"""Phase 0 answer key. The deck engine is not changed here.

Open cases are still missed by the current engine and xfail. A case that starts
passing fails this file until its ``open`` flag is set to false in the answer key.
"""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from unittest.mock import patch

import pytest
import yaml
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from agetic_cdd_api.report_builder_base import BuildContext, DealProfile
from agetic_cdd_api.report_cdd_evidence import EvidenceBook, _absorb_vdr, build_agent_evidence
from agetic_cdd_api.report_cdd_factcheck import contradicts_profit_return, drop_mislabeled_contract_revenue
from agetic_cdd_api.report_cdd_skill import _as_waterfall, _margin_title, build_skill_content
from agetic_cdd_api.report_storyline import default_storyline

KEY_PATH = Path(__file__).parent / "fixtures" / "cdd_phase0_answer_key.yaml"
ANSWER_KEY = yaml.safe_load(KEY_PATH.read_text())
DEALS = {deal["id"]: deal for deal in ANSWER_KEY["deals"]}
CASES = {case["id"]: case for case in ANSWER_KEY["cases"]}


def _money(amount: float) -> str:
    return f"${amount:.2f}M"


def _vdr_fact(fid: str, section: str, row: str, cells: dict, doc: str) -> dict:
    return {
        "id": fid,
        "module": "HPF",
        "module_title": "Data room",
        "section_title": section,
        "row_label": row,
        "text": row + " " + " ".join(f"{key} {value}" for key, value in cells.items()),
        "cells": {str(key): str(value) for key, value in cells.items()},
        "citations": {"doc": [doc], "web": [], "computed": False},
        "source": "vdr_direct",
        "basis": "data_room_doc",
        "grade": "established",
        "class": "fact",
        "flags": [],
        "period_status": {},
    }


def _descr(shape, text: str) -> None:
    for node in shape._element.iter():
        if node.tag.endswith("cNvPr"):
            node.set("descr", text)
            return


def _write_cim(path: Path, deal: dict) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    history = prs.slides.add_slide(blank)
    box = history.shapes.add_textbox(Inches(0.4), Inches(0.4), Inches(12), Inches(2))
    box.text_frame.text = (
        "FOUNDED 2014 Started collection in the home city "
        "ACQUISITION 2018 Bought the local routes "
        "DIVERSIFICATION 2020 Added processing "
        "ACCELERATION Close Series B investment and open a new city "
        "EXPANSION 2024 Opened two further routes"
    )
    claims = prs.slides.add_slide(blank)
    claim = claims.shapes.add_textbox(Inches(0.4), Inches(0.4), Inches(12), Inches(2))
    rate = f"{deal['monthly_retention'] * 100:.0f}"
    claim.text_frame.text = (
        f"The CIM states {rate}% monthly retention and 20 municipal partners. "
        f"The model prices homes at ${deal['model_price_year']:.0f} a year."
    )
    for shape, descr in (
        (box, "River Haul logo"),
        (claim, "Bay Organics brand"),
    ):
        _descr(shape, descr)
    extra = history.shapes.add_textbox(Inches(0.4), Inches(3), Inches(4), Inches(0.4))
    extra.text_frame.text = "Local names"
    _descr(extra, "North Cart logo")
    prs.save(path)


def _outputs(deal: dict) -> dict:
    decoy = deal["decoy"]
    return {
        "historical_performance": {
            "spec": {
                "period_facts": [
                    {"metric": "Revenue", "period": str(decoy["year"]), "period_type": "actual", "value": decoy["amount"]},
                    {"metric": "Revenue", "period": "2024", "period_type": "actual", "value": deal["accounts"][2024]},
                    {"metric": "Revenue", "period": "2025", "period_type": "actual", "value": deal["accounts"][2025]},
                ],
                "findings": [
                    deal["stale_margin"],
                    "8 period(s) from accounts. 0 CIM gap(s)",
                    "Databook promoted · 4 P&L line(s)",
                    "EBITDA returned to profit by FY2025.",
                ],
                "primary_sources": ["Management accounts"],
            }
        },
        "market_volume_and_growth": {
            "spec": {
                "tam": {"value": deal["tam"], "scale": "million"},
                "sam": {"value": deal["sam"], "scale": "million"},
                "som": {"value": deal["som"], "scale": "million"},
                "primary_sources": ["Company workbook"],
            }
        },
        "competitive_differentiation": {
            "spec": {
                "self_assessment": {
                    "columns": ["Company", "National hauler"],
                    "rows": [
                        {"criterion": "Local service", "Company": "Yes", "National hauler": "No"},
                        {"criterion": "Own processing", "Company": "Yes", "National hauler": "Yes"},
                    ],
                }
            }
        },
    }


def _facts(deal: dict) -> list[dict]:
    doc = "management_accounts.xlsx"
    accounts = {year: _money(amount) for year, amount in deal["accounts"].items()}
    accounts[f"{min(deal['ytd'])} YTD"] = _money(deal["ytd"][min(deal["ytd"])])
    accounts[f"{max(deal['ytd'])} YTD"] = _money(deal["ytd"][max(deal["ytd"])])
    facts = [
        _vdr_fact(
            f"{deal['id']}-rev",
            "Statement of operations (management accounts)",
            "Total Revenues",
            accounts,
            doc,
        )
    ]
    for name, amount in deal["costs"].items():
        facts.append(_vdr_fact(
            f"{deal['id']}-cost-{len(facts)}",
            "Statement of operations (management accounts)",
            name,
            {"2025": _money(amount)},
            doc,
        ))
    facts.append(_vdr_fact(
        f"{deal['id']}-ebitda",
        "Statement of operations (management accounts)",
        "EBITDA",
        {"2024": _money(deal["ebitda"][2024]), "2025": _money(deal["ebitda_line"])},
        doc,
    ))
    for metric, series in (("Gross margin", deal["gross_margin"]), ("EBITDA margin", deal["ebitda_margin"])):
        facts.append(_vdr_fact(
            f"{deal['id']}-margin-{metric.split()[0].lower()}",
            "Statement of operations (management accounts)",
            metric,
            {year: f"{value:.1f}%" for year, value in series.items()},
            doc,
        ))
    facts.append(_vdr_fact(
        f"{deal['id']}-wb",
        "Financial summary (company workbook)",
        "Revenue",
        {year: _money(amount) for year, amount in deal["workbook"].items()},
        "company_workbook.xlsx",
    ))
    decoy = deal["decoy"]
    facts.append(_vdr_fact(
        f"{deal['id']}-contract",
        "Municipal contracts DC contract upside spectrum",
        "Modeled",
        {str(decoy["contract_year"]): _money(decoy["amount"])},
        "company_workbook.xlsx",
    ))
    for month, subscriber, revenue in deal["retention"]:
        facts.append(_vdr_fact(
            f"{deal['id']}-ret-{month}",
            "Customer retention summary",
            str(month),
            {"Subscriber retention": f"{subscriber:.1f}%", "Revenue retention": f"{revenue:.1f}%"},
            "company_workbook.xlsx",
        ))
    prices = deal["prices"]
    facts.append(_vdr_fact(
        f"{deal['id']}-arpa",
        "Unit economics",
        "ARPA ($ per month)",
        {
            "Commercial": f"${prices['commercial']:.2f}",
            "Municipal": f"${prices['municipal']:.2f}",
            "Residential": f"${prices['residential']:.2f}",
        },
        "company_workbook.xlsx",
    ))
    facts.append(_vdr_fact(
        f"{deal['id']}-locations",
        "Unit economics",
        "Active locations",
        {"Commercial": "100", "Municipal": "40", "Residential": "800"},
        "company_workbook.xlsx",
    ))
    segment = {
        "2025": _money(deal["accounts"][2025] * 0.2),
        "2030": _money(deal["workbook"][2030] * 0.3),
        "2025 Mix": deal["mix"],
    }
    facts.append(_vdr_fact(
        f"{deal['id']}-segment",
        "Revenue by segment",
        "Municipal - DC",
        segment,
        "company_workbook.xlsx",
    ))
    facts.append(_vdr_fact(
        f"{deal['id']}-utica",
        "Revenue by segment",
        "Utica",
        {"2026": _money(deal["accounts"][2025] * 0.1), "2030": _money(deal["workbook"][2030] * 0.1)},
        "company_workbook.xlsx",
    ))
    return facts


def _ctx(deal: dict) -> BuildContext:
    return BuildContext(
        deal_slug=deal["slug"],
        report_type="cdd_deck",
        profile=DealProfile(deal_slug=deal["slug"], company=deal["company"], sector="waste"),
        storyline=default_storyline("cdd_deck"),
        agent_outputs=_outputs(deal),
    )


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    root = tmp_path_factory.mktemp("phase0-deals")
    prepared = {}
    for deal in ANSWER_KEY["deals"]:
        docs = root / deal["slug"] / "documents"
        docs.mkdir(parents=True)
        _write_cim(docs / "CIM_memorandum.pptx", deal)
        Workbook().save(docs / "management_accounts.xlsx")
        ctx = _ctx(deal)
        book = build_agent_evidence(ctx)
        _absorb_vdr(book, _facts(deal))
        prepared[deal["id"]] = (ctx, book)
    with patch("agetic_cdd_api.services_deals.deals_root", return_value=root):
        decks = {}
        for deal_id, (ctx, book) in prepared.items():
            content, _withheld = build_skill_content(ctx, book)
            decks[deal_id] = {"content": content, "book": book, "deal": DEALS[deal_id]}
        yield decks


def _page(deck: dict, sid: str) -> dict | None:
    page = deck["content"]["slides"].get(sid)
    return page if isinstance(page, dict) else None


def _text(page: dict | None) -> str:
    return json.dumps(page or {}, ensure_ascii=False)


def _all_text(deck: dict) -> str:
    return json.dumps(deck["content"]["slides"], ensure_ascii=False)


def _table(page: dict, name: str) -> tuple[list, list]:
    spec = (page.get("elements") or {}).get(name) or {}
    return list(spec.get("columns") or []), list(spec.get("rows") or [])


def _money_in(text: str) -> list[float]:
    return [float(match) for match in re.findall(r"\d+(?:\.\d+)?", text.replace(",", ""))]


def _check(case_id: str, deck: dict) -> tuple[str, str]:
    deal = deck["deal"]
    slides = deck["content"]["slides"]
    if case_id == "N-diverge":
        page = _page(deck, "S048_downturn_resilience")
        if page is None:
            return "missing", "retention page was not built"
        gap = max(abs(sub - rev) for _month, sub, rev in deal["retention"])
        if gap > 1:
            return "missing", "fixture curves are more than 1 point apart"
        if re.search(r"(?i)diverge", _text(page)):
            return "fail", "title says the curves diverge"
        return "pass", ""
    if case_id == "N-returned-to-profit":
        page = _page(deck, "S009_company_snapshot")
        if page is None:
            return "missing", "snapshot was not built"
        if min(deal["ebitda"].values()) <= 0:
            return "missing", "fixture EBITDA is not already positive"
        if re.search(r"(?i)returned to profit", _text(page)):
            return "fail", "snapshot says EBITDA returned to profit"
        return "pass", ""
    if case_id == "N-evidenced-model":
        page = _page(deck, "S021_h1_market_size")
        if page is None:
            return "missing", "market page was not built"
        if re.search(r"(?i)evidenced model", _text(page)):
            return "fail", "workbook sizing is called an evidenced model"
        return "pass", ""
    if case_id == "N-cost-tie":
        page = _page(deck, "S011_value_chain")
        if page is None:
            return "missing", "cost page was not built"
        _columns, rows = _table(page, "table")
        shown = 0.0
        for row in rows:
            amounts = _money_in(row[1])
            if not amounts:
                return "fail", f"cost row has no amount: {row}"
            shown += amounts[0]
        total = shown + deal["ebitda_line"]
        if abs(total - deal["accounts"][2025]) > 0.01:
            return "fail", f"costs {shown:.2f} plus EBITDA {deal['ebitda_line']:.2f} = {total:.2f}, revenue is {deal['accounts'][2025]:.2f}"
        return "pass", ""
    if case_id == "N-decoy-revenue":
        year = deal["decoy"]["year"]
        amount = f"{deal['decoy']['amount']:.2f}"
        for fact in deck["book"].facts:
            text = str(fact.get("text") or "")
            if re.search(rf"Revenue {year} \(actual\).*{re.escape(amount)}", text):
                return "fail", f"fact store has Revenue {year} (actual) for the modelled contract amount"
        return "pass", ""
    if case_id == "N-wrong-column":
        page = _page(deck, "S072_growth_levers_overview")
        if page is None:
            return "missing", "levers page was not built"
        columns, rows = _table(page, "rows")
        if "Not established" not in columns:
            return "pass", ""
        index = columns.index("Not established")
        for row in rows:
            cell = row[index] if index < len(row) else ""
            if re.search(r"\d", cell) and "%" in cell:
                return "fail", f"{cell} is under Not established"
        return "pass", ""
    if case_id == "N-date-column":
        page = _page(deck, "S010_platform_evolution")
        if page is None:
            return "missing", "history page was not built"
        _columns, rows = _table(page, "flow")
        bad = [row[0] for row in rows if not re.fullmatch(r"(19|20)\d\d", row[0])]
        if bad:
            return "fail", f"date column holds {bad}"
        return "pass", ""
    if case_id == "N-plan-as-milestone":
        page = _page(deck, "S010_platform_evolution")
        if page is None:
            return "missing", "history page was not built"
        if re.search(r"Close Series B", _text(page)):
            return "fail", "Close Series B is listed as a milestone"
        return "pass", ""
    if case_id == "N-picture-title":
        page = _page(deck, "S062_criteria_vs_archetypes")
        if page is None:
            return "missing", "competitor matrix was not built"
        title = str((page.get("title") or {}).get("text") or "")
        if re.search(r"(?i)pictures on the CIM|are described as", title):
            return "fail", "title is a picture description"
        return "pass", ""
    if case_id == "N-mixed-price-axis":
        page = _page(deck, "S064_price_benchmarking")
        if page is None:
            return "missing", "price page was not built"
        chart = (page.get("elements") or {}).get("chart") or {}
        values = [value for series in chart.get("series") or [] for value in series.get("values") or []]
        if len(values) >= 2 and min(values) > 0 and max(values) / min(values) > 8:
            return "fail", f"one axis plots {min(values)} and {max(values)}"
        return "pass", ""
    if case_id == "N-period-debug":
        blob = _all_text(deck)
        if re.search(r"period\(s\) from accounts|CIM gap\(s\)", blob):
            return "fail", "debug period text is on a page"
        return "pass", ""
    if case_id == "N-stale-margin-lead":
        page = _page(deck, "S009_company_snapshot")
        if page is None:
            return "missing", "snapshot was not built"
        bullets = ((page.get("elements") or {}).get("left") or {}).get("items") or []
        lead = str((bullets[0] or {}).get("text") or "") if bullets else ""
        if "2023" in lead and re.search(r"(?i)margin", lead):
            return "fail", "snapshot leads with the 2023 margin"
        return "pass", ""
    if case_id == "N-databook-promoted":
        if re.search(r"(?i)databook promoted", _all_text(deck)):
            return "fail", "databook promoted is on a page"
        return "pass", ""
    if case_id == "N-margin-chart":
        page = _page(deck, "S045_gf4_pricing_stability")
        if page is None:
            return "missing", "margin page was not built"
        chart = (page.get("elements") or {}).get("chart") or {}
        categories = " ".join(str(item) for item in chart.get("categories") or [])
        series = chart.get("series") or []
        if (
            chart.get("chart_type", "column") == "column"
            and len(series) == 1
            and re.search(r"(?i)gross", categories)
            and re.search(r"(?i)ebitda", categories)
        ):
            return "fail", "gross and EBITDA margins are one column series"
        return "pass", ""
    if case_id == "N-bridge-waterfall":
        page = _page(deck, "S046_sam_growth_build")
        if page is None:
            return "missing", "bridge page was not built"
        chart = (page.get("elements") or {}).get("right") or {}
        if chart.get("type") == "chart" and chart.get("chart_type") != "waterfall":
            return "fail", f"bridge chart is {chart.get('chart_type')}"
        return "pass", ""
    if case_id == "N-bridge-basis":
        page = _page(deck, "S046_sam_growth_build")
        if page is None:
            return "missing", "bridge page was not built"
        blob = _text(page)
        actual = f"{deal['accounts'][2025]:.2f}"
        workbook_actual = f"{deal['workbook'][2025]:.2f}"
        plan = f"{deal['workbook'][2030]:.2f}"
        same_accounts = actual in blob and plan not in blob
        same_workbook = workbook_actual in blob and plan in blob and actual not in blob
        if not same_accounts and not same_workbook:
            return "fail", f"bridge mixes accounts actual {actual} with workbook plan {plan}"
        return "pass", ""
    if case_id == "P-ytd":
        page = _page(deck, "S009_company_snapshot")
        if page is None:
            return "missing", "snapshot was not built"
        head = str((page.get("title") or {}).get("text") or "") + " " + str((page.get("subtitle") or {}).get("text") or "")
        if re.search(r"(?i)year-to-date|same months", head):
            return "pass", ""
        return "fail", "snapshot does not compare year-to-date with the same months last year"
    if case_id == "P-margin-compare":
        page = _page(deck, "S045_gf4_pricing_stability")
        if page is None:
            return "missing", "margin page was not built"
        title = str((page.get("title") or {}).get("text") or "")
        if re.search(r"(?i)gross", title) and re.search(r"(?i)ebitda", title):
            return "pass", ""
        return "fail", "margin title does not compare gross and EBITDA"
    if case_id == "P-trading":
        page = _page(deck, "S047_scenario_analysis")
        if page is None:
            return "missing", "scenario page was not built"
        blob = _text(page).lower()
        plan = f"{deal['workbook'][2030]:.2f}"
        if plan in _text(page) and ("trading" in blob or "growth forward" in blob):
            return "pass", ""
        return "fail", "scenario page does not set a trading case against the plan"
    if case_id == "P-model-vs-billed":
        annual = f"{deal['model_price_year']:.0f}"
        monthly = f"{deal['prices']['residential']:.2f}".rstrip("0").rstrip(".")
        for page in slides.values():
            if not isinstance(page, dict):
                continue
            blob = _text(page)
            if annual in blob and monthly in blob and re.search(r"(?i)against|versus|\bvs\b|below|above", blob):
                return "pass", ""
        return "fail", "model price per year is not set against the billed monthly price"
    if case_id == "P-annualize":
        annual = round((deal["monthly_retention"] ** 12) * 100, 1)
        if re.search(rf"{annual:.1f}", _all_text(deck)):
            return "pass", ""
        return "fail", f"monthly retention is not shown as {annual:.1f}% a year"
    if case_id == "P-competitors":
        blob = _all_text(deck)
        if re.search(r"No outside competitor research was loaded", blob):
            return "pass", ""
        web = len(re.findall(r"outside the data room", blob, re.I))
        names = re.findall(r"[A-Z][A-Za-z]+(?: [A-Z][A-Za-z]+){0,2}", blob)
        if web and len(set(names)) >= 3 and re.search(r"\$\d+", blob):
            # The price and the outside-the-room label must be on the competitor page.
            page = _page(deck, "S060_competitor_archetypes")
            page_text = _text(page)
            named = re.findall(r"[A-Z][A-Za-z]+(?: [A-Z][A-Za-z]+)?", page_text)
            if page and re.search(r"outside the data room", page_text, re.I) and len(set(named)) >= 3 and re.search(r"\$\d+", page_text):
                return "pass", ""
        return "fail", "no outside competitors and no explicit research gap"
    if case_id == "P-graded-summary":
        page = _page(deck, "S020_highlights_summary")
        blob = _text(page)
        columns, _rows = _table(page or {}, "rows")
        if page and "Grade" in columns and re.search(r"(?i)test before signing", blob):
            return "pass", ""
        return "fail", "highlights do not open with graded findings"
    if case_id == "P-insight-title":
        page = _page(deck, "S045_gf4_pricing_stability")
        if page is None:
            return "missing", "margin page was not built"
        title = str((page.get("title") or {}).get("text") or "")
        if re.fullmatch(r"Margin on record is \d+(?:\.\d+)?% in 20\d\d", title):
            return "fail", "title only restates one margin"
        if re.search(r"(?i)gross", title) and re.search(r"(?i)ebitda", title):
            return "pass", ""
        return "fail", "title does not state the margin finding"
    return "missing", f"no checker for {case_id}"


def _open(case: dict, deal_id: str) -> bool:
    if not case.get("open"):
        return False
    deals = case.get("open_deals")
    return deal_id in deals if deals else True


def test_answer_key_covers_three_deals_and_every_case():
    assert [deal["id"] for deal in ANSWER_KEY["deals"]] == ["harbor", "northwind", "compost-crew"]
    assert [deal["company"] for deal in ANSWER_KEY["deals"]] == [
        "Harbor Bend",
        "Northwind Organics",
        "Compost Crew",
    ]
    assert len(CASES) == 24
    for case in ANSWER_KEY["cases"]:
        assert case["fix"]
        assert case["rule"]
        assert case["polarity"] in {"negative", "positive"}


@pytest.mark.parametrize("deal_id", list(DEALS))
def test_answer_key_cost_fixtures_balance(deal_id: str):
    deal = DEALS[deal_id]
    total = sum(deal["costs"].values()) + deal["ebitda_line"]
    assert abs(total - deal["accounts"][2025]) <= 0.01


@pytest.mark.parametrize("deal_id", list(DEALS))
@pytest.mark.parametrize("case_id", list(CASES))
def test_phase0_case(case_id: str, deal_id: str, built):
    case = CASES[case_id]
    status, detail = _check(case_id, built[deal_id])
    if status == "missing":
        pytest.fail(f"{case_id} fixture for {deal_id}: {detail}")
    if _open(case, deal_id):
        if status == "pass":
            pytest.fail(f"{case_id} now passes for {deal_id}. Set open to false in the answer key.")
        pytest.xfail(f"{case['fix']}: {detail}")
    assert status == "pass", f"{case_id} for {deal_id}: {detail}"


def test_profit_check_keeps_the_sign_on_thousands():
    book = EvidenceBook()
    book.facts = [{
        "id": "E1",
        "row_label": "EBITDA",
        "cells": {"2024": "-$722k", "2025": "$0.57M"},
    }]
    assert contradicts_profit_return("EBITDA returned to profit by FY2025.", book) is False
    book.facts[0]["cells"]["2024"] = "572k"
    assert contradicts_profit_return("EBITDA returned to profit by FY2025.", book) is True


def test_drop_mislabeled_contract_revenue_logs_the_fact(caplog):
    book = EvidenceBook()
    book.facts = [
        {
            "id": "REV",
            "row_label": "Total Revenues",
            "section_title": "Statement of operations",
            "text": "Total Revenues",
            "cells": {"2024": "$11.87M", "2025": "$13.41M"},
        },
        {
            "id": "SCEN",
            "row_label": "Modeled",
            "section_title": "DC contract upside spectrum",
            "text": "Modeled",
            "cells": {"2030": "10420k"},
        },
        {
            "id": "BAD",
            "row_label": "Revenue",
            "section_title": "agent",
            "text": "Revenue 2020 (actual): $10.42M",
            "cells": {},
        },
    ]
    book.by_id = {fact["id"]: fact for fact in book.facts}
    book.buckets = {"revenue_actual": ["BAD", "REV"]}
    with caplog.at_level(logging.INFO, logger="agetic_cdd_api.report_cdd_factcheck"):
        drop_mislabeled_contract_revenue(book)
    assert [fact["id"] for fact in book.facts] == ["REV", "SCEN"]
    assert "BAD" not in book.by_id
    assert book.buckets["revenue_actual"] == ["REV"]
    assert "BAD" in caplog.text
    assert "2020" in caplog.text


def test_waterfall_hook_converts_only_a_start_and_an_end(caplog):
    start_end = {
        "type": "chart",
        "chart_type": "column",
        "categories": ["2025", "2030P"],
        "series": [{"name": "Revenue", "values": [13.28, 50.17]}],
        "point_refs": [["HPF-A", "HPF-B"]],
    }
    longer = {
        "type": "chart",
        "chart_type": "column",
        "categories": ["2024", "2025", "2030P"],
        "series": [{"name": "Revenue", "values": [11.87, 13.28, 50.17]}],
    }
    with caplog.at_level(logging.INFO, logger="agetic_cdd_api.report_cdd_skill"):
        _as_waterfall(start_end)
        _as_waterfall(longer)
    assert start_end["chart_type"] == "waterfall"
    assert start_end["categories"] == ["2025", "Bridge", "2030P"]
    assert start_end["series"][0]["values"] == [13.28, 36.89, 50.17]
    assert start_end["totals"] == [0, 2]
    assert start_end["point_refs"] == [["HPF-A", "CALC-BRIDGE-DELTA", "HPF-B"]]
    assert longer["chart_type"] == "column"
    assert longer["categories"] == ["2024", "2025", "2030P"]
    assert "not a start and an end" in caplog.text


def test_margin_title_skips_a_cell_that_is_not_a_number(caplog):
    margins = [
        {"metric": "Gross margin", "label": "2024", "value": 31.8},
        {"metric": "Gross margin", "label": "2025", "value": "N/A"},
        {"metric": "Gross margin", "label": "2026", "value": 30.0},
        {"metric": "EBITDA margin", "label": "2024", "value": 9.6},
        {"metric": "EBITDA margin", "label": "2025", "value": "—"},
        {"metric": "EBITDA margin", "label": "2026", "value": 4.3},
    ]
    with caplog.at_level(logging.INFO, logger="agetic_cdd_api.report_cdd_skill"):
        title = _margin_title(margins)
    assert title == (
        "Gross margin fell from 31.8% to 30.0% while EBITDA margin fell from 9.6% to 4.3%"
    )
    assert "N/A" in caplog.text
