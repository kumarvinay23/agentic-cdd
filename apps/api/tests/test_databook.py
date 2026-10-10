"""Databook Phase 0/1 unit tests — mapping, conflicts, data-quality shape."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from agetic_cdd_api.app import app
from agetic_cdd_api.services_databook_extract import rows_from_table
from agetic_cdd_api.services_databook_map import map_caption, same_family_allowed
from agetic_cdd_api.services_databook_models import (
    DealDatabookParams,
    ExtractedRow,
    MetricFamily,
    ProofLevel,
    RowStatus,
)
from agetic_cdd_api.services_databook_resolve import (
    apply_series_drops,
    detect_conflicts,
)


FIXTURE = (
    Path(__file__).resolve().parents[3]
    / "docs"
    / "reference"
    / "diligenceiq-data-quality-test2-fixture.json"
)


def test_atomic_write_and_skip_invalid_rows(tmp_path: Path, caplog) -> None:
    import logging

    from agetic_cdd_api.services_databook_store import _load_model_list, _read_json, _write_json

    path = tmp_path / "extract.json"
    _write_json(
        path,
        {
            "rows": [
                {
                    "row_id": "ok1",
                    "doc_id": "d",
                    "source_name": "a.pdf",
                    "caption": "Revenue",
                    "value": 100.0,
                },
                {"row_id": "bad", "caption": "missing required fields"},
            ]
        },
    )
    assert path.is_file()
    assert list(tmp_path.glob(".*.tmp")) == []
    assert _read_json(path, {})["rows"][0]["row_id"] == "ok1"

    deal = type("Deal", (), {"slug": "tmp-deal", "id": "x"})()
    # Point loader at tmp_path by monkeypatching databook_dir via writing under a fake deal folder is heavy;
    # exercise _load_model_list through a thin wrap: write then read via helper with patched path.
    from agetic_cdd_api import services_databook_store as store

    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        with caplog.at_level(logging.WARNING):
            rows = _load_model_list(deal, "extract.json", "rows", ExtractedRow)  # type: ignore[arg-type]
        assert len(rows) == 1
        assert rows[0].row_id == "ok1"
        assert any("Failed to validate ExtractedRow" in r.message for r in caplog.records)
    finally:
        store.databook_dir = original


def _headers(client: TestClient) -> dict[str, str]:
    res = client.post(
        "/api/v1/auth/login",
        json={"email": "admin@ageticcdd.com", "password": "adminpass"},
    )
    token = res.json()["data"]["tokens"]["accessToken"]
    return {"Authorization": f"Bearer {token}"}


def test_map_family_gates() -> None:
    assert map_caption("Revenue (INR Cr)").metric_key == "revenue"
    assert map_caption("Units Sold (000s)").metric_key == "units_sold"
    assert map_caption("Net Revenue Retention (NRR)").metric_key == "nrr"
    assert map_caption("Gross Revenue Retention (GRR)").metric_key == "grr"
    assert map_caption("Revenue Churn").metric_key == "revenue_churn"
    assert map_caption("YoY Growth").metric_key == "yoy_growth"

    assert same_family_allowed("Revenue (INR Cr)", "revenue")
    assert not same_family_allowed("Units Sold", "revenue")
    assert not same_family_allowed("Net Revenue Retention (NRR)", "revenue")
    assert not same_family_allowed("Revenue Churn", "revenue")


def test_series_drop_ratio_like_values() -> None:
    rows = [
        ExtractedRow(
            row_id=f"r{i}",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=v,
        )
        for i, v in enumerate([31_400_000_000.0, 49_000_000_000.0, 14_680_000_000.0, 81.0, 74.0])
    ]
    out, issues = apply_series_drops(rows, params=DealDatabookParams())
    dropped_vals = {r.value for r in out if r.status == RowStatus.DROPPED}
    assert 81.0 in dropped_vals
    assert 74.0 in dropped_vals
    assert 31_400_000_000.0 not in dropped_vals
    assert any(i["kind"] == "dropped" for i in issues)
    assert "orders of magnitude" in issues[0]["reason"]


def test_units_sold_cannot_win_revenue_conflict() -> None:
    """Anti-goal from live DIQ: Units Sold 320 must not be a Revenue candidate."""
    rows = [
        ExtractedRow(
            row_id="u1",
            doc_id="a",
            source_name="01_Executive_Summary.pdf",
            caption="Units Sold",
            metric_key="units_sold",
            metric_family=MetricFamily.UNITS,
            fiscal_year=2024,
            value=320.0,
        ),
        ExtractedRow(
            row_id="u2",
            doc_id="b",
            source_name="04_Commercial_Due_Diligence.pdf",
            caption="Units Sold",
            metric_key="units_sold",
            metric_family=MetricFamily.UNITS,
            fiscal_year=2024,
            value=320.0,
        ),
        ExtractedRow(
            row_id="r1",
            doc_id="b",
            source_name="04_Commercial_Due_Diligence.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=31_400_000_000.0,
        ),
        ExtractedRow(
            row_id="r2",
            doc_id="a",
            source_name="01_Executive_Summary.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=49_000_000_000.0,
        ),
        ExtractedRow(
            row_id="n1",
            doc_id="c",
            source_name="13_Valuation.pdf",
            caption="Net Revenue Retention (NRR)",
            metric_key="nrr",
            metric_family=MetricFamily.RETENTION,
            fiscal_year=2024,
            value=81.0,
        ),
    ]
    issues, hold_ids = detect_conflicts(rows)
    revenue_conflicts = [i for i in issues if i.get("metric") in {"Revenue", "revenue"} and i.get("fiscal_year") == 2024]
    assert revenue_conflicts, "expected Revenue FY2024 conflict across docs"
    candidates = revenue_conflicts[0]["candidates"]
    values = {c["value"] for c in candidates}
    assert 320.0 not in values
    assert 81.0 not in values
    assert 31_400_000_000.0 in values
    assert 49_000_000_000.0 in values
    chosen = next(c for c in candidates if c["chosen"])
    assert chosen["value"] in {31_400_000_000.0, 49_000_000_000.0}
    assert hold_ids
    assert sum(1 for c in candidates if c["chosen"]) == 1


def test_conflict_rank_tie_and_float_chosen() -> None:
    from agetic_cdd_api.services_databook_resolve import apply_statuses_and_promote

    # Identical ranking dimensions except value centrality — must not TypeError
    rows = [
        ExtractedRow(
            row_id="a",
            doc_id="d1",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0000001,
        ),
        ExtractedRow(
            row_id="b",
            doc_id="d2",
            source_name="b.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=200.0,
        ),
    ]
    issues, hold_ids = detect_conflicts(rows)
    assert len(issues) == 1
    assert sum(1 for c in issues[0]["candidates"] if c["chosen"]) == 1
    assert hold_ids == {"a", "b"}


def test_promote_merges_identical_multi_source_rows() -> None:
    from agetic_cdd_api.services_databook_models import ProofLevel
    from agetic_cdd_api.services_databook_resolve import apply_statuses_and_promote

    rows = [
        ExtractedRow(
            row_id="a",
            doc_id="d1",
            source_name="zzz.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=3140.0,
            assumption=False,
            dual_agree=True,
            proof_level=ProofLevel.L2,
            proof_checks=["dual_agree"],
        ),
        ExtractedRow(
            row_id="b",
            doc_id="d2",
            source_name="aaa.pdf",
            caption="Revenue (INR Cr)",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=3140.0,
            currency="INR",
            scale="Cr",
            assumption=False,
            dual_agree=True,
            proof_level=ProofLevel.L2,
            proof_checks=["dual_agree"],
        ),
    ]
    out, promoted = apply_statuses_and_promote(rows, hold_ids=set())
    assert len(promoted) == 1
    assert set(promoted[0].sources) == {"zzz.pdf", "aaa.pdf"}
    assert promoted[0].currency == "INR"
    assert promoted[0].scale == "Cr"
    assert promoted[0].proof_level == ProofLevel.L3  # multi-source bump
    assert all(r.status == RowStatus.PROMOTED for r in out)


def test_conflict_emitted_when_plainness_zero() -> None:
    import agetic_cdd_api.services_databook_resolve as resolve_mod

    original = resolve_mod.caption_plainness
    resolve_mod.caption_plainness = lambda caption, metric_key: 0.0  # type: ignore[assignment]
    try:
        issues, hold = detect_conflicts(
            [
                ExtractedRow(
                    row_id="a",
                    doc_id="d1",
                    source_name="a.pdf",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=100.0,
                ),
                ExtractedRow(
                    row_id="b",
                    doc_id="d2",
                    source_name="b.pdf",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=200.0,
                ),
            ]
        )
        assert len(issues) == 1
        assert hold == {"a", "b"}
        assert sum(1 for c in issues[0]["candidates"] if c["chosen"]) == 1
    finally:
        resolve_mod.caption_plainness = original


def test_series_drop_negative_centre() -> None:
    rows = [
        ExtractedRow(
            row_id=f"r{i}",
            doc_id="d",
            source_name="a.pdf",
            caption="EBITDA",
            metric_key="ebitda",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=v,
        )
        for i, v in enumerate([-1000.0, -1100.0, -900.0, -1.0])
    ]
    out, issues = apply_series_drops(rows, params=DealDatabookParams())
    dropped = {r.value for r in out if r.status == RowStatus.DROPPED}
    assert -1.0 in dropped
    assert issues


def test_series_drop_skips_opposite_sign_adjustments() -> None:
    """Small reclass/adjustment must not be series-dropped vs a positive centre."""
    rows = [
        ExtractedRow(
            row_id=f"r{i}",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2020 + i,
            value=v,
        )
        for i, v in enumerate([-0.05, 100.0, 105.0, 110.0])
    ]
    out, issues = apply_series_drops(rows, params=DealDatabookParams())
    dropped = {r.value for r in out if r.status == RowStatus.DROPPED}
    assert -0.05 not in dropped
    assert not any(abs(float((i.get("detail") or {}).get("value") or 0)) < 1 for i in issues)


def test_assumption_reason_flags_not_substring_gated() -> None:
    from agetic_cdd_api.services_databook_resolve import (
        _dedupe_issues,
        assumption_issues_from_rows,
    )

    rows = [
        ExtractedRow(
            row_id="a1",
            doc_id="1",
            source_name="fin.pdf",
            caption="Other income",
            metric_key="other_income",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=12.0,
            assumption=True,
            status=RowStatus.CANDIDATE,
        )
    ]
    assumed = assumption_issues_from_rows(rows)
    assert len(assumed) == 1
    reason = assumed[0]["reason"]
    assert "no currency stated" in reason
    assert "unit/currency inferred or missing" in reason

    # Dedupe must collide on metric_key even when one side omits row_id.
    a = {
        "kind": "scale_step",
        "metric_key": "revenue",
        "metric": "Revenue",
        "fiscal_year": 2024,
        "value": 100.0,
        "row_id": "r1",
    }
    b = {
        "kind": "scale_step",
        "metric_key": "revenue",
        "fiscal_year": 2024,
        "value": 100.0,
    }
    assert len(_dedupe_issues([a, b])) == 1


def test_rows_from_year_table() -> None:
    table = {
        "name": "pnl",
        "rows": [
            ["Metric", "FY2022", "FY2023", "FY2024"],
            ["Revenue (INR Cr)", "190", "1468", "3140"],
            ["Units Sold (000s)", "20.1", "152.1", "320"],
        ],
    }
    rows = rows_from_table(source_name="cdd.pdf", doc_id="cdd", table=table)
    metrics = {(r.metric_key, r.fiscal_year, r.value) for r in rows}
    assert ("revenue", 2024, 3140.0) in metrics
    assert ("units_sold", 2022, 20.1) in metrics
    rev = next(r for r in rows if r.metric_key == "revenue" and r.fiscal_year == 2024)
    assert rev.currency == "INR"
    assert rev.scale == "Cr"
    # Caption-sourced money scale is kept but flagged (header wins when present).
    assert rev.scale_source == "caption"
    assert rev.assumption is True
    assert rev.period_length == "FY"
    assert rev.period_end == "2024-12-31"
    assert rev.source_ref is not None


def test_parse_number_negatives_and_percents() -> None:
    from agetic_cdd_api.services_databook_extract import (
        _detect_year_headers,
        _infer_unit,
        _iter_all_text_facts,
        _iter_text_facts,
        _parse_number,
    )

    assert _parse_number("(12.5%)") == -12.5
    assert _parse_number("-150.00") == -150.0
    assert _parse_number(" - 100 ") == -100.0
    assert _parse_number("(1,234.5)") == -1234.5
    assert _parse_number("81.0") == 81.0
    assert _parse_number("N/M") is None
    assert _parse_number("100.50.") == 100.5
    assert _parse_number("100.50%") == 100.5
    # Hyphenated ranges must not become concatenated negatives.
    assert _parse_number("10-12") is None
    assert _parse_number("10–12") is None

    prose = "Revenue in 2024 was 100.50."
    facts = _iter_text_facts(prose)
    assert facts == [("Revenue", 2024, 100.5)]
    # Long line without a trailing value must not hang / explode combinations
    long_noise = "Revenue " + ("x" * 200) + " and more prose without a year-value pair"
    assert _iter_text_facts(long_noise) == []

    # Flattened P&L: years precede the metric name (common in pypdf diligence extracts).
    flat = (
        "Line Item FY2021 FY2022 FY2023 FY2024E FY2025E "
        "Revenue 8 456 2,630 4,900 8,200 Cost of Goods Sold (11) (420) (2,443) (4,288) (6,560)"
    )
    series = _iter_all_text_facts(flat)
    assert ("Revenue", 2024, 4900.0) in series
    assert ("Revenue", 2021, 8.0) in series
    units = "FY2022 FY2023 FY2024E YoY Growth Units Sold (000s) ~20 ~152 ~320 +110.5%"
    unit_facts = [(c, y, v) for c, y, v in _iter_all_text_facts(units) if "units" in c.lower()]
    assert (2022, 20.0) in {(y, v) for _, y, v in unit_facts}
    assert (2024, 320.0) in {(y, v) for _, y, v in unit_facts}

    years = _detect_year_headers(["Metric", "FY2024", "Notes (2001)", "Amount 4500", "2023"])
    assert years[1] == 2024
    assert 2 not in years  # note reference, not a column year
    assert 3 not in years  # 4500 is not 19xx/20xx
    assert years[4] == 2023  # bare year-only cell OK

    _, currency, scale = _infer_unit("Revenue ($ in Millions)")
    assert currency == "USD"
    assert scale == "M"
    _, _, scale_k = _infer_unit("ARR in thousands")
    assert scale_k == "K"


def test_fallback_multi_column_and_map_hit_meta() -> None:
    from agetic_cdd_api.services_databook_extract import _resolve_unit_meta
    from agetic_cdd_api.services_databook_map import MetricMapHit
    from agetic_cdd_api.services_databook_models import MetricFamily

    # No year headers — previously only first numeric cell was kept.
    table = {
        "name": "wide",
        "rows": [
            ["Metric", "2022", "2023", "Note"],
            ["Revenue", "100", "200", "see CIM"],
        ],
    }
    rows = rows_from_table(source_name="a.pdf", doc_id="a", table=table)
    # Header cells "2022"/"2023" are bare years → year_cols path, not fallback.
    # Force true fallback with non-year column labels:
    table2 = {
        "name": "wide2",
        "rows": [
            ["Metric", "Actual", "Budget", "Note"],
            ["Revenue (USD M)", "100", "200", "n/a"],
        ],
    }
    rows2 = rows_from_table(source_name="a.pdf", doc_id="a", table=table2)
    values = sorted(r.value for r in rows2)
    assert values == [100.0, 200.0]
    assert all(r.currency == "USD" and r.scale == "M" for r in rows2)

    hit = MetricMapHit(
        "revenue",
        MetricFamily.REVENUE,
        currency="INR",
        scale="Cr",
    )
    unit, currency, scale = _resolve_unit_meta("Top line", hit)
    assert currency == "INR"
    assert scale == "Cr"
    assert unit is None


def test_two_row_grid_second_row_year_headers() -> None:
    table = {
        "name": "tiny",
        "rows": [
            ["Company financials"],
            ["Metric", "FY2023", "FY2024"],
        ],
    }
    # Only headers — no body values, but must not crash and must detect years path
    assert rows_from_table(source_name="a.pdf", doc_id="a", table=table) == []

    table2 = {
        "name": "tiny2",
        "rows": [
            ["Company financials"],
            ["Metric", "FY2023", "FY2024"],
            ["Revenue (USD M)", "10", "(2.5)"],
        ],
    }
    rows = rows_from_table(source_name="a.pdf", doc_id="a", table=table2)
    assert len(rows) == 2
    by_year = {r.fiscal_year: r.value for r in rows}
    assert by_year[2023] == 10.0
    assert by_year[2024] == -2.5
    assert rows[0].scale == "M"
    assert rows[0].currency == "USD"


def test_fixture_shape_fields_present() -> None:
    import json

    assert FIXTURE.is_file(), FIXTURE
    data = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert "item_kinds" in data
    assert "dropped" in data["item_kinds"]
    assert "conflict" in data["item_kinds"]
    for item in data.get("representative_items") or []:
        assert "kind" in item


def test_databook_endpoints_empty_deal() -> None:
    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "Databook Deal", "slug": "databook-deal", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]

        summary = client.get(f"/api/v1/portfolios/{deal_id}/cdd/databook", headers=headers)
        assert summary.status_code == 200
        assert summary.json()["success"] is True
        assert "meta" in summary.json()["data"]

        dq = client.get(f"/api/v1/portfolios/{deal_id}/cdd/data-quality", headers=headers)
        assert dq.status_code == 200
        body = dq.json()
        assert "items" in body and "summary" in body
        assert body["summary"]["total"] == len(body["items"])
        assert body["summary"]["needs_review"] == body["summary"]["total"]

        rescan = client.post(f"/api/v1/portfolios/{deal_id}/cdd/databook/rescan", headers=headers)
        assert rescan.status_code == 200

        rows = client.get(f"/api/v1/portfolios/{deal_id}/cdd/databook/rows", headers=headers)
        assert rows.status_code == 200
        findings = client.get(f"/api/v1/portfolios/{deal_id}/cdd/databook/findings", headers=headers)
        assert findings.status_code == 200
        promoted = client.get(f"/api/v1/portfolios/{deal_id}/cdd/databook/promoted", headers=headers)
        assert promoted.status_code == 200


def test_databook_rescan_from_library_tables(tmp_path_factory=None) -> None:
    """Upload + synthetic library document → conflicts surface in data-quality."""
    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "DQ Pack", "slug": "dq-pack", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]
        # Minimal PDF upload so deal has a VDR file; we'll write library JSON directly
        client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/vdr/upload",
            headers=headers,
            files={"file": ("01_Executive_Summary.pdf", BytesIO(b"%PDF-1.4 exec"), "application/pdf")},
        )
        client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/vdr/upload",
            headers=headers,
            files={"file": ("04_Commercial_Due_Diligence.pdf", BytesIO(b"%PDF-1.4 cdd"), "application/pdf")},
        )

        # Seed library index + docs with conflicting revenue tables
        from agetic_cdd_api.db import SessionLocal
        from agetic_cdd_api.models import Deal
        from agetic_cdd_api.services_databook import rescan_databook
        from agetic_cdd_api.services_vdr import library_dir

        db = SessionLocal()
        try:
            deal = db.get(Deal, deal_id)
            assert deal is not None
            lib = library_dir(deal)
            doc_dir = lib / "documents"
            doc_dir.mkdir(parents=True, exist_ok=True)

            def write_doc(filename: str, revenue_row: list[str]) -> dict:
                stem = filename.replace(".pdf", "").replace(" ", "_")[:40] + "_deadbeef"
                # use real safe stem via library helper
                from agetic_cdd_api.services_library import _safe_stem

                stem = _safe_stem(filename)
                payload = {
                    "filename": filename,
                    "library_stem": stem,
                    "status": "ready",
                    "tables": [
                        {
                            "name": "financials",
                            "rows": [
                                ["Metric", "FY2022", "FY2023", "FY2024"],
                                revenue_row,
                                ["Units Sold (000s)", "20.0", "152.0", "320.0"],
                            ],
                        }
                    ],
                    "text": "",
                    "chunks": [],
                }
                (doc_dir / f"{stem}.json").write_text(
                    __import__("json").dumps(payload, indent=2),
                    encoding="utf-8",
                )
                return {"filename": filename, "library_stem": stem, "status": "ready"}

            entries = [
                write_doc("01_Executive_Summary.pdf", ["Revenue (INR Cr)", "456", "2630", "4900"]),
                write_doc("04_Commercial_Due_Diligence.pdf", ["Revenue (INR Cr)", "190", "1468", "3140"]),
            ]
            index = {
                "deal_id": deal.id,
                "documents": entries,
                "document_count": 2,
                "status": "ready",
            }
            (lib / "index.json").write_text(__import__("json").dumps(index, indent=2), encoding="utf-8")

            summary = rescan_databook(deal)
            assert summary.meta.row_count > 0
            assert summary.flags.get("conflict", 0) >= 1

            from agetic_cdd_api.services_databook import get_data_quality

            dq = get_data_quality(deal, ensure=False)
            assert dq.summary.total == len(dq.items)
            conflicts = [i for i in dq.items if i.get("kind") == "conflict"]
            assert conflicts
            # Units Sold must not appear inside Revenue conflict candidates
            for item in conflicts:
                if "Revenue" not in str(item.get("metric")):
                    continue
                for cand in item.get("candidates") or []:
                    for cap in cand.get("captions") or []:
                        assert "Units Sold" not in cap
                        assert "NRR" not in cap
        finally:
            db.close()


def test_hitl_accept_drop_vouch_and_reason_required() -> None:
    from agetic_cdd_api.db import SessionLocal
    from agetic_cdd_api.models import Deal
    from agetic_cdd_api.services_databook import get_data_quality, rescan_databook
    from agetic_cdd_api.services_databook_decisions import accept_conflict, drop_row, vouch_row
    from agetic_cdd_api.services_databook_store import load_promoted, load_rows
    from agetic_cdd_api.services_library import _safe_stem
    from agetic_cdd_api.services_vdr import library_dir
    from fastapi import HTTPException

    with TestClient(app) as client:
        headers = _headers(client)
        import uuid

        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "HITL Deal", "slug": f"hitl-{uuid.uuid4().hex[:8]}", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]
        client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/vdr/upload",
            headers=headers,
            files={"file": ("01_Executive_Summary.pdf", BytesIO(b"%PDF-1.4 a"), "application/pdf")},
        )
        client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/vdr/upload",
            headers=headers,
            files={"file": ("04_Commercial_Due_Diligence.pdf", BytesIO(b"%PDF-1.4 b"), "application/pdf")},
        )

        db = SessionLocal()
        try:
            deal = db.get(Deal, deal_id)
            assert deal is not None
            lib = library_dir(deal)
            doc_dir = lib / "documents"
            doc_dir.mkdir(parents=True, exist_ok=True)

            def write_doc(filename: str, revenue_row: list[str]) -> dict:
                stem = _safe_stem(filename)
                payload = {
                    "filename": filename,
                    "library_stem": stem,
                    "status": "ready",
                    "tables": [
                        {
                            "name": "financials",
                            "rows": [
                                ["Metric", "FY2024"],
                                revenue_row,
                            ],
                        }
                    ],
                    "text": "",
                    "chunks": [],
                }
                (doc_dir / f"{stem}.json").write_text(
                    __import__("json").dumps(payload, indent=2), encoding="utf-8"
                )
                return {"filename": filename, "library_stem": stem, "status": "ready"}

            entries = [
                write_doc("01_Executive_Summary.pdf", ["Revenue (INR Cr)", "4900"]),
                write_doc("04_Commercial_Due_Diligence.pdf", ["Revenue (INR Cr)", "3140"]),
            ]
            (lib / "index.json").write_text(
                __import__("json").dumps(
                    {"deal_id": deal.id, "documents": entries, "document_count": 2, "status": "ready"},
                    indent=2,
                ),
                encoding="utf-8",
            )
            rescan_databook(deal)
            dq = get_data_quality(deal, ensure=False)
            conflicts = [i for i in dq.items if i.get("kind") == "conflict"]
            assert conflicts

            # reason required
            try:
                accept_conflict(
                    deal,
                    metric_key="Revenue",
                    fiscal_year=2024,
                    value=3140.0,
                    reason="x",
                    actor="tester",
                )
                assert False, "expected HTTPException"
            except HTTPException as exc:
                assert exc.status_code == 400

            result = accept_conflict(
                deal,
                metric_key="Revenue",
                fiscal_year=2024,
                value=3140.0,
                reason="Commercial DD governs FY2024 revenue",
                actor="tester@example.com",
            )
            assert result["action"] == "accept"
            assert result["value"] == 3140.0
            promoted = load_promoted(deal)
            assert any(p.metric_key == "revenue" and p.value == 3140.0 and not p.auto for p in promoted)
            dq2 = get_data_quality(deal, ensure=False)
            assert not any(
                i.get("kind") == "conflict" and i.get("fiscal_year") == 2024 and "Revenue" in str(i.get("metric"))
                for i in dq2.items
            )

            # API path: drop / vouch on synthetic rows
            from agetic_cdd_api.services_databook_models import ExtractedRow, MetricFamily, RowStatus
            from agetic_cdd_api.services_databook_store import save_rows

            noise = ExtractedRow(
                row_id="drop1",
                doc_id="x",
                source_name="noise.pdf",
                caption="Page number",
                metric_key=None,
                metric_family=MetricFamily.UNKNOWN,
                fiscal_year=None,
                value=12.0,
                status=RowStatus.CANDIDATE,
            )
            held = ExtractedRow(
                row_id="vouch1",
                doc_id="x",
                source_name="x.pdf",
                caption="Revenue",
                metric_key="revenue",
                metric_family=MetricFamily.REVENUE,
                fiscal_year=2023,
                value=100.0,
                status=RowStatus.HELD_OUT,
            )
            save_rows(deal, load_rows(deal) + [noise, held])

            drop_res = client.post(
                f"/api/v1/portfolios/{deal_id}/cdd/databook/rows/{noise.row_id}/drop",
                headers=headers,
                json={"reason": "Not a governing figure for this review"},
            )
            assert drop_res.status_code == 200, drop_res.text
            assert drop_res.json()["data"]["action"] == "drop"

            vouch_res = client.post(
                f"/api/v1/portfolios/{deal_id}/cdd/databook/rows/{held.row_id}/vouch",
                headers=headers,
                json={"reason": "Reviewed source; releasing on my authority"},
            )
            assert vouch_res.status_code == 200, vouch_res.text
            assert vouch_res.json()["data"]["action"] == "vouch"
        finally:
            db.close()


def test_statement_block_pass_fail_and_hold_out() -> None:
    from agetic_cdd_api.services_databook_blocks import annotate_and_detect_blocks, reconcile_blocks
    from agetic_cdd_api.services_databook_models import DealDatabookParams, RowStatus
    from agetic_cdd_api.services_databook_resolve import apply_statuses_and_promote

    table = {
        "name": "income",
        "rows": [
            ["Metric", "FY2024"],
            ["Revenue", "100"],
            ["Other income", "20"],
            ["Total income", "120"],
        ],
    }
    rows, blocks = annotate_and_detect_blocks(
        source_name="pnl.pdf",
        doc_id="pnl",
        table=table,
    )
    assert any(r.caption.lower().startswith("revenue") for r in rows)
    assert len(blocks) == 1
    checked, fail_ids, issues = reconcile_blocks(blocks, params=DealDatabookParams())
    assert checked[0].outcome == "pass"
    assert checked[0].ties is True
    assert not fail_ids
    assert not issues

    bad = {
        "name": "income",
        "rows": [
            ["Metric", "FY2024"],
            ["Revenue", "100"],
            ["Other income", "20"],
            ["Total income", "999"],
        ],
    }
    rows2, blocks2 = annotate_and_detect_blocks(source_name="pnl.pdf", doc_id="pnl", table=bad)
    checked2, fail_ids2, issues2 = reconcile_blocks(blocks2, params=DealDatabookParams())
    assert checked2[0].outcome == "fail"
    assert fail_ids2
    assert issues2 and issues2[0]["kind"] == "failed_check"
    out, promoted = apply_statuses_and_promote(rows2, hold_ids=set(), block_fail_ids=fail_ids2)
    assert not promoted
    assert all(
        r.status == RowStatus.HELD_OUT for r in out if r.row_id in fail_ids2
    )


def test_nested_grand_total_carry_and_duplicate_lines() -> None:
    from agetic_cdd_api.services_databook_blocks import annotate_and_detect_blocks, reconcile_blocks
    from agetic_cdd_api.services_databook_models import DealDatabookParams

    # Nested: Total Revenue closes product/services; Total income = prior total + Other.
    nested = {
        "name": "pnl",
        "rows": [
            ["Metric", "FY2024"],
            ["Product revenue", "60"],
            ["Services revenue", "40"],
            ["Total revenue", "100"],
            ["Other income", "20"],
            ["Total income", "120"],
        ],
    }
    rows, blocks = annotate_and_detect_blocks(
        source_name="pnl.pdf", doc_id="pnl", table=nested, params=DealDatabookParams()
    )
    checked, fail_ids, issues = reconcile_blocks(blocks, params=DealDatabookParams())
    assert len(checked) == 2
    assert all(b.outcome == "pass" for b in checked)
    assert not fail_ids
    assert not issues
    # Independent expense section must NOT absorb prior revenue total.
    split = {
        "name": "pnl",
        "rows": [
            ["Metric", "FY2024"],
            ["Product revenue", "60"],
            ["Services revenue", "40"],
            ["Total revenue", "100"],
            ["COGS", "30"],
            ["Total expenses", "30"],
        ],
    }
    _, blocks2 = annotate_and_detect_blocks(
        source_name="pnl.pdf", doc_id="pnl", table=split, params=DealDatabookParams()
    )
    checked2, fail2, _ = reconcile_blocks(blocks2, params=DealDatabookParams())
    assert all(b.outcome == "pass" for b in checked2)
    assert not fail2

    # Duplicate identical lines claim distinct ExtractedRows / unique row_ids.
    dup = {
        "name": "seg",
        "rows": [
            ["Metric", "FY2024"],
            ["Segment A", "10"],
            ["Segment A", "10"],
            ["Total", "20"],
        ],
    }
    rows3, blocks3 = annotate_and_detect_blocks(
        source_name="seg.pdf", doc_id="seg", table=dup, params=DealDatabookParams()
    )
    line_rows = [r for r in rows3 if r.caption == "Segment A"]
    assert len(line_rows) == 2
    assert line_rows[0].row_id != line_rows[1].row_id
    checked3, fail3, _ = reconcile_blocks(blocks3, params=DealDatabookParams())
    assert checked3[0].outcome == "pass"
    assert checked3[0].computed_sum == 20.0
    assert not fail3
    assert len({r.block_id for r in line_rows if r.block_id}) == 1


def test_findings_trust_ledger_and_triage(tmp_path: Path, monkeypatch) -> None:
    from types import SimpleNamespace

    import agetic_cdd_api.services_databook as databook
    import agetic_cdd_api.services_databook_store as store
    from agetic_cdd_api.services_databook_models import (
        ExtractedRow,
        MetricFamily,
        RowStatus,
        StatementBlock,
    )

    monkeypatch.setattr(store, "databook_dir", lambda _deal: tmp_path)
    monkeypatch.setattr(
        databook,
        "load_library_index",
        lambda _deal: {
            "documents": [
                {"filename": "good.pdf", "library_stem": "good"},
                {"filename": "bad.pdf", "library_stem": "bad"},
                {"filename": "empty.pdf", "library_stem": "empty"},
            ]
        },
    )

    deal = SimpleNamespace(id="d3", slug="p3-findings")
    store.save_rows(
        deal,  # type: ignore[arg-type]
        [
            ExtractedRow(
                row_id="r1",
                doc_id="bad",
                source_name="bad.pdf",
                caption="Revenue",
                metric_key="revenue",
                metric_family=MetricFamily.REVENUE,
                fiscal_year=2024,
                value=100.0,
                status=RowStatus.HELD_OUT,
                assumption=True,
            ),
            ExtractedRow(
                row_id="r2",
                doc_id="good",
                source_name="good.pdf",
                caption="Revenue",
                metric_key="revenue",
                metric_family=MetricFamily.REVENUE,
                fiscal_year=2024,
                value=100.0,
                currency="INR",
                status=RowStatus.PROMOTED,
            ),
        ],
    )
    store.save_blocks(
        deal,  # type: ignore[arg-type]
        [
            StatementBlock(
                block_id="b1",
                doc_id="bad",
                source_name="bad.pdf",
                period="FY2024",
                fiscal_year=2024,
                line_row_ids=["r1"],
                printed_subtotal=120.0,
                computed_sum=100.0,
                ties=False,
                outcome="fail",
            ),
            StatementBlock(
                block_id="b2",
                doc_id="good",
                source_name="good.pdf",
                period="FY2024",
                fiscal_year=2024,
                line_row_ids=["r2"],
                printed_subtotal=100.0,
                computed_sum=100.0,
                ties=True,
                outcome="pass",
            ),
            StatementBlock(
                block_id="b3",
                doc_id="empty",
                source_name="empty.pdf",
                outcome="no_table",
            ),
        ],
    )
    store.save_issues(deal, [])  # type: ignore[arg-type]
    findings = databook.list_databook_findings(deal)  # type: ignore[arg-type]
    by_name = {f["source_name"]: f for f in findings}
    assert by_name["bad.pdf"]["checks_failed"] == 1
    assert by_name["bad.pdf"]["checks_total"] == 1
    assert "failed_check" in by_name["bad.pdf"]["flags"] or "assumption" in by_name["bad.pdf"]["flags"]
    assert by_name["good.pdf"]["checks_passed"] == 1
    assert by_name["empty.pdf"]["no_table"] is True
    assert by_name["empty.pdf"]["triage"] == "unread"
    # Triage order: assumption/failed_check before unread
    assert findings[0]["source_name"] == "bad.pdf"


def test_json_primitive_and_correct_patch_serializes_enum(tmp_path: Path, monkeypatch) -> None:
    from types import SimpleNamespace

    import agetic_cdd_api.services_databook_decisions as decisions
    import agetic_cdd_api.services_databook_store as store

    monkeypatch.setattr(store, "databook_dir", lambda _deal: tmp_path)

    deal = SimpleNamespace(id="d1", slug="hitl-enum")
    row = ExtractedRow(
        row_id="r1",
        doc_id="d",
        source_name="a.pdf",
        caption="Mystery line",
        metric_key=None,
        metric_family=MetricFamily.UNKNOWN,
        fiscal_year=2024,
        value=10.0,
        status=RowStatus.CANDIDATE,
    )
    store.save_rows(deal, [row])  # type: ignore[arg-type]
    store.save_promoted(deal, [])  # type: ignore[arg-type]
    store.save_meta(deal, store.load_meta(deal))  # type: ignore[arg-type]

    decisions.correct_row(
        deal,  # type: ignore[arg-type]
        "r1",
        reason="Map to custom metric key",
        actor="tester",
        metric_key="custom_metric",
        value=42.0,
    )
    raw = (tmp_path / "decisions.jsonl").read_text(encoding="utf-8")
    assert '"metric_family": "other"' in raw
    assert "MetricFamily" not in raw


def test_accept_overlay_merges_sources_and_held_out_excluded(tmp_path: Path, monkeypatch, caplog) -> None:
    import logging
    from types import SimpleNamespace

    import agetic_cdd_api.services_databook_decisions as decisions
    import agetic_cdd_api.services_databook_store as store

    monkeypatch.setattr(store, "databook_dir", lambda _deal: tmp_path)

    deal = SimpleNamespace(id="d2", slug="hitl-overlay")
    rows = [
        ExtractedRow(
            row_id="a",
            doc_id="1",
            source_name="exec.pdf",
            caption="Revenue (INR Cr)",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=3140.0,
            status=RowStatus.CANDIDATE,
        ),
        ExtractedRow(
            row_id="b",
            doc_id="2",
            source_name="cdd.pdf",
            caption="Revenue (INR Cr)",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=3140.0,
            status=RowStatus.CANDIDATE,
        ),
        ExtractedRow(
            row_id="c",
            doc_id="3",
            source_name="cim.pdf",
            caption="Revenue (INR Cr)",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=4900.0,
            status=RowStatus.CANDIDATE,
        ),
    ]
    store.append_decision(
        deal,  # type: ignore[arg-type]
        {
            "decision_id": "acc1",
            "action": "accept",
            "reason": "Commercial DD governs",
            "actor": "tester",
            "at": "2026-01-01T00:00:00Z",
            "row_id": "a",
            "match": {
                "source_name": "exec.pdf",
                "caption": "Revenue (INR Cr)",
                "fiscal_year": 2024,
                "metric_key": "revenue",
            },
            "patch": {"value": 3140.0, "metric_key": "revenue", "fiscal_year": 2024},
            "accepted_value": 3140.0,
            "metric_key": "revenue",
            "fiscal_year": 2024,
        },
    )
    # Corrupt trailing line should warn, not crash
    with caplog.at_level(logging.WARNING, logger="agetic_cdd_api.services_databook_decisions"):
        (tmp_path / "decisions.jsonl").write_text(
            (tmp_path / "decisions.jsonl").read_text(encoding="utf-8") + "{not-json\n",
            encoding="utf-8",
        )
        working, promoted = decisions.apply_decision_overlays(deal, rows)  # type: ignore[arg-type]
    assert any("malformed decision" in r.message.lower() for r in caplog.records)
    assert len(promoted) == 1
    assert set(promoted[0].sources) == {"cdd.pdf", "exec.pdf"}
    assert set(promoted[0].captions) == {"Revenue (INR Cr)"}
    by_id = {r.row_id: r for r in working}
    assert by_id["a"].status == RowStatus.PROMOTED
    assert by_id["b"].status == RowStatus.PROMOTED
    assert by_id["c"].status == RowStatus.DROPPED

    # HELD_OUT siblings must not re-enter conflict detection after vouch
    store.save_promoted(
        deal,  # type: ignore[arg-type]
        [
            *promoted,
            decisions._upsert_promoted([], row=ExtractedRow(
                row_id="v1",
                doc_id="1",
                source_name="a.pdf",
                caption="ARR",
                metric_key="arr",
                metric_family=MetricFamily.REVENUE,
                fiscal_year=2023,
                value=100.0,
                status=RowStatus.VOUCHED,
            ), decision_id="v", auto=False)[0],
        ],
    )
    vouch_rows = [
        ExtractedRow(
            row_id="v1",
            doc_id="1",
            source_name="a.pdf",
            caption="ARR",
            metric_key="arr",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2023,
            value=100.0,
            status=RowStatus.VOUCHED,
        ),
        ExtractedRow(
            row_id="v2",
            doc_id="2",
            source_name="b.pdf",
            caption="ARR",
            metric_key="arr",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2023,
            value=200.0,
            status=RowStatus.HELD_OUT,
        ),
    ]
    issues = decisions._rebuild_issues_from_rows(deal, vouch_rows)  # type: ignore[arg-type]
    assert not any(
        i.get("kind") == "conflict" and i.get("fiscal_year") == 2023 for i in issues
    )


# ---------------------------------------------------------------------------
# Phase 4 — promoted consumers, materiality, CSV, ops WARN demotion
# ---------------------------------------------------------------------------


def test_assumption_issues_in_data_quality() -> None:
    from agetic_cdd_api.services_databook_resolve import (
        assumption_issues_from_rows,
        build_data_quality_items,
    )

    rows = [
        ExtractedRow(
            row_id="a1",
            doc_id="1",
            source_name="fin.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=3140.0,
            scale="Cr",
            assumption=True,
            status=RowStatus.CANDIDATE,
        ),
        ExtractedRow(
            row_id="a2",
            doc_id="1",
            source_name="fin.pdf",
            caption="Units Sold",
            metric_key="units_sold",
            metric_family=MetricFamily.UNITS,
            fiscal_year=2024,
            value=20.0,
            assumption=False,
            status=RowStatus.CANDIDATE,
        ),
    ]
    assumed = assumption_issues_from_rows(rows)
    assert len(assumed) == 1
    assert assumed[0]["kind"] == "assumed"
    assert assumed[0]["metric"] == "revenue"
    items = build_data_quality_items(
        dropped=[],
        conflicts=[],
        assumed=assumed,
    )
    assert any(i["kind"] == "assumed" for i in items)
    by_kind: dict[str, int] = {}
    for i in items:
        by_kind[i["kind"]] = by_kind.get(i["kind"], 0) + 1
    assert by_kind.get("assumed") == 1

    from agetic_cdd_api.services_databook_consume import (
        _display_unit,
        filter_material_promoted,
        promoted_csv,
        promoted_to_pl_lines,
    )
    from agetic_cdd_api.services_databook_models import PromotedMetric

    metrics = [
        PromotedMetric(
            metric_key="revenue",
            fiscal_year=2024,
            value=3140.0,
            currency="INR",
            scale="Cr",
            row_id="r1",
            sources=["a.xlsx"],
            captions=["Revenue"],
            auto=True,
        ),
        PromotedMetric(
            metric_key="revenue",
            fiscal_year=2023,
            value=2800.0,
            currency="INR",
            scale="Cr",
            row_id="r2",
            sources=["a.xlsx"],
            captions=["Revenue"],
            auto=True,
        ),
        PromotedMetric(
            metric_key="revenue",
            fiscal_year=2025,
            value=3500.0,
            currency="INR",
            scale="Cr",
            row_id="r2b",
            sources=["a.xlsx"],
            captions=["Revenue"],
            auto=True,
        ),
        PromotedMetric(
            metric_key="units_sold",
            fiscal_year=2024,
            value=20.0,
            unit="000s",
            scale="M",
            row_id="r3",
            sources=["b.xlsx"],
            captions=["Units Sold"],
            auto=True,
        ),
    ]
    lines = promoted_to_pl_lines(metrics)
    rev = next(l for l in lines if l["line_item"] == "Revenue")
    assert rev["fy2024_value"] == 3140.0
    assert rev["fy2023_value"] == 2800.0
    assert rev["fy2025_value"] == 3500.0
    assert rev["databook"] is True

    units = next(l for l in lines if l["line_item"] == "Units Sold")
    assert units["unit"] == "M 000s"
    assert _display_unit(metrics[-1]) == "M 000s"
    assert _display_unit(metrics[0]) == "INR Cr"

    material = filter_material_promoted(metrics, material_only=True)
    assert all(m.metric_key != "units_sold" for m in material)
    assert len(material) == 3

    csv_text = promoted_csv(material)
    assert "metric_key,fiscal_year,value" in csv_text
    assert "revenue,2024,3140.0" in csv_text
    assert "units_sold" not in csv_text


def test_merge_promoted_does_not_mutate_input_spec(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_consume import merge_promoted_into_historical_spec
    from agetic_cdd_api.services_databook_models import PromotedMetric
    from agetic_cdd_api import services_databook_store as store

    class _Deal:
        id = "d1"
        slug = "consume-iso"
        name = "Consume"
        company = "Consume Co"

    deal = _Deal()  # type: ignore[assignment]
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        store.save_promoted(
            deal,  # type: ignore[arg-type]
            [
                PromotedMetric(
                    metric_key="revenue",
                    fiscal_year=2022,
                    value=2000.0,
                    currency="INR",
                    scale="Cr",
                    row_id="r1",
                    sources=["a.xlsx"],
                    captions=["Revenue"],
                    auto=True,
                    proof_level=ProofLevel.L2,
                ),
            ],
        )
        nested_note = ["keep"]
        original_spec = {
            "pl_lines": [{"line_item": "Revenue", "fy2024_value": 4900.0}],
            "bridge_notes": nested_note,
            "performance_metrics": {"x": 1.0},
            "empty": False,
        }
        merged = merge_promoted_into_historical_spec(deal, original_spec)  # type: ignore[arg-type]
        merged["bridge_notes"].append("mutated")
        merged["performance_metrics"]["x"] = 99.0
        assert nested_note == ["keep"]
        assert original_spec["performance_metrics"]["x"] == 1.0
        assert merged["performance_metrics"]["revenue_fy2022"] == 2000.0
        assert merged["pl_lines"][0]["fy2022_value"] == 2000.0
    finally:
        store.databook_dir = original



def test_ops_dashboard_skips_warn_when_databook_promoted() -> None:
    from agetic_cdd_api.report_ops_dashboard import MetricFact, detect_figure_conflicts

    facts = [
        MetricFact(
            family="revenue",
            label="Revenue",
            value=3140.0,
            year=2024,
            agent_key="databook",
            agent_name="Databook",
            sources=["a.xlsx"],
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=4900.0,
            year=2024,
            agent_key="historical_performance",
            agent_name="Historical Performance",
            sources=["b.pdf"],
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=2800.0,
            year=2024,
            agent_key="executive_summary",
            agent_name="Executive Summary",
            sources=["c.pdf"],
        ),
    ]
    warnings = detect_figure_conflicts(facts)
    assert warnings == []

    # Without databook, conflict still fires
    no_db = [f for f in facts if f.agent_key != "databook"]
    assert len(detect_figure_conflicts(no_db)) == 1


def test_merge_promoted_into_historical_spec(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_consume import merge_promoted_into_historical_spec
    from agetic_cdd_api.services_databook_models import PromotedMetric
    from agetic_cdd_api import services_databook_store as store

    class _Deal:
        id = "d1"
        slug = "consume-test"
        name = "Consume"
        company = "Consume Co"

    deal = _Deal()  # type: ignore[assignment]
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        store.save_promoted(
            deal,  # type: ignore[arg-type]
            [
                PromotedMetric(
                    metric_key="revenue",
                    fiscal_year=2024,
                    value=3140.0,
                    currency="INR",
                    scale="Cr",
                    row_id="r1",
                    sources=["commercial.xlsx"],
                    captions=["Revenue"],
                    auto=False,
                    decision_id="dec1",
                ),
            ],
        )
        merged = merge_promoted_into_historical_spec(
            deal,  # type: ignore[arg-type]
            {"pl_lines": [{"line_item": "Revenue", "fy2024_value": 4900.0}], "empty": False},
        )
        assert merged["databook_promoted"] is True
        assert merged["pl_lines"][0]["fy2024_value"] == 3140.0
        assert "commercial.xlsx" in merged["sources"]
    finally:
        store.databook_dir = original


def test_databook_promoted_csv_endpoint() -> None:
    with TestClient(app) as client:
        headers = _headers(client)
        deal = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "CSV Deal", "slug": "csv-deal-p4", "industry": "generic"},
        )
        deal_id = deal.json()["data"]["id"]
        csv_res = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/promoted.csv",
            headers=headers,
        )
        assert csv_res.status_code == 200
        assert "text/csv" in csv_res.headers.get("content-type", "")
        assert "metric_key,fiscal_year,value" in csv_res.text
        deep = client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/deep",
            headers=headers,
        )
        assert deep.status_code == 200
        assert deep.json()["success"] is True
        rows = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/rows?material=true",
            headers=headers,
        )
        assert rows.status_code == 200


def test_excel_export_import_roundtrip(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from agetic_cdd_api.services_databook_excel import (
        build_editable_workbook,
        databook_freshness,
        import_edited_workbook,
    )
    from agetic_cdd_api.services_databook_models import ExtractedRow, MetricFamily, RowStatus
    from agetic_cdd_api.services_databook_store import load_promoted, load_rows, save_meta, save_rows
    from agetic_cdd_api import services_databook_store as store
    from agetic_cdd_api.services_databook_models import DatabookMeta
    from agetic_cdd_api.services_ingestion import utc_now_iso

    class _Deal:
        id = "excel-deal"
        slug = "excel-deal"
        name = "Excel"
        company = "Excel Co"

    deal = _Deal()  # type: ignore[assignment]
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        save_rows(
            deal,  # type: ignore[arg-type]
            [
                ExtractedRow(
                    row_id="r1",
                    doc_id="1",
                    source_name="a.xlsx",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=100.0,
                    currency="INR",
                    scale="Cr",
                    status=RowStatus.PROMOTED,
                ),
            ],
        )
        save_meta(deal, DatabookMeta(last_rescan_at=utc_now_iso(), row_count=1, promoted_count=1))  # type: ignore[arg-type]
        fresh = databook_freshness(deal)  # type: ignore[arg-type]
        assert fresh["up_to_date"] is True
        assert fresh["label"] == "Databook is up to date"

        xlsx = build_editable_workbook(deal)  # type: ignore[arg-type]
        assert xlsx[:2] == b"PK"
        wb = load_workbook(BytesIO(xlsx))
        assert "Databook" in wb.sheetnames
        sheet = wb["Databook"]
        # header + 1 data row
        assert sheet.cell(2, 1).value == "r1"
        sheet.cell(2, 5).value = 3140.0  # value column
        out = BytesIO()
        wb.save(out)
        result = import_edited_workbook(
            deal,  # type: ignore[arg-type]
            out.getvalue(),
            actor="tester@example.com",
        )
        assert result["applied"] == 1
        rows = load_rows(deal)  # type: ignore[arg-type]
        assert rows[0].value == 3140.0
        assert rows[0].status == RowStatus.PROMOTED
        promoted = load_promoted(deal)  # type: ignore[arg-type]
        assert any(p.value == 3140.0 for p in promoted)
    finally:
        store.databook_dir = original


def test_databook_excel_endpoints() -> None:
    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "Excel API Deal", "slug": "excel-api-deal", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]
        summary = client.get(f"/api/v1/portfolios/{deal_id}/cdd/databook", headers=headers)
        assert summary.status_code == 200
        body = summary.json()["data"]
        assert "freshness" in body
        assert body["freshness"]["up_to_date"] is False
        assert "Update databook" in body["freshness"]["label"]

        export = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/export.xlsx",
            headers=headers,
        )
        assert export.status_code == 200
        assert "spreadsheetml" in export.headers.get("content-type", "")
        assert export.content[:2] == b"PK"

        files = {"file": ("databook.xlsx", export.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        imported = client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/import.xlsx",
            headers=headers,
            files=files,
        )
        assert imported.status_code == 200
        assert imported.json()["success"] is True



def test_release_create_and_hard_consume(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_consume import (
        load_promoted_metrics_for_slug,
        prefer_released_pl_lines,
    )
    from agetic_cdd_api.services_databook_models import (
        ExtractedRow,
        MetricFamily,
        PromotedMetric,
        ProofLevel,
        ReleaseCellStatus,
        RowStatus,
    )
    from agetic_cdd_api.services_databook_release import create_release
    from agetic_cdd_api import services_databook_store as store
    from agetic_cdd_api.services_databook_store import load_current_release, save_promoted, save_rows

    class _Deal:
        id = "rel-deal"
        slug = "rel-deal"
        name = "Release"
        company = "Release Co"

    deal = _Deal()  # type: ignore[assignment]
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]

    # Point deals_root slug loader at tmp parent so for_slug helpers work.
    from agetic_cdd_api import services_deals as deals_mod
    from agetic_cdd_api import services_databook_consume as consume_mod
    from agetic_cdd_api import services_databook_release as release_mod

    deals_root_original = deals_mod.deals_root
    fake_root = tmp_path.parent / "deals_root"
    deal_dir = fake_root / "rel-deal" / "databook"
    deal_dir.mkdir(parents=True)

    def _deals_root() -> Path:
        return fake_root

    try:
        store.databook_dir = lambda _deal: deal_dir  # type: ignore[assignment]
        deals_mod.deals_root = _deals_root  # type: ignore[assignment]
        consume_mod.deals_root = _deals_root  # type: ignore[assignment]
        store.deals_root = _deals_root  # type: ignore[assignment]

        save_rows(
            deal,  # type: ignore[arg-type]
            [
                ExtractedRow(
                    row_id="r1",
                    doc_id="1",
                    source_name="a.xlsx",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=3140.0,
                    currency="INR",
                    scale="Cr",
                    status=RowStatus.PROMOTED,
                    proof_level=ProofLevel.L2,
                    dual_agree=True,
                ),
                ExtractedRow(
                    row_id="r2",
                    doc_id="1",
                    source_name="a.xlsx",
                    caption="EBITDA",
                    metric_key="ebitda",
                    metric_family=MetricFamily.MARGIN,
                    fiscal_year=2024,
                    value=500.0,
                    currency="INR",
                    scale="Cr",
                    status=RowStatus.HELD_OUT,
                    proof_level=ProofLevel.L1,
                ),
            ],
        )
        save_promoted(
            deal,  # type: ignore[arg-type]
            [
                PromotedMetric(
                    metric_key="revenue",
                    fiscal_year=2024,
                    value=3140.0,
                    currency="INR",
                    scale="Cr",
                    row_id="r1",
                    sources=["a.xlsx"],
                    captions=["Revenue"],
                    auto=True,
                    proof_level=ProofLevel.L2,
                ),
            ],
        )
        # Fix slug loader shape regression
        loaded = load_promoted_metrics_for_slug("rel-deal")
        assert len(loaded) == 1
        assert loaded[0].value == 3140.0

        release = create_release(deal, source="manual", note="test")  # type: ignore[arg-type]
        assert release.release_id == "v1"
        assert release.counts.get("proven") == 1
        assert release.counts.get("doubtful") == 1
        current = load_current_release(deal)  # type: ignore[arg-type]
        assert current is not None
        assert current.release_id == "v1"
        statuses = {c.metric_key: c.status for c in current.cells}
        assert statuses["revenue"] == ReleaseCellStatus.PROVEN
        assert statuses["ebitda"] == ReleaseCellStatus.DOUBTFUL

        agent_lines = [
            {"line_item": "Revenue", "metric_key": "revenue", "fy2024_value": 9999.0, "unit": "INR Cr"},
            {"line_item": "Other note", "fy2024_value": 1.0},
        ]
        preferred = prefer_released_pl_lines("rel-deal", agent_lines)
        rev = next(l for l in preferred if l.get("metric_key") == "revenue" or l.get("line_item") == "Revenue")
        assert rev["fy2024_value"] == 3140.0
        assert rev["databook_status"] == "proven"
        ebitda = next(l for l in preferred if l.get("metric_key") == "ebitda")
        assert "fy2024_value" not in ebitda
        assert ebitda["databook_status"] == "doubtful"
        assert ebitda.get("databook_provisional") is True
        assert 2024 in (ebitda.get("doubtful_years") or [])

        # Immutable: second identical release_id must fail via versioning (v2 ok)
        release2 = create_release(deal, source="manual", note="v2")  # type: ignore[arg-type]
        assert release2.release_id == "v2"
        assert load_current_release(deal).release_id == "v2"  # type: ignore[union-attr]
    finally:
        store.databook_dir = original
        deals_mod.deals_root = deals_root_original
        consume_mod.deals_root = deals_root_original
        store.deals_root = deals_root_original


def test_hard_consume_strips_agent_when_no_release(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_consume import prefer_released_pl_lines
    from agetic_cdd_api import services_deals as deals_mod
    from agetic_cdd_api import services_databook_consume as consume_mod
    from agetic_cdd_api import services_databook_store as store

    fake_root = tmp_path / "deals"
    fake_root.mkdir()
    original = deals_mod.deals_root
    store_root = store.deals_root
    consume_root = consume_mod.deals_root
    try:
        deals_mod.deals_root = lambda: fake_root  # type: ignore[assignment]
        store.deals_root = lambda: fake_root  # type: ignore[assignment]
        consume_mod.deals_root = lambda: fake_root  # type: ignore[assignment]
        lines = prefer_released_pl_lines(
            "missing-deal",
            [{"line_item": "Revenue", "metric_key": "revenue", "fy2024_value": 4900.0}],
        )
        assert lines
        assert "fy2024_value" not in lines[0]
        assert lines[0]["databook_status"] == "missing"
    finally:
        deals_mod.deals_root = original
        store.deals_root = store_root
        consume_mod.deals_root = consume_root


def test_g1_material_agent_fy_guard_and_display_rows() -> None:
    from agetic_cdd_api.services_databook_consume import (
        material_agent_fy_values_present,
        released_pl_display_rows,
    )

    agent_only = [{"line_item": "Revenue", "metric_key": "revenue", "fy2024_value": 99.0}]
    assert material_agent_fy_values_present(agent_only) is True
    stamped = [
        {
            "line_item": "Revenue",
            "metric_key": "revenue",
            "fy2024_value": 99.0,
            "databook": True,
            "databook_release_id": "v1",
            "databook_status": "proven",
        }
    ]
    assert material_agent_fy_values_present(stamped) is False

    rows, pl, footnote = released_pl_display_rows(
        "__no_such_deal_g1__",
        agent_only,
        limit=5,
    )
    assert footnote is None
    assert pl
    assert pl[0]["databook_status"] == "missing"
    assert rows
    assert rows[0][3] == "missing"


def test_g1_release_marks_reports_stale(tmp_path: Path) -> None:
    from agetic_cdd_api.report_store import (
        finish_report,
        get_report,
        mark_reports_stale_for_databook_release,
    )
    from agetic_cdd_api import services_deals as deals_mod

    fake_root = tmp_path / "deals"
    (fake_root / "g1-deal" / "reports").mkdir(parents=True)
    original = deals_mod.deals_root
    try:
        deals_mod.deals_root = lambda: fake_root  # type: ignore[assignment]
        finish_report("g1-deal", "ic_memo", artifact_path=None)
        finish_report("g1-deal", "market_deck", artifact_path=None)
        meta = get_report("g1-deal", "ic_memo")
        assert meta and meta["status"] == "ready"

        marked = mark_reports_stale_for_databook_release(
            "g1-deal", release_id="v3", release_version=3
        )
        assert "ic_memo" in marked and "market_deck" in marked
        stale = get_report("g1-deal", "ic_memo")
        assert stale is not None
        assert stale["status"] == "stale"
        assert stale["stale_databook"] is True
        assert stale["stale_release_id"] == "v3"

        finish_report("g1-deal", "ic_memo", artifact_path=None)
        fresh = get_report("g1-deal", "ic_memo")
        assert fresh is not None
        assert fresh["status"] == "ready"
        assert not fresh.get("stale_databook")
    finally:
        deals_mod.deals_root = original


def test_databook_release_endpoint() -> None:
    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "Release API Deal", "slug": "release-api-deal", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]
        released = client.post(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/release",
            headers=headers,
            json={"note": "manual"},
        )
        assert released.status_code == 200
        payload = released.json()["data"]
        assert payload["release"]["release_id"].startswith("v")
        assert int(payload["release"]["version"]) >= 1
        assert payload["summary"]["release"]["release_id"] == payload["release"]["release_id"]
        current = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/release",
            headers=headers,
        )
        assert current.status_code == 200
        assert current.json()["data"]["release_id"] == payload["release"]["release_id"]
        listing = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/releases",
            headers=headers,
        )
        assert listing.status_code == 200
        assert listing.json()["data"]["total"] >= 1


def test_conflict_display_label_joins_metric_key_cell(tmp_path: Path) -> None:
    """Legacy conflict tickets keyed as 'Total Revenue' must attach to revenue."""
    from agetic_cdd_api.services_databook_models import (
        ExtractedRow,
        MetricFamily,
        ReleaseCellStatus,
        RowStatus,
    )
    from agetic_cdd_api.services_databook_release import build_release_cells
    from agetic_cdd_api import services_databook_store as store

    class _Deal:
        id = "label-join"
        slug = "label-join"

    deal = _Deal()
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        store.save_promoted(deal, [])  # type: ignore[arg-type]
        store.save_rows(
            deal,  # type: ignore[arg-type]
            [
                ExtractedRow(
                    row_id="r1",
                    doc_id="1",
                    source_name="CC_Investor_Workbook_EXTERNAL.xlsx",
                    caption="Total Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2025,
                    value=13.282,
                    status=RowStatus.HELD_OUT,
                ),
                ExtractedRow(
                    row_id="r2",
                    doc_id="2",
                    source_name="Compost Crew Monthly Financials 23YTD26.xlsx",
                    caption="Total Revenues",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2025,
                    value=13.4115,
                    status=RowStatus.HELD_OUT,
                ),
            ],
        )
        store.save_issues(
            deal,  # type: ignore[arg-type]
            [
                {
                    "kind": "conflict",
                    "metric": "Total Revenue",
                    "fiscal_year": 2025,
                    "chosen": 13.4115,
                    "candidates": [
                        {
                            "value": 13.282,
                            "sources": ["CC_Investor_Workbook_EXTERNAL.xlsx"],
                            "captions": ["Total Revenue"],
                            "chosen": False,
                            "stated_by": 1,
                        },
                        {
                            "value": 13.4115,
                            "sources": [
                                "Compost Crew Monthly Financials 23YTD26.xlsx"
                            ],
                            "captions": ["Total Revenues"],
                            "chosen": True,
                            "stated_by": 1,
                        },
                    ],
                    "scope": "across documents",
                    "rule": "test",
                }
            ],
        )
        cells = build_release_cells(deal)  # type: ignore[arg-type]
        rev = next(c for c in cells if c.metric_key == "revenue")
        assert rev.status == ReleaseCellStatus.DOUBTFUL
        assert rev.value == pytest.approx(13.4115)
        assert len(rev.alternatives) == 2
        assert {round(float(a["value"]), 4) for a in rev.alternatives} == {
            13.282,
            13.4115,
        }
        # Display-label key must not create a second orphan cell.
        assert sum(1 for c in cells if c.metric_key.lower() == "revenue") == 1
    finally:
        store.databook_dir = original  # type: ignore[assignment]


def test_proven_retains_conflict_alternatives(tmp_path: Path, caplog) -> None:
    import logging

    from agetic_cdd_api.services_databook_models import (
        ExtractedRow,
        MetricFamily,
        PromotedMetric,
        ReleaseCellStatus,
        RowStatus,
    )
    from agetic_cdd_api.services_databook_release import build_release_cells
    from agetic_cdd_api import services_databook_store as store

    class _Deal:
        id = "alt-deal"
        slug = "alt-deal"

    deal = _Deal()
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        store.save_promoted(
            deal,  # type: ignore[arg-type]
            [
                PromotedMetric(
                    metric_key="revenue",
                    fiscal_year=2024,
                    value=3140.0,
                    row_id="r1",
                    sources=["audit.xlsx"],
                    captions=["Revenue"],
                    auto=False,
                    decision_id="dec1",
                ),
            ],
        )
        store.save_rows(
            deal,  # type: ignore[arg-type]
            [
                ExtractedRow(
                    row_id="r1",
                    doc_id="1",
                    source_name="audit.xlsx",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=3140.0,
                    status=RowStatus.PROMOTED,
                ),
                ExtractedRow(
                    row_id="r2",
                    doc_id="2",
                    source_name="draft.xlsx",
                    caption="Revenue",
                    metric_key="revenue",
                    metric_family=MetricFamily.REVENUE,
                    fiscal_year=2024,
                    value=2800.0,
                    status=RowStatus.HELD_OUT,
                ),
            ],
        )
        store.save_issues(
            deal,  # type: ignore[arg-type]
            [
                {
                    "kind": "conflict",
                    "metric": "revenue",
                    "fiscal_year": 2024,
                    "chosen": 3140.0,
                    "candidates": [
                        {
                            "value": 3140.0,
                            "sources": ["audit.xlsx"],
                            "captions": ["Revenue"],
                            "chosen": True,
                            "stated_by": 1,
                        },
                        {
                            "value": "not-a-number",
                            "sources": ["draft.xlsx"],
                            "captions": ["Revenue"],
                            "chosen": False,
                            "stated_by": 1,
                        },
                    ],
                    "scope": "across documents",
                    "rule": "test",
                }
            ],
        )
        with caplog.at_level(logging.WARNING):
            cells = build_release_cells(deal)  # type: ignore[arg-type]
        rev = next(c for c in cells if c.metric_key == "revenue")
        assert rev.status == ReleaseCellStatus.PROVEN
        assert rev.value == 3140.0
        assert len(rev.alternatives) == 2
        assert "competing candidates retained" in (rev.reason or "")
        # Unparseable alternative values are retained raw; chosen parse path logs only when needed
        assert any(a.get("value") == "not-a-number" for a in rev.alternatives)
    finally:
        store.databook_dir = original


def test_slug_deal_protocol_bootstrap(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_models import PromotedMetric, ProofLevel
    from agetic_cdd_api.services_databook_release import SlugDeal, create_release, ensure_current_release
    from agetic_cdd_api import services_databook_store as store

    deal = SlugDeal(id="slug-proto", slug="slug-proto")
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        store.save_promoted(
            deal,
            [
                PromotedMetric(
                    metric_key="ebitda",
                    fiscal_year=2023,
                    value=100.0,
                    row_id="e1",
                    sources=["a.xlsx"],
                    captions=["EBITDA"],
                    proof_level=ProofLevel.L2,
                ),
            ],
        )
        release = create_release(deal, source="manual")
        assert release.counts["proven"] == 1
        again = ensure_current_release(deal, bootstrap=False)
        assert again is not None
        assert again.release_id == release.release_id
    finally:
        store.databook_dir = original


def test_release_path_rejects_traversal(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_release import SlugDeal
    from agetic_cdd_api import services_databook_store as store

    deal = SlugDeal(id="sec", slug="sec")
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        for bad in ("../../etc/passwd", "..\\..\\secret", "v1/../../x", "", ".", ".."):
            try:
                store._release_path(deal, bad)
                raised = False
            except ValueError:
                raised = True
            assert raised, f"expected ValueError for {bad!r}"
        ok = store._release_path(deal, "v12")
        assert ok.name == "v12.json"
        assert ok.is_relative_to((tmp_path / "releases").resolve())
    finally:
        store.databook_dir = original


def test_save_release_immutable_under_lock(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_models import DatabookRelease, ReleasedCell, ReleaseCellStatus
    from agetic_cdd_api.services_databook_release import SlugDeal
    from agetic_cdd_api import services_databook_store as store

    deal = SlugDeal(id="imm", slug="imm")
    original = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    try:
        release = DatabookRelease(
            release_id="v1",
            version=1,
            created_at="2026-01-01T00:00:00Z",
            deal_slug="imm",
            cells=[
                ReleasedCell(
                    metric_key="revenue",
                    fiscal_year=2024,
                    status=ReleaseCellStatus.PROVEN,
                    value=1.0,
                )
            ],
            counts={"proven": 1, "doubtful": 0, "missing": 0},
        )
        store.save_release(deal, release, set_current=True)
        assert store.load_meta(deal).current_release_id == "v1"
        try:
            store.save_release(deal, release, set_current=True)
            assert False, "expected immutable conflict"
        except ValueError as exc:
            assert "already exists" in str(exc)
    finally:
        store.databook_dir = original


def test_is_material_item_no_substring_false_positives() -> None:
    from agetic_cdd_api.services_databook_consume import (
        _is_material_item,
        _normalize_metric_label,
        _strip_material_agent_values,
    )

    assert _is_material_item("revenue", "Revenue") is True
    assert _is_material_item("gross_margin", "Gross Margin") is True
    assert _is_material_item("net_revenue_share", "Net Revenue Share") is False
    assert _is_material_item("gross_margin_bps", "Gross Margin bps") is False
    assert _is_material_item("", "gross margin percentage") is False
    # Non-string / missing fields must not raise.
    assert _normalize_metric_label(None) == ""
    assert _normalize_metric_label(42) == "42"
    assert _is_material_item(None, None) is False
    assert _is_material_item(True, False) is False

    stripped = _strip_material_agent_values(
        [
            {"line_item": "Revenue", "metric_key": "revenue", "fy2024_value": 1.0},
            {"line_item": "Net Revenue Share", "metric_key": "net_revenue_share", "fy2024_value": 2.0},
            {"line_item": "Other", "fy2024_value": 3.0},
            {"line_item": None, "metric_key": None, "fy2024_value": 9.0},
        ]
    )
    by_key = {r.get("metric_key") or r.get("line_item"): r for r in stripped}
    assert "fy2024_value" not in by_key["revenue"]
    assert by_key["revenue"]["databook_status"] == "missing"
    assert by_key["net_revenue_share"]["fy2024_value"] == 2.0
    assert by_key["Other"]["fy2024_value"] == 3.0
    assert by_key[None]["fy2024_value"] == 9.0


def test_released_to_pl_lines_mixed_status_and_custom_order() -> None:
    from agetic_cdd_api.services_databook_consume import (
        _collapse_line_status,
        _display_unit_parts,
        released_to_pl_lines,
    )
    from agetic_cdd_api.services_databook_models import (
        DatabookRelease,
        ReleaseCellStatus,
        ReleasedCell,
    )

    assert _collapse_line_status(["proven", "missing"]) == "mixed"
    assert _collapse_line_status(["proven", "doubtful", "missing"]) == "doubtful"
    assert _display_unit_parts("USD", "USD", "m") == "USD m"
    assert _display_unit_parts(None, "k", None) == "k"

    release = DatabookRelease(
        release_id="v1",
        version=1,
        created_at="2026-01-01T00:00:00Z",
        deal_slug="u",
        cells=[
            ReleasedCell(
                metric_key="zebra_kpi",
                fiscal_year=2024,
                status=ReleaseCellStatus.PROVEN,
                value=1.0,
            ),
            ReleasedCell(
                metric_key="alpha_kpi",
                fiscal_year=2024,
                status=ReleaseCellStatus.PROVEN,
                value=2.0,
            ),
            ReleasedCell(
                metric_key="revenue",
                fiscal_year=2023,
                status=ReleaseCellStatus.PROVEN,
                value=100.0,
            ),
            ReleasedCell(
                metric_key="revenue",
                fiscal_year=2024,
                status=ReleaseCellStatus.MISSING,
                value=None,
            ),
        ],
        counts={"proven": 3, "doubtful": 0, "missing": 1},
    )
    lines = released_to_pl_lines(release)
    keys = [l["metric_key"] for l in lines]
    # Material first, then custom in first-seen (document) order — not alpha.
    assert keys == ["revenue", "zebra_kpi", "alpha_kpi"]
    rev = next(l for l in lines if l["metric_key"] == "revenue")
    assert rev["databook_status"] == "mixed"
    assert rev["fy2023_value"] == 100.0
    assert "fy2024_value" not in rev
    assert 2024 in rev["missing_years"]


def test_released_to_pl_lines_keeps_fy_units(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_consume import released_to_pl_lines
    from agetic_cdd_api.services_databook_models import (
        DatabookRelease,
        ReleaseCellStatus,
        ReleasedCell,
    )

    release = DatabookRelease(
        release_id="v1",
        version=1,
        created_at="2026-01-01T00:00:00Z",
        deal_slug="u",
        cells=[
            ReleasedCell(
                metric_key="revenue",
                fiscal_year=2021,
                status=ReleaseCellStatus.PROVEN,
                value=100.0,
                currency="EUR",
            ),
            ReleasedCell(
                metric_key="revenue",
                fiscal_year=2022,
                status=ReleaseCellStatus.PROVEN,
                value=110.0,
                currency="EUR",
                scale="k",
            ),
        ],
        counts={"proven": 2, "doubtful": 0, "missing": 0},
    )
    lines = released_to_pl_lines(release)
    assert len(lines) == 1
    assert lines[0]["fy_units"]["2021"] == "EUR"
    assert lines[0]["fy_units"]["2022"] == "EUR k"
    assert "EUR" in lines[0]["unit"] and "EUR k" in lines[0]["unit"]


def test_prefer_promoted_emits_deprecation_warning() -> None:
    import warnings

    from agetic_cdd_api.services_databook_consume import prefer_promoted_pl_lines

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always", DeprecationWarning)
        prefer_promoted_pl_lines("no-such-deal-xyz", [])
    assert any(issubclass(w.category, DeprecationWarning) for w in caught)


def test_forecast_period_headers_isolated() -> None:
    from agetic_cdd_api.services_databook_classify import is_forecast_period_header
    from agetic_cdd_api.services_databook_extract import _detect_year_headers

    assert is_forecast_period_header("Year 1") is True
    assert is_forecast_period_header("Y2") is True
    assert is_forecast_period_header("Year 3 (FY2026)") is True
    assert is_forecast_period_header("FY2024") is False
    years = _detect_year_headers(["Line", "FY2022", "FY2023", "Year 1", "Year 2"])
    assert set(years.values()) == {2022, 2023}
    assert 3 not in years and 4 not in years


def test_classify_set_aside_and_ladder(tmp_path: Path) -> None:
    from agetic_cdd_api.services_databook_classify import (
        classify_entry,
        deal_identity_tokens,
        infer_basis,
    )
    from agetic_cdd_api.services_databook_models import FileRelevance, SourceBasis

    class _Deal:
        id = "d1"
        slug = "acme-co"
        name = "Acme Co"
        company = "Acme Industries"

    tokens = deal_identity_tokens(_Deal())  # type: ignore[arg-type]
    assert "acme" in tokens

    personal = classify_entry(
        {"filename": "Personal_Brokerage_Letter.pdf", "library_stem": "p1", "excerpt": ""},
        deal_tokens=tokens,
    )
    assert personal.relevance == FileRelevance.SET_ASIDE

    audited = classify_entry(
        {
            "filename": "Acme_Audited_Financial_Statements_2024.pdf",
            "library_stem": "a1",
            "excerpt": "Independent auditor report",
            "cdl_category": "financial",
        },
        deal_tokens=tokens,
    )
    assert audited.basis == SourceBasis.AUDITED
    assert audited.ladder_score == 100

    draft = classify_entry(
        {
            "filename": "Acme_Draft_Accounts.xlsx",
            "library_stem": "d1",
            "excerpt": "Draft management pack",
        },
        deal_tokens=tokens,
    )
    assert draft.basis == SourceBasis.DRAFT
    assert draft.ladder_score < audited.ladder_score
    assert infer_basis("QoE_Grant_Thornton.pdf", "quality of earnings") == SourceBasis.ADVISER


def test_source_ladder_prefers_audited_in_conflicts() -> None:
    from agetic_cdd_api.services_databook_models import (
        ExtractedRow,
        MetricFamily,
        RowStatus,
        SourceBasis,
    )
    from agetic_cdd_api.services_databook_resolve import detect_conflicts

    rows = [
        ExtractedRow(
            row_id="a",
            doc_id="1",
            source_name="audit.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0,
            status=RowStatus.CANDIDATE,
            source_basis=SourceBasis.AUDITED,
            ladder_score=100,
        ),
        ExtractedRow(
            row_id="b",
            doc_id="2",
            source_name="draft.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=90.0,
            status=RowStatus.CANDIDATE,
            source_basis=SourceBasis.DRAFT,
            ladder_score=40,
        ),
    ]
    issues, hold = detect_conflicts(rows)
    assert hold == {"a", "b"}
    assert len(issues) == 1
    chosen = next(c for c in issues[0]["candidates"] if c["chosen"])
    assert chosen["value"] == 100.0
    assert "audit.pdf" in chosen["sources"]


def test_file_register_endpoint() -> None:
    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={"name": "Classify Deal", "slug": "classify-api-deal", "industry": "generic"},
        )
        deal_id = created.json()["data"]["id"]
        reg = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/file-register",
            headers=headers,
        )
        assert reg.status_code == 200
        body = reg.json()["data"]
        assert body is not None
        assert "entries" in body
        assert "counts" in body


def test_extract_row_id_includes_col_and_absolute_grid_row() -> None:
    """Review fixes: distinct col_idx in row_id; SourceRef.row is absolute grid index."""
    # Fallback path: identical values across columns must not collide.
    table = {
        "name": "wide2",
        "rows": [
            ["Metric", "Actual", "Budget", "Note"],
            ["Operating Expenses", "100", "100", "n/a"],
        ],
    }
    rows = rows_from_table(source_name="a.pdf", doc_id="a", table=table)
    values = [r for r in rows if r.value == 100.0]
    assert len(values) == 2
    assert values[0].row_id != values[1].row_id
    assert {r.source_ref.col for r in values if r.source_ref} == {1, 2}
    assert all(r.source_ref and r.source_ref.row == 1 for r in values)

    # Title row + year header → first body line is grid row 2.
    titled = {
        "name": "pnl",
        "rows": [
            ["Company financials"],
            ["Metric", "FY2023", "FY2024"],
            ["Revenue (USD M)", "10", "12"],
        ],
    }
    titled_rows = rows_from_table(source_name="a.pdf", doc_id="a", table=titled)
    assert titled_rows
    assert all(r.source_ref and r.source_ref.row == 2 for r in titled_rows)


def test_phase3_header_scale_and_eight_labels() -> None:
    from agetic_cdd_api.services_databook_header import parse_table_header, period_end_for_year

    meta = parse_table_header(
        table_name="pnl",
        title_row=["Consolidated income statement — Figures in INR Cr"],
        header_row=["Metric", "FY2022", "FY2023", "FY2024"],
        fy_convention="ending_march",
    )
    assert meta.currency == "INR"
    assert meta.scale == "Cr"
    assert meta.scope == "consolidated"
    assert meta.statement == "income_statement"
    assert meta.period_length == "FY"
    assert period_end_for_year(2024, period_length="FY", fy_convention="ending_march") == "2024-03-31"

    table = {
        "name": "pnl",
        "rows": [
            ["Consolidated P&L (INR Crore)"],
            ["Metric", "FY2022", "FY2023", "FY2024"],
            ["Revenue", "190", "1468", "3140"],
        ],
    }
    rows = rows_from_table(
        source_name="audit.pdf",
        doc_id="audit",
        table=table,
        fy_convention="ending_march",
    )
    assert rows
    rev = next(r for r in rows if r.metric_key == "revenue" and r.fiscal_year == 2024)
    assert rev.currency == "INR"
    assert rev.scale == "Cr"
    assert rev.currency_source == "header"
    assert rev.scale_source == "header"
    assert rev.assumption is False
    assert rev.scope == "consolidated"
    assert rev.statement == "income_statement"
    assert rev.period_length == "FY"
    assert rev.period_end == "2024-03-31"
    assert rev.source_ref is not None
    assert rev.source_ref.table == "pnl"
    assert rev.source_ref.col is not None
    assert rev.source_ref.rule == "year_column"


def test_phase3_caption_scale_is_assumption_not_header() -> None:
    table = {
        "name": "kpi",
        "rows": [
            ["Metric", "FY2023", "FY2024"],
            ["Revenue (USD M)", "10", "12"],
        ],
    }
    rows = rows_from_table(source_name="a.pdf", doc_id="a", table=table)
    assert rows
    assert all(r.scale == "M" and r.currency == "USD" for r in rows)
    assert all(r.scale_source == "caption" for r in rows)
    assert all(r.assumption for r in rows)  # money scale not from header


def test_phase3_prose_goes_to_notes_not_rows() -> None:
    from agetic_cdd_api.services_databook_extract import notes_from_prose

    text = "Revenue in 2024 was 100.50. FY2021 FY2022 Revenue 8 456"
    notes = notes_from_prose(source_name="cim.pdf", doc_id="cim", text=text)
    assert notes
    assert any(n.fiscal_year == 2024 and n.value == 100.5 for n in notes)
    # Notes must never be ExtractedRow statement lines
    assert all(n.reason == "prose" for n in notes)


def test_phase3_scale_step_guard_holds_thousand_misread() -> None:
    from agetic_cdd_api.services_databook_resolve import apply_scale_step_guard

    rows = [
        ExtractedRow(
            row_id=f"r{i}",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2020 + i,
            value=v,
            currency="INR",
            scale="Cr",
        )
        for i, v in enumerate([100.0, 110.0, 90.0, 100_000.0])
    ]
    _, hold_ids, issues = apply_scale_step_guard(rows, params=DealDatabookParams())
    assert "r3" in hold_ids
    assert any(i.get("kind") == "scale_step" for i in issues)

    # Volatile peers [1, 1e6] must not invent a baseline that traps a mid value.
    volatile = [
        ExtractedRow(
            row_id=f"v{i}",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2020 + i,
            value=v,
        )
        for i, v in enumerate([1.0, 1_000_000.0, 1_000.0])
    ]
    _, volatile_hold, _ = apply_scale_step_guard(volatile, params=DealDatabookParams())
    assert "v2" not in volatile_hold


def test_map_wrapper_plainness_and_display_label() -> None:
    from agetic_cdd_api.services_databook_map import caption_plainness, display_metric_label

    # Word-boundary plainness: "arr" must not score inside unrelated words.
    assert caption_plainness("Gross Margin", "arr", sector_pack="saas") == 0.0
    assert caption_plainness("arrival fees", "arr", sector_pack="saas") == 0.0
    assert caption_plainness("ARR", "arr", sector_pack="saas") == 1.0
    # revenue_share_pct matches "Revenue Share (%)" via phrase / alias, not "_pct" literal.
    assert caption_plainness("Revenue Share (%)", "revenue_share_pct") == 1.0

    # Hardcoded labels work even when CoA display helper would be unused.
    assert display_metric_label("revenue") == "Revenue"
    assert display_metric_label("nrr") == "Net Revenue Retention (NRR)"


def test_phase4_coa_dual_agree_and_section_gate() -> None:
    from agetic_cdd_api.services_databook_coa import get_sector_pack, map_caption_coa
    from agetic_cdd_api.services_databook_map import map_caption

    hit = map_caption("Revenue (INR Cr)")
    assert hit is not None
    assert hit.metric_key == "revenue"
    assert hit.coa_id == "is.revenue.total"
    assert hit.dual_agree is True
    assert hit.section == "income_statement"

    # Section gate: cash maps on BS, not on IS (incl. synonyms).
    assert map_caption("Cash", statement="balance_sheet") is not None
    assert map_caption("Cash", statement="income_statement") is None
    assert map_caption("Cash", statement="P&L") is None
    assert map_caption("Cash", statement="bs") is not None

    # SaaS pack adds ARR; generic does not.
    assert map_caption("ARR", sector_pack="generic") is None
    arr = map_caption("Annual Recurring Revenue", sector_pack="saas")
    assert arr is not None and arr.metric_key == "arr"
    # Short alias must not match compound captions.
    assert map_caption("ARR growth", sector_pack="saas") is None
    # Soft qualifier beside a single-token alias is allowed ("total cogs").
    total_cogs = map_caption("Total COGS")
    assert total_cogs is not None and total_cogs.metric_key == "cogs"
    # Ratio / P&L hitch-ons must not map to BS debt or COGS.
    assert map_caption("Direct costs percentage") is None
    assert map_caption("Term loan interest expense", statement="balance_sheet") is None
    assert map_caption("Term loan", statement="balance_sheet") is not None

    pack = get_sector_pack("saas")
    assert any(n.metric_key == "logo_churn" for n in pack.nodes)

    # Dual methods agree on NRR.
    coa = map_caption_coa("Net Revenue Retention (NRR)")
    assert coa is not None and coa.dual_agree and coa.method_a and coa.method_b

    # Meta captions rejected by both methods (no soft single-method hit).
    assert map_caption("CIM page(s)") is None
    assert map_caption("CIM Crosswalk") is None
    # revenue shareholding must not map to revenue_share_pct
    assert map_caption("revenue shareholding") is None
    assert map_caption("Revenue Share") is not None


def test_phase4_derived_not_auto_promoted_and_unmapped_surfaced() -> None:
    from agetic_cdd_api.services_databook_resolve import (
        apply_statuses_and_promote,
        unmapped_issues_from_rows,
    )

    rows = [
        ExtractedRow(
            row_id="e1",
            doc_id="d",
            source_name="a.pdf",
            caption="EBITDA",
            metric_key="ebitda",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=50.0,
            assumption=False,
            derived=True,
            dual_agree=True,
            coa_id="is.ebitda",
        ),
        ExtractedRow(
            row_id="u1",
            doc_id="d",
            source_name="a.pdf",
            caption="Widget adjustment",
            metric_key=None,
            metric_family=MetricFamily.UNKNOWN,
            fiscal_year=2024,
            value=3.0,
        ),
    ]
    out, promoted = apply_statuses_and_promote(rows, hold_ids=set())
    assert promoted == []
    assert out[0].status == RowStatus.CANDIDATE
    unmapped = unmapped_issues_from_rows(rows)
    assert any(i.get("row_id") == "u1" for i in unmapped)


def test_phase4_extract_stamps_coa_fields() -> None:
    table = {
        "name": "pnl",
        "rows": [
            ["Consolidated income statement — Figures in INR Cr"],
            ["Metric", "FY2024"],
            ["Revenue", "3140"],
            ["EBITDA", "400"],
        ],
    }
    rows = rows_from_table(
        source_name="audit.pdf",
        doc_id="a",
        table=table,
        fy_convention="ending_march",
        sector_pack="generic",
    )
    rev = next(r for r in rows if r.metric_key == "revenue")
    assert rev.coa_id == "is.revenue.total"
    assert rev.dual_agree is True
    assert rev.derived is False
    ebitda = next(r for r in rows if r.metric_key == "ebitda")
    assert ebitda.derived is True
    assert ebitda.coa_section == "income_statement"


def test_phase5_proof_levels_bs_balance_and_release_gate() -> None:
    from agetic_cdd_api.services_databook_models import (
        DealDatabookParams,
        ProofLevel,
        ReleaseCellStatus,
    )
    from agetic_cdd_api.services_databook_prove import (
        assign_proof_levels,
        check_balance_sheet,
        meets_min_proof,
        run_prove_pipeline,
    )
    from agetic_cdd_api.services_databook_resolve import apply_statuses_and_promote

    # Dual-agree unique row → L2 → promote
    rows = [
        ExtractedRow(
            row_id="r1",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0,
            assumption=False,
            dual_agree=True,
        )
    ]
    annotated = assign_proof_levels(rows, block_fail_ids=set(), hold_ids=set())
    assert annotated[0].proof_level == ProofLevel.L2
    _, promoted = apply_statuses_and_promote(annotated, hold_ids=set())
    assert len(promoted) == 1
    assert promoted[0].proof_level == ProofLevel.L2

    # Assumption → L1 → no auto-promote
    weak = [
        ExtractedRow(
            row_id="w1",
            doc_id="d",
            source_name="a.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0,
            assumption=True,
            dual_agree=True,
        )
    ]
    weak_a = assign_proof_levels(weak, block_fail_ids=set(), hold_ids=set())
    assert weak_a[0].proof_level == ProofLevel.L1
    _, no_promo = apply_statuses_and_promote(weak_a, hold_ids=set())
    assert no_promo == []

    # BS balance check holds on imbalance
    bs_rows = [
        ExtractedRow(
            row_id="a1",
            doc_id="d",
            source_name="bs.pdf",
            caption="Total assets",
            metric_key="total_assets",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=100.0,
            coa_section="balance_sheet",
            statement="balance_sheet",
            dual_agree=True,
        ),
        ExtractedRow(
            row_id="l1",
            doc_id="d",
            source_name="bs.pdf",
            caption="Total liabilities",
            metric_key="total_liabilities",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=40.0,
            coa_section="balance_sheet",
            statement="balance_sheet",
            dual_agree=True,
        ),
        ExtractedRow(
            row_id="e1",
            doc_id="d",
            source_name="bs.pdf",
            caption="Total equity",
            metric_key="total_equity",
            metric_family=MetricFamily.OTHER,
            fiscal_year=2024,
            value=50.0,  # 40+50 ≠ 100
            coa_section="balance_sheet",
            statement="balance_sheet",
            dual_agree=True,
        ),
    ]
    issues, holds = check_balance_sheet(bs_rows, params=DealDatabookParams())
    assert holds == {"a1", "l1", "e1"}
    assert any(i.get("check_id") == "bs_balance" for i in issues)

    annotated_bs, extra, prove_issues = run_prove_pipeline(
        bs_rows,
        blocks=[],
        block_fail_ids=set(),
        hold_ids=set(),
        params=DealDatabookParams(),
    )
    assert extra
    assert all(r.proof_level == ProofLevel.L1 for r in annotated_bs)
    assert any(i.get("check_id") == "bs_balance" for i in prove_issues)

    assert meets_min_proof(ProofLevel.L2, ProofLevel.L2)
    assert not meets_min_proof(ProofLevel.L1, ProofLevel.L2)
    assert ReleaseCellStatus.PROVEN.value == "proven"


def test_conflict_sources_ordered_by_ladder_not_alpha() -> None:
    """Candidate.sources[0] must be ladder-best, not alphabetical."""
    rows = [
        ExtractedRow(
            row_id="a",
            doc_id="d1",
            source_name="adviser_report.pdf",  # alphabetically first
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0,
            ladder_score=55,
        ),
        ExtractedRow(
            row_id="b",
            doc_id="d2",
            source_name="audited_financials.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=100.0,
            ladder_score=100,
        ),
        ExtractedRow(
            row_id="c",
            doc_id="d3",
            source_name="cim.pdf",
            caption="Revenue",
            metric_key="revenue",
            metric_family=MetricFamily.REVENUE,
            fiscal_year=2024,
            value=200.0,
            ladder_score=25,
        ),
    ]
    issues, _ = detect_conflicts(rows)
    assert len(issues) == 1
    chosen = next(c for c in issues[0]["candidates"] if c["chosen"])
    assert chosen["value"] == 100.0
    assert chosen["sources"][0] == "audited_financials.pdf"
    assert chosen["sources"][0] != "adviser_report.pdf"


def test_classify_review_refinements() -> None:
    from agetic_cdd_api.services_databook_classify import (
        classify_set_aside,
        infer_basis,
        is_forecast_period_header,
        ladder_score_for_source,
        register_by_filename,
    )
    from agetic_cdd_api.services_databook_models import (
        FileRegister,
        FileRegisterEntry,
        FileRelevance,
        FileRole,
        SourceBasis,
    )

    # A: standard doc tokens must not false-positive set-aside
    aside, _ = classify_set_aside(
        filename="Quarterly_Operating_Breakdown.xlsx",
        excerpt="x" * 250,
        deal_tokens={"acme"},
    )
    assert aside is False

    # B: 2-digit FY keeps history; Year N still isolates
    assert is_forecast_period_header("Forecast vs Actual FY24") is False
    assert is_forecast_period_header("Projected FY25") is False
    assert is_forecast_period_header("Forecast column") is True
    assert is_forecast_period_header("Year 1 (FY2026)") is True

    # C: adviser draft packs are ADVISER; audited still wins over adviser brand
    assert infer_basis("QoE_Draft_Report_v2.pdf", "quality of earnings draft") == SourceBasis.ADVISER
    assert infer_basis("KPMG_Audited_Financial_Statements.pdf", "independent auditor") == SourceBasis.AUDITED

    # Path resilience for ladder lookup
    reg = FileRegister(
        deal_slug="x",
        generated_at="t",
        entries=[
            FileRegisterEntry(
                filename="audit.pdf",
                doc_id="a",
                relevance=FileRelevance.IN_SCOPE,
                role=FileRole.HISTORY_SOURCE,
                basis=SourceBasis.AUDITED,
                ladder_score=100,
            )
        ],
    )
    by_file = register_by_filename(reg)
    assert ladder_score_for_source("/tmp/vdr/audit.pdf", by_file) == 100
    assert ladder_score_for_source("audit.PDF", by_file) == 100


def test_revenue_cohort_and_pct_captions_do_not_map() -> None:
    assert map_caption("Revenue by cohort ($)") is None
    assert map_caption("Top-10 accounts % of revenue") is None
    assert map_caption("Revenue") is not None
    assert map_caption("Revenue").metric_key == "revenue"


def test_year_as_value_and_raw_dollar_guards() -> None:
    from agetic_cdd_api.services_databook_extract import (
        is_implausible_money_magnitude,
        is_year_as_value,
        rows_from_table,
    )

    assert is_year_as_value("revenue", 2025, 2025.0) is True
    assert is_year_as_value("revenue", 2025, 13.282) is False
    assert is_implausible_money_magnitude("revenue", 4_161_985.57) is True
    assert is_implausible_money_magnitude("revenue", 7.528) is False

    rows = rows_from_table(
        source_name="wb.xlsx",
        doc_id="d1",
        table={
            "name": "cohort",
            "rows": [
                ["Metric", "FY2023", "FY2024", "FY2025"],
                ["Revenue by cohort ($)", "2023", "2024", "2025"],
                ["Revenue", "7.528", "11.741", "13.282"],
                ["Revenue ($)", "4161985.57", "4889359.66", "5612996.64"],
            ],
        },
    )
    rev = [r for r in rows if r.metric_key == "revenue"]
    assert all(r.caption != "Revenue by cohort ($)" for r in rev)
    assert all(not is_year_as_value(r.metric_key, r.fiscal_year, r.value) for r in rev)
    assert all(not is_implausible_money_magnitude(r.metric_key, r.value) for r in rev)
    assert {(r.fiscal_year, r.value) for r in rev} == {
        (2023, 7.528),
        (2024, 11.741),
        (2025, 13.282),
    }


def test_doubtful_release_omits_numeric_values() -> None:
    from agetic_cdd_api.services_databook_consume import (
        released_as_metric_fact_dicts,
        released_to_pl_lines,
    )
    from agetic_cdd_api.services_databook_models import (
        DatabookRelease,
        ReleaseCellStatus,
        ReleasedCell,
    )

    release = DatabookRelease(
        release_id="v1",
        version=1,
        created_at="2026-01-01T00:00:00Z",
        deal_slug="u",
        cells=[
            ReleasedCell(
                metric_key="revenue",
                fiscal_year=2025,
                status=ReleaseCellStatus.DOUBTFUL,
                value=2025.0,
            ),
            ReleasedCell(
                metric_key="ebitda",
                fiscal_year=2025,
                status=ReleaseCellStatus.PROVEN,
                value=0.572,
            ),
        ],
        counts={"proven": 1, "doubtful": 1, "missing": 0},
    )
    lines = {l["metric_key"]: l for l in released_to_pl_lines(release)}
    assert "fy2025_value" not in lines["revenue"]
    assert lines["revenue"]["databook_status"] == "doubtful"
    assert lines["ebitda"]["fy2025_value"] == 0.572
    facts = released_as_metric_fact_dicts(release)
    assert [(f["family"], f["value"]) for f in facts] == [("ebitda", 0.572)]


def test_ops_kpi_prefers_near_term_not_plan_horizon() -> None:
    from agetic_cdd_api.report_ops_dashboard import MetricFact, derive_kpi_scorecard

    facts = [
        MetricFact(
            family="revenue",
            label="Revenue",
            value=10.5,
            year=2023,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=11.869,
            year=2024,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=13.412,
            year=2025,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        # Full Metric Grid / model years must not become headline "Latest".
        MetricFact(
            family="revenue",
            label="Revenue",
            value=14.3,
            year=2026,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=50.15,
            year=2030,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
    ]
    pl_lines = [
        {
            "line_item": "Revenue",
            "fy2024_value": 11.741,
            "fy2024_period_type": "actual",
            "fy2025_value": 13.282,
            "fy2025_period_type": "actual",
            "fy2026_value": 14.3,
            "fy2026_period_type": "forecast",
            "fy2030_value": 50.15,
            "fy2030_period_type": "forecast",
        }
    ]
    kpis = derive_kpi_scorecard(facts, pl_lines)
    yoy = next(k for k in kpis if k.name == "Revenue Growth (YoY)")
    assert "FY2030" not in str(yoy.latest)
    assert "FY2026" not in str(yoy.latest)
    assert "13.0%" in str(yoy.yoy) or "13%" in str(yoy.yoy)
    assert "FY2025" in str(yoy.latest)


def test_ops_kpi_caps_latest_even_when_plan_years_mis_tagged_actual() -> None:
    """Regression: mis-tagged FY2030 'actual' previously flipped headline YoY."""
    from agetic_cdd_api.report_ops_dashboard import MetricFact, derive_kpi_scorecard

    facts = [
        MetricFact(
            family="revenue",
            label="Revenue",
            value=11.869,
            year=2024,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=13.412,
            year=2025,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
        MetricFact(
            family="revenue",
            label="Revenue",
            value=50.15,
            year=2030,
            agent_key="historical_performance",
            agent_name="Historical",
        ),
    ]
    pl_lines = [
        {
            "line_item": "Revenue",
            "fy2024_value": 11.869,
            "fy2024_period_type": "actual",
            "fy2025_value": 13.412,
            "fy2025_period_type": "actual",
            "fy2030_value": 50.15,
            "fy2030_period_type": "actual",  # wrong tag — still must not win
        }
    ]
    kpis = derive_kpi_scorecard(facts, pl_lines)
    yoy = next(k for k in kpis if k.name == "Revenue Growth (YoY)")
    assert "FY2030" not in str(yoy.latest)
    assert "FY2025" in str(yoy.latest)
    assert "13.0%" in str(yoy.yoy) or "13%" in str(yoy.yoy)


def test_sanitize_dc_contract_framing_rewrites_4m_as_current() -> None:
    from agetic_cdd_api.report_builder_base import sanitize_dc_contract_framing

    raw = (
        "from 12,000 active homes today to 30,000-60,000 homes by 2030, "
        "increasing municipal contract revenue potential from $4.0M to $10.4M-$22.2M."
    )
    fixed = sanitize_dc_contract_framing(raw)
    assert "$4.0M to $10.4M" not in fixed
    assert "$3.0M (FY2025A)" in fixed
    assert "$10.4M-$22.2M" in fixed

    assert "FY2025A" in sanitize_dc_contract_framing(
        "representing $4.0M in current annual revenue under the DC contract"
    )
    untouched = "Plan bridge uses $4.0M FY2026P municipal revenue toward $10.4M."
    assert sanitize_dc_contract_framing(untouched) == untouched


def test_sanitize_dc_homes_expansion_injects_baseline_and_potential() -> None:
    from agetic_cdd_api.report_builder_base import sanitize_dc_contract_framing

    raw = (
        "Expansion of the anchor D.C. municipal contract from ~12,000 homes today "
        "to 30,000 homes by 2030, contributing up to $10.41M in recurring collection revenue."
    )
    fixed = sanitize_dc_contract_framing(raw)
    assert "$3.0M FY2025A" in fixed
    assert "$22.2M full-contract potential" in fixed
    assert "$10.41M" in fixed


def test_sanitize_report_prose_strips_boilerplate_and_inc() -> None:
    from agetic_cdd_api.report_builder_base import sanitize_report_prose

    garbled = (
        "Supply Chain Resilience for Compost Crew "
        "(registered under Maryland and District of Columbia securities exemptions)."
    )
    assert "registered under" not in sanitize_report_prose(garbled, company="Compost Crew")
    assert "Inc." not in sanitize_report_prose(
        "Compost Crew Inc. is a D.C. Metro region organics platform.",
        company="Compost Crew",
    )
    owned = (
        "Ownership: Total evidenced common equity/option shares: 10,202,561 "
        "(10,000,010 + 200,000 + 2,070 + 481). "
        "Ownership: Total evidenced common equity/option shares: 10,202,561 "
        "(10,000,010 + 200,000 + 2,070 + 481)."
    )
    cleaned = sanitize_report_prose(owned)
    assert cleaned.count("Total evidenced") == 1
    assert not cleaned.startswith("Ownership:")


def test_upgrade_generic_doc_cites() -> None:
    from agetic_cdd_api.report_builder_base import upgrade_generic_doc_cites

    raw = "Food scrap collection services. (DOC: data room financials)"
    fixed = upgrade_generic_doc_cites(raw, ["a.pdf", "b.xlsx", "c.pptx"])
    assert "(DOC: [1], DOC: [2], DOC: [3])" in fixed
    assert "data room financials" not in fixed


def test_approx_equal_flags_qbo_model_variance() -> None:
    from agetic_cdd_api.report_ops_dashboard import _approx_equal

    assert _approx_equal(7.558, 7.558) is True
    assert _approx_equal(7.528, 7.558) is False
    assert _approx_equal(11.741, 11.869) is False
    assert _approx_equal(13.282, 13.282) is True


def test_compost_company_name_infers_organics_sector() -> None:
    from agetic_cdd_api.services_accounts_extract import infer_sector_id_from_corpus

    assert infer_sector_id_from_corpus("Compost Crew") == "waste_organics"
    assert (
        infer_sector_id_from_corpus("Financial Services\nCompost Crew CIM")
        == "waste_organics"
    )


def test_phase6_validation_pack_and_learning(tmp_path: Path) -> None:
    """OL-1…OL-7: pack after release, confirm/remap/exclude, memory, traps, lineage."""
    from agetic_cdd_api import services_databook_store as store
    from agetic_cdd_api.services_databook_decisions import confirm_row, remap_row
    from agetic_cdd_api.services_databook_learn import (
        load_lineage,
        load_mapping_memory,
        load_traps,
        lookup_mapping_memory,
    )
    from agetic_cdd_api.services_databook_map import map_caption
    from agetic_cdd_api.services_databook_models import (
        DatabookRelease,
        ExtractedRow,
        MetricFamily,
        PromotedMetric,
        ProofLevel,
        ReleaseCellStatus,
        ReleasedCell,
        RowStatus,
        SourceBasis,
    )
    from agetic_cdd_api.services_databook_pack import build_validation_pack, select_calibration_sample
    from agetic_cdd_api.services_databook_store import save_promoted, save_release, save_rows

    class _Deal:
        id = "deal-p6"
        slug = "phase6-pack"

    deal = _Deal()
    original_dir = store.databook_dir
    store.databook_dir = lambda _deal: tmp_path  # type: ignore[assignment]
    tmp_path.mkdir(parents=True, exist_ok=True)
    try:
        proven = ReleasedCell(
            metric_key="revenue",
            fiscal_year=2024,
            status=ReleaseCellStatus.PROVEN,
            value=100.0,
            row_id="r_rev",
            sources=["a.pdf"],
            captions=["Revenue"],
            statement="IS",
            scope="consolidated",
            period_end="2024-12-31",
            period_length="FY",
            source_basis=SourceBasis.MANAGEMENT,
            proof_level=ProofLevel.L2,
            proof_checks=["tie"],
        )
        proven2 = ReleasedCell(
            metric_key="units_sold",
            fiscal_year=2023,
            status=ReleaseCellStatus.PROVEN,
            value=50.0,
            row_id="r_units",
            sources=["a.pdf"],
            captions=["Units sold"],
            statement="KPI",
        )
        doubtful = ReleasedCell(
            metric_key="ebitda",
            fiscal_year=2024,
            status=ReleaseCellStatus.DOUBTFUL,
            value=None,
            row_id="r_ebitda",
            sources=["a.pdf"],
            captions=["EBITDA"],
            reason="failed prove check",
            statement="IS",
        )
        conflict = ReleasedCell(
            metric_key="gross_profit",
            fiscal_year=2024,
            status=ReleaseCellStatus.PROVEN,
            value=40.0,
            row_id="r_gp",
            sources=["a.pdf"],
            captions=["Gross profit"],
            alternatives=[{"value": 42.0, "source_name": "b.pdf", "row_id": "r_gp_alt"}],
            statement="IS",
        )
        missing = ReleasedCell(
            metric_key="net_income",
            fiscal_year=2024,
            status=ReleaseCellStatus.MISSING,
            value=None,
            reason="required line absent",
        )

        empty = build_validation_pack(deal, persist_timestamp=False)  # type: ignore[arg-type]
        assert empty.ready_for_review is False
        assert empty.counts["total"] == 0

        rows = [
            ExtractedRow(
                row_id="r_rev",
                doc_id="a",
                source_name="a.pdf",
                caption="Total turnover",
                metric_key="revenue",
                metric_family=MetricFamily.REVENUE,
                fiscal_year=2024,
                value=100.0,
                status=RowStatus.PROMOTED,
            ),
            ExtractedRow(
                row_id="r_ebitda",
                doc_id="a",
                source_name="a.pdf",
                caption="EBITDA",
                metric_key="ebitda",
                metric_family=MetricFamily.MARGIN,
                fiscal_year=2024,
                value=20.0,
                status=RowStatus.HELD_OUT,
            ),
            ExtractedRow(
                row_id="r_gp",
                doc_id="a",
                source_name="a.pdf",
                caption="Gross profit",
                metric_key="gross_profit",
                metric_family=MetricFamily.OTHER,
                fiscal_year=2024,
                value=40.0,
                status=RowStatus.PROMOTED,
            ),
        ]
        save_rows(deal, rows)  # type: ignore[arg-type]
        save_promoted(
            deal,  # type: ignore[arg-type]
            [
                PromotedMetric(
                    metric_key="revenue",
                    fiscal_year=2024,
                    value=100.0,
                    row_id="r_rev",
                    sources=["a.pdf"],
                    captions=["Total turnover"],
                ),
                PromotedMetric(
                    metric_key="gross_profit",
                    fiscal_year=2024,
                    value=40.0,
                    row_id="r_gp",
                    sources=["a.pdf"],
                    captions=["Gross profit"],
                ),
            ],
        )
        save_release(
            deal,  # type: ignore[arg-type]
            DatabookRelease(
                release_id="v1",
                version=1,
                created_at="2026-10-04T12:00:00Z",
                deal_slug=deal.slug,
                source="manual",
                note="p6 test",
                cells=[proven, proven2, doubtful, conflict, missing],
                counts={"proven": 2, "doubtful": 1, "missing": 1},
            ),
        )

        pack = build_validation_pack(deal, calibration_limit=3)  # type: ignore[arg-type]
        assert pack.ready_for_review is True
        kinds = {c.kind.value for c in pack.cards}
        assert "doubtful" in kinds
        assert "conflict" in kinds
        assert "missing" in kinds
        assert "calibration" in kinds
        cal = [c for c in pack.cards if c.kind.value == "calibration"]
        assert len(cal) >= 1
        assert any(c.dependents for c in pack.cards if c.kind.value == "doubtful")
        assert any(c.alternatives for c in pack.cards if c.kind.value == "conflict")

        sample = select_calibration_sample([proven, proven2], limit=5)
        assert len(sample) == 2

        remap = remap_row(
            deal,  # type: ignore[arg-type]
            "r_rev",
            reason="Caption is revenue not other",
            actor="tester@example.com",
            metric_key="revenue",
            value=101.0,
        )
        assert remap["action"] == "remap"
        mem = load_mapping_memory(deal)  # type: ignore[arg-type]
        assert mem
        assert lookup_mapping_memory(deal, "Total turnover") is not None  # type: ignore[arg-type]
        hit = map_caption("Total turnover", deal=deal)
        assert hit is not None and hit.metric_key == "revenue"
        assert load_traps(deal)  # type: ignore[arg-type]
        assert load_lineage(deal)  # type: ignore[arg-type]

        confirm = confirm_row(
            deal,  # type: ignore[arg-type]
            "r_ebitda",
            reason="Reviewed source; releasing on my authority",
            actor="tester@example.com",
        )
        assert confirm["action"] == "confirm"
    finally:
        store.databook_dir = original_dir  # type: ignore[assignment]

    with TestClient(app) as client:
        headers = _headers(client)
        created = client.post(
            "/api/v1/deals",
            headers=headers,
            json={
                "name": "P6 Pack",
                "slug": f"p6-{__import__('uuid').uuid4().hex[:8]}",
                "industry": "generic",
            },
        )
        deal_id = created.json()["data"]["id"]
        pack_res = client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/validation-pack",
            headers=headers,
        )
        assert pack_res.status_code == 200, pack_res.text
        assert pack_res.json()["data"]["ready_for_review"] is False
        assert client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/mapping-memory",
            headers=headers,
        ).status_code == 200
        assert client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/traps",
            headers=headers,
        ).status_code == 200
        assert client.get(
            f"/api/v1/portfolios/{deal_id}/cdd/databook/decisions",
            headers=headers,
        ).status_code == 200
