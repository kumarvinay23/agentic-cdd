"""v4 CDD deck: chapter shape, charts, and no invented reference-deal figures."""

from __future__ import annotations

from pptx import Presentation

from agetic_cdd_api.report_builder_base import BuildContext, DealProfile
from agetic_cdd_api.report_cdd_deck_v4 import plan_cdd_deck, render_cdd_deck
from agetic_cdd_api.report_storyline import default_storyline


def _ctx(outputs: dict) -> BuildContext:
    return BuildContext(
        deal_slug="sample-deal",
        report_type="cdd_deck",
        profile=DealProfile(deal_slug="sample-deal", company="Northwind Organics", sector="waste"),
        storyline=default_storyline("cdd_deck"),
        agent_outputs=outputs,
    )


def _rich() -> dict:
    return {
        "historical_performance": {
            "spec": {
                "period_facts": [
                    {"metric": "Revenue", "period": "2024", "period_type": "actual", "value": 8.5},
                    {"metric": "Revenue", "period": "2025", "period_type": "actual", "value": 13.41},
                    {"metric": "Revenue", "period": "2030", "period_type": "plan", "value": 50.2},
                ],
                "growth_and_margins": [
                    {"metric": "EBITDA margin", "period": "2025", "value": 4.3},
                ],
                "quality_verdict": "PASS",
                "reliance_verdict": "LIMITED",
                "primary_sources": ["QBO statement of operations"],
            }
        },
        "company_background": {
            "spec": {
                "history_events": [{"year": "2019", "event": "First municipal award"}],
                "primary_sources": ["CIM"],
            }
        },
        "customer_segmentation": {
            "spec": {
                "segments": [
                    {"name": "Municipal", "share_pct": 36, "key_driver": "One contract"},
                    {"name": "Commercial", "share_pct": 34, "key_driver": "Location ARPU"},
                    {"name": "Residential", "share_pct": 30, "key_driver": "Subscriptions"},
                ],
                "geo_mix": [
                    {"name": "District core", "share_pct": 60},
                    {"name": "Adjacent county", "share_pct": 40},
                ],
            }
        },
        "market_volume_and_growth": {
            "spec": {
                "tam": {"value": 400, "scale": "million"},
                "sam": {"value": 139, "scale": "million"},
                "som": {"value": 40, "scale": "million"},
                "bottom_up": {
                    "units": "homes x price",
                    "capture_share": "30% haircut",
                    "realised_price": "company price",
                    "arithmetic": "count x price x haircuts",
                    "bottom_up_size": "see SOM",
                },
                "plan_vs_market": {
                    "plan_rate": "30%",
                    "market_rate": "8%",
                    "plain_english": "The plan outruns the market rate.",
                    "what_must_be_true": "Win rate holds.",
                },
                "lines": [
                    {"name": "Hauling", "value": 100, "scale": "million"},
                    {"name": "Products", "value": 39, "scale": "million"},
                ],
                "primary_sources": ["Company market workbook"],
            }
        },
        "competitive_differentiation": {
            "spec": {
                "differentiators": [
                    {"claim": "Retention", "evidence": "Workbook GRR"},
                    {"claim": "Brand", "evidence": ""},
                ],
                "findings": ["Protection rests on one contract."],
            }
        },
        "competitor_identification": {
            "spec": {
                "competitors": [
                    {"name": "Local Hauler A", "geography": "Same city", "overlap": "Residential"},
                    {"name": "Local Hauler B", "geography": "Same city", "overlap": "Commercial"},
                ]
            }
        },
        "customer_stickiness": {
            "spec": {
                "retention_on_same_population": {
                    "population": "commercial accounts",
                    "grr": "87.3%",
                    "nrr": "106.6%",
                }
            }
        },
        "growth_opportunities": {
            "spec": {"levers": [{"name": "Density", "effect": "+locations"}]}
        },
        "executive_summary": {
            "spec": {"insight_snapshot": "A thinly profitable collector; the plan is untested."}
        },
        "demand_drivers": {"spec": {"findings": ["Municipal programmes drive volume."]}},
        "market_definition": {
            "spec": {
                "perimeter": "Residential and commercial collection in the home region.",
                "excluded": ["National roll-up"],
            }
        },
    }


def _slide_text(prs: Presentation) -> str:
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def test_rich_deck_follows_v4_chapters_and_charts():
    ctx = _ctx(_rich())
    prs, plan = render_cdd_deck(ctx)
    text = _slide_text(prs)
    assert "DiligenceIQ" in text
    assert "Agentic CDD" not in text
    assert "Compost Crew" not in text
    assert "Northwind Organics" in text
    assert "BUSINESS INTRODUCTION" in text
    assert "INVESTMENT HIGHLIGHTS" in text
    assert "MARKET DYNAMICS" in text
    assert "COMPETITIVE LANDSCAPE" in text
    assert "GROWTH OPPORTUNITIES" in text
    assert "APPENDIX A" in text
    assert "APPENDIX B" in text
    assert any(
        "2030P" in page.title or any("2030P" in cell for row in page.rows for cell in row)
        for page in plan.pages
    )
    charts = sum(1 for slide in prs.slides for shape in slide.shapes if shape.has_chart)
    assert charts >= 3
    assert len(prs.slides) >= 20
    assert any(page.kind == "divider" for page in plan.pages)


def test_thin_deck_withholds_pages_instead_of_inventing_figures():
    ctx = _ctx({"company_background": {"spec": {"insight_snapshot": "A regional operator."}}})
    plan = plan_cdd_deck(ctx)
    text_bits = " ".join(page.title for page in plan.pages)
    assert "$50.2" not in text_bits
    assert "Compost Crew" not in text_bits
    assert any("Sizing" in why or "TAM" in why for _name, why in plan.withheld)
    blob = " ".join(page.title + page.subtitle for page in plan.pages)
    assert "withheld" in blob.lower() or any(page.section.startswith("APPENDIX B") for page in plan.pages)


def test_preview_uses_cdd_sections_not_market_intel_label():
    from pathlib import Path
    import tempfile

    from agetic_cdd_api.report_preview import preview_pptx

    ctx = _ctx(_rich())
    prs, _plan = render_cdd_deck(ctx)
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "deck.pptx"
        prs.save(str(path))
        payload = preview_pptx(path)
    titles = [slide["title"] for slide in payload["slides"]]
    eyebrows = [slide.get("eyebrow") or "" for slide in payload["slides"]]
    blob = " ".join(titles + eyebrows)
    assert "Market Intel" not in blob
    assert any("ENGAGEMENT OVERVIEW" in item for item in eyebrows)
    assert any("Commercial Due Diligence: Northwind Organics" in title for title in titles)
    assert any(title.startswith("Revenue was") for title in titles)


def test_skill_template_renders_cdd_deck():
    from pathlib import Path
    import tempfile

    from pptx import Presentation

    from agetic_cdd_api.report_cdd_skill import render_cdd_with_skill

    ctx = _ctx(_rich())
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "Northwind_CDD_Deck.pptx"
        render_cdd_with_skill(ctx, path)
        prs = Presentation(str(path))
        text = _slide_text(prs)
    assert "Market Intel Deck" not in text
    assert "Commercial Due Diligence: Northwind Organics" in text
    assert "DiligenceIQ" in text
    assert "ENGAGEMENT OVERVIEW" in text.upper() or "Engagement Overview" in text
    assert len(prs.slides) >= 12
    assert any(shape.has_chart for slide in prs.slides for shape in slide.shapes)


def test_skill_drops_plan_years_and_wrong_exhibits():
    from agetic_cdd_api.report_cdd_skill import build_skill_content

    outputs = _rich()
    outputs["historical_performance"]["spec"]["period_facts"] = [
        {"metric": "Revenue", "period": "2023", "period_type": "actual", "value": 7.56},
        {"metric": "Revenue", "period": "2024", "period_type": "actual", "value": 11.87},
        {"metric": "Revenue", "period": "2025", "period_type": "actual", "value": 13.41},
        {"metric": "Revenue", "period": "2026", "period_type": "actual", "value": 14.3},
        {"metric": "Revenue", "period": "2028", "period_type": "actual", "value": 24.5},
        {"metric": "Revenue", "period": "2029", "period_type": "actual", "value": 36.0},
        {"metric": "Revenue", "period": "TTM to Jun 2026", "period_type": "ttm", "value": 13.61},
    ]
    outputs["historical_performance"]["spec"]["findings"] = [
        "Quality PASS · Reliance BLOCKED",
        "Revenue grew on the municipal contract.",
    ]
    outputs["company_background"]["spec"]["history_events"] = [
        {"year": "2018", "event": "Founder acquires the operating company"},
        {"year": "2023", "event": "CS-01 10,000,010 shares issued to the founder"},
        {"year": "2023", "event": "Option plan GR-2023-11"},
    ]
    outputs["market_definition"]["spec"]["perimeter"] = (
        "Total evidenced common equity/option shares: 1,662,075 (1,661,465 + 310)."
    )
    outputs["market_volume_and_growth"]["spec"]["bottom_up"] = {
        "units": "23,000 (DOC: securities-ledger)",
        "capture_share": "Information request",
        "realised_price": "company price",
        "arithmetic": "1,661,465 homes x $300",
        "bottom_up_size": "see SOM",
    }
    outputs["competitive_differentiation"]["spec"]["findings"] = [
        "Test UNTESTED: Information request",
        "QA 2/7 PASS Quality REWORK · Reliance BLOCKED",
    ]
    content, _withheld = build_skill_content(_ctx(outputs))
    slides = content["slides"]
    snap = slides["S009_company_snapshot"]
    title = snap["title"]["text"]
    assert "2025" in title
    assert "2028" not in title
    assert "2029" not in title
    assert "$36.0M" not in title
    categories = snap["elements"]["right"]["categories"]
    assert categories == ["2023", "2024", "2025"]
    snap_blob = str(snap)
    assert "Quality PASS" not in snap_blob
    assert "TTM to Jun 2026" in snap_blob
    history = str(slides["S010_platform_evolution"])
    assert "CS-01" not in history
    assert "GR-2023" not in history
    assert "acquires the operating company" in history
    assert "S034_market_definitions" not in slides
    method = str(slides["S089_sizing_method_core"])
    assert "securities" not in method.lower()
    assert "Information request" not in method
    assert "homes" in method
    assert "S015_barriers" not in slides


def test_evidence_register_cites_the_file_not_deal_modules():
    import sys
    from pathlib import Path

    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content, skill_dir

    outputs = _rich()
    outputs["historical_performance"]["spec"]["period_facts"] = [
        {"metric": "Revenue", "period": "2024", "period_type": "actual", "value": 11.87},
        {"metric": "Revenue", "period": "2025", "period_type": "actual", "value": 13.41},
        {"metric": "Revenue", "period": "2029", "period_type": "actual", "value": 36.0},
    ]
    outputs["historical_performance"]["spec"]["primary_sources"] = ["QBO statement of operations.xlsx"]
    ctx = _ctx(outputs)
    book = build_agent_evidence(ctx)
    content, _withheld = build_skill_content(ctx, book)
    refs = content["slides"]["S009_company_snapshot"]["title"]["refs"]
    assert "DEAL-1" not in refs
    assert "RUN-1" not in refs
    fact = book.by_id[refs[0]]
    assert fact["module_title"] == "Historical Performance"
    assert fact["citations"]["doc"] == ["QBO statement of operations.xlsx"]
    assert fact["period_status"].get("actual") == "actual"
    assert fact["grade"] == "established"
    assert all("2029" not in book.by_id[ref]["text"] for ref in refs)
    sys.path.insert(0, str(Path(skill_dir()) / "scripts"))
    from render_deck import source_line

    line = source_line(refs, book.by_id, {}, False)
    assert "Deal modules" not in line
    assert "Historical Performance" in line
    assert "QBO statement of operations" in line


def test_vdr_extract_merges_accounts_into_the_register(tmp_path):
    from openpyxl import Workbook

    from agetic_cdd_api.report_cdd_evidence import prepare_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content

    vdr = tmp_path / "vdr"
    vdr.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "P&L"
    sheet.append(["Line", 2023, 2024, 2025, 2029])
    sheet.append(["Total Revenue", 7_560_000, 11_870_000, 13_410_000, 36_000_000])
    workbook.save(vdr / "Management accounts.xlsx")
    ctx = _ctx({"historical_performance": {"spec": {"primary_sources": ["Management accounts.xlsx"]}}})
    book = prepare_evidence(ctx, tmp_path / "work", vdr)
    assert any(fact.get("source") == "vdr_direct" and fact.get("module") == "HPF" for fact in book.facts)
    assert any("2025" in label for label, _value, _fid, status in book.vdr_points if status == "actual")
    assert not any(status == "actual" and label in {"2029", "2029P"} for label, _v, _f, status in book.vdr_points)
    content, _withheld = build_skill_content(ctx, book)
    snap = content["slides"]["S009_company_snapshot"]
    assert "2029" not in snap["title"]["text"]
    assert "2025" in snap["title"]["text"]
    cited = [book.by_id[ref] for ref in snap["elements"]["right"]["refs"] if ref in book.by_id]
    assert any(fact.get("source") == "vdr_direct" for fact in cited)
    assert any("Management accounts" in doc for fact in cited for doc in fact["citations"]["doc"])


def test_step3_calculations_are_traced_and_plan_stays_off_the_snapshot(tmp_path):
    import json
    import subprocess
    import sys

    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content, skill_dir

    ctx = _ctx(_rich())
    book = build_agent_evidence(ctx)
    content, _withheld = build_skill_content(ctx, book)
    ids = {calc["id"] for calc in content["calculations"]}
    assert "CALC-YOY-2025" in ids
    assert "CALC-BRIDGE-DELTA" in ids
    assert "CALC-MOMENTUM" in ids
    assert "CALC-NRR-GRR" in ids
    assert "CALC-HAIRCUT-1" not in ids
    slides = content["slides"]
    assert "S023_h1_growth_drivers" in slides
    assert "S046_sam_growth_build" in slides
    assert "S047_scenario_analysis" in slides
    assert "S068_churn_triggers" in slides
    assert "S042_gf2_mix_shift" in slides
    assert "Hauling" in slides["S037_tam_walkdown"]["title"]["text"]
    assert "2029" not in slides["S009_company_snapshot"]["title"]["text"]
    assert "2030" not in slides["S009_company_snapshot"]["title"]["text"]

    outputs = _rich()
    outputs["market_volume_and_growth"]["spec"]["bottom_up"]["bottom_up_size"] = "$40 million"
    outputs["market_volume_and_growth"]["spec"]["bottom_up"]["capture_share"] = "two 30% haircuts"
    content_hair, _ = build_skill_content(_ctx(outputs))
    hair = [calc["output"] for calc in content_hair["calculations"] if calc["id"].startswith("CALC-HAIRCUT")]
    assert hair == [28.0, 19.6]

    work = tmp_path / "work"
    work.mkdir()
    (work / "content.json").write_text(json.dumps(content), encoding="utf-8")
    (work / "evidence.json").write_text(json.dumps({"facts": book.facts}), encoding="utf-8")
    script = skill_dir() / "scripts" / "check_content.py"
    result = subprocess.run(
        [sys.executable, str(script), "--content", str(work / "content.json"), "--work", str(work)],
        cwd=str(skill_dir()),
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_step4_competitive_and_growth_exhibits_are_traced(tmp_path):
    import json
    import subprocess
    import sys

    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content, skill_dir

    outputs = _rich()
    outputs["competitor_identification"]["spec"]["competitors"] = [
        {"name": "Civic Cart", "geography": "Same city", "overlap": "Residential", "classification": "Local", "in_geography": True},
        {"name": "National Hauler", "geography": "National", "overlap": "Municipal", "classification": "National", "in_geography": False},
    ]
    outputs["competitive_differentiation"]["spec"]["self_assessment"] = {
        "columns": ["Northwind Organics", "National hauler", "Local operator"],
        "rows": [
            {"criterion": "Collection is the core business", "scores": {"Northwind Organics": "yes", "National hauler": "no", "Local operator": "yes"}},
            {"criterion": "Owns processing", "scores": {"Northwind Organics": "no", "National hauler": "yes", "Local operator": "no"}},
        ],
    }
    outputs["market_pricing"] = {"spec": {"price_points": [
        {"name": "Company book", "price": 20, "unit": "$ per month", "role": "company"},
        {"name": "Civic Cart", "price": 25, "unit": "$ per month", "role": "peer"},
    ], "primary_sources": ["Price sheet"]}}
    outputs["customer_stickiness"]["spec"]["retention_by_year"] = [
        {"year": "2024", "grr": 80, "nrr": 110},
        {"year": "2025", "grr": 90, "nrr": 120},
    ]
    outputs["customer_stickiness"]["spec"]["logo_history"] = [
        {"year": "2024", "new": 20, "churned": 6},
        {"year": "2025", "new": 14, "churned": 9},
    ]
    outputs["growth_opportunities"]["spec"]["levers"] = [
        {"name": "Density", "actual": 12, "plan": 18, "actual_label": "2025", "plan_label": "2030P"},
        {"name": "Contracts", "actual": 8, "plan": 14, "actual_label": "2025", "plan_label": "2030P"},
    ]
    outputs["growth_opportunities"]["spec"]["sales_quotas"] = [
        {"driver": "Commercial", "baseline": "7 accounts a year", "growth_case": "flat", "quota": 7},
    ]
    outputs["growth_opportunities"]["spec"]["zip_scores"] = [
        {"territory": "North", "area": "Arlington", "county": "Arlington", "zip": "10001", "score": 7.25},
        {"territory": "South", "area": "Fairfax", "county": "Fairfax", "zip": "10002", "score": 6.40},
    ]
    outputs["growth_opportunities"]["spec"]["initiatives"] = [
        {"name": "Riverside plant", "status": "Permitted", "actual": 2, "plan": 4, "actual_label": "2024", "plan_label": "2027"},
        {"name": "Bagged goods", "status": "Trading", "actual": 3, "plan": 6, "actual_label": "2024", "plan_label": "2027"},
    ]
    outputs["growth_opportunities"]["spec"]["revenue_total"] = 20
    outputs["synergies"] = {"spec": {"acquisitions": [
        {"name": "Fieldstone", "status": "Proposed", "scale": "One city route", "price": 4, "multiple": 8, "synergy_ebitda": 1.2},
    ]}}

    ctx = _ctx(outputs)
    book = build_agent_evidence(ctx)
    content, _withheld = build_skill_content(ctx, book)
    slides = content["slides"]
    calcs = {calc["id"]: calc["output"] for calc in content["calculations"]}
    assert "Outside geography" in slides["S060_competitor_archetypes"]["elements"]["table"]["rows"][1][1]
    assert "self-assessment" in slides["S062_criteria_vs_archetypes"]["subtitle"]["text"]
    assert "S063_feature_benchmarking" not in slides
    assert "S065_incumbent_nps" not in slides
    assert calcs["CALC-PRICE-GAP"] == -20.0
    assert slides["S079_l3_expand_accounts"]["elements"]["chart"]["categories"] == ["2024", "2025"]
    assert len(slides["S079_l3_expand_accounts"]["elements"]["chart"]["series"]) == 2
    assert calcs["CALC-LEVER-TOTAL"] == 12.0
    assert calcs["CALC-LOGO-NET-2025"] == 5.0
    assert calcs["CALC-LOGO-QUOTA"] == 2.0
    assert "7.25" in slides["S081_l5_adjacent_segments"]["elements"]["table"]["rows"][0][4]
    assert calcs["CALC-INIT-SHARE"] == 25.0
    assert calcs["CALC-INIT-CAGR-1"] == 26.0
    assert calcs["CALC-DEAL-EBITDA-Fieldstone"] == 0.5
    assert calcs["CALC-DEAL-RESIDUAL-Fieldstone"] == 0.7
    assert "2030" not in slides["S009_company_snapshot"]["title"]["text"]

    work = tmp_path / "work"
    work.mkdir()
    (work / "content.json").write_text(json.dumps(content), encoding="utf-8")
    (work / "evidence.json").write_text(json.dumps({"facts": book.facts}), encoding="utf-8")
    script = skill_dir() / "scripts" / "check_content.py"
    result = subprocess.run(
        [sys.executable, str(script), "--content", str(work / "content.json"), "--work", str(work)],
        cwd=str(skill_dir()),
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_step5_appendix_adjusts_verdicts_and_names_real_gaps(tmp_path):
    import json
    import subprocess
    import sys

    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content, skill_dir

    outputs = _rich()
    outputs["market_risk"] = {
        "spec": {
            "quality_verdict": "PASS",
            "risks": [{"name": "Weather", "severity": "High"}],
        }
    }
    ctx = _ctx(outputs)
    book = build_agent_evidence(ctx)
    content, _withheld = build_skill_content(ctx, book)
    slides = content["slides"]

    def rows_of(slide: dict | list) -> list:
        pages = slide if isinstance(slide, list) else [slide]
        return [row for page in pages for row in page["elements"]["table"]["rows"]]

    verdicts = rows_of(slides["S099_module_verdicts"])
    historical = next(row for row in verdicts if row[0] == "Historical Performance")
    assert "LIMITED" in historical[1]
    assert historical[2] == "Limited"
    risk = next(row for row in verdicts if row[0] == "Market Risk")
    assert risk[1] == "PASS"
    assert risk[2] == "Asserted"
    assert "Lowered" in risk[3]

    gaps = json.dumps(slides["S101_exceptions"])
    assert "No exhibit in the loaded modules" not in gaps
    assert "Plan years" in gaps
    assert "survey" in gaps.lower()
    assert "Cited only" in json.dumps(slides["S100_source_register"])

    glossary = slides["S007_glossary"]["elements"]["table"]["rows"]
    terms = [row[0] for row in glossary]
    assert "TAM" in terms
    assert "NRR" in terms
    assert "ARR" not in terms

    work = tmp_path / "work"
    work.mkdir()
    (work / "content.json").write_text(json.dumps(content), encoding="utf-8")
    (work / "evidence.json").write_text(json.dumps({"facts": book.facts}), encoding="utf-8")
    script = skill_dir() / "scripts" / "check_content.py"
    result = subprocess.run(
        [sys.executable, str(script), "--content", str(work / "content.json"), "--work", str(work)],
        cwd=str(skill_dir()),
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_room_exhibits_use_workbook_and_cim_tables():
    import sys
    from pathlib import Path

    from agetic_cdd_api.report_cdd_calcs import DeckMath, _cim_deals, _room_cohorts, _room_ties
    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence

    sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "tools" / "cdd-report-builder" / "scripts"))
    from vdr_extract import fmt_value

    assert fmt_value(0.753846, "auto", "2021 cohort", "Logo retention", False, "2025") == "75.4%"
    assert fmt_value(1, "auto", "2021 cohort", "Logo retention", False, "2021") == "100.0%"
    assert "$" in fmt_value(851452.6, "auto", "2021 cohort", "Revenue by cohort ($)", False, "2025")

    math = DeckMath()
    _room_ties(math, [
        {
            "id": "HPF-A",
            "section_title": "Statement of operations (management accounts, annual roll-up)",
            "row_label": "Total Revenues",
            "cells": {"2024": "$8.00M", "2025": "$10.00M"},
        },
        {
            "id": "HPF-W",
            "section_title": "Financial summary (company workbook)",
            "row_label": "Revenue",
            "cells": {"2024": "$7.90M", "2025": "$9.80M"},
        },
        {
            "id": "CBG-C",
            "section_title": "CIM crosswalk: CIM figures traced to support",
            "row_label": "Revenue 2025",
            "cells": {"Value in the CIM": "$9.8mn"},
        },
    ])
    assert math.ties["rows"][-1] == ["2025", "$10.00M", "$9.80M", "$9.8mn"]
    _room_cohorts(math, [
        {
            "id": "CST-1",
            "section_title": "Customer cohorts",
            "row_label": "2024 cohort",
            "cells": {"Accounts": "10", "2024": "100.0%", "2025": "90.0%"},
        },
        {
            "id": "CST-2",
            "section_title": "Customer cohorts",
            "row_label": "2025 cohort",
            "cells": {"Accounts": "4", "2024": "100.0%", "2025": "95.0%"},
        },
        {
            "id": "CST-3",
            "section_title": "Customer cohorts",
            "row_label": "2024 cohort",
            "cells": {"Accounts": "10", "2025": "1200000"},
        },
    ])
    assert [row[0] for row in math.cohorts["rows"]] == ["2024 cohort", "2025 cohort"]
    assert "90.0%" in math.cohorts["rows"][0]

    book = build_agent_evidence(_ctx({}))
    deals = DeckMath()
    _cim_deals(
        book,
        deals,
        [[
            ["Target", "Status", "Description & Rationale"],
            ["Northwind Rollup", "Closed (2024)", "About 1.5 million of revenue"],
        ]],
        ["Structured terms. 4.0 purchase price at 3.0x EBITDA."],
        ["cim.pptx"],
    )
    assert deals.deals["rows"][0]["name"] == "Northwind Rollup"
    assert "1.5 million" in deals.deals["rows"][0]["terms_text"]
    assert "4.0" in deals.deals["note"]


def test_step6_rendered_deck_covers_the_reference_pages(tmp_path):
    import json
    import subprocess
    import sys
    from pathlib import Path

    from agetic_cdd_api.report_cdd_diff import diff_decks
    from agetic_cdd_api.report_cdd_evidence import build_agent_evidence
    from agetic_cdd_api.report_cdd_skill import build_skill_content, render_cdd_with_skill, skill_dir

    outputs = _rich()
    outputs["historical_performance"]["spec"]["growth_and_margins"].append(
        {"metric": "Gross margin", "period": "2024", "value": 28.0}
    )
    outputs["customer_segmentation"]["spec"]["segments"][0]["price"] = "18 per stop"
    outputs["customer_segmentation"]["spec"]["segments"][0]["locations"] = "12 sites"
    outputs["customer_segmentation"]["spec"]["segments"][1]["price"] = "40 per site"
    outputs["customer_segmentation"]["spec"]["segments"][1]["locations"] = "9 sites"
    outputs["cost_structure"] = {"spec": {"cost_lines": [
        {"name": "Disposal", "value": 4.2, "share": "31%"},
        {"name": "Labour", "value": 3.1, "share": "23%"},
    ]}}
    outputs["market_risk"] = {"spec": {
        "risks": [{"name": "Weather", "severity": "High", "note": "Seasonal volume"}],
        "quality_verdict": "PASS",
    }}
    outputs["competitor_identification"]["spec"]["competitors"] = [
        {"name": "Civic Cart", "geography": "Same city", "overlap": "Residential", "classification": "Local", "in_geography": True},
        {"name": "National Hauler", "geography": "National", "overlap": "Municipal", "classification": "National", "in_geography": False},
    ]
    outputs["competitive_differentiation"]["spec"]["self_assessment"] = {
        "columns": ["Northwind Organics", "National hauler", "Local operator"],
        "rows": [
            {"criterion": "Collection is the core business", "scores": {"Northwind Organics": "yes", "National hauler": "no", "Local operator": "yes"}},
            {"criterion": "Owns processing", "scores": {"Northwind Organics": "no", "National hauler": "yes", "Local operator": "no"}},
        ],
    }
    outputs["market_pricing"] = {"spec": {
        "price_points": [
            {"name": "Company book", "price": 20, "unit": "$ per month", "role": "company"},
            {"name": "Civic Cart", "price": 25, "unit": "$ per month", "role": "peer"},
        ],
        "price_levers": [
            {"name": "Municipal", "prior": 20, "current": 21},
            {"name": "Commercial", "prior": 40, "current": 42},
        ],
        "primary_sources": ["Price sheet"],
    }}
    outputs["customer_stickiness"]["spec"]["retention_by_year"] = [
        {"year": "2024", "grr": 80, "nrr": 110},
        {"year": "2025", "grr": 90, "nrr": 120},
    ]
    outputs["customer_stickiness"]["spec"]["logo_history"] = [
        {"year": "2024", "new": 20, "churned": 6},
        {"year": "2025", "new": 14, "churned": 9},
    ]
    outputs["growth_opportunities"]["spec"]["levers"] = [
        {"name": "Density", "actual": 12, "plan": 18, "actual_label": "2025", "plan_label": "2030P"},
        {"name": "Contracts", "actual": 8, "plan": 14, "actual_label": "2025", "plan_label": "2030P"},
    ]
    outputs["growth_opportunities"]["spec"]["sales_quotas"] = [
        {"driver": "Commercial", "baseline": "7 accounts a year", "growth_case": "flat", "quota": 7},
    ]
    outputs["growth_opportunities"]["spec"]["zip_scores"] = [
        {"territory": "North", "area": "Arlington", "county": "Arlington", "zip": "10001", "score": 7.25},
        {"territory": "South", "area": "Fairfax", "county": "Fairfax", "zip": "10002", "score": 6.40},
    ]
    outputs["growth_opportunities"]["spec"]["initiatives"] = [
        {"name": "Riverside plant", "status": "Permitted", "actual": 2, "plan": 4, "actual_label": "2024", "plan_label": "2027"},
        {"name": "Bagged goods", "status": "Trading", "actual": 3, "plan": 6, "actual_label": "2024", "plan_label": "2027"},
    ]
    outputs["growth_opportunities"]["spec"]["revenue_total"] = 20
    outputs["synergies"] = {"spec": {"acquisitions": [
        {"name": "Fieldstone", "status": "Proposed", "scale": "One city route", "price": 4, "multiple": 8, "synergy_ebitda": 1.2},
    ]}}
    outputs["market_volume_and_growth"]["spec"]["adjacent_lines"] = [
        {"name": "Bagged goods", "basis": "volume times a stated price", "som": 10, "actual": 2},
        {"name": "Bulk amendments", "basis": "order count times a stated price", "som": 8, "actual": 1},
    ]

    ctx = _ctx(outputs)
    book = build_agent_evidence(ctx)
    content, _withheld = build_skill_content(ctx, book)
    calcs = {calc["id"]: calc["output"] for calc in content["calculations"]}
    assert calcs["CALC-PRICE-MOVE-1"] == 5.0
    assert calcs["CALC-ADJ-CAPTURE-1"] == 20.0
    assert calcs["CALC-ADJ-CAPTURE-2"] == 12.5
    assert "S080_l4_pricing" in content["slides"]
    assert "S090_sizing_method_adjacent" in content["slides"]
    assert "S012_business_model_flow" in content["slides"]
    assert "S013_business_model_comparison" not in content["slides"]
    assert "2030" not in content["slides"]["S009_company_snapshot"]["title"]["text"]

    work = tmp_path / "work"
    work.mkdir()
    (work / "content.json").write_text(json.dumps(content), encoding="utf-8")
    (work / "evidence.json").write_text(json.dumps({"facts": book.facts}), encoding="utf-8")
    script = skill_dir() / "scripts" / "check_content.py"
    checked = subprocess.run(
        [sys.executable, str(script), "--content", str(work / "content.json"), "--work", str(work)],
        cwd=str(skill_dir()),
        capture_output=True,
        text=True,
    )
    assert checked.returncode == 0, checked.stdout + checked.stderr

    out = tmp_path / "Northwind_CDD_Deck.pptx"
    render_cdd_with_skill(ctx, out)
    reference = Path(__file__).resolve().parents[3] / "Compost_Crew_CDD_v4 (2).pptx"
    gap = diff_decks(out, reference)
    assert gap["missing"] == [], gap


def test_vdr_total_revenue_collapses_to_one_series():
    from agetic_cdd_api.report_cdd_calcs import derive
    from agetic_cdd_api.report_cdd_evidence import EvidenceBook, _vdr_plan_points, _vdr_revenue_points

    def fact(fid, label, section, cells, status, module="HPF", measures=None):
        return {
            "id": fid,
            "source": "vdr_direct",
            "module": module,
            "module_title": section,
            "section_title": section,
            "row_label": label,
            "text": " | ".join([label, *[f"{key}: {value}" for key, value in cells.items()]]),
            "cells": {"Line": label, **cells},
            "period_status": status,
            "measures": measures or [],
            "citations": {"doc": ["Management accounts.xlsx"], "web": []},
        }

    book = EvidenceBook()
    book.facts = [
        fact("HPF-A", "Total Revenues", "Statement of operations (management accounts)",
             {"2024": "$8.00M", "2025": "$10.00M", "2026 YTD (Jul)": "$4.00M"},
             {"2024": "actual", "2025": "actual", "2026 YTD (Jul)": "ytd"}, measures=["fin.revenue"]),
        fact("HPF-B", "Total Revenue", "Financial summary (company workbook)",
             {"2024": "$7.50M", "2025": "$9.00M", "2030P": "$20.00M"},
             {"2024": "actual", "2025": "actual", "2030P": "plan"}, measures=["fin.revenue"]),
        fact("HPF-C", "Other Revenue", "Statement of operations (management accounts)",
             {"2024": "$0.10M", "2025": "$0.20M"},
             {"2024": "actual", "2025": "actual"}, measures=["fin.revenue"]),
        fact("RVQ-1", "Hauling", "Revenue by segment (company workbook)",
             {"2025": "$6.00M"}, {"2025": "actual"}, module="RVQ", measures=["fin.revenue_mix"]),
        fact("RVQ-2", "Products", "Revenue by segment (company workbook)",
             {"2025": "$4.00M"}, {"2025": "actual"}, module="RVQ", measures=["fin.revenue_mix"]),
        fact("CST-1", "Gross revenue retention (GRR)", "Commercial summary",
             {"2024": "80.0%", "2025": "90.0%"}, {"2024": "actual", "2025": "actual"}, module="CST"),
        fact("CST-2", "Net revenue retention (NRR)", "Commercial summary",
             {"2024": "110.0%", "2025": "120.0%"}, {"2024": "actual", "2025": "actual"}, module="CST"),
        fact("CST-3", "New logos", "Commercial summary",
             {"2024": "10", "2025": "12"}, {"2024": "actual", "2025": "actual"}, module="CST"),
        fact("CST-4", "Churned logos", "Commercial summary",
             {"2024": "4", "2025": "5"}, {"2024": "actual", "2025": "actual"}, module="CST"),
        fact("CST-5", "Gross revenue retention (GRR)", "Residential summary",
             {"2022": "70.0%", "2023": "71.0%", "2024": "72.0%", "2025": "73.0%"},
             {"2022": "actual", "2023": "actual", "2024": "actual", "2025": "actual"}, module="CST"),
        fact("CST-6", "Net revenue retention (NRR)", "Residential summary",
             {"2022": "74.0%", "2023": "75.0%", "2024": "76.0%", "2025": "77.0%"},
             {"2022": "actual", "2023": "actual", "2024": "actual", "2025": "actual"}, module="CST"),
    ]
    book.by_id = {item["id"]: item for item in book.facts}
    book.vdr_points = _vdr_revenue_points(book.facts)
    book.vdr_plan = _vdr_plan_points(book.facts)
    labels = [point[0] for point in book.vdr_points if point[3] == "actual"]
    assert labels == ["2024", "2025"]
    assert book.vdr_points[1][1] == 10.0
    assert book.vdr_plan[-1][0] == "2030P"
    math = derive(_ctx({}), book)
    assert math.yoy[-1]["earlier"] == "2024"
    assert math.yoy[-1]["later"] == "2025"
    assert math.yoy[-1]["pct"] == 25.0
    assert math.bridge["actual"] == 9.0
    assert math.bridge["actual_id"] == "HPF-B"
    assert math.bridge["plan"] == 20.0
    assert math.bridge["plan_id"] == "HPF-B"
    assert math.bridge["plan_label"] == "2030P"
    assert math.mix_calculated is True
    assert {item["name"] for item in math.mix} == {"Hauling", "Products"}
    assert math.retention["grr"] == 90.0
    assert math.retention["nrr"] == 120.0
    assert math.retention["population"] == "Commercial summary"
    assert math.retention_curve["categories"] == ["2024", "2025"]


def test_dollar_revenue_is_shown_in_millions():
    ctx = _ctx({
        "historical_performance": {
            "spec": {
                "pl_lines": [{
                    "line_item": "Total revenue",
                    "fy2024_value": 8_500_000,
                    "fy2025_value": 13_410_000,
                }]
            }
        }
    })
    plan = plan_cdd_deck(ctx)
    titles = " ".join(page.title for page in plan.pages)
    assert "$13.4M" in titles
    assert "$8.5M" in titles
    assert "13410000" not in titles


def test_preview_keeps_chart_kpis_and_columns():
    from pathlib import Path
    import tempfile

    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Inches

    from agetic_cdd_api.report_preview import preview_pptx

    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    cover = prs.slides.add_slide(blank)
    cover.shapes.add_textbox(Inches(0.45), Inches(0.45), Inches(8), Inches(0.4)).text_frame.text = (
        "Commercial Due Diligence: Northwind"
    )
    slide = prs.slides.add_slide(blank)

    def box(text, left, top, width, height):
        shape = slide.shapes.add_textbox(left, top, width, height)
        shape.text_frame.text = text
        return shape

    box("BUSINESS INTRODUCTION", Inches(0.45), Inches(0.12), Inches(8), Inches(0.28))
    box("Revenue was $10.0M in 2025", Inches(0.45), Inches(0.45), Inches(10), Inches(0.45))
    box("Actuals from the accounts extract.", Inches(0.45), Inches(1.05), Inches(10), Inches(0.28))
    for value, label, left in (("$8.0M", "Revenue 2024", 0.45), ("$10.0M", "Revenue 2025", 2.8)):
        tile = slide.shapes.add_textbox(Inches(left), Inches(1.7), Inches(2.1), Inches(0.9))
        tile.text_frame.text = value
        tile.text_frame.add_paragraph().text = label
    box("Key facts", Inches(0.45), Inches(2.9), Inches(5.2), Inches(0.3))
    body = slide.shapes.add_textbox(Inches(0.45), Inches(3.3), Inches(5.2), Inches(2.2))
    body.text_frame.text = "EBITDA returned to profit in 2025."
    data = CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Revenue", (8.0, 10.0))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.2), Inches(2.9), Inches(5.5), Inches(2.8), data
    )

    work = prs.slides.add_slide(blank)
    def workbox(text, left, top, width, height):
        shape = work.shapes.add_textbox(left, top, width, height)
        shape.text_frame.text = text
    workbox("ENGAGEMENT OVERVIEW", Inches(0.45), Inches(0.12), Inches(8), Inches(0.28))
    workbox("The mandate for this engagement", Inches(0.45), Inches(0.45), Inches(10), Inches(0.45))
    workbox("Business & Financials", Inches(0.45), Inches(2.5), Inches(2.2), Inches(0.35))
    questions = work.shapes.add_textbox(Inches(0.45), Inches(3.0), Inches(2.2), Inches(2.0))
    questions.text_frame.text = "What has the business delivered?"
    questions.text_frame.add_paragraph().text = "How do margins trend?"
    workbox("Market & Geography", Inches(3.0), Inches(2.5), Inches(2.2), Inches(0.35))
    market = work.shapes.add_textbox(Inches(3.0), Inches(3.0), Inches(2.2), Inches(2.0))
    market.text_frame.text = "How large is the market?"

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "preview.pptx"
        prs.save(str(path))
        payload = preview_pptx(path)

    snapshot = payload["slides"][1]
    assert [item["value"] for item in snapshot["kpis"]] == ["$8.0M", "$10.0M"]
    assert snapshot["chart"]["categories"] == ["2024", "2025"]
    assert snapshot["chart"]["series"][0]["values"] == [8.0, 10.0]
    assert snapshot["columns"][0]["heading"] == "Key facts"
    assert "$8.0M" not in snapshot["paragraphs"]

    mandate = payload["slides"][2]
    assert [column["heading"] for column in mandate["columns"]] == [
        "Business & Financials",
        "Market & Geography",
    ]

