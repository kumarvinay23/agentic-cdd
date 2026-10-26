"""In-browser report preview helpers (xlsx sheets + docx/pdf document blocks)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def _iter_block_items(doc: DocumentObject):
    """Yield paragraphs and tables in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _has_page_break(paragraph: Paragraph) -> bool:
    for run in paragraph.runs:
        for br in run._element.findall(qn("w:br")):
            if br.get(qn("w:type")) == "page":
                return True
    if paragraph._element.findall(".//" + qn("w:lastRenderedPageBreak")):
        return True
    return False


def _style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name if paragraph.style is not None else "Normal"
    except Exception:
        return "Normal"


def _para_blocks(paragraph: Paragraph) -> list[dict[str, Any]]:
    """Convert a paragraph into preview blocks (inline page breaks + text)."""
    text = (paragraph.text or "").strip()
    blocks: list[dict[str, Any]] = []
    if _has_page_break(paragraph):
        blocks.append({"type": "page_break"})

    if not text:
        return blocks

    style = _style_name(paragraph)
    if style.startswith("Heading"):
        try:
            level = int(style.split()[-1])
        except ValueError:
            level = 1
        blocks.append({"type": "heading", "level": max(1, min(level, 3)), "text": text})
        return blocks

    if style.startswith("List"):
        blocks.append({"type": "bullet", "text": text})
        return blocks

    if text.lower().startswith("insight:"):
        body = text[8:].strip()
        blocks.append({"type": "insight", "text": body or text})
        return blocks

    blocks.append({"type": "paragraph", "text": text})
    return blocks


def _table_kind(rows: list[list[str]]) -> str:
    """Classify cover / divider / data tables for richer preview chrome."""
    if not rows:
        return "table"
    flat = " ".join(" ".join(r) for r in rows).strip()
    flat_l = flat.lower()
    # Single-cell hero / divider
    if len(rows) == 1 and len(rows[0]) == 1:
        cell = rows[0][0]
        if "investment due diligence" in cell.lower() and "strategy report" in cell.lower():
            return "cover_hero"
        if cell.strip()[:2].isdigit() and any(
            x in cell for x in ("Deal Framing", "Company & Management", "Legal, IP")
        ):
            return "divider"
        if "strictly private" in cell.lower() and len(cell) < 600:
            return "callout"
    if "document reference" in flat_l or (
        "prepared by" in flat_l and "sector" in flat_l and len(rows) <= 4
    ):
        return "meta"
    return "table"


def _table_block(table: Table) -> dict[str, Any] | None:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            # Collapse whitespace; avoid runaway cell text
            t = " ".join((cell.text or "").split())
            if len(t) > 800:
                t = t[:799] + "…"
            cells.append(t)
        # Skip fully empty rows
        if any(c.strip() for c in cells):
            rows.append(cells)
    if not rows:
        return None
    # Cap preview size
    if len(rows) > 60:
        rows = rows[:60]
    kind = _table_kind(rows)
    return {"type": kind, "rows": rows}


def preview_docx(path: Path, *, max_blocks: int = 400) -> dict[str, Any]:
    """Convert a Strategy Report .docx into ordered preview blocks."""
    doc = Document(str(path))
    blocks: list[dict[str, Any]] = []
    for item in _iter_block_items(doc):
        if len(blocks) >= max_blocks:
            blocks.append({"type": "paragraph", "text": "…preview truncated…", "muted": True})
            break
        if isinstance(item, Paragraph):
            blocks.extend(_para_blocks(item))
        elif isinstance(item, Table):
            block = _table_block(item)
            if block:
                blocks.append(block)

    # Drop leading empty page breaks
    while blocks and blocks[0].get("type") == "page_break":
        blocks.pop(0)

    return {
        "format": "docx",
        "blocks": blocks,
        "block_count": len(blocks),
    }


def preview_xlsx(path: Path, sheet: str | None = None, *, max_rows: int = 120) -> dict[str, Any]:
    """Return workbook sheet data for Excel-style preview."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        active = sheet if sheet in sheet_names else sheet_names[0]
        ws = wb[active]
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append([str(cell) if cell is not None else "" for cell in row])
        return {
            "format": "xlsx",
            "sheets": sheet_names,
            "active_sheet": active,
            "rows": rows,
        }
    finally:
        wb.close()


def preview_pdf(path: Path, *, max_pages: int = 40) -> dict[str, Any]:
    """Convert a PDF report into ordered preview blocks (text extraction)."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[dict[str, Any]] = []
    page_count = len(reader.pages)
    for i, page in enumerate(reader.pages[:max_pages]):
        if i > 0:
            blocks.append({"type": "page_break"})
        text = page.extract_text() or ""
        # Heuristic split into paragraphs / headings
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        buf: list[str] = []
        for ln in lines:
            upperish = ln.isupper() and len(ln) > 8
            looks_heading = (
                ln.startswith(("1.", "2.", "3.", "01", "02", "03", "Table of Contents", "IC Memo", "Insight:"))
                or upperish
            )
            if looks_heading and buf:
                blocks.append({"type": "paragraph", "text": " ".join(buf)})
                buf = []
            if ln.lower().startswith("insight:"):
                blocks.append({"type": "insight", "text": ln[8:].strip() or ln})
            elif ln.startswith(("1.0 ", "2.0 ", "3.0 ", "1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "2.3", "3.1", "3.2", "3.3")):
                level = 1 if ln[1:3] == ".0" else 2
                blocks.append({"type": "heading", "level": level, "text": ln})
            elif ln in ("Risk & Opportunity", "Valuation", "Executive Synthesis") or ln.startswith("IC Memo"):
                blocks.append({"type": "heading", "level": 1, "text": ln})
            elif ln.startswith(("STRICTLY PRIVATE", "INVESTMENT DUE DILIGENCE")):
                blocks.append({"type": "callout" if "STRICTLY" in ln else "cover_hero", "text": ln, "rows": [[ln]]})
            else:
                buf.append(ln)
                if len(" ".join(buf)) > 420:
                    blocks.append({"type": "paragraph", "text": " ".join(buf)})
                    buf = []
        if buf:
            blocks.append({"type": "paragraph", "text": " ".join(buf)})
    if page_count > max_pages:
        blocks.append({
            "type": "paragraph",
            "text": f"…preview truncated after {max_pages} of {page_count} pages…",
            "muted": True,
        })
    return {
        "format": "pdf",
        "blocks": blocks,
        "block_count": len(blocks),
        "page_count": page_count,
    }


_V4_SECTIONS = (
    "ENGAGEMENT OVERVIEW",
    "BUSINESS INTRODUCTION",
    "INVESTMENT HIGHLIGHTS",
    "MARKET DYNAMICS",
    "COMPETITIVE LANDSCAPE",
    "GROWTH OPPORTUNITIES",
    "APPENDIX A",
    "APPENDIX B",
)


def _is_v4_section(text: str) -> bool:
    up = text.strip().upper()
    return up.startswith(_V4_SECTIONS)


_KPI_VALUE = re.compile(r"^[+-]?(?:\$[\d,.]+[MBK]?|\d+(?:\.\d+)?%)$")


def _chart_payload(shape) -> dict[str, Any] | None:
    try:
        chart = shape.chart
        categories = [str(cat) for cat in chart.plots[0].categories]
        series = []
        for item in chart.series:
            values = []
            for raw in item.values:
                try:
                    values.append(round(float(raw), 2) if raw is not None else 0.0)
                except (TypeError, ValueError):
                    values.append(0.0)
            series.append({"name": str(item.name or "Series"), "values": values})
        if not categories or not series:
            return None
        return {"categories": categories[:12], "series": series[:4]}
    except Exception:
        return None


def _layout_from_shapes(bodies: list[dict[str, Any]]) -> tuple[list[dict[str, str]], list[dict[str, Any]], set[str]]:
    """KPI tiles and side-by-side columns. Returns the lines those layouts already show."""
    used: set[str] = set()
    kpis: list[dict[str, str]] = []
    for body in bodies:
        lines = body["lines"]
        if len(lines) != 2 or int(body["width"]) > 3_600_000:
            continue
        if not _KPI_VALUE.match(lines[0].replace(" ", "")) and not _KPI_VALUE.match(lines[0]):
            continue
        if len(lines[1]) > 48:
            continue
        kpis.append({"value": lines[0], "label": lines[1]})
        used.update(lines)
    kpis = kpis[:6]

    columns: list[dict[str, Any]] = []
    groups: dict[int, list[dict[str, Any]]] = {}
    for body in bodies:
        if any(line in used for line in body["lines"]):
            continue
        if int(body["top"]) < 1_500_000:
            continue
        if str(body["lines"][0]).startswith("Source:"):
            used.add(body["lines"][0])
            continue
        bucket = int(body["left"]) // 500_000
        groups.setdefault(bucket, []).append(body)
    if 1 <= len(groups) <= 6:
        for bucket in sorted(groups):
            parts = sorted(groups[bucket], key=lambda item: item["top"])
            heading = ""
            items: list[str] = []
            for part in parts:
                line = part["lines"][0]
                if not heading and len(part["lines"]) == 1 and len(line) <= 42:
                    heading = line
                    continue
                if heading:
                    items.extend(part["lines"])
            if heading and items:
                columns.append({
                    "heading": heading,
                    "items": items[:6],
                    "left": int(parts[0]["left"]),
                })
                used.add(heading)
                used.update(items)
        if not columns:
            used = {line for kpi in kpis for line in (kpi["value"], kpi["label"])}
            for body in bodies:
                if str(body["lines"][0]).startswith("Source:"):
                    used.add(body["lines"][0])
    return kpis, columns[:6], used


def preview_pptx(path: Path, *, max_slides: int = 80) -> dict[str, Any]:
    """Structured slide preview for CDD and Market Intel decks.

    Returns both:
    - ``slides``: per-slide payloads for the deck-frame UI
    - ``blocks``: flattened fallback for generic document preview
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    slide_count = len(prs.slides)
    slides_out: list[dict[str, Any]] = []
    blocks: list[dict[str, Any]] = []

    chrome_prefixes = (
        "MARKET INTEL DECK ·",
        "CDD DECK ·",
        "STRICTLY PRIVATE",
        "DILIGENCEIQ ·",
        "Workflow:",
        "Source:",
    )

    def _is_chrome(text: str) -> bool:
        t = text.strip()
        if not t:
            return True
        if re.match(r"^\d{1,2}\s*/\s*\d{1,2}$", t):
            return True
        upper = t.upper()
        if upper.startswith("DILIGENCEIQ"):
            return True
        if "© DILIGENCEIQ" in upper or "© AGENTIC CDD" in upper:
            return True
        if t.startswith(chrome_prefixes):
            return True
        if "Investment Due Diligence · ©" in t:
            return True
        if t.startswith("Section ") and "Investment Due Diligence" in t:
            return True
        return False

    def _shape_top(shape) -> int:
        try:
            return int(shape.top or 0)
        except Exception:
            return 0

    for i, slide in enumerate(prs.slides):
        if i >= max_slides:
            break

        texts_pos: list[tuple[int, int, str]] = []
        tables: list[list[list[str]]] = []
        bodies: list[dict[str, Any]] = []
        chart: dict[str, Any] | None = None
        has_full_navy = False

        for shape in slide.shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    continue
            except Exception:
                pass

            try:
                if (
                    shape.width and shape.height
                    and shape.width >= int(prs.slide_width * 0.9)
                    and shape.height >= int(prs.slide_height * 0.9)
                    and hasattr(shape, "fill")
                ):
                    fill = shape.fill
                    if fill.type is not None:
                        try:
                            rgb = fill.fore_color.rgb
                            if rgb and str(rgb).upper() in {"1F3864", "0F2A4A", "002B5B"}:
                                has_full_navy = True
                        except Exception:
                            pass
            except Exception:
                pass

            if getattr(shape, "has_chart", False) and chart is None:
                chart = _chart_payload(shape)
                if chart is not None:
                    try:
                        chart["left"] = int(shape.left or 0)
                    except Exception:
                        chart["left"] = 0
                continue

            if getattr(shape, "has_table", False):
                tbl = shape.table
                rows: list[list[str]] = []
                for r in tbl.rows:
                    rows.append([(c.text or "").strip() for c in r.cells])
                if rows:
                    tables.append(rows)
                continue

            if not getattr(shape, "has_text_frame", False):
                continue
            top = _shape_top(shape)
            try:
                left = int(shape.left or 0)
            except Exception:
                left = 0
            lines = []
            for para in shape.text_frame.paragraphs:
                t = (para.text or "").strip()
                if t:
                    texts_pos.append((top, left, t))
                    lines.append(t)
            if lines and not any(_is_chrome(line) for line in lines):
                try:
                    width = int(shape.width or 0)
                except Exception:
                    width = 0
                bodies.append({"top": top, "left": left, "width": width, "lines": lines})

        texts_pos.sort(key=lambda x: (x[0], x[1]))
        ordered = [t for _, _, t in texts_pos]
        footer_top = int(prs.slide_height * 0.9)
        content = [
            t for top, _, t in texts_pos
            if not _is_chrome(t) and not (re.fullmatch(r"\d{2}", t) and top >= footer_top)
        ]

        kind = "content"
        joined = " | ".join(content[:8])
        joined_u = joined.upper()
        if has_full_navy and any(
            t in {"01", "02", "03", "04", "05", "06", "07"} or t.startswith("SECTION")
            for t in content[:5]
        ):
            kind = "divider"
        elif i == 0 or (has_full_navy and (
            "INVESTMENT DUE DILIGENCE" in joined_u or "COMMERCIAL CDD" in joined_u
        )):
            kind = "cover"
        elif any(t == "Agenda" for t in content[:4]):
            kind = "agenda"
        elif any("agents' read" in t for t in content[:3]):
            kind = "workflow"
        elif any(t.startswith("Key Findings") for t in content[:4]) and any(
            x in joined_u for x in ("IN-SCOPE", "VERDICT", "OUT-OF-SCOPE")
        ):
            kind = "key_findings"
        elif any(
            "CLOSING SUMMARY" in t.upper()
            or t.startswith("Market Intel Deck — Summary")
            or t.startswith("Market Intel Deck - Summary")
            or t.startswith("Recommendation & monitoring")
            or t.startswith("Executive summary")
            for t in content[:5]
        ):
            kind = "summary"
        elif any(t.startswith("Appendix") for t in content[:3]):
            kind = "appendix"

        eyebrow = ""
        title = f"Slide {i + 1}"
        subtitle = ""
        chapter_titles = (
            "Business Introduction & Market Primer",
            "Investment Highlights",
            "Market Dynamics & Opportunity",
            "Competitive Landscape",
            "Growth Opportunities",
            "Appendix A - Market Sizing Methodology",
            "Appendix B - Evidence & Audit Trail",
        )
        v4_section = next((t for t in content if _is_v4_section(t)), "")
        title_locked = False
        if any("Commercial Due Diligence" in t for t in content[:6]):
            kind = "cover"
            title = next(t for t in content if "Commercial Due Diligence" in t)
            subtitle = next(
                (t for t in content if t not in {title, "DiligenceIQ", "Commercial due diligence"}),
                "",
            )
            title_locked = True
        elif any(t in {"01", "02", "03", "04", "05", "A", "B"} for t in content[:3]) and any(
            t in chapter_titles for t in content[:6]
        ):
            kind = "divider"
            title = next(t for t in content if t in chapter_titles)
            subtitle = next(
                (
                    t for t in content
                    if t != title and t not in {"01", "02", "03", "04", "05", "A", "B"} and not _is_v4_section(t)
                ),
                "",
            )
            for t in content[:3]:
                if t in {"01", "02", "03", "04", "05", "A", "B"}:
                    eyebrow = t
                    break
            title_locked = True
        elif v4_section:
            eyebrow = v4_section
            rest = [t for t in content if t != v4_section and not _is_v4_section(t)]
            if rest:
                title = rest[0]
            if len(rest) > 1 and not rest[1].startswith("Source:"):
                subtitle = rest[1]
            title_locked = True
        if not title_locked:
            for t in content:
                up = t.upper()
                if up.startswith(("SECTION 0", "DOCUMENT OVERVIEW", "WORKFLOW ANALYSIS", "CLOSING SUMMARY")):
                    eyebrow = t
                    continue
                if t in {"01", "02", "03", "04", "05", "06", "07"} and kind == "divider":
                    continue
                if (
                    t.startswith((
                        "1.", "2.", "3.", "4.", "Agenda", "Appendix", "Key Findings",
                        "Market Intel", "INVESTMENT", "DEAL POSITION", "DILIGENCE WORKFLOW",
                        "Engagement overview", "Basis of preparation", "Scope & research",
                        "Business introduction", "Market primer", "Market sizing",
                        "Growth trajectory", "Competitor universe", "Positioning",
                        "Strategic differentiators", "SWOT", "Customer segmentation",
                        "Stickiness", "Operations & supply", "Risk register",
                        "Historical performance", "Capital structure", "Executive summary",
                        "Recommendation", "CDD Deck",
                    ))
                    or "MARKET DEFINITION" in t
                    or t in {"Market Analysis", "Competitive Landscape", "Summary", "Key Findings",
                             "Engagement & Scope", "Business & Market", "Customer & Commercial",
                             "Operations & Risk", "Financials & Valuation", "IC Decision"}
                    or any(
                        k in t
                        for k in (
                            "TAM", "SWOT", "Competitive", "CAGR", "Pricing", "Demand",
                            "battlecard", "moat", "taxonomy", "Historical", "Sustainability",
                            "positioning", "capability", "Differentiation", "Direct &",
                            "valuation", "go/no-go", "monitoring",
                        )
                    )
                ):
                    title = t
                    break

        if not title_locked:
            seen_title = False
            for t in content:
                if t == title:
                    seen_title = True
                    continue
                if not seen_title:
                    continue
                if t.upper().startswith(("SECTION", "DOCUMENT", "WORKFLOW", "CLOSING", "DEAL POSITION", "DILIGENCE")):
                    continue
                if t.startswith(("•", "—", "-", "Insight")):
                    break
                if len(t) < 160:
                    subtitle = t
                break

        insight = ""
        bullets: list[str] = []
        cards: list[dict[str, str]] = []
        body_paras: list[str] = []
        meta_rows: list[list[str]] = []

        for t in content:
            if t in {title, subtitle, eyebrow} or t in {"01", "02", "03", "04"}:
                continue
            low = t.lower()
            if low.startswith("insight snapshot:"):
                insight = t.split(":", 1)[-1].strip()
                continue
            if low.startswith("insight:"):
                insight = t.split(":", 1)[-1].strip()
                continue
            if t.startswith(("•", "—", "- ")):
                bullets.append(t.lstrip("•—- ").strip())
                continue
            body_paras.append(t)

        labels = {
            "IN-SCOPE DRIVERS", "OUT-OF-SCOPE", "ENGAGEMENT RATIONALE",
            "VERDICT", "RELATIVE STRENGTH", "PRIMARY COMP. THREAT",
            "STRENGTHS", "WEAKNESSES", "OPPORTUNITIES", "THREATS",
        }
        for idx_t, t in enumerate(content):
            if t.upper() in labels:
                body = ""
                for nxt in content[idx_t + 1: idx_t + 4]:
                    if nxt.upper() in labels or nxt == title:
                        break
                    if nxt.startswith(("•", "—")):
                        body = nxt.lstrip("•—- ").strip()
                        break
                    if nxt not in {eyebrow, subtitle}:
                        body = nxt
                        break
                cards.append({"label": t.upper(), "body": body or "—"})

        if kind == "cover":
            cover_labels = {"DOCUMENT", "TARGET", "SECTOR", "REPORT DATE", "PREPARED BY"}
            label_hits = [
                (top, left, t.upper())
                for top, left, t in texts_pos
                if t.upper() in cover_labels and not _is_chrome(t)
            ]
            value_hits = [
                (top, left, t)
                for top, left, t in texts_pos
                if t.upper() not in cover_labels and not _is_chrome(t)
            ]
            for ltop, lleft, label in label_hits:
                best = None
                best_score = None
                for vtop, vleft, val in value_hits:
                    if vtop < ltop - 20000:
                        continue
                    score = abs(vleft - lleft) + max(0, vtop - ltop) // 20
                    if best_score is None or score < best_score:
                        best_score = score
                        best = val
                if best:
                    meta_rows.append([label, best])
                    value_hits = [v for v in value_hits if v[2] != best]

        agenda_items: list[dict[str, str]] = []
        if kind == "agenda":
            for n in ("01", "02", "03", "04"):
                if n not in content:
                    continue
                ni = content.index(n)
                agenda_items.append({
                    "num": n,
                    "title": content[ni + 1] if ni + 1 < len(content) else "",
                    "blurb": content[ni + 2] if ni + 2 < len(content) else "",
                })

        primary_table = tables[0] if tables else []
        paragraphs = [p[:420] for p in body_paras if p not in {title, subtitle}][:10]
        kpis, columns, laid_out = _layout_from_shapes(bodies)
        if kind in {"cover", "divider"} or any(
            re.match(r"^\d{2}\b", column["heading"]) for column in columns
        ):
            columns = []
            laid_out = {line for kpi in kpis for line in (kpi["value"], kpi["label"])}
            for body in bodies:
                if str(body["lines"][0]).startswith("Source:"):
                    laid_out.add(body["lines"][0])
        if laid_out:
            paragraphs = [p for p in paragraphs if p not in laid_out]
            bullets = [b for b in bullets if b not in laid_out]
        if kind == "cover":
            cover_skip = {
                "DOCUMENT", "TARGET", "SECTOR", "REPORT DATE", "PREPARED BY",
                *(r[0].upper() for r in meta_rows),
                *(r[1] for r in meta_rows),
                *(r[1].upper() for r in meta_rows),
            }
            paragraphs = [
                p for p in paragraphs
                if p not in cover_skip and p.upper() not in cover_skip
            ][:4]

        slide_payload: dict[str, Any] = {
            "index": i + 1,
            "kind": kind,
            "eyebrow": eyebrow,
            "title": title[:200],
            "subtitle": subtitle[:240],
            "insight": insight[:400],
            "bullets": bullets[:12],
            "paragraphs": paragraphs,
            "cards": cards[:8],
            "agenda_items": agenda_items,
            "meta_rows": meta_rows,
            "table": primary_table[:12],
            "kpis": kpis,
            "chart": chart,
            "columns": columns,
            "page": f"{i + 1:02d} / {slide_count:02d}",
        }
        if kind == "divider":
            for t in content[:4]:
                if t in {"01", "02", "03", "04", "05", "06", "07"}:
                    slide_payload["number"] = t
                    break
            if not slide_payload.get("number") and eyebrow.upper().startswith("SECTION"):
                slide_payload["number"] = eyebrow.replace("SECTION", "").strip()

        slides_out.append(slide_payload)

        if i > 0:
            blocks.append({"type": "page_break"})
        if kind in {"cover", "divider"}:
            blocks.append({
                "type": "cover_hero" if kind == "cover" else "divider",
                "text": f"{slide_payload.get('number', '')} {title}".strip(),
                "rows": [[f"{eyebrow}\n{title}\n{subtitle}".strip()]],
            })
        else:
            if eyebrow:
                blocks.append({"type": "heading", "level": 3, "text": eyebrow})
            blocks.append({"type": "heading", "level": 1 if i == 0 else 2, "text": title})
            if subtitle:
                blocks.append({"type": "paragraph", "text": subtitle, "muted": True})
            if insight:
                blocks.append({"type": "insight", "text": insight})
            for b in bullets[:8]:
                blocks.append({"type": "bullet", "text": b})
            for c in cards[:6]:
                blocks.append({"type": "paragraph", "text": f"{c['label']}: {c['body']}"})
            if primary_table:
                blocks.append({"type": "table", "rows": primary_table})

    if slide_count > max_slides:
        blocks.append({
            "type": "paragraph",
            "text": f"…preview truncated after {max_slides} of {slide_count} slides…",
            "muted": True,
        })

    return {
        "format": "pptx",
        "slides": slides_out,
        "blocks": blocks,
        "block_count": len(blocks),
        "page_count": slide_count,
        "slide_count": slide_count,
    }
