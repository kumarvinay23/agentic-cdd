"""Operations Dashboard builder — 5-sheet Excel workbook from Deep Dive stores.

Sheets:
  1. Executive Dashboard — KPI scorecard + RAG + figure validation + workflow findings
  2. Customer Analysis — segments, concentration, stickiness, satisfaction
  3. Operational & Risk — supplier landscape, ops risk register, cost breakdown
  4. Financial Analysis — multi-year metric grid, P&L, capital structure
  5. Source Data — provenance fact rows (metric × year × value × agent)

KPI selection, RAG status, figure validation, and year columns are derived from
agent-output shapes (keys, units, overlapping values) — not hardcoded deal values.
"""

from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Iterator
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agetic_cdd_api.report_builder_base import BuildContext, ReportBuilder, _evt, sanitize_report_prose
from agetic_cdd_api.report_store import report_artifact_dir
from agetic_cdd_api.routers_reports import register_builder
from agetic_cdd_api.services_deals import deals_root

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
_SUBHEADER_FONT = Font(bold=True, size=11)
_WARN_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
_ONTRACK_FILL = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
_ATRISK_FILL = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

_FY_KEY_RE = re.compile(r"^fy(\d{4})_value$", re.I)
_FY_IN_NAME_RE = re.compile(r"(?:^|_)(?:fy)?(\d{4})(?:_|$)", re.I)
_SKIP_SPEC_KEYS = frozenset({
    "document", "dd_code", "sources", "empty", "slug", "track", "coverage",
    "extractor", "extract_source", "vdr_backed", "llm_refined", "metrics",
    "raw", "as_of", "scale", "unit", "label", "note", "notes",
})


def _rag_fill(status: str) -> PatternFill | None:
    s = (status or "").lower()
    if s in ("on-track", "on track", "green"):
        return _ONTRACK_FILL
    if s in ("at-risk", "at risk", "red", "below target", "critical", "high"):
        return _ATRISK_FILL
    if s in ("monitor", "warn", "amber", "medium"):
        return _WARN_FILL
    return None


def _safe_num(v: Any) -> Any:
    if v is None:
        return "-"
    if isinstance(v, (int, float)):
        return v
    try:
        return float(v)
    except (ValueError, TypeError):
        return v


def _approx_equal(a: float, b: float, *, rel: float = 0.001, abs_tol: float = 0.005) -> bool:
    """True when databook and Full Metric Grid agree within a tight band.

    QBO-vs-model gaps like 7.528 vs 7.558 (~40 bps) must flag as disagreement.
    """
    if a == b:
        return True
    scale = max(abs(a), abs(b), 1.0)
    return abs(a - b) <= max(abs_tol, rel * scale)


def _write_header_row(ws, row: int, headers: list[str]) -> None:
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _write_row(ws, row: int, values: list, bold: bool = False) -> None:
    for ci, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=ci, value=v)
        cell.border = _THIN_BORDER
        if bold:
            cell.font = Font(bold=True)


def _auto_width(ws, min_width: int = 10, max_width: int = 50) -> None:
    for col in ws.columns:
        mx = min_width
        for cell in col:
            if cell.value:
                mx = max(mx, min(len(str(cell.value)), max_width))
        ws.column_dimensions[get_column_letter(col[0].column)].width = mx + 2


def _get_agent(ctx: BuildContext, key: str) -> dict:
    return ctx.agent_outputs.get(key, {})


def _get_spec(ctx: BuildContext, key: str) -> dict:
    agent = _get_agent(ctx, key)
    return agent.get("spec", {}) if isinstance(agent, dict) else {}


def _humanize(key: str) -> str:
    """Turn snake/camel metric keys into readable labels with unit hints."""
    k = key.strip()
    suffix = ""
    if k.endswith("_pct"):
        k, suffix = k[:-4], " (%)"
    elif k.endswith("_inr_cr"):
        k, suffix = k[:-7], " (INR Cr)"
    elif k.endswith("_usd_m"):
        k, suffix = k[:-6], " (USD M)"
    elif k.endswith("_inr"):
        k, suffix = k[:-4], " (INR)"
    elif k.endswith("_days"):
        k, suffix = k[:-5], " (Days)"
    elif k.endswith("_ratio"):
        k, suffix = k[:-6], " (x)"
    k = _FY_IN_NAME_RE.sub("_", k)
    k = re.sub(r"_+", " ", k).strip()
    label = k.replace("_", " ").title()
    for acro in ("Ebitda", "Ebit", "Nps", "Grr", "Nrr", "Cac", "Ltv", "Wacc", "Asp", "Bom", "Capex", "Fcf"):
        label = label.replace(acro, acro.upper())
    return label + suffix


def _canonical_family(name: str) -> str:
    """Collapse metric names into comparable families for conflict detection."""
    s = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    s = _FY_IN_NAME_RE.sub("_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    for suf in ("_pct", "_inr_cr", "_usd_m", "_inr", "_days", "_value", "_share"):
        if s.endswith(suf):
            s = s[: -len(suf)]
    aliases = {
        "net_revenue": "revenue",
        "gross_margin": "gross_margin",
        "adj_ebitda": "ebitda",
        "adjusted_ebitda": "ebitda",
        "liquidity": "cash",
        "cash_equivalents": "cash",
        "cash_and_equivalents": "cash",
        "avg_selling_price": "asp",
        "selling_price": "asp",
        "grr": "grr",
        "nrr": "nrr",
        "gross_revenue_retention": "grr",
        "net_revenue_retention": "nrr",
        "customer_acquisition_cost": "cac",
        "cac_inr": "cac",
        "ltv_inr": "ltv",
        "estimated_3_year_ltv": "ltv",
        "ltv_cac_ratio": "ltv_cac",
        "wacc_pct": "wacc",
    }
    return aliases.get(s, s)


@dataclass
class MetricFact:
    family: str
    label: str
    value: float
    year: int | None = None
    unit: str = ""
    agent_key: str = ""
    agent_name: str = ""
    sources: list[str] = field(default_factory=list)


@dataclass
class KpiRow:
    name: str
    latest: Any
    yoy: Any
    rag: str
    period: str = ""


def _period_tag(year: int, period_type: str | None) -> str:
    """FY2025A / FY2030P — never leave plan years looking like current run-rate."""
    pt = (period_type or "").strip().lower()
    if pt in {"actual", "ltm", "last twelve months"}:
        suffix = "A"
    elif pt in {"forecast", "plan", "budget", "projected", "estimate"}:
        suffix = "P"
    elif "e" in pt:
        suffix = "E"
    else:
        suffix = ""
    return f"FY{year}{suffix}" if year else ""


def _as_float(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", "").strip())
        except ValueError:
            return None
    return None


def _year_from_key(key: str) -> int | None:
    m = _FY_IN_NAME_RE.search(key)
    return int(m.group(1)) if m else None


def _unit_from_path(path: tuple[str, ...], explicit: str = "") -> str:
    if explicit:
        return explicit
    joined = "_".join(path).lower()
    if joined.endswith("_pct") or ("margin" in joined and "pct" in joined):
        return "%"
    if "inr_cr" in joined:
        return "INR Cr"
    if "usd_m" in joined:
        return "USD M"
    if joined.endswith("_inr") or "_inr" in joined:
        return "INR"
    if joined.endswith("_days"):
        return "Days"
    return ""


def harvest_metric_facts(ctx: BuildContext) -> list[MetricFact]:
    """Collect numeric facts across all agent specs for scorecard / validation / grid."""
    facts: list[MetricFact] = []

    # Prefer Databook *released* metrics first (canonical FY history — P0).
    try:
        from agetic_cdd_api.services_databook_consume import (
            load_current_release_metrics_for_slug,
            released_as_metric_fact_dicts,
        )

        release = load_current_release_metrics_for_slug(ctx.deal_slug)
        if release is not None:
            for d in released_as_metric_fact_dicts(release):
                status = str(d.get("databook_status") or "proven")
                label = str(d["label"])
                if status == "doubtful":
                    label = f"{label} (doubtful)"
                facts.append(
                    MetricFact(
                        family=str(d["family"]),
                        label=label,
                        value=float(d["value"]),
                        year=d.get("year"),
                        unit=str(d.get("unit") or ""),
                        agent_key="databook",
                        agent_name="Databook",
                        sources=list(d.get("sources") or []),
                    )
                )
    except Exception:
        pass

    for agent_key, agent in ctx.agent_outputs.items():
        if not isinstance(agent, dict):
            continue
        spec = agent.get("spec")
        if not isinstance(spec, dict):
            continue
        agent_name = agent.get("agentName") or agent_key.replace("_", " ").title()
        sources = list(agent.get("sources") or spec.get("sources") or [])

        for list_key in ("pl_lines", "segments", "price_points", "kpi_gaps", "bom_components"):
            rows = spec.get(list_key)
            if not isinstance(rows, list):
                continue
            for row in rows:
                if not isinstance(row, dict):
                    continue
                name = (
                    row.get("line_item")
                    or row.get("name")
                    or row.get("label")
                    or row.get("kpi")
                    or row.get("category")
                    or ""
                )
                if not name:
                    continue
                unit = str(row.get("unit") or "")
                year_vals: dict[int, float] = {}
                for rk, rv in row.items():
                    m = _FY_KEY_RE.match(str(rk))
                    if m:
                        num = _as_float(rv)
                        if num is not None:
                            year_vals[int(m.group(1))] = num
                if year_vals:
                    for yr, num in year_vals.items():
                        facts.append(MetricFact(
                            family=_canonical_family(str(name)),
                            label=str(name),
                            value=num,
                            year=yr,
                            unit=unit,
                            agent_key=agent_key,
                            agent_name=agent_name,
                            sources=sources,
                        ))
                for vk in ("value", "share_pct"):
                    if vk not in row:
                        continue
                    num = _as_float(row.get(vk))
                    if num is None:
                        continue
                    label = str(name)
                    if label.lower() == "asp":
                        label = "Avg Selling Price"
                    facts.append(MetricFact(
                        family=_canonical_family(str(name)),
                        label=label,
                        value=num,
                        year=None,
                        unit=unit or _unit_from_path((str(name), vk)),
                        agent_key=agent_key,
                        agent_name=agent_name,
                        sources=sources,
                    ))

        metric_dict_keys = (
            "retention_metrics", "cost_metrics", "cash_metrics", "buying_metrics",
            "leverage_ratios", "dcf_metrics", "performance_metrics", "revenue_mix",
        )
        for sk, sv in spec.items():
            if sk in _SKIP_SPEC_KEYS or not isinstance(sv, dict):
                continue
            is_metric_dict = sk.endswith("_metrics") or sk.endswith("_ratios") or sk in metric_dict_keys
            if not is_metric_dict:
                if not all(v is None or isinstance(v, (int, float)) for v in sv.values()):
                    continue
            for mk, mv in sv.items():
                num = _as_float(mv)
                if num is None:
                    continue
                facts.append(MetricFact(
                    family=_canonical_family(str(mk)),
                    label=_humanize(str(mk)),
                    value=num,
                    year=_year_from_key(str(mk)),
                    unit=_unit_from_path((str(mk),)),
                    agent_key=agent_key,
                    agent_name=agent_name,
                    sources=sources,
                ))

        for sk, sv in spec.items():
            if sk in _SKIP_SPEC_KEYS or isinstance(sv, (dict, list)):
                continue
            num = _as_float(sv)
            if num is None:
                continue
            facts.append(MetricFact(
                family=_canonical_family(str(sk)),
                label=_humanize(str(sk)),
                value=num,
                year=_year_from_key(str(sk)),
                unit=_unit_from_path((str(sk),)),
                agent_key=agent_key,
                agent_name=agent_name,
                sources=sources,
            ))

    return facts


def _format_display(value: float, unit: str = "", family: str = "") -> Any:
    if unit == "%" or family in ("gross_margin", "ebitda_margin", "grr", "nrr", "wacc"):
        return f"{value:g}%"
    if unit == "INR Cr":
        return f"{value:g} INR Cr"
    if unit == "INR" or family in ("cac", "ltv", "asp"):
        return f"INR {value:,.0f}"
    if family == "ltv_cac" or unit == "x":
        return f"{value:g}x"
    if abs(value) >= 1_000_000:
        return f"{value:,.0f}"
    return value


def _rag_for_metric(family: str, value: float, *, yoy: float | None = None) -> str:
    """Family-based RAG rules (thresholds are rule families, not deal-specific values)."""
    f = family.lower()
    if yoy is not None and ("growth" in f or f == "revenue"):
        return "On-Track" if yoy > 0 else "At-Risk"
    if "margin" in f or f in ("gross_margin", "ebitda_margin"):
        if value < 0:
            return "At-Risk"
        if value < 15:
            return "At-Risk"
        return "On-Track"
    if f in ("ebitda", "ebit", "net_income", "ufcf"):
        return "At-Risk" if value < 0 else "On-Track"
    if "churn" in f:
        return "At-Risk" if value > 20 else "Monitor"
    if f in ("nrr", "grr", "retention"):
        return "At-Risk" if value < 90 else "On-Track"
    if f == "ltv_cac" or "ltv_cac" in f:
        return "On-Track" if value >= 3 else "At-Risk"
    if f in ("cac", "wacc", "asp", "cash", "liquidity"):
        return "Monitor"
    if "burn" in f:
        return "At-Risk" if value > 50 else "Monitor"
    return "Monitor"


def derive_kpi_scorecard(facts: list[MetricFact], pl_lines: list[dict]) -> list[KpiRow]:
    """Build scorecard from P&L YoY logic + latest facts by family priority.

    Headline P&L KPIs prefer agent-confirmed (non-databook) year series when
    available, then tagged actuals, then near-term years — never the terminal
    plan year when earlier actuals exist.
    """
    from datetime import date

    kpis: list[KpiRow] = []
    seen: set[str] = set()

    rev_by_year: dict[int, float] = {}
    ebitda_by_year: dict[int, float] = {}
    gm_by_year: dict[int, float] = {}
    type_by_year: dict[int, str] = {}

    # Prefer Full Metric Grid / agent facts for headline growth (confirmed series).
    agent_rev_years: set[int] = set()
    for fact in facts:
        if fact.year is None or fact.agent_key == "databook":
            continue
        if fact.family == "revenue":
            rev_by_year[fact.year] = fact.value
            agent_rev_years.add(fact.year)
        elif fact.family == "ebitda":
            ebitda_by_year[fact.year] = fact.value
        elif fact.family == "gross_margin":
            gm_by_year[fact.year] = fact.value

    for line in pl_lines:
        if not isinstance(line, dict):
            continue
        item = str(line.get("line_item") or "")
        years: dict[int, float] = {}
        for k, v in line.items():
            m = _FY_KEY_RE.match(str(k))
            if m:
                num = _as_float(v)
                if num is not None:
                    yr = int(m.group(1))
                    years[yr] = num
                    pt = line.get(f"fy{yr}_period_type")
                    if isinstance(pt, str) and pt.strip():
                        type_by_year.setdefault(yr, pt.strip().lower())
        fam = _canonical_family(item)
        # Only fill gaps — do not overwrite agent-confirmed grid figures.
        # When agent revenue history exists, do not extend with plan-only databook years.
        if fam == "revenue":
            for yr, val in years.items():
                if agent_rev_years and yr not in agent_rev_years and yr > max(agent_rev_years):
                    continue
                rev_by_year.setdefault(yr, val)
        elif fam == "ebitda" and "%" not in item:
            for yr, val in years.items():
                ebitda_by_year.setdefault(yr, val)
        elif fam == "gross_margin" or "gross margin" in item.lower():
            for yr, val in years.items():
                gm_by_year.setdefault(yr, val)

    def _prefer_actual_years(by_year: dict[int, float], *, lock_years: set[int] | None = None) -> list[int]:
        """Prefer tagged actuals over plan/forecast — never terminal plan year as 'Latest'.

        Hard-cap at the latest closed calendar FY so Full Metric Grid plan years
        (FY2026–FY2030) cannot flip the headline YoY even when mis-tagged actual
        or when period_type tags were stripped by databook release overlays.
        """
        closed = date.today().year - 1
        forecastish = {"forecast", "plan", "budget", "projected", "estimate"}
        actualish = {"actual", "ltm", "last twelve months", "year to date", "ytd"}

        def _is_forecast(y: int) -> bool:
            return type_by_year.get(y) in forecastish

        eligible = [y for y in by_year if not _is_forecast(y)]
        # Prefer history through closed FY; only admit current calendar year if
        # explicitly tagged actual/YTD (never bare untagged future plan years).
        capped = [y for y in eligible if y <= closed]
        if len(capped) < 2:
            cur = date.today().year
            capped = [
                y for y in eligible
                if y <= closed or (y == cur and type_by_year.get(y) in actualish)
            ]
        pool = capped if len(capped) >= 2 else (eligible if eligible else list(by_year))

        actuals = sorted(y for y in pool if type_by_year.get(y) in actualish)
        if len(actuals) >= 2:
            return actuals
        if len(actuals) == 1:
            prior = sorted(y for y in pool if y < actuals[0])
            if prior:
                return [prior[-1], actuals[0]]
            return actuals

        if lock_years:
            near = sorted(y for y in pool if y in lock_years)
            if len(near) >= 2:
                return near
        near = sorted(pool)
        if len(near) >= 2:
            if near[-1] - near[0] > 4:
                cluster_end = near[0] + 3
                cluster = [y for y in near if y <= cluster_end]
                if len(cluster) >= 2:
                    return cluster
            return near
        return sorted(by_year)

    if len(rev_by_year) >= 2:
        prefer = _prefer_actual_years(rev_by_year, lock_years=agent_rev_years or None)
        if len(prefer) >= 2:
            y0, y1 = prefer[-2], prefer[-1]
        else:
            ys = sorted(rev_by_year)
            y0, y1 = ys[-2], ys[-1]
        if rev_by_year[y0] != 0:
            yoy = round((rev_by_year[y1] - rev_by_year[y0]) / abs(rev_by_year[y0]) * 100, 1)
            period = _period_tag(y1, type_by_year.get(y1))
            kpis.append(KpiRow(
                "Revenue Growth (YoY)",
                f"{yoy}% ({period})" if period else f"{yoy}%",
                f"{yoy}%",
                _rag_for_metric("revenue_growth", yoy, yoy=yoy),
                period=period,
            ))
            seen.add("revenue_growth")

    if gm_by_year and "gross_margin" not in seen:
        prefer = _prefer_actual_years(gm_by_year)
        yr = prefer[-1] if prefer else max(gm_by_year)
        val = gm_by_year[yr]
        period = _period_tag(yr, type_by_year.get(yr))
        kpis.append(KpiRow(
            "Gross Margin",
            f"{val:g}% ({period})" if period else f"{val:g}%",
            "-",
            _rag_for_metric("gross_margin", val),
            period=period,
        ))
        seen.add("gross_margin")

    if ebitda_by_year:
        prefer = _prefer_actual_years(ebitda_by_year)
        yr = prefer[-1] if prefer else max(ebitda_by_year)
        ebitda = ebitda_by_year[yr]
        period = _period_tag(yr, type_by_year.get(yr))
        kpis.append(KpiRow(
            "EBITDA",
            f"{ebitda:g} ({period})" if period else ebitda,
            "-",
            _rag_for_metric("ebitda", ebitda),
            period=period,
        ))
        seen.add("ebitda")
        if yr in rev_by_year and rev_by_year[yr] != 0:
            margin = round(ebitda / rev_by_year[yr] * 100, 1)
            kpis.append(KpiRow(
                "EBITDA Margin",
                f"{margin}% ({period})" if period else f"{margin}%",
                "-",
                _rag_for_metric("ebitda_margin", margin),
                period=period,
            ))
            seen.add("ebitda_margin")

    latest_by_family: dict[str, MetricFact] = {}
    for fact in facts:
        prev = latest_by_family.get(fact.family)
        if prev is None:
            latest_by_family[fact.family] = fact
            continue
        if fact.year and (prev.year is None or fact.year > prev.year):
            latest_by_family[fact.family] = fact

    # Order only — presence is data-driven from harvested facts
    priority = [
        "gross_margin", "asp", "cash", "liquidity", "ltv", "cac", "ltv_cac", "wacc", "nrr", "grr",
    ]
    label_overrides = {
        "asp": "Avg Selling Price",
        "cash": "Cash & Equivalents",
        "liquidity": "Cash & Equivalents",
        "ltv": "Estimated 3-Year LTV",
        "cac": "Customer Acquisition Cost (CAC)",
        "ltv_cac": "LTV / CAC Ratio",
        "wacc": "WACC",
        "nrr": "Net Revenue Retention",
        "grr": "Gross Revenue Retention",
        "gross_margin": "Gross Margin",
    }

    for fam in priority:
        display_fam = "cash" if fam == "liquidity" else fam
        if display_fam in seen or fam in seen:
            continue
        fact = latest_by_family.get(fam)
        if not fact:
            continue
        name = label_overrides.get(fam, fact.label)
        display = _format_display(fact.value, fact.unit, fam)
        if fam == "gross_margin":
            display = f"{fact.value:g}%"
        if fam == "wacc":
            display = f"{fact.value:g}%"
        if fam == "ltv_cac":
            display = f"{fact.value:g}x"
        kpis.append(KpiRow(name, display, "-", _rag_for_metric(fam, fact.value)))
        seen.add(display_fam)

    return kpis


# Families too generic to treat as figure conflicts (valuation bands, booleans, etc.)
_CONFLICT_SKIP_FAMILIES = frozenset({
    "high", "low", "mid", "base", "empty", "slug", "count", "true", "false",
    "order", "idx", "id", "score",
})


def detect_figure_conflicts(facts: list[MetricFact], *, rel_tol: float = 0.05) -> list[tuple[str, str, str]]:
    """Compare same-family facts across agents; emit WARN rows when values disagree.

    When Databook has promoted a (family, year), skip agent WARNs for that key —
    the promoted figure is authoritative.
    """
    by_key: dict[tuple[str, int | None], list[MetricFact]] = defaultdict(list)
    for f in facts:
        if f.family in _CONFLICT_SKIP_FAMILIES or len(f.family) < 3:
            continue
        by_key[(f.family, f.year)].append(f)

    warnings: list[tuple[str, str, str]] = []
    for (family, year), group in sorted(by_key.items(), key=lambda x: (x[0][0], x[0][1] or 0)):
        databook_facts = [f for f in group if f.agent_key == "databook"]
        if databook_facts:
            # Canonical value present — do not surface cross-agent WARN noise.
            continue
        by_agent: dict[str, MetricFact] = {}
        for f in group:
            by_agent.setdefault(f.agent_key, f)
        if len(by_agent) < 2:
            continue
        values = list(by_agent.values())
        base = values[0].value
        has_conflict = False
        for other in values[1:]:
            denom = max(abs(base), abs(other.value), 1e-9)
            if abs(base - other.value) / denom > rel_tol:
                has_conflict = True
                break
        if not has_conflict:
            continue
        scale_note = ""
        nums = [v.value for v in values]
        nonzero = [abs(n) for n in nums if n != 0]
        if nonzero and max(nonzero) / min(nonzero) >= 100:
            scale_note = f" Possible unit/scale mismatch (ratio ~{max(nonzero) / min(nonzero):,.0f}x)."
        year_bit = f" for FY{year}" if year else ""
        detail_parts = [
            f"{f.value:g} ({f.agent_name}" + (f"; {f.sources[0]}" if f.sources else "") + ")"
            for f in values
        ]
        finding = (
            f"Agents disagree on {family}{year_bit}: "
            + "; ".join(detail_parts)
            + "."
            + scale_note
            + " Confirm which figure is authoritative."
        )
        warnings.append(("WARN", "source conflict", finding))
    return warnings


def build_year_metric_grid(facts: list[MetricFact]) -> tuple[list[int], list[tuple[str, dict[int, float]]]]:
    """Pivot year-stamped facts into Metric × Year grid."""
    years = sorted({f.year for f in facts if f.year is not None})
    series: dict[str, dict[int, float]] = defaultdict(dict)
    labels: dict[str, str] = {}
    for f in facts:
        if f.year is None:
            continue
        series[f.family][f.year] = f.value
        labels[f.family] = f.label or _humanize(f.family)
    rows = [(labels[k], series[k]) for k in sorted(series.keys())]
    return years, rows


def section_coverage(ctx: BuildContext) -> list[tuple[str, str, str]]:
    """Mark workbook sections Covered only when grounding agents have content."""
    groups = [
        ("01 — Customer Analysis", ["customer_segmentation", "customer_stickiness", "customer_satisfaction"]),
        ("02 — Operational & Risk", ["operational_risk", "supplier_dependence", "supply_chain_resilience", "cost_structure"]),
        ("03 — Financial Analysis", ["historical_performance", "revenue_quality", "cost_structure", "capital_structure"]),
    ]
    rows = []
    for title, agents in groups:
        present = []
        for ak in agents:
            agent = _get_agent(ctx, ak)
            if isinstance(agent, dict) and (agent.get("findings") or agent.get("spec")):
                present.append(ak.replace("_", " "))
        if present:
            rows.append((title, "grounded in: " + ", ".join(present), "Covered"))
        else:
            rows.append((title, "no workflow agent analysed this area", "Not covered"))
    return rows


def _finding_rag(text: str) -> str:
    t = (text or "").lower()
    if any(w in t for w in ("below target", "at risk", "high risk", "critical", "negative", "fail")):
        return "At-Risk"
    if any(w in t for w in ("monitor", "medium", "warn")):
        return "Monitor"
    if any(w in t for w in ("on-track", "on track", "strong", "ready")):
        return "On-Track"
    return "Monitor"


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_exec_dashboard(wb: Workbook, ctx: BuildContext, facts: list[MetricFact]) -> None:
    ws = wb.active
    ws.title = "Executive Dashboard"

    r = 1
    ws.cell(row=r, column=1, value="Operations Dashboard").font = Font(bold=True, size=14)
    r += 1
    company = ctx.profile.company if ctx.profile else ctx.deal_slug
    ws.cell(row=r, column=1, value=f"{company} · Executive Dashboard · Strictly Private & Confidential")
    r += 1
    ws.cell(row=r, column=1, value="Prepared by Agentic CDD")
    r += 2

    pl_lines = _get_spec(ctx, "historical_performance").get("pl_lines", [])
    if not isinstance(pl_lines, list):
        pl_lines = []
    # Keep agent period_type tags for KPI YoY (actual vs forecast). Databook release
    # overlays strip those tags and previously pulled FY2030 into "Latest".
    agent_pl_for_kpi = [dict(l) for l in pl_lines if isinstance(l, dict)]
    try:
        from agetic_cdd_api.services_databook_consume import prefer_released_pl_lines

        pl_lines = prefer_released_pl_lines(ctx.deal_slug, pl_lines)
    except Exception:
        pass

    ws.cell(row=r, column=1, value="Executive Scorecard — Key Performance Indicators").font = _SUBHEADER_FONT
    r += 1
    _write_header_row(ws, r, ["KPI", "Latest (period)", "YoY", "RAG Status"])
    r += 1

    for kpi in derive_kpi_scorecard(facts, agent_pl_for_kpi or pl_lines):
        _write_row(ws, r, [kpi.name, kpi.latest, kpi.yoy, kpi.rag])
        fill = _rag_fill(kpi.rag)
        if fill:
            ws.cell(row=r, column=4).fill = fill
            ws.cell(row=r, column=4).font = Font(bold=True, color="FFFFFF")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Workbook Contents").font = _SUBHEADER_FONT
    r += 1
    _write_header_row(ws, r, ["Section", "Coverage", "Status"])
    r += 1
    for name, coverage, status in section_coverage(ctx):
        _write_row(ws, r, [name, coverage, status])
        r += 1

    warnings = detect_figure_conflicts(facts)
    if warnings:
        r += 1
        ws.cell(row=r, column=1, value="Figure Validation — what was contested, corrected or qualified").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Level", "Check", "Finding"])
        r += 1
        for level, check, finding in warnings[:25]:
            _write_row(ws, r, [level, check, finding])
            ws.cell(row=r, column=1).fill = _WARN_FILL
            r += 1

    r += 1
    ws.cell(row=r, column=1, value="Diligence Workflow Analysis — findings from this deal's agents").font = _SUBHEADER_FONT
    r += 1
    _write_header_row(ws, r, ["Agent", "Finding", "Source document(s)", "Status"])
    r += 1

    preferred = [
        "historical_performance", "operational_risk", "internal_risk",
        "deal_context_and_objectives", "capital_structure", "cash_flow",
    ]
    agent_keys = [k for k in preferred if k in ctx.agent_outputs] or sorted(ctx.agent_outputs.keys())
    for agent_key in agent_keys:
        agent = _get_agent(ctx, agent_key)
        if not isinstance(agent, dict):
            continue
        findings = agent.get("findings") or []
        if not findings:
            continue
        sources = agent.get("sources") or []
        src_str = ", ".join(sources[:3]) if sources else "-"
        name = agent.get("agentName", agent_key.replace("_", " ").title())
        for finding in findings[:2]:
            rag = _finding_rag(str(finding))
            cleaned = sanitize_report_prose(
                str(finding),
                company=company if isinstance(company, str) else None,
                sources=sources if isinstance(sources, list) else None,
            )
            _write_row(ws, r, [name, cleaned, src_str, rag])
            fill = _rag_fill(rag)
            if fill:
                ws.cell(row=r, column=4).fill = fill
                ws.cell(row=r, column=4).font = Font(bold=True, color="FFFFFF")
            r += 1

    _auto_width(ws)


def _build_customer_analysis(wb: Workbook, ctx: BuildContext) -> None:
    ws = wb.create_sheet("Customer Analysis")

    r = 1
    ws.cell(row=r, column=1, value="Operations Dashboard").font = Font(bold=True, size=14)
    r += 1
    company = ctx.profile.company if ctx.profile else ctx.deal_slug
    ws.cell(row=r, column=1, value=f"{company} · Customer Analysis · Strictly Private & Confidential")
    r += 2

    seg_spec = _get_spec(ctx, "customer_segmentation")
    segments = seg_spec.get("segments", []) if isinstance(seg_spec.get("segments"), list) else []

    if segments:
        col_map = [("name", "Segment")]
        for key, label in (
            ("share_pct", "Share (%)"),
            ("avg_age", "Avg Age"),
            ("use_case", "Use Case"),
            ("key_driver", "Key Driver"),
        ):
            if any(isinstance(s, dict) and s.get(key) not in (None, "") for s in segments):
                col_map.append((key, label))
        col_map.append(("_source", "Source"))

        ws.cell(row=r, column=1, value="Revenue by Segment").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, [c[1] for c in col_map])
        r += 1
        for seg in segments:
            if not isinstance(seg, dict):
                continue
            vals = []
            for key, _ in col_map:
                if key == "_source":
                    vals.append("VDR segment rows")
                elif key.endswith("_pct"):
                    vals.append(_safe_num(seg.get(key)))
                else:
                    vals.append(seg.get(key) or "-")
            _write_row(ws, r, vals)
            r += 1

        shares = [
            s.get("share_pct") for s in segments
            if isinstance(s, dict) and isinstance(s.get("share_pct"), (int, float))
        ]
        if shares:
            total = sum(shares)
            if abs(total - 100) > 5:
                r += 1
                ws.cell(
                    row=r, column=1,
                    value=f"⚠ Segment shares sum to {total:g}% (expected ~100%) — confirm segmental note",
                )
                ws.cell(row=r, column=1).fill = _WARN_FILL
                r += 1
        r += 1

        ws.cell(row=r, column=1, value="Customer / Segment Concentration").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Segment", "Share (%)", "Source"])
        r += 1
        for seg in sorted(
            [s for s in segments if isinstance(s, dict)],
            key=lambda s: s.get("share_pct") or 0,
            reverse=True,
        ):
            _write_row(ws, r, [seg.get("name", ""), _safe_num(seg.get("share_pct")), "VDR segment rows"])
            r += 1
        r += 1

    stick_spec = _get_spec(ctx, "customer_stickiness")
    retention = stick_spec.get("retention_metrics")
    if isinstance(retention, dict) and retention:
        ws.cell(row=r, column=1, value="Retention & Stickiness Metrics").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Metric", "Value", "Source"])
        r += 1
        for mk, mv in retention.items():
            if _as_float(mv) is None:
                continue
            _write_row(ws, r, [_humanize(str(mk)), _safe_num(mv), "from the data room"])
            r += 1
        r += 1

    cohorts = stick_spec.get("cohorts")
    if isinstance(cohorts, list) and cohorts:
        keys: list[str] = []
        for c in cohorts:
            if isinstance(c, dict):
                for k in c.keys():
                    if k not in keys:
                        keys.append(k)
        ws.cell(row=r, column=1, value="Cohort Retention Analysis").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, [_humanize(k) for k in keys])
        r += 1
        for c in cohorts:
            if not isinstance(c, dict):
                continue
            _write_row(
                ws, r,
                [
                    _safe_num(c.get(k)) if isinstance(c.get(k), (int, float)) else (c.get(k) or "-")
                    for k in keys
                ],
            )
            r += 1
        r += 1

    sat_spec = _get_spec(ctx, "customer_satisfaction")
    sat_nums = [(k, v) for k, v in sat_spec.items() if k not in _SKIP_SPEC_KEYS and _as_float(v) is not None]
    if sat_nums:
        ws.cell(row=r, column=1, value="Satisfaction Metrics").font = _SUBHEADER_FONT
        r += 1
        for k, v in sat_nums:
            _write_row(ws, r, [_humanize(str(k)), _safe_num(v)])
            r += 1
        r += 1

    churn = sat_spec.get("churn_drivers")
    if isinstance(churn, list) and churn:
        ws.cell(row=r, column=1, value="Churn Drivers").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Driver", "Contribution (%)", "Severity"])
        r += 1
        for cd in churn:
            if not isinstance(cd, dict):
                continue
            sev = str(cd.get("severity") or "")
            _write_row(ws, r, [cd.get("driver", ""), _safe_num(cd.get("contribution_pct")), sev])
            fill = _rag_fill(sev)
            if fill:
                ws.cell(row=r, column=3).fill = fill
                ws.cell(row=r, column=3).font = Font(bold=True, color="FFFFFF")
            r += 1

    _auto_width(ws)


def _build_ops_risk(wb: Workbook, ctx: BuildContext) -> None:
    ws = wb.create_sheet("Operational & Risk")

    r = 1
    ws.cell(row=r, column=1, value="Operations Dashboard").font = Font(bold=True, size=14)
    r += 1
    company = ctx.profile.company if ctx.profile else ctx.deal_slug
    ws.cell(row=r, column=1, value=f"{company} · Operational & Risk · Strictly Private & Confidential")
    r += 2

    vendors = _get_spec(ctx, "supplier_dependence").get("vendors", [])
    if isinstance(vendors, list) and vendors:
        sample = next((v for v in vendors if isinstance(v, dict)), {})
        preferred = [
            ("name", "Vendor"), ("component", "Component"), ("country", "Country"),
            ("annual_value_inr_cr", "Annual Value (INR Cr)"), ("risk_level", "Risk Level"),
        ]
        cols = [
            (k, lab) for k, lab in preferred
            if k in sample or any(isinstance(v, dict) and k in v for v in vendors)
        ]
        if not cols and isinstance(sample, dict):
            cols = [(k, _humanize(k)) for k in sample.keys()]
        ws.cell(row=r, column=1, value="Primary Supplier Landscape").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, [c[1] for c in cols])
        r += 1
        risk_col = next((i for i, (k, _) in enumerate(cols, 1) if k == "risk_level"), None)
        for v in vendors:
            if not isinstance(v, dict):
                continue
            _write_row(
                ws, r,
                [
                    _safe_num(v.get(k)) if isinstance(v.get(k), (int, float)) else (v.get(k) or "")
                    for k, _ in cols
                ],
            )
            if risk_col:
                fill = _rag_fill(str(v.get("risk_level") or ""))
                if fill:
                    ws.cell(row=r, column=risk_col).fill = fill
                    ws.cell(row=r, column=risk_col).font = Font(bold=True, color="FFFFFF")
            r += 1
        r += 1

    kpi_gaps = _get_spec(ctx, "operational_risk").get("kpi_gaps", [])
    if isinstance(kpi_gaps, list) and kpi_gaps:
        ws.cell(row=r, column=1, value="Operational Risk Register").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["KPI", "Status", "RAG"])
        r += 1
        for gap in kpi_gaps:
            if not isinstance(gap, dict):
                continue
            status = str(gap.get("status") or "")
            _write_row(ws, r, [gap.get("kpi", ""), status, status])
            fill = _rag_fill(status)
            if fill:
                ws.cell(row=r, column=3).fill = fill
                ws.cell(row=r, column=3).font = Font(bold=True, color="FFFFFF")
            r += 1
        r += 1

    bom = _get_spec(ctx, "cost_structure").get("bom_components", [])
    if isinstance(bom, list) and bom:
        ws.cell(row=r, column=1, value="Cost Component Breakdown").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Category", "% of Revenue"])
        r += 1
        for comp in bom:
            if not isinstance(comp, dict):
                continue
            _write_row(ws, r, [comp.get("category", ""), _safe_num(comp.get("share_pct"))])
            r += 1
        r += 1

    cash_metrics = _get_spec(ctx, "cash_flow").get("cash_metrics")
    if isinstance(cash_metrics, dict) and cash_metrics:
        ws.cell(row=r, column=1, value="Cash Flow & Liquidity").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Metric", "Value"])
        r += 1
        for mk, mv in cash_metrics.items():
            if _as_float(mv) is None:
                continue
            _write_row(ws, r, [_humanize(str(mk)), _safe_num(mv)])
            r += 1
        r += 1

    ws.cell(row=r, column=1, value="Diligence Workflow Analysis — findings from this deal's agents").font = _SUBHEADER_FONT
    r += 1
    _write_header_row(ws, r, ["Agent", "Finding", "Source document(s)", "Status"])
    r += 1
    for agent_key in [
        "internal_risk", "operational_risk", "supplier_dependence",
        "supply_chain_resilience", "capital_structure",
    ]:
        agent = _get_agent(ctx, agent_key)
        if not isinstance(agent, dict):
            continue
        findings = agent.get("findings") or []
        sources = agent.get("sources") or []
        src_str = ", ".join(sources[:3]) if sources else "-"
        name = agent.get("agentName", agent_key.replace("_", " ").title())
        for finding in findings[:2]:
            rag = _finding_rag(str(finding))
            cleaned = sanitize_report_prose(
                str(finding),
                company=company if isinstance(company, str) else None,
                sources=sources if isinstance(sources, list) else None,
            )
            _write_row(ws, r, [name, cleaned, src_str, rag])
            fill = _rag_fill(rag)
            if fill:
                ws.cell(row=r, column=4).fill = fill
                ws.cell(row=r, column=4).font = Font(bold=True, color="FFFFFF")
            r += 1

    _auto_width(ws)


def _build_financial(wb: Workbook, ctx: BuildContext, facts: list[MetricFact]) -> None:
    ws = wb.create_sheet("Financial Analysis")

    r = 1
    ws.cell(row=r, column=1, value="Operations Dashboard").font = Font(bold=True, size=14)
    r += 1
    company = ctx.profile.company if ctx.profile else ctx.deal_slug
    ws.cell(row=r, column=1, value=f"{company} · Financial Analysis · Strictly Private & Confidential")
    r += 2

    years, grid_rows = build_year_metric_grid(facts)
    grid_by_family: dict[str, dict[int, float]] = {}
    if years and grid_rows:
        ws.cell(row=r, column=1, value="Financial Series — Full Metric Grid").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Metric", *[str(y) for y in years]])
        r += 1
        for label, by_year in grid_rows:
            fam = _canonical_family(str(label))
            grid_by_family[fam] = dict(by_year)
            _write_row(ws, r, [label, *[_safe_num(by_year.get(y)) for y in years]])
            r += 1
        r += 1

    pl_lines = _get_spec(ctx, "historical_performance").get("pl_lines", [])
    if not isinstance(pl_lines, list):
        pl_lines = []
    try:
        from agetic_cdd_api.services_databook_consume import prefer_released_pl_lines

        pl_lines = prefer_released_pl_lines(ctx.deal_slug, pl_lines)
    except Exception:
        pass
    if isinstance(pl_lines, list) and pl_lines:
        fy_cols: list[tuple[int, str]] = []
        seen_cols: set[tuple[int, str]] = set()
        for line in pl_lines:
            if not isinstance(line, dict):
                continue
            for k in line.keys():
                m = _FY_KEY_RE.match(str(k))
                if m:
                    yr = int(m.group(1))
                    key = (yr, str(k))
                    if key not in seen_cols:
                        seen_cols.add(key)
                        fy_cols.append(key)
        fy_cols = sorted(fy_cols, key=lambda x: x[0], reverse=True)
        headers = ["Line Item", *[f"FY{y}" for y, _ in fy_cols], "Unit", "Status"]
        ws.cell(row=r, column=1, value="Financial Series — Databook release (P0)").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, headers)
        r += 1
        for line in pl_lines:
            if not isinstance(line, dict):
                continue
            label = str(line.get("line_item", ""))
            fam = _canonical_family(label)
            fy_status = line.get("fy_status") if isinstance(line.get("fy_status"), dict) else {}
            vals: list[Any] = [label]
            for yr, col in fy_cols:
                raw = line.get(col)
                if raw is None and fy_status.get(str(yr)) == "missing":
                    vals.append("missing")
                    continue
                if raw is None and (
                    fy_status.get(str(yr)) == "doubtful" or yr in (line.get("doubtful_years") or [])
                ):
                    vals.append("doubtful †")
                    continue
                num = _safe_num(raw)
                grid_val = (grid_by_family.get(fam) or {}).get(yr)
                disagrees = (
                    isinstance(num, (int, float))
                    and isinstance(grid_val, (int, float))
                    and not _approx_equal(float(num), float(grid_val))
                )
                if (
                    fy_status.get(str(yr)) == "doubtful"
                    or yr in (line.get("doubtful_years") or [])
                    or disagrees
                ):
                    vals.append(f"{num} †" if num not in (None, "") else "doubtful †")
                else:
                    vals.append(num)
            vals.append(line.get("unit", ""))
            vals.append(str(line.get("databook_status") or ("doubtful" if line.get("databook_provisional") else "proven")))
            _write_row(ws, r, vals)
            r += 1
        r += 1
        ws.cell(
            row=r,
            column=1,
            value=(
                "† Doubtful — provisional, missing, or disagrees with Full Metric Grid. "
                "Trust Financial Series — Full Metric Grid for confirmed figures. "
                "Missing = no releasable evidence."
            ),
        ).font = Font(name="Calibri", size=9, italic=True, color="666666")
        r += 2

        rev = next(
            (l for l in pl_lines if isinstance(l, dict) and _canonical_family(str(l.get("line_item", ""))) == "revenue"),
            None,
        )
        ebitda = next(
            (
                l for l in pl_lines
                if isinstance(l, dict)
                and _canonical_family(str(l.get("line_item", ""))) == "ebitda"
                and "%" not in str(l.get("line_item", ""))
            ),
            None,
        )
        if rev and fy_cols:
            ws.cell(row=r, column=1, value="Revenue & EBITDA Summary").font = _SUBHEADER_FONT
            r += 1
            _write_header_row(ws, r, ["Period", "Revenue", "EBITDA", "Rev Growth", "EBITDA Margin"])
            r += 1
            prev_rev = None
            for yr, col in sorted(fy_cols, key=lambda x: x[0]):
                rv = _as_float(rev.get(col))
                ev = _as_float(ebitda.get(col)) if ebitda else None
                growth = "-"
                if isinstance(rv, float) and isinstance(prev_rev, float) and prev_rev != 0:
                    growth = f"{round((rv - prev_rev) / abs(prev_rev) * 100, 1)}%"
                margin = "-"
                if isinstance(rv, float) and isinstance(ev, float) and rv != 0:
                    margin = f"{round(ev / rv * 100, 1)}%"
                _write_row(ws, r, [f"FY{yr}", _safe_num(rv), _safe_num(ev), growth, margin])
                r += 1
                if isinstance(rv, float):
                    prev_rev = rv
            r += 1

    cs_spec = _get_spec(ctx, "capital_structure")
    cap_blocks = []
    for key in ("leverage_ratios", "dcf_metrics"):
        block = cs_spec.get(key)
        if isinstance(block, dict) and block:
            cap_blocks.append(block)
    if cap_blocks:
        ws.cell(row=r, column=1, value="Capital Structure & Valuation Metrics").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Metric", "Value"])
        r += 1
        for block in cap_blocks:
            for mk, mv in block.items():
                num = _as_float(mv)
                if num is None:
                    continue
                _write_row(
                    ws, r,
                    [
                        _humanize(str(mk)),
                        _format_display(num, _unit_from_path((str(mk),)), _canonical_family(str(mk))),
                    ],
                )
                r += 1
        r += 1

    flags = _get_spec(ctx, "revenue_quality").get("quality_flags", [])
    if isinstance(flags, list) and flags:
        ws.cell(row=r, column=1, value="Revenue Quality Flags").font = _SUBHEADER_FONT
        r += 1
        for flag in flags:
            ws.cell(row=r, column=1, value=flag if isinstance(flag, str) else str(flag))
            ws.cell(row=r, column=1).fill = _WARN_FILL
            r += 1

    _auto_width(ws)


def _build_source_data(wb: Workbook, ctx: BuildContext, facts: list[MetricFact]) -> None:
    ws = wb.create_sheet("Source Data")

    r = 1
    ws.cell(row=r, column=1, value="Operations Dashboard").font = Font(bold=True, size=14)
    r += 1
    company = ctx.profile.company if ctx.profile else ctx.deal_slug
    ws.cell(row=r, column=1, value=f"{company} · Source Data · Strictly Private & Confidential")
    r += 2

    year_facts = [f for f in facts if f.year is not None]
    if year_facts:
        ws.cell(row=r, column=1, value="Underlying Reference Data — provenance fact rows").font = _SUBHEADER_FONT
        r += 1
        _write_header_row(ws, r, ["Metric", "Fiscal Year", "Value", "Unit", "Agent", "Source"])
        r += 1
        for f in sorted(year_facts, key=lambda x: (x.family, x.year or 0, x.agent_key)):
            _write_row(ws, r, [
                f.label,
                f.year,
                f.value,
                f.unit or "-",
                f.agent_name,
                (f.sources[0] if f.sources else "-"),
            ])
            r += 1
        r += 1

    ws.cell(row=r, column=1, value="Agent Catalog").font = _SUBHEADER_FONT
    r += 1
    _write_header_row(ws, r, ["Agent", "DD Code", "Document", "Sources", "Coverage", "Fact Count"])
    r += 1
    for agent_key, agent_data in sorted(ctx.agent_outputs.items()):
        if not isinstance(agent_data, dict):
            continue
        spec = agent_data.get("spec", {})
        if not isinstance(spec, dict):
            continue
        sources = agent_data.get("sources", [])
        findings = agent_data.get("findings", [])
        _write_row(ws, r, [
            agent_data.get("agentName", agent_key.replace("_", " ").title()),
            spec.get("dd_code", "-"),
            spec.get("document", "-"),
            ", ".join(sources[:3]) if sources else "-",
            agent_data.get("source_coverage", "-"),
            len(findings),
        ])
        r += 1

    _auto_width(ws)


class OpsDashboardBuilder(ReportBuilder):
    report_type = "ops_dashboard"

    def build_sections(self, ctx: BuildContext) -> Iterator[dict[str, Any]]:
        facts = harvest_metric_facts(ctx)
        warnings = detect_figure_conflicts(facts)
        pl = _get_spec(ctx, "historical_performance").get("pl_lines", []) or []
        kpis = derive_kpi_scorecard(facts, pl if isinstance(pl, list) else [])
        yield _evt("section", f"Building Executive Dashboard ({len(kpis)} KPIs, {len(warnings)} validation warnings)")
        yield _evt("section", "Building Customer Analysis sheet")
        yield _evt("section", "Building Operational & Risk sheet")
        years, grid = build_year_metric_grid(facts)
        yield _evt("section", f"Building Financial Analysis ({len(grid)} metrics × {len(years)} years)")
        yield _evt("section", "Building Source Data sheet")
        yield _evt("validation", f"{len(warnings)} warning(s) from cross-agent figure checks", warnings=len(warnings))

    def export_artifact(self, ctx: BuildContext) -> str:
        facts = harvest_metric_facts(ctx)
        wb = Workbook()

        _build_exec_dashboard(wb, ctx, facts)
        _build_customer_analysis(wb, ctx)
        _build_ops_risk(wb, ctx)
        _build_financial(wb, ctx, facts)
        _build_source_data(wb, ctx, facts)

        out_dir = report_artifact_dir(ctx.deal_slug, "ops_dashboard")
        raw_company = ctx.profile.company if ctx.profile else ctx.deal_slug
        company = re.sub(r"[^\w\s-]", "", raw_company).strip().replace(" ", "_") or ctx.deal_slug
        filename = f"{company}_Operations_Dashboard.xlsx"
        path = out_dir / filename
        wb.save(str(path))
        deal_root = deals_root() / ctx.deal_slug
        return str(path.relative_to(deal_root))


register_builder("ops_dashboard", OpsDashboardBuilder)
