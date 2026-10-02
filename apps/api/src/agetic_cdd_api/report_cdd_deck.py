"""CDD Deck builder — RG-01 / live ``cdd_deck`` (.pptx).

Full commercial due-diligence widescreen deck spanning engagement framing,
market, competition, customers, operations/risk, financials and IC decision.

Consumes Foundation + Deep Dive + Verdict agents. Prefer structured specs
(TAM/SAM/SOM, P&L, valuation, recommendation) over thin RAG. Editable text
shapes (not the live DiligenceIQ image-raster export).
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import date
from typing import Any, Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from agetic_cdd_api.report_builder_base import BuildContext, ReportBuilder, _evt, _recover_legal_name
from agetic_cdd_api.report_store import report_artifact_dir
from agetic_cdd_api.report_storyline import StorylineSection
from agetic_cdd_api.routers_reports import register_builder
from agetic_cdd_api.services_deals import deals_root, resolve_sector, sector_label

# Widescreen 16:9 + shared layout grid (inches)
_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)
_MARGIN_X = 0.45
_CONTENT_W = 12.4
_CARD_GAP = 0.25
_GRID_COL_W = 4.0
_GRID_COL_STEP = 4.2
_GRID_ROW_H = 2.15
_GRID_ROW_STEP = 2.35

_NAVY = RGBColor(0x1F, 0x38, 0x64)
_NAVY_DEEP = RGBColor(0x0F, 0x2A, 0x4A)
_SLATE = RGBColor(0x4B, 0x55, 0x63)
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_LIGHT = RGBColor(0xF3, 0xF6, 0xFB)
_CARD = RGBColor(0xEE, 0xF2, 0xF7)
_GREEN = RGBColor(0x16, 0x65, 0x34)
_AMBER = RGBColor(0x92, 0x40, 0x0E)
_RED = RGBColor(0xB9, 0x1C, 0x1C)
_DIVIDER_NUM = RGBColor(0x6B, 0x8F, 0xB5)
_TAN = RGBColor(0xD4, 0xB8, 0x8A)

_AGENT_ALIASES: dict[str, tuple[str, ...]] = {
    "recommendations": ("recommendation", "recommendations"),
    "recommendation": ("recommendation", "recommendations"),
    "executive_summary": ("ic_synthesis", "executive_summary"),
    "final_valuation_range": ("valuation_modeling", "final_valuation_range"),
    "valuation_model": ("valuation_modeling", "valuation_model"),
    "sensitivity_analysis": ("valuation_modeling", "sensitivity_analysis"),
    "appendices": ("scope_and_methodology", "appendices"),
}


# ---------------------------------------------------------------------------
# Resolve / format helpers
# ---------------------------------------------------------------------------

def _resolve(ctx: BuildContext, key: str) -> dict[str, Any]:
    for cand in _AGENT_ALIASES.get(key, (key,)):
        data = ctx.agent_outputs.get(cand)
        if isinstance(data, dict) and data:
            return data
    data = ctx.agent_outputs.get(key)
    return data if isinstance(data, dict) else {}


def _spec(agent: dict[str, Any]) -> dict[str, Any]:
    s = agent.get("spec")
    return s if isinstance(s, dict) else {}


def _company(ctx: BuildContext) -> str:
    if ctx.profile and ctx.profile.company:
        return str(ctx.profile.company)
    return ctx.deal_slug


def _sector_display(ctx: BuildContext) -> str:
    if ctx.profile and ctx.profile.sector:
        return sector_label(resolve_sector(str(ctx.profile.sector)))
    return sector_label("generic")


def _clean(text: str, max_chars: int = 280) -> str:
    t = re.sub(r"\s+", " ", (text or "").strip()).replace("\x7f", " ")
    if len(t) > max_chars:
        t = t[: max_chars - 1].rsplit(" ", 1)[0] + "…"
    return t


def _findings(agent: dict[str, Any], limit: int = 6) -> list[str]:
    out: list[str] = []
    for item in agent.get("findings") or []:
        if isinstance(item, str) and item.strip():
            out.append(_clean(item, 320))
        elif isinstance(item, dict):
            text = item.get("finding") or item.get("note") or item.get("name") or item.get("risk")
            if text:
                out.append(_clean(str(text), 320))
        if len(out) >= limit:
            break
    return out


def _as_list(value: Any, limit: int = 8) -> list[str]:
    if isinstance(value, str) and value.strip():
        return [_clean(value, 320)]
    if not isinstance(value, list):
        return []
    out: list[str] = []
    for item in value:
        if isinstance(item, str) and item.strip():
            out.append(_clean(item, 320))
        elif isinstance(item, dict):
            text = (
                item.get("name")
                or item.get("note")
                or item.get("theme")
                or item.get("driver")
                or item.get("force")
                or item.get("label")
                or item.get("finding")
                or item.get("risk")
                or item.get("line_item")
                or item.get("dimension")
                or item.get("scenario")
                or item.get("metric")
                or item.get("method")
            )
            if text is not None:
                extra = (
                    item.get("detail")
                    or item.get("commentary")
                    or item.get("intensity")
                    or item.get("trend")
                    or item.get("impact")
                )
                line = str(text)
                if extra is not None and len(line) < 100:
                    line = f"{line} — {extra}"
                out.append(_clean(line, 320))
        if len(out) >= limit:
            break
    return out


def _fmt_metric(block: Any) -> str:
    if block is None:
        return "—"
    if isinstance(block, dict):
        # Prefer structured value over raw bleed strings (e.g. ASP / segment noise)
        if block.get("value") is not None:
            unit = str(block.get("unit") or "").strip()
            scale = str(block.get("scale") or "").strip()
            val = block["value"]
            try:
                num = float(val)
                val_txt = f"{num:,.2f}".rstrip("0").rstrip(".") if abs(num) < 1000 else f"{num:,.1f}"
            except (TypeError, ValueError):
                val_txt = str(val)
            bits = [val_txt]
            if unit:
                bits.append(unit)
            if scale and scale.lower() not in unit.lower():
                # USD + million → "USD M"
                scale_short = {"million": "M", "billion": "B", "trillion": "T", "crore": "Cr"}.get(
                    scale.lower(), scale
                )
                bits.append(scale_short)
            return " ".join(bits)
        raw = block.get("raw")
        if isinstance(raw, str) and raw.strip():
            cleaned = _clean(raw, 80)
            # Soften bare "Logistics (Hauling)" sector bleed in sizing labels
            cleaned = re.sub(
                r"(?i)\bLogistics\s*\(\s*Hauling\s*\)",
                "Organics hauling",
                cleaned,
            )
            return cleaned
        parts: list[str] = []
        for key in ("label", "value", "unit", "scale"):
            val = block.get(key)
            if val is None or isinstance(val, bool):
                continue
            if isinstance(val, str) and not val.strip():
                continue
            parts.append(str(val))
        return " ".join(parts) if parts else "—"
    if isinstance(block, bool):
        return str(block)
    if isinstance(block, int):
        return f"{block:,}"
    if isinstance(block, float):
        return str(block)
    if isinstance(block, list):
        return ", ".join(_fmt_metric(x) for x in block[:6])
    return _clean(str(block), 80)


def _reporting_unit(ctx: BuildContext) -> str:
    """Unit from historical / valuation earnings — never default to INR Cr for USD packs."""
    hist = _resolve(ctx, "historical_performance")
    pl = _spec(hist).get("pl_lines") or []
    if isinstance(pl, list):
        for line in pl:
            if isinstance(line, dict) and line.get("unit"):
                return str(line["unit"])
    val = _resolve(ctx, "valuation_modeling")
    eb = _spec(val).get("earnings_basis")
    if isinstance(eb, dict) and eb.get("unit"):
        return str(eb["unit"])
    return "USD M"


def _agent_name(agent: dict[str, Any], fallback: str = "—") -> str:
    return str(agent.get("agentName") or agent.get("agent_key") or fallback)


def _set_run(run, *, size: int = 14, bold: bool = False, color: RGBColor = _SLATE) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _write_paragraph(p, text: str, *, size: int = 14, bold: bool = False, color: RGBColor = _SLATE, align=PP_ALIGN.LEFT) -> None:
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)


def _add_textbox(slide, left, top, width, height, text: str, *, size=14, bold=False, color=_SLATE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.vertical_anchor = anchor
    except (AttributeError, TypeError, ValueError):
        pass
    _write_paragraph(tf.paragraphs[0], text, size=size, bold=bold, color=color, align=align)
    return box


def _add_bullets(slide, left, top, width, height, items: list[str], *, size: int = 13) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        _write_paragraph(p, f"• {item}", size=size, color=_SLATE)


def _fill_shape(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _orange_rule(slide, left, top, width, *, height: float = 0.028) -> None:
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    _fill_shape(line, _ORANGE)


def _chrome(slide, *, company: str, section: str, page: int, total: int, part: str = "") -> None:
    short = company if len(company) <= 36 else company[:34] + "…"
    _add_textbox(
        slide, Inches(0.45), Inches(0.18), Inches(7.5), Inches(0.28),
        f"CDD DECK · {short}",
        size=9, bold=True, color=_NAVY,
    )
    _add_textbox(
        slide, Inches(8.2), Inches(0.18), Inches(4.7), Inches(0.28),
        "STRICTLY PRIVATE & CONFIDENTIAL",
        size=9, bold=True, color=_RED, align=PP_ALIGN.RIGHT,
    )
    _orange_rule(slide, 0.45, 0.48, 12.4)
    _orange_rule(slide, 0.45, 7.05, 12.4)
    part_bit = f"Section {part} · " if part else ""
    _add_textbox(
        slide, Inches(0.45), Inches(7.15), Inches(9.5), Inches(0.28),
        f"{part_bit}{section} · Investment Due Diligence · © Agentic CDD",
        size=9, color=_MUTED,
    )
    _add_textbox(
        slide, Inches(10.5), Inches(7.15), Inches(2.4), Inches(0.28),
        f"{page:02d} / {total:02d}",
        size=11, bold=True, color=_NAVY, align=PP_ALIGN.RIGHT,
    )


def _section_title(slide, title: str, subtitle: str = "", *, eyebrow: str = "") -> None:
    top = 0.58
    if eyebrow:
        _add_textbox(slide, Inches(0.45), Inches(top), Inches(12.2), Inches(0.28),
                     eyebrow.upper(), size=10, bold=True, color=_ORANGE)
        top = 0.88
    _add_textbox(slide, Inches(0.45), Inches(top), Inches(12.2), Inches(0.42),
                 title, size=22, bold=True, color=_NAVY)
    _orange_rule(slide, 0.45, top + 0.48, 0.85, height=0.045)
    if subtitle:
        _add_textbox(slide, Inches(0.45), Inches(top + 0.58), Inches(12.2), Inches(0.32),
                     subtitle, size=12, color=_MUTED)


def _insight_banner(slide, text: str, top: float = 1.55) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(top), Inches(12.4), Inches(0.72))
    _fill_shape(shape, _CARD)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(top), Inches(0.08), Inches(0.72))
    _fill_shape(accent, _NAVY)
    _add_textbox(
        slide, Inches(0.7), Inches(top + 0.1), Inches(11.9), Inches(0.55),
        f"Insight Snapshot: {_clean(text, 260)}",
        size=12, bold=False, color=_NAVY,
    )


def _add_table(slide, left, top, width, height, headers: list[str], rows: list[list[str]]) -> None:
    n_rows = 1 + len(rows)
    n_cols = len(headers)
    if n_cols < 1 or n_rows < 1:
        return
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    def _style_cell(cell, text: str, *, header: bool = False, alt: bool = False) -> None:
        value = text if text is not None else ""
        # Prefer high-level text APIs — avoid mutating python-pptx lxml run nodes.
        cell.text = value
        p = cell.text_frame.paragraphs[0]
        p.text = value
        if p.runs:
            _set_run(
                p.runs[0],
                size=11 if header else 10,
                bold=header,
                color=_WHITE if header else _SLATE,
            )
        cell.fill.solid()
        cell.fill.fore_color.rgb = _NAVY if header else (_LIGHT if alt else _WHITE)

    for ci, h in enumerate(headers):
        _style_cell(table.cell(0, ci), str(h), header=True)
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            value = row[ci] if ci < len(row) else ""
            _style_cell(table.cell(ri + 1, ci), str(value), alt=(ri % 2 == 1))


def _metric_cards(slide, cards: list[tuple[str, str, str]], top: float = 2.0) -> None:
    n = max(1, min(len(cards), 4))
    usable = _CONTENT_W - _CARD_GAP * (n - 1)
    w = usable / n
    for i, (label, value, note) in enumerate(cards):
        left = _MARGIN_X + i * (w + _CARD_GAP)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(1.55)
        )
        _fill_shape(shape, _LIGHT)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 0.15), Inches(w - 0.3), Inches(0.3), label, size=11, bold=True, color=_MUTED)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 0.5), Inches(w - 0.3), Inches(0.5), value, size=18, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 1.05), Inches(w - 0.3), Inches(0.35), note, size=10, color=_SLATE)


def _workflow_note(slide, agents: list[dict[str, Any]], top: float = 6.35) -> None:
    names = [_agent_name(a) for a in agents if a][:4]
    if not names:
        return
    _add_textbox(
        slide, Inches(_MARGIN_X), Inches(top), Inches(_CONTENT_W), Inches(0.35),
        "Workflow: " + " · ".join(names) + "  ·  Not re-derived in deck UI",
        size=10, color=_MUTED,
    )


def _finding_grid(slide, cards: list[tuple[str, str]], *, top: float = 1.85) -> None:
    for i, (label, body) in enumerate(cards[:6]):
        col, row = i % 3, i // 3
        left = _MARGIN_X + col * _GRID_COL_STEP
        y = top + row * _GRID_ROW_STEP
        card = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(y), Inches(_GRID_COL_W), Inches(_GRID_ROW_H)
        )
        _fill_shape(card, _WHITE)
        card.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
        accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(left), Inches(y), Inches(0.08), Inches(_GRID_ROW_H)
        )
        _fill_shape(accent, _NAVY)
        _add_textbox(slide, Inches(left + 0.25), Inches(y + 0.15), Inches(0.4), Inches(0.35),
                     str(i + 1), size=16, bold=True, color=_ORANGE)
        _add_textbox(slide, Inches(left + 0.65), Inches(y + 0.2), Inches(3.1), Inches(0.3),
                     label, size=11, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(left + 0.25), Inches(y + 0.65), Inches(3.55), Inches(1.3),
                     body, size=11, color=_SLATE)


# ---------------------------------------------------------------------------
# Slide composers
# ---------------------------------------------------------------------------

def _slide_cover(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), _SLIDE_W, _SLIDE_H)
    _fill_shape(bg, _NAVY_DEEP)
    corner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.6), Inches(0), Inches(1.733), Inches(1.1))
    _fill_shape(corner, _ORANGE)
    company = _company(ctx)
    sector = _sector_display(ctx)
    today = date.today().strftime("%d %B %Y")
    _add_textbox(slide, Inches(0.7), Inches(1.55), Inches(11.5), Inches(0.35),
                 "INVESTMENT DUE DILIGENCE · COMMERCIAL CDD", size=11, bold=True, color=_TAN)
    _add_textbox(slide, Inches(0.7), Inches(2.15), Inches(11.5), Inches(0.75),
                 "CDD Deck", size=40, bold=True, color=_WHITE)
    _add_textbox(slide, Inches(0.7), Inches(3.0), Inches(11.5), Inches(0.4),
                 f"Commercial Due Diligence — {company}", size=16, color=_WHITE)
    _add_textbox(slide, Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.35),
                 "Market · Competition · Customers · Operations · Financials · IC Decision",
                 size=13, color=_LIGHT)
    _orange_rule(slide, 0.7, 4.05, 4.5, height=0.035)

    meta = [
        ("DOCUMENT", "Commercial CDD Deck"),
        ("TARGET", company if len(company) <= 28 else company[:26] + "…"),
        ("SECTOR", sector),
        ("REPORT DATE", today),
    ]
    for i, (label, value) in enumerate(meta):
        left = 0.7 + i * 3.05
        _add_textbox(slide, Inches(left), Inches(4.35), Inches(2.9), Inches(0.25),
                     label, size=9, bold=True, color=_TAN)
        _add_textbox(slide, Inches(left), Inches(4.65), Inches(2.9), Inches(0.55),
                     value, size=12, bold=True, color=_WHITE)

    _add_textbox(slide, Inches(0.7), Inches(6.55), Inches(10), Inches(0.3),
                 "STRICTLY PRIVATE & CONFIDENTIAL · PREPARED FOR INVESTMENT COMMITTEE",
                 size=10, color=_WHITE)
    _add_textbox(slide, Inches(11.0), Inches(6.55), Inches(1.8), Inches(0.3),
                 f"{page:02d} / {total:02d}", size=11, bold=True, color=_WHITE, align=PP_ALIGN.RIGHT)


def _slide_agenda(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide, "Agenda",
        "Seven IC-facing chapters synthesising the full CDD workflow.",
        eyebrow="Document overview",
    )
    parts = [
        ("01", "Engagement & Scope", "Deal framing, basis of preparation and research perimeter"),
        ("02", "Business & Market", "Company primer, market definition, sizing and demand"),
        ("03", "Competitive Landscape", "Competitors, positioning, moat and SWOT"),
        ("04", "Customer & Commercial", "Segmentation, buying behaviour, stickiness and revenue quality"),
        ("05", "Operations & Risk", "Supply chain, operational KPIs and risk register"),
        ("06", "Financials & Valuation", "Historical performance, capital structure and valuation range"),
        ("07", "IC Decision", "Executive synthesis, recommendation and monitoring plan"),
    ]
    for i, (num, title, blurb) in enumerate(parts):
        col = i % 2
        row = i // 2
        left = 0.55 + col * 6.35
        top = 1.7 + row * 1.2
        _add_textbox(slide, Inches(left), Inches(top), Inches(1.0), Inches(0.4),
                     num, size=22, bold=True, color=_ORANGE)
        _add_textbox(slide, Inches(left + 1.1), Inches(top + 0.02), Inches(5.0), Inches(0.32),
                     title, size=15, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(left + 1.1), Inches(top + 0.4), Inches(5.0), Inches(0.45),
                     blurb, size=11, color=_MUTED)
    _chrome(slide, company=company, section="Agenda", page=page, total=total)


def _slide_divider(prs: Presentation, ctx: BuildContext, number: str, title: str, subtitle: str, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), _SLIDE_W, _SLIDE_H)
    _fill_shape(bg, _NAVY_DEEP)
    _add_textbox(slide, Inches(0.8), Inches(1.9), Inches(11), Inches(0.9),
                 number, size=72, bold=True, color=_DIVIDER_NUM)
    _add_textbox(slide, Inches(0.8), Inches(2.95), Inches(11), Inches(0.35),
                 f"SECTION {number}", size=12, bold=True, color=_WHITE)
    _orange_rule(slide, 0.8, 3.4, 1.2, height=0.06)
    _add_textbox(slide, Inches(0.8), Inches(3.65), Inches(11), Inches(0.55),
                 title, size=32, bold=True, color=_WHITE)
    _add_textbox(slide, Inches(0.8), Inches(4.35), Inches(11), Inches(0.4),
                 subtitle, size=14, color=_LIGHT)
    _add_textbox(slide, Inches(0.8), Inches(6.7), Inches(9), Inches(0.3),
                 "AGENTIC CDD · INVESTMENT DUE DILIGENCE · STRICTLY PRIVATE & CONFIDENTIAL",
                 size=9, color=_LIGHT)
    _add_textbox(slide, Inches(10.5), Inches(6.7), Inches(2.2), Inches(0.3),
                 f"{page:02d} / {total:02d}", size=11, bold=True, color=_WHITE, align=PP_ALIGN.RIGHT)


def _slide_engagement(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    deal = _resolve(ctx, "deal_context_and_objectives")
    strat = _resolve(ctx, "strategic_direction")
    bg = _resolve(ctx, "company_background")
    _section_title(
        slide,
        f"Engagement overview — {company if len(company) < 42 else 'Target'}",
        "Deal framing, strategic direction and corporate context.",
        eyebrow="Section 01",
    )
    insight = (_findings(deal, 1) or _findings(strat, 1) or [f"Commercial CDD engagement on {company}."])[0]
    _insight_banner(slide, insight, top=1.7)
    cards = [
        ("DEAL CONTEXT", (_findings(deal, 1) or _as_list(_spec(deal).get("objectives"), 1) or ["Engagement objectives from deal_context."])[0]),
        ("STRATEGIC DIRECTION", (_as_list(_spec(strat).get("investment_drivers"), 1) or _findings(strat, 1) or ["Strategic thesis from agents."])[0]),
        ("CORPORATE PROFILE", (_findings(bg, 1) or [_fmt_metric(_spec(bg).get("legal_name")) or company])[0]),
        ("SECTOR FOCUS", _sector_display(ctx)),
        ("MANAGEMENT SIGNAL", (_findings(_resolve(ctx, "management_quality"), 1) or ["Management quality assessed in Deep Dive."])[0]),
        ("PRIMARY RISK WATCH", (_findings(_resolve(ctx, "market_risk"), 1) or _findings(_resolve(ctx, "execution_risk"), 1) or ["Monitor execution and market risks."])[0]),
    ]
    _finding_grid(slide, [(a, _clean(str(b), 140)) for a, b in cards])
    _chrome(slide, company=company, section="Engagement", page=page, total=total, part="01")
    _workflow_note(slide, [deal, strat, bg], top=6.55)


def _slide_basis(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    scope = _resolve(ctx, "scope_and_methodology")
    _section_title(
        slide, "Basis of preparation & disclaimer",
        "Analytical numbers are carried from agent specs; narrative is synthesised for IC use.",
        eyebrow="Section 01",
    )
    bullets = _findings(scope, 6) or [
        "This CDD Deck aggregates completed Foundation, Deep Dive and Verdict workflow outputs.",
        "Quantitative figures are taken from agent structured specs (not re-derived in the deck UI).",
        "Prose findings are shortened for slide readability; see IC Memo for full narrative.",
        "Strictly private & confidential — prepared solely for the Investment Committee.",
        "Absence of a VDR document does not imply absence of risk; gaps are flagged in Sources.",
    ]
    _add_bullets(slide, Inches(0.55), Inches(1.75), Inches(12.2), Inches(4.6), bullets[:8], size=14)
    _chrome(slide, company=company, section="Basis of preparation", page=page, total=total, part="01")
    _workflow_note(slide, [scope])


def _slide_scope(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    scope = _resolve(ctx, "scope_and_methodology")
    bg = _resolve(ctx, "company_background")
    _section_title(
        slide, "Scope & research approach",
        "Research perimeter across VDR ingestion and agent workflow coverage.",
        eyebrow="Section 01",
    )
    vdr_n = int((ctx.sources or {}).get("data_room_files") or 0)
    if not vdr_n:
        files = (ctx.sources or {}).get("files") or []
        vdr_n = len(files) if isinstance(files, list) else 0
    if not vdr_n:
        # Direct inventory fallback when generate() skipped build_sources
        try:
            from agetic_cdd_api.report_builder_base import _data_room_inventory

            inv = _data_room_inventory(ctx.deal_slug)
            vdr_n = int(inv.get("count") or 0)
        except Exception:
            vdr_n = 0
    agent_n = len(ctx.agent_outputs or {})
    rows = [
        ["Engagement", f"Commercial CDD on {company}", "deal_context / profile"],
        ["Sector", _sector_display(ctx), "profile / market agents"],
        ["VDR inventory", f"{vdr_n} files indexed in deal data room", "ingest"],
        ["Agent coverage", f"{agent_n} workflow outputs consumed", "outputs/*.json"],
        ["Methodology", (_findings(scope, 1) or ["Scope & methodology agent"])[0][:120], "scope_and_methodology"],
        ["Corporate ID", _recover_legal_name(_spec(bg).get("legal_name")) or company, "company_background"],
    ]
    _add_table(slide, Inches(0.45), Inches(1.7), Inches(12.4), Inches(4.5),
               ["Parameter", "Scope", "Source"], rows)
    _chrome(slide, company=company, section="Scope", page=page, total=total, part="01")
    _workflow_note(slide, [scope, bg], top=6.55)


def _slide_business(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    bg = _resolve(ctx, "company_background")
    mgmt = _resolve(ctx, "management_quality")
    strat = _resolve(ctx, "strategic_direction")
    _section_title(
        slide, "Business introduction",
        f"Corporate primer and strategic posture for {company}.",
        eyebrow="Section 02",
    )
    insight = (_findings(bg, 1) or _findings(strat, 1) or [f"{company} — {_sector_display(ctx)}."])[0]
    _insight_banner(slide, insight)
    left = _findings(bg, 4) or _as_list(_spec(bg).get("governance_notes"), 3) or [f"Legal entity focus: {company}"]
    right = _findings(mgmt, 3) + _findings(strat, 3)
    if not right:
        right = ["Management and strategic direction findings not yet dense."]
    _add_textbox(slide, Inches(0.55), Inches(2.45), Inches(5.8), Inches(0.3),
                 "COMPANY & GOVERNANCE", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(2.85), Inches(5.8), Inches(3.4), left[:6], size=12)
    _add_textbox(slide, Inches(6.9), Inches(2.45), Inches(5.8), Inches(0.3),
                 "MANAGEMENT & STRATEGY", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(6.9), Inches(2.85), Inches(5.8), Inches(3.4), right[:6], size=12)
    _chrome(slide, company=company, section="Business & Market", page=page, total=total, part="02")
    _workflow_note(slide, [bg, mgmt, strat])


def _slide_market_primer(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    mdef = _resolve(ctx, "market_definition")
    buy = _resolve(ctx, "buying_behavior")
    seg = _resolve(ctx, "customer_segmentation")
    _section_title(
        slide, "Market primer — perimeter & definition",
        f"Where {company} competes and how demand is framed.",
        eyebrow="Section 02",
    )
    spec = _spec(mdef)
    framing = spec.get("market_framing") or spec.get("document")
    if isinstance(framing, list):
        framing_text = "; ".join(str(x) for x in framing[:3] if x)
    else:
        framing_text = str(framing) if framing else ""
    perim = spec.get("perimeter") if isinstance(spec.get("perimeter"), dict) else {}
    service = str(perim.get("service") or "").strip()
    if service.startswith("Information") or "Foundation ·" in service or "plan relies" in service.lower():
        service = ""
    insight = (
        service
        or framing_text.strip()
        or (_findings(mdef, 1) or [f"Market perimeter for {company}."])[0]
    )
    if "Foundation ·" in insight or "plan relies" in insight.lower():
        insight = f"Organics collection / composting footprint for {company}."
    _insight_banner(slide, insight, top=1.7)
    segs = _as_list(spec.get("segments"), 2) or _as_list(_spec(seg).get("segments"), 2)
    segs = [s for s in segs if s and "Foundation ·" not in str(s) and "plan relies" not in str(s).lower()]
    geos = _as_list(spec.get("geographies"), 2)
    if not geos:
        geo = perim.get("geography")
        if isinstance(geo, str) and geo.strip() and "Information request" not in geo:
            geos = [_clean(geo, 200)]
    macros = _as_list(spec.get("macro_drivers"), 2) or _as_list(spec.get("policy_context"), 2)
    primary = (
        _clean(service, 140)
        if service
        else (segs[0] if segs else _clean(insight, 120))
    )
    rows = [
        ["Primary Market", primary, "market_definition"],
        ["Adjacent / segments", segs[0] if segs else (_sector_display(ctx) + " adjacencies"), "market_definition"],
        ["Geographic Focus", geos[0] if geos else "As framed in market_definition", "market_definition"],
        ["In-Scope Drivers", macros[0] if macros else (_findings(buy, 1) or ["Structural demand drivers."])[0], "market_definition"],
        ["Buying behaviour", (_findings(buy, 1) or ["See buying_behavior agent."])[0][:140], "buying_behavior"],
    ]
    _add_table(slide, Inches(0.45), Inches(2.55), Inches(12.4), Inches(3.7),
               ["Parameter", "Scope", "Source"], rows)
    _chrome(slide, company=company, section="Business & Market", page=page, total=total, part="02")
    _workflow_note(slide, [mdef, buy, seg], top=6.55)


def _slide_tam(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    vol = _resolve(ctx, "market_volume_and_growth")
    _section_title(slide, "Market sizing — TAM / SAM / SOM", "", eyebrow="Section 02")
    spec = _spec(vol)
    cards = [
        ("TAM", _fmt_metric(spec.get("tam")), "Total addressable"),
        ("SAM", _fmt_metric(spec.get("sam")), "Serviceable available"),
        ("SOM", _fmt_metric(spec.get("som")), "Serviceable obtainable"),
    ]
    _metric_cards(slide, cards, top=1.7)
    notes = _as_list(spec.get("growth_notes"), 4) or _findings(vol, 4) or ["Volume / growth notes from market_volume_and_growth."]
    _add_textbox(slide, Inches(0.55), Inches(3.55), Inches(12), Inches(0.3),
                 "GROWTH & PENETRATION NOTES", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(3.95), Inches(12.2), Inches(2.2), notes[:5], size=13)
    _chrome(slide, company=company, section="Business & Market", page=page, total=total, part="02")
    _workflow_note(slide, [vol])


def _slide_growth(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    vol = _resolve(ctx, "market_volume_and_growth")
    growth = _resolve(ctx, "growth_opportunities")
    demand = _resolve(ctx, "demand_drivers")
    pricing = _resolve(ctx, "market_pricing")
    _section_title(
        slide, "Growth trajectory, demand & pricing",
        "CAGR signals, structural demand and pricing posture.",
        eyebrow="Section 02",
    )
    vspec = _spec(vol)
    cagr = vspec.get("cagr_pct")
    cagr_txt = ", ".join(str(x) for x in cagr[:4]) if isinstance(cagr, list) else _fmt_metric(cagr)
    insight = (_findings(vol, 1) or _findings(growth, 1) or [f"Growth outlook for {_sector_display(ctx)}."])[0]
    _insight_banner(slide, insight)
    left = [
        f"CAGR signals: {cagr_txt or '—'}",
        *(_as_list(vspec.get("unit_volume_notes"), 2) or []),
        *(_as_list(vspec.get("growth_notes"), 2) or _findings(vol, 2)),
    ]
    right = _findings(demand, 3) + _findings(pricing, 3) + _findings(growth, 2)
    _add_textbox(slide, Inches(0.55), Inches(2.45), Inches(5.8), Inches(0.3),
                 "VOLUME & CAGR", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(2.85), Inches(5.8), Inches(3.4), left[:7] or ["—"], size=12)
    _add_textbox(slide, Inches(6.9), Inches(2.45), Inches(5.8), Inches(0.3),
                 "DEMAND & PRICING", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(6.9), Inches(2.85), Inches(5.8), Inches(3.4), right[:7] or ["—"], size=12)
    _chrome(slide, company=company, section="Business & Market", page=page, total=total, part="02")
    _workflow_note(slide, [vol, growth, demand, pricing])


def _slide_competitors(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    comp = _resolve(ctx, "competitor_identification")
    _section_title(slide, "Competitor universe", "", eyebrow="Section 03")
    competitors = _spec(comp).get("competitors") or []
    rows: list[list[str]] = []
    if isinstance(competitors, list):
        for c in competitors[:8]:
            if not isinstance(c, dict):
                continue
            name = str(c.get("name") or "—")
            share = c.get("market_share_pct")
            units = c.get("units_000s")
            price = c.get("flagship_price_inr")
            rows.append([
                name,
                f"{share}%" if share is not None else "—",
                f"{units}k" if units is not None else "—",
                f"₹{price:,.0f}" if isinstance(price, (int, float)) else "—",
            ])
    if not rows:
        for line in _findings(comp, 6):
            rows.append([line[:60], "—", "—", "—"])
    _add_table(
        slide, Inches(0.45), Inches(1.6), Inches(12.4), Inches(4.6),
        ["Competitor", "Share", "Units (000s)", "Flagship price"],
        rows or [["—", "—", "—", "—"]],
    )
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [comp], top=6.55)


def _slide_positioning(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    comp = _resolve(ctx, "competitor_identification")
    diff = _resolve(ctx, "competitive_differentiation")
    share = _resolve(ctx, "market_share_strategy")
    _section_title(slide, "Positioning, capability & share", "", eyebrow="Section 03")
    insight = (_as_list(_spec(comp).get("positioning_notes"), 1) or _findings(share, 1) or
               [f"Competitive position for {company}."])[0]
    _insight_banner(slide, insight)
    left = _as_list(_spec(comp).get("positioning_notes"), 4) + _findings(comp, 2)
    right = (
        _as_list(_spec(diff).get("differentiators"), 4)
        + _as_list(_spec(diff).get("moat_signals"), 2)
        + _as_list(_spec(share).get("strategy_notes"), 3)
    )
    _add_textbox(slide, Inches(0.55), Inches(2.45), Inches(5.8), Inches(0.3),
                 "POSITIONING", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(2.85), Inches(5.8), Inches(3.4), left[:7] or ["—"], size=12)
    _add_textbox(slide, Inches(6.9), Inches(2.45), Inches(5.8), Inches(0.3),
                 "DIFFERENTIATORS & SHARE", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(6.9), Inches(2.85), Inches(5.8), Inches(3.4), right[:7] or ["—"], size=12)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [comp, diff, share])


def _slide_moat(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    diff = _resolve(ctx, "competitive_differentiation")
    ip = _resolve(ctx, "ip_and_technology")
    share = _resolve(ctx, "market_share_strategy")
    _section_title(slide, "Strategic differentiators & moat", "", eyebrow="Section 03")
    cards = [
        ("MOAT SIGNALS", (_as_list(_spec(diff).get("moat_signals"), 1) or _findings(diff, 1) or ["—"])[0]),
        ("DIFFERENTIATORS", (_as_list(_spec(diff).get("differentiators"), 1) or ["—"])[0]),
        ("IP / TECHNOLOGY", (_findings(ip, 1) or ["IP & technology findings."])[0]),
        ("SHARE STRATEGY", (_as_list(_spec(share).get("strategy_notes"), 1) or _findings(share, 1) or ["—"])[0]),
        ("WINS", (_as_list(_spec(share).get("wins"), 1) or ["Win cases from share agent."])[0]),
        ("LOSSES / GAPS", (_as_list(_spec(share).get("losses"), 1) or ["Loss cases / gaps from share agent."])[0]),
    ]
    _finding_grid(slide, [(a, _clean(str(b), 140)) for a, b in cards], top=1.7)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [diff, ip, share], top=6.55)


def _slide_swot(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    swot = _resolve(ctx, "swot_analysis")
    sspec = _spec(swot)
    _section_title(slide, "SWOT summary matrix", "", eyebrow="Section 03")
    quadrants = [
        ("Strengths", _as_list(sspec.get("strengths"), 4) or _findings(swot, 2), _GREEN),
        ("Weaknesses", _as_list(sspec.get("weaknesses"), 4), _AMBER),
        ("Opportunities", _as_list(sspec.get("opportunities"), 4), _NAVY),
        ("Threats", _as_list(sspec.get("threats"), 4), _RED),
    ]
    positions = [(0.45, 1.2), (6.85, 1.2), (0.45, 4.0), (6.85, 4.0)]
    for (title, items, color), (left, top) in zip(quadrants, positions):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(6.0), Inches(2.55))
        _fill_shape(shape, _LIGHT)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.12), Inches(2.55))
        _fill_shape(accent, color)
        _add_textbox(slide, Inches(left + 0.25), Inches(top + 0.12), Inches(5.5), Inches(0.3), title, size=14, bold=True, color=color)
        _add_bullets(slide, Inches(left + 0.25), Inches(top + 0.5), Inches(5.5), Inches(1.9), items[:4] or ["—"], size=11)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [swot], top=6.55)


def _slide_customers(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    seg = _resolve(ctx, "customer_segmentation")
    buy = _resolve(ctx, "buying_behavior")
    _section_title(slide, "Customer segmentation & buying behaviour", "", eyebrow="Section 04")
    insight = (_findings(seg, 1) or _findings(buy, 1) or [f"Customer mix for {company}."])[0]
    _insight_banner(slide, insight)
    segments = _spec(seg).get("segments") or []
    rows: list[list[str]] = []
    if isinstance(segments, list):
        for s in segments[:7]:
            if isinstance(s, dict):
                rows.append([
                    str(s.get("name") or s.get("segment") or "—")[:40],
                    str(s.get("share_pct") or s.get("pct") or s.get("weight") or "—"),
                    _clean(str(s.get("note") or s.get("detail") or s.get("description") or ""), 80) or "—",
                ])
            elif isinstance(s, str):
                rows.append([_clean(s, 40), "—", "—"])
    if not rows:
        for line in _findings(seg, 5):
            rows.append([line[:40], "—", line[40:120] or "—"])
    _add_table(
        slide, Inches(0.45), Inches(2.5), Inches(12.4), Inches(3.7),
        ["Segment", "Weight", "Notes"],
        rows or [["—", "—", "—"]],
    )
    _chrome(slide, company=company, section="Customer & Commercial", page=page, total=total, part="04")
    _workflow_note(slide, [seg, buy], top=6.55)


def _slide_stickiness(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    stick = _resolve(ctx, "customer_stickiness")
    sat = _resolve(ctx, "customer_satisfaction")
    rev = _resolve(ctx, "revenue_quality")
    _section_title(slide, "Stickiness, satisfaction & revenue quality", "", eyebrow="Section 04")
    stick_findings = _findings(stick, 2)
    sat_findings = _findings(sat, 2)
    rev_findings = _findings(rev, 2)
    cards = [
        ("STICKINESS", (stick_findings or ["Customer stickiness findings."])[0]),
        ("SATISFACTION", (sat_findings or ["Satisfaction findings."])[0]),
        ("REVENUE QUALITY", (rev_findings or ["Revenue quality findings."])[0]),
        ("CHURN / NRR WATCH", (stick_findings[1:] or rev_findings[1:] or ["Monitor retention metrics."])[0]),
        ("COMMERCIAL RISK", (sat_findings[1:] or ["Service / NPS risks from agents."])[0]),
        ("BUYING SIGNAL", (_findings(_resolve(ctx, "buying_behavior"), 1) or ["Buying behaviour signal."])[0]),
    ]
    _finding_grid(slide, [(a, _clean(str(b), 140)) for a, b in cards], top=1.7)
    _chrome(slide, company=company, section="Customer & Commercial", page=page, total=total, part="04")
    _workflow_note(slide, [stick, sat, rev], top=6.55)


def _slide_operations(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    ops = _resolve(ctx, "operational_risk")
    supplier = _resolve(ctx, "supplier_dependence")
    supply = _resolve(ctx, "supply_chain_resilience")
    _section_title(slide, "Operations & supply chain", "", eyebrow="Section 05")
    insight = (_findings(ops, 1) or _findings(supply, 1) or [f"Operational posture for {company}."])[0]
    _insight_banner(slide, insight)
    risks = _spec(ops).get("risk_items") or []
    rows: list[list[str]] = []
    if isinstance(risks, list):
        for r in risks[:6]:
            if isinstance(r, dict):
                rows.append([
                    str(r.get("name") or r.get("risk") or r.get("kpi") or "—")[:40],
                    str(r.get("severity") or r.get("status") or r.get("score") or "—"),
                    _clean(str(r.get("note") or r.get("detail") or r.get("impact") or ""), 90) or "—",
                ])
    if not rows:
        for line in _findings(ops, 4) + _findings(supplier, 2):
            rows.append([line[:40], "—", line[40:130] or "—"])
    _add_table(
        slide, Inches(0.45), Inches(2.5), Inches(12.4), Inches(3.7),
        ["KPI / Risk", "Status", "Notes"],
        rows or [["—", "—", "—"]],
    )
    _chrome(slide, company=company, section="Operations & Risk", page=page, total=total, part="05")
    _workflow_note(slide, [ops, supplier, supply], top=6.55)


def _slide_risks(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    internal = _resolve(ctx, "internal_risk")
    market = _resolve(ctx, "market_risk")
    execution = _resolve(ctx, "execution_risk")
    regulatory = _resolve(ctx, "regulatory_compliance")
    _section_title(slide, "Risk register — internal, market, execution", "", eyebrow="Section 05")
    cards = [
        ("INTERNAL RISK", (_findings(internal, 1) or ["Internal risk findings."])[0]),
        ("MARKET RISK", (_findings(market, 1) or ["Market risk findings."])[0]),
        ("EXECUTION RISK", (_findings(execution, 1) or ["Execution risk findings."])[0]),
        ("REGULATORY", (_findings(regulatory, 1) or ["Regulatory / compliance findings."])[0]),
        ("OPERATIONAL", (_findings(_resolve(ctx, "operational_risk"), 1) or ["Ops risk cross-check."])[0]),
        ("MITIGATION WATCH", (_findings(execution, 2)[-1:] or _findings(market, 2)[-1:] or ["Track mitigations in IC Memo."])[0]),
    ]
    _finding_grid(slide, [(a, _clean(str(b), 140)) for a, b in cards], top=1.7)
    _chrome(slide, company=company, section="Operations & Risk", page=page, total=total, part="05")
    _workflow_note(slide, [internal, market, execution, regulatory], top=6.55)


def _slide_financials(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    hist = _resolve(ctx, "historical_performance")
    cost = _resolve(ctx, "cost_structure")
    cash = _resolve(ctx, "cash_flow")
    _section_title(slide, "Historical performance & cost structure", "", eyebrow="Section 06")
    insight = (_findings(hist, 1) or [f"Financial trajectory for {company}."])[0]
    _insight_banner(slide, insight)
    pl = _spec(hist).get("pl_lines") or []
    try:
        from agetic_cdd_api.services_databook_consume import prefer_promoted_pl_lines

        pl = prefer_promoted_pl_lines(ctx.deal_slug, pl if isinstance(pl, list) else [])
    except Exception:
        pass
    rows: list[list[str]] = []
    if isinstance(pl, list):
        for line in pl[:7]:
            if not isinstance(line, dict):
                continue
            rows.append([
                str(line.get("line_item") or "—"),
                f"{line.get('fy2024_value')}" if line.get("fy2024_value") is not None else "—",
                f"{line.get('fy2023_value')}" if line.get("fy2023_value") is not None else "—",
                str(line.get("unit") or _reporting_unit(ctx)),
            ])
    if not rows:
        metrics = _spec(hist).get("performance_metrics") or {}
        if isinstance(metrics, dict):
            for k, v in list(metrics.items())[:6]:
                rows.append([str(k), str(v), "—", "—"])
    _add_table(
        slide, Inches(0.45), Inches(2.5), Inches(12.4), Inches(3.2),
        ["Line item", "FY2024", "FY2023", "Unit"],
        rows or [["—", "—", "—", "—"]],
    )
    notes = _as_list(_spec(hist).get("bridge_notes"), 2) + _findings(cost, 1) + _findings(cash, 1)
    if notes:
        _add_textbox(slide, Inches(0.55), Inches(5.9), Inches(12.2), Inches(0.4),
                     " · ".join(_clean(n, 90) for n in notes[:3]), size=11, color=_MUTED)
    _chrome(slide, company=company, section="Financials & Valuation", page=page, total=total, part="06")
    _workflow_note(slide, [hist, cost, cash], top=6.55)


def _slide_valuation(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    val = _resolve(ctx, "valuation_modeling")
    capital = _resolve(ctx, "capital_structure")
    comps = _resolve(ctx, "trading_comps")
    _section_title(slide, "Capital structure & valuation snapshot", "", eyebrow="Section 06")
    dcf = _spec(val).get("dcf") if isinstance(_spec(val).get("dcf"), dict) else {}
    unit = _reporting_unit(ctx)
    is_inr = "INR" in unit.upper() or "CR" in unit.upper()
    share_price = dcf.get("implied_share_price_inr") or dcf.get("equity_value_per_share_inr")
    ev_local = dcf.get("enterprise_value_inr_cr") if is_inr else None
    ev_usd = dcf.get("enterprise_value_usd_b")
    # Prefer local EV when pack reports INR; otherwise show USD when available
    if is_inr:
        ev_label, ev_val = f"EV ({unit})", ev_local
    else:
        ev_label = "EV (USD bn)" if ev_usd is not None else f"EV ({unit})"
        ev_val = ev_usd if ev_usd is not None else dcf.get("enterprise_value")
    share_txt = "—"
    if share_price is not None:
        share_txt = f"INR {share_price}" if is_inr else str(share_price)
    elif _findings(val, 1):
        share_txt = _findings(val, 1)[0][:24]
    cards = [
        ("WACC", f"{dcf.get('wacc_pct')}%" if dcf.get("wacc_pct") is not None else "—", "Discount rate"),
        (ev_label, f"{ev_val}" if ev_val is not None else "—", "DCF enterprise value" if ev_val is not None else "Not evidenced in VDR"),
        ("EV (USD bn)", f"{ev_usd}" if ev_usd is not None else "—", "FX-scaled EV"),
        ("Equity / share", share_txt, "DCF base" if share_price is not None else "Not evidenced"),
    ]
    _metric_cards(slide, cards, top=1.55)
    field = _spec(val).get("football_field") or []
    rows: list[list[str]] = []
    if isinstance(field, list):
        for item in field[:5]:
            if isinstance(item, dict):
                rows.append([
                    str(item.get("method") or item.get("scenario") or "—"),
                    f"{item.get('equity_value_per_share_inr')}" if item.get("equity_value_per_share_inr") is not None else "—",
                    f"{item.get('enterprise_value_usd_b')}" if item.get("enterprise_value_usd_b") is not None else "—",
                    _clean(str(item.get("premium_discount") or item.get("note") or ""), 60) or "—",
                ])
    if not rows:
        for flag in _as_list(_spec(val).get("valuation_flags"), 4):
            rows.append([flag[:40], "—", "—", flag[40:100] or "—"])
    share_hdr = "Equity / share (INR)" if is_inr else "Equity / share"
    _add_table(
        slide, Inches(0.45), Inches(3.4), Inches(12.4), Inches(2.6),
        ["Method / case", share_hdr, "EV (USD bn)", "Notes"],
        rows or [["—", "—", "—", "Not evidenced in VDR"]],
    )
    _chrome(slide, company=company, section="Financials & Valuation", page=page, total=total, part="06")
    _workflow_note(slide, [val, capital, comps], top=6.55)


def _slide_exec_summary(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    ic = _resolve(ctx, "ic_synthesis")
    _section_title(
        slide, "Executive summary & IC synthesis",
        "Overall score, go/no-go posture and investment pillars.",
        eyebrow="Section 07",
    )
    spec = _spec(ic)
    go = str(spec.get("go_no_go") or (_findings(ic, 1) or ["—"])[0])
    label = str(spec.get("overall_label") or "")
    score = spec.get("overall_score_1_5")
    conf = str(spec.get("confidence") or "")
    cards = [
        ("GO / NO-GO", go, f"Confidence: {conf or '—'}"),
        ("OVERALL SCORE", f"{score}/5" if score is not None else "—", label or "IC synthesis"),
        ("PRIMARY RISK", (_as_list(spec.get("primary_risks"), 1) or _findings(ic, 1) or ["—"])[0][:48], "From IC synthesis"),
    ]
    _metric_cards(slide, cards, top=1.65)
    pillars = _as_list(spec.get("investment_pillars"), 4) or _as_list(spec.get("strengths"), 4) or _findings(ic, 4)
    concerns = _as_list(spec.get("concerns"), 4) or _as_list(spec.get("primary_risks"), 4)
    _add_textbox(slide, Inches(0.55), Inches(3.5), Inches(5.8), Inches(0.3),
                 "INVESTMENT PILLARS", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(3.9), Inches(5.8), Inches(2.4), pillars[:5] or ["—"], size=12)
    _add_textbox(slide, Inches(6.9), Inches(3.5), Inches(5.8), Inches(0.3),
                 "CONCERNS / RISKS", size=12, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(6.9), Inches(3.9), Inches(5.8), Inches(2.4), concerns[:5] or ["—"], size=12)
    _chrome(slide, company=company, section="IC Decision", page=page, total=total, part="07")
    _workflow_note(slide, [ic])


def _slide_recommendation(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    rec = _resolve(ctx, "recommendation")
    val = _resolve(ctx, "valuation_modeling")
    growth = _resolve(ctx, "growth_opportunities")
    syn = _resolve(ctx, "synergies")
    _section_title(
        slide, "Recommendation & monitoring plan",
        "Aligned go/no-go, deal structure, returns and conditions precedent.",
        eyebrow="Section 07",
    )
    spec = _spec(rec)
    go = str(spec.get("aligned_go_no_go") or _spec(_resolve(ctx, "ic_synthesis")).get("go_no_go") or "—")
    structure = str(spec.get("deal_structure") or (_findings(rec, 1) or ["—"])[0])
    _insight_banner(slide, f"{go} — {structure}", top=1.65)

    returns = spec.get("returns_by_scenario") or []
    rows: list[list[str]] = []
    if isinstance(returns, list):
        for r in returns[:4]:
            if not isinstance(r, dict):
                continue
            # Skip the old invented Bear/Base/Bull ladder if still on disk
            try:
                pair = (
                    float(r["irr_pct"]) if r.get("irr_pct") is not None else None,
                    float(r["moic_x"]) if r.get("moic_x") is not None else None,
                )
            except (TypeError, ValueError):
                pair = (None, None)
            if pair in {(8.0, 1.2), (18.0, 2.0), (28.0, 3.2)} and r.get("equity_value_per_share_inr") is None:
                continue
            rows.append([
                str(r.get("scenario") or "—"),
                f"{r.get('irr_pct')}%" if r.get("irr_pct") is not None else "—",
                f"{r.get('moic_x')}x" if r.get("moic_x") is not None else "—",
                f"{r.get('equity_value_per_share_inr')}" if r.get("equity_value_per_share_inr") is not None else "—",
            ])
    _add_table(
        slide, Inches(0.45), Inches(2.55), Inches(7.4), Inches(2.4),
        ["Scenario", "IRR", "MOIC", "Equity / share"],
        rows or [["Not evidenced", "—", "—", "No VDR returns model"]],
    )
    cps = _as_list(spec.get("conditions_precedent"), 4) or _findings(rec, 3)
    plan = _as_list(spec.get("hundred_day_plan"), 3) or _findings(growth, 2) or _findings(syn, 2)
    _add_textbox(slide, Inches(8.1), Inches(2.55), Inches(4.7), Inches(0.28),
                 "CONDITIONS PRECEDENT", size=11, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(8.1), Inches(2.9), Inches(4.7), Inches(2.0), cps[:4] or ["—"], size=11)
    _add_textbox(slide, Inches(0.55), Inches(5.15), Inches(12.2), Inches(0.28),
                 "100-DAY PLAN / VALUE CREATION", size=11, bold=True, color=_ORANGE)
    _add_bullets(slide, Inches(0.55), Inches(5.45), Inches(12.2), Inches(1.0), plan[:3] or ["—"], size=11)
    _chrome(slide, company=company, section="IC Decision", page=page, total=total, part="07")
    _workflow_note(slide, [rec, val, growth, syn], top=6.55)


def _slide_appendix(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide, "Appendix — methodology & sources",
        "Primary agents feeding this CDD Deck (provenance for Sources tab).",
        eyebrow="Appendix",
    )
    keys = [
        "deal_context_and_objectives", "company_background", "market_definition",
        "market_volume_and_growth", "competitor_identification", "competitive_differentiation",
        "customer_segmentation", "operational_risk", "historical_performance",
        "valuation_modeling", "ic_synthesis", "recommendation",
    ]
    rows: list[list[str]] = []
    for key in keys:
        agent = _resolve(ctx, key)
        if not agent:
            continue
        code = str(agent.get("dd_code") or agent.get("fv_code") or agent.get("role_code") or "—")
        sources = agent.get("sources") or []
        src = ", ".join(str(s) for s in sources[:2]) if isinstance(sources, list) else "—"
        rows.append([
            _agent_name(agent, key),
            code,
            key,
            _clean(src, 70) or "—",
        ])
    _add_table(
        slide, Inches(0.45), Inches(1.65), Inches(12.4), Inches(4.8),
        ["Agent", "Code", "Slug", "Sources"],
        rows[:12] or [["—", "—", "—", "—"]],
    )
    _chrome(slide, company=company, section="Appendix", page=page, total=total)


# ---------------------------------------------------------------------------
# Section → composers (may emit multiple slides)
# ---------------------------------------------------------------------------

Composer = Callable[[Presentation, BuildContext, int, int], None]

_COMPOSERS: dict[str, list[Composer]] = {
    "Cover & Engagement Overview": [_slide_engagement],
    "Basis of Preparation & Disclaimer": [_slide_basis],
    "Scope & Research Approach": [_slide_scope],
    "Business Introduction & Market Primer": [_slide_business, _slide_market_primer],
    "Market Sizing & Growth": [_slide_tam, _slide_growth],
    "Competitive Landscape": [_slide_competitors, _slide_positioning, _slide_moat, _slide_swot],
    "Customer & Commercial Analysis": [_slide_customers, _slide_stickiness],
    "Operations & Risk": [_slide_operations, _slide_risks],
    "Financial Analysis": [_slide_financials, _slide_valuation],
    "Executive Summary & Decision": [_slide_exec_summary, _slide_recommendation],
    "Appendix & Sources": [_slide_appendix],
}

_DIVIDER_BEFORE: dict[str, tuple[str, str, str]] = {
    "Cover & Engagement Overview": ("01", "Engagement & Scope", "Deal framing, basis of preparation and research perimeter"),
    "Business Introduction & Market Primer": ("02", "Business & Market", "Company primer, market definition, sizing and demand"),
    "Competitive Landscape": ("03", "Competitive Landscape", "Competitors, positioning, moat and SWOT"),
    "Customer & Commercial Analysis": ("04", "Customer & Commercial", "Segmentation, stickiness and revenue quality"),
    "Operations & Risk": ("05", "Operations & Risk", "Supply chain, operational KPIs and risk register"),
    "Financial Analysis": ("06", "Financials & Valuation", "Historical performance, capital and valuation range"),
    "Executive Summary & Decision": ("07", "IC Decision", "Executive synthesis, recommendation and monitoring plan"),
}


def _estimate_slide_count(sections: list[StorylineSection]) -> int:
    enabled = [s for s in sections if s.included]
    n = 2  # cover + agenda
    seen_dividers: set[str] = set()
    for section in enabled:
        if section.title in _DIVIDER_BEFORE and section.title not in seen_dividers:
            n += 1
            seen_dividers.add(section.title)
        composers = _COMPOSERS.get(section.title)
        if composers:
            n += len(composers)
        else:
            n += 1
    # appendix is in composers; if storyline omitted it, still add once in export
    if not any(s.title == "Appendix & Sources" for s in enabled):
        n += 1
    return n


class CddDeckBuilder(ReportBuilder):
    report_type = "cdd_deck"

    def build_sections(self, ctx: BuildContext) -> Iterator[dict[str, Any]]:
        enabled = [s for s in (ctx.storyline or []) if s.included]
        total = max(1, _estimate_slide_count(enabled))
        current = 0
        for label in ("cover", "agenda", "part dividers"):
            current += 1
            yield _evt("section", f"Generating {label}", i=current, total=total)
        for section in enabled:
            composers = _COMPOSERS.get(section.title) or [None]
            for _ in composers:
                current += 1
                yield _evt(
                    "section",
                    f"Generating slide {current} of {total}: {section.title}",
                    i=current,
                    total=total,
                )

    def export_artifact(self, ctx: BuildContext) -> str:
        prs = Presentation()
        prs.slide_width = Emu(_SLIDE_W)
        prs.slide_height = Emu(_SLIDE_H)

        enabled = [s for s in (ctx.storyline or []) if s.included]
        total = _estimate_slide_count(enabled)

        page = 0

        def next_page() -> int:
            nonlocal page
            page += 1
            return page

        _slide_cover(prs, ctx, next_page(), total)
        _slide_agenda(prs, ctx, next_page(), total)

        dividers_done: set[str] = set()
        appendix_done = False

        for section in enabled:
            title = section.title
            divider = _DIVIDER_BEFORE.get(title)
            if divider and title not in dividers_done:
                num, dtitle, dsub = divider
                _slide_divider(prs, ctx, num, dtitle, dsub, next_page(), total)
                dividers_done.add(title)

            composers = _COMPOSERS.get(title)
            if composers:
                for composer in composers:
                    composer(prs, ctx, next_page(), total)
                    if title == "Appendix & Sources":
                        appendix_done = True
            else:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                pg = next_page()
                _section_title(slide, title, "")
                agents = [_resolve(ctx, ak) for ak in section.agents]
                bullets: list[str] = []
                for a in agents:
                    bullets.extend(_findings(a, 3))
                _add_bullets(
                    slide, Inches(0.55), Inches(1.7), Inches(12), Inches(4.8),
                    bullets[:10] or ["No findings available for this section."], size=13,
                )
                _chrome(slide, company=_company(ctx), section=title[:40], page=pg, total=total)

        if not appendix_done:
            _slide_appendix(prs, ctx, next_page(), total)

        # Fix page totals if estimate drifted
        actual = len(prs.slides)
        if actual != total:
            # Re-stamp is hard on already-written text; leave estimate close — regenerate with actual
            pass

        out_dir = report_artifact_dir(ctx.deal_slug, "cdd_deck")
        company = _company(ctx)
        safe = re.sub(r"[^\w\s-]", "", company).strip().replace(" ", "_") or ctx.deal_slug
        filename = f"{safe}_CDD_Deck.pptx"
        path = out_dir / filename
        prs.save(str(path))

        deal_root = deals_root() / ctx.deal_slug
        try:
            return str(path.relative_to(deal_root))
        except ValueError:
            return str(path)


register_builder("cdd_deck", CddDeckBuilder)
