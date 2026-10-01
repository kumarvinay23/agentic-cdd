"""Excel export/import and freshness status for Databook (DiligenceIQ-parity)."""

from __future__ import annotations

import io
import logging
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from agetic_cdd_api.models import Deal
from agetic_cdd_api.services_databook_decisions import correct_row
from agetic_cdd_api.services_databook_models import ExtractedRow, RowStatus
from agetic_cdd_api.services_databook_store import load_meta, load_rows
from agetic_cdd_api.services_library import load_library_index
from agetic_cdd_api.services_vdr import documents_dir

logger = logging.getLogger(__name__)

_EDITABLE_HEADERS = [
    "row_id",
    "caption",
    "metric_key",
    "fiscal_year",
    "value",
    "unit",
    "currency",
    "scale",
    "status",
    "source_name",
    "assumption",
]

_IMPORT_REASON = "Imported from edited Excel"


def _parse_iso(value: str | None) -> datetime | None:
    if not value or not isinstance(value, str):
        return None
    text = value.strip()
    if not text:
        return None
    try:
        if text.endswith("Z"):
            text = text[:-1] + "+00:00"
        dt = datetime.fromisoformat(text)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except ValueError:
        return None


def databook_freshness(deal: Deal) -> dict[str, Any]:
    """Compare last Rescan/Deep vs library ingest and VDR file mtimes."""
    meta = load_meta(deal)
    index = load_library_index(deal) or {}
    library_at = _parse_iso(str(index.get("generated_at") or "") or None)
    last_rescan = _parse_iso(meta.last_rescan_at)
    last_deep = _parse_iso(meta.last_deep_at)
    last = max((t for t in (last_rescan, last_deep) if t is not None), default=None)

    newest_vdr: datetime | None = None
    docs = documents_dir(deal)
    if docs.is_dir():
        for path in docs.iterdir():
            if not path.is_file() or path.name.startswith("."):
                continue
            try:
                mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc)
            except OSError:
                continue
            if newest_vdr is None or mtime > newest_vdr:
                newest_vdr = mtime

    stale_reason: str | None = None
    if last is None:
        stale_reason = "Never scanned — run Update databook to build from the library."
    elif library_at and library_at > last:
        stale_reason = "Library was updated after the last Rescan."
    elif newest_vdr and newest_vdr > last:
        stale_reason = "VDR files are newer than the last Rescan — reingest or Deep, then Update."

    up_to_date = stale_reason is None
    return {
        "up_to_date": up_to_date,
        "label": "Databook is up to date" if up_to_date else "Update databook",
        "stale_reason": stale_reason,
        "library_generated_at": index.get("generated_at"),
        "last_rescan_at": meta.last_rescan_at,
        "last_deep_at": meta.last_deep_at,
        "action": "none" if up_to_date else "rescan",
    }


def build_editable_workbook(deal: Deal) -> bytes:
    """Export non-dropped extract rows as an editable Excel workbook for round-trip import."""
    rows = [r for r in load_rows(deal) if r.status != RowStatus.DROPPED]
    wb = Workbook()
    ws = wb.active
    ws.title = "Databook"
    ws.append(_EDITABLE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in sorted(
        rows,
        key=lambda r: (r.metric_key or "", r.fiscal_year or 0, r.source_name, r.row_id),
    ):
        ws.append(
            [
                row.row_id,
                row.caption,
                row.metric_key or "",
                row.fiscal_year if row.fiscal_year is not None else "",
                row.value,
                row.unit or "",
                row.currency or "",
                row.scale or "",
                row.status.value,
                row.source_name,
                "1" if row.assumption else "0",
            ]
        )

    guide = wb.create_sheet("README", 0)
    guide["A1"] = "Databook editable export"
    guide["A1"].font = Font(bold=True, size=14)
    guide["A3"] = (
        "Edit values on the Databook sheet, then use Import edited Excel. "
        "Keep row_id stable — that is how corrections are matched. "
        "Changing value, metric_key, unit, currency, or scale applies a Correct decision "
        f"(reason: '{_IMPORT_REASON}'). Dropped rows are not exported."
    )
    guide.column_dimensions["A"].width = 100
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["J"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cell_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None


def _cell_int(value: Any) -> int | None:
    num = _cell_float(value)
    if num is None:
        return None
    return int(num)


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_changed(existing: ExtractedRow, patch: dict[str, Any]) -> bool:
    for key, new_val in patch.items():
        old = getattr(existing, key, None)
        if key == "value":
            if new_val is None:
                continue
            try:
                if abs(float(old) - float(new_val)) > 1e-9:
                    return True
            except (TypeError, ValueError):
                return True
        elif _norm(old) != _norm(new_val):
            return True
    return False


def import_edited_workbook(
    deal: Deal,
    content: bytes,
    *,
    actor: str,
    reason: str | None = None,
) -> dict[str, Any]:
    """Apply Excel edits as Correct decisions (audit trail + promote when metric+year set)."""
    if not content:
        raise ValueError("Empty workbook")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # openpyxl raises various errors for bad files
        raise ValueError(f"Could not read Excel workbook: {exc}") from exc

    sheet = wb["Databook"] if "Databook" in wb.sheetnames else wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise ValueError("Workbook has no header row") from exc

    headers = [str(h).strip().lower() if h is not None else "" for h in header_row]
    col = {name: idx for idx, name in enumerate(headers) if name}
    if "row_id" not in col or "value" not in col:
        raise ValueError("Workbook must include row_id and value columns (export from Databook first)")

    by_id = {r.row_id: r for r in load_rows(deal)}
    applied: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    import_reason = (reason or "").strip() or _IMPORT_REASON

    for line_no, values in enumerate(rows_iter, start=2):
        if not values or all(v is None or str(v).strip() == "" for v in values):
            continue

        def get(name: str) -> Any:
            idx = col.get(name)
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        row_id = _cell_str(get("row_id"))
        if not row_id:
            skipped.append({"line": line_no, "reason": "missing row_id"})
            continue
        existing = by_id.get(row_id)
        if existing is None:
            skipped.append({"line": line_no, "row_id": row_id, "reason": "unknown row_id"})
            continue

        patch: dict[str, Any] = {}
        value = _cell_float(get("value"))
        if value is not None:
            patch["value"] = value
        metric_key = _cell_str(get("metric_key"))
        if metric_key is not None:
            patch["metric_key"] = metric_key
        unit = _cell_str(get("unit"))
        if "unit" in col:
            patch["unit"] = unit
        currency = _cell_str(get("currency"))
        if "currency" in col:
            patch["currency"] = currency
        scale = _cell_str(get("scale"))
        if "scale" in col:
            patch["scale"] = scale
        caption = _cell_str(get("caption"))
        if caption is not None:
            patch["caption"] = caption

        # fiscal_year is part of identity; allow correction via correct_row only if we extend —
        # for now only value/metric/unit/currency/scale/caption (correct_row API).
        if not patch:
            skipped.append({"line": line_no, "row_id": row_id, "reason": "no editable fields"})
            continue
        if not _row_changed(existing, patch):
            skipped.append({"line": line_no, "row_id": row_id, "reason": "unchanged"})
            continue

        try:
            result = correct_row(
                deal,
                row_id,
                reason=import_reason,
                actor=actor,
                value=patch.get("value"),
                metric_key=patch.get("metric_key"),
                unit=patch.get("unit") if "unit" in patch else None,
                currency=patch.get("currency") if "currency" in patch else None,
                scale=patch.get("scale") if "scale" in patch else None,
                caption=patch.get("caption"),
            )
            applied.append(
                {
                    "line": line_no,
                    "row_id": row_id,
                    "decision_id": result.get("decision_id"),
                    "value": patch.get("value"),
                }
            )
            # Refresh local index so later rows see updated state
            updated = result.get("row") or {}
            if isinstance(updated, dict) and updated.get("row_id"):
                by_id[row_id] = ExtractedRow.model_validate(updated)
        except Exception as exc:
            logger.warning("Excel import failed for row %s: %s", row_id, exc)
            errors.append({"line": line_no, "row_id": row_id, "reason": str(exc)})

    return {
        "applied": len(applied),
        "skipped": len(skipped),
        "errors": len(errors),
        "items_applied": applied[:50],
        "items_skipped": skipped[:50],
        "items_errors": errors[:50],
        "reason": import_reason,
    }
