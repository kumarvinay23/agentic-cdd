"""Market Intel Deck builder — RG-02 / live ``market_deck`` (.pptx).

Live DiligenceIQ storyline: 17 content sections in 4 agenda parts
(Key Findings · Market Analysis · Competitive Landscape · Summary).

Consumes Foundation (F-01 / strategic_direction) + market / competition
Deep Dive agents. Prefer structured specs (TAM/SAM/SOM, competitors, SWOT)
over thin RAG. Editable text shapes (not raster slides).
"""

from __future__ import annotations

import re
from collections.abc import Iterator
from datetime import date
from typing import Any, Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from agetic_cdd_api.report_builder_base import BuildContext, ReportBuilder, _evt, sanitize_report_prose
from agetic_cdd_api.report_store import report_artifact_dir
from agetic_cdd_api.report_storyline import StorylineSection
from agetic_cdd_api.routers_reports import register_builder
from agetic_cdd_api.services_deals import deals_root, resolve_sector, sector_label

# Widescreen 16:9
_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)

_NAVY = RGBColor(0x1F, 0x38, 0x64)
_NAVY_DEEP = RGBColor(0x0F, 0x2A, 0x4A)
_SLATE = RGBColor(0x4B, 0x55, 0x63)
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_ORANGE = RGBColor(0xE8, 0x7A, 0x2E)
_ACCENT = _ORANGE
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_LIGHT = RGBColor(0xF3, 0xF6, 0xFB)
_CARD = RGBColor(0xEE, 0xF2, 0xF7)
_GREEN = RGBColor(0x16, 0x65, 0x34)
_AMBER = RGBColor(0x92, 0x40, 0x0E)
_RED = RGBColor(0xB9, 0x1C, 0x1C)
_DIVIDER_NUM = RGBColor(0x6B, 0x8F, 0xB5)
_TAN = RGBColor(0xD4, 0xB8, 0x8A)

_AGENT_ALIASES: dict[str, tuple[str, ...]] = {
    "recommendations": ("recommendation", "recommendations"),
    "recommendation": ("recommendation", "recommendations"),
    "market_intel_deck": ("market_definition", "market_volume_and_growth", "market_intel_deck"),
    "sensitivity_analysis": ("valuation_modeling", "sensitivity_analysis"),
}


# ---------------------------------------------------------------------------
# Resolve / format helpers
# ---------------------------------------------------------------------------

def _resolve(ctx: BuildContext, key: str) -> dict[str, Any]:
    for cand in _AGENT_ALIASES.get(key, (key,)):
        data = ctx.agent_outputs.get(cand)
        if isinstance(data, dict) and data:
            return data
    data = ctx.agent_outputs.get(key)
    return data if isinstance(data, dict) else {}


def _spec(agent: dict[str, Any]) -> dict[str, Any]:
    s = agent.get("spec")
    return s if isinstance(s, dict) else {}


def _company(ctx: BuildContext) -> str:
    if ctx.profile and ctx.profile.company:
        return str(ctx.profile.company)
    return ctx.deal_slug


def _sector_display(ctx: BuildContext) -> str:
    if ctx.profile and ctx.profile.sector:
        return sector_label(resolve_sector(str(ctx.profile.sector)))
    return sector_label("generic")


def _clean(text: str, max_chars: int = 280, *, sources: list[Any] | None = None) -> str:
    t = re.sub(r"\s+", " ", (text or "").strip()).replace("\x7f", " ")
    t = sanitize_report_prose(t, sources=sources)
    if len(t) > max_chars:
        t = t[: max_chars - 1].rsplit(" ", 1)[0] + "…"
    return t


def _findings(agent: dict[str, Any], limit: int = 6) -> list[str]:
    out: list[str] = []
    for item in agent.get("findings") or []:
        if isinstance(item, str) and item.strip():
            out.append(_clean(item, 320))
        elif isinstance(item, dict):
            text = item.get("finding") or item.get("note") or item.get("name") or item.get("risk")
            if text:
                out.append(_clean(str(text), 320))
        if len(out) >= limit:
            break
    return out


def _as_list(value: Any, limit: int = 8) -> list[str]:
    if isinstance(value, str) and value.strip():
        return [_clean(value, 320)]
    if not isinstance(value, list):
        return []
    out: list[str] = []
    for item in value:
        if isinstance(item, str) and item.strip():
            out.append(_clean(item, 320))
        elif isinstance(item, dict):
            text = (
                item.get("name")
                or item.get("note")
                or item.get("theme")
                or item.get("driver")
                or item.get("force")
                or item.get("label")
                or item.get("finding")
            )
            if text is not None:
                extra = item.get("intensity") or item.get("detail") or item.get("trend")
                line = str(text)
                if extra is not None and len(line) < 100:
                    line = f"{line} ({extra})"
                out.append(_clean(line, 320))
        if len(out) >= limit:
            break
    return out


def _fmt_metric(block: Any) -> str:
    if block is None:
        return "—"
    if isinstance(block, dict):
        # Prefer structured value over raw (raw may still cite a single segment row)
        if block.get("value") is not None:
            unit = str(block.get("unit") or "").strip()
            scale = str(block.get("scale") or "").strip()
            val = block["value"]
            try:
                num = float(val)
                val_txt = f"{num:,.2f}".rstrip("0").rstrip(".") if abs(num) < 1000 else f"{num:,.1f}"
            except (TypeError, ValueError):
                val_txt = str(val)
            bits = [val_txt]
            if unit:
                bits.append(unit)
            if scale and scale.lower() not in unit.lower():
                scale_short = {"million": "M", "billion": "B", "trillion": "T", "crore": "Cr"}.get(
                    scale.lower(), scale
                )
                bits.append(scale_short)
            return " ".join(bits)
        raw = block.get("raw")
        if isinstance(raw, str) and raw.strip():
            return _clean(raw, 80)
        value = block.get("value")
        unit = block.get("unit") or ""
        scale = block.get("scale") or ""
        label = block.get("label") or ""
        bits = []
        if label:
            bits.append(str(label))
        if value is not None and not isinstance(value, bool):
            bits.append(str(value))
        if unit:
            bits.append(str(unit))
        if scale:
            bits.append(str(scale))
        return " ".join(bits) if bits else "—"
    if isinstance(block, (int, float)) and not isinstance(block, bool):
        return str(block)
    if isinstance(block, list):
        return ", ".join(_fmt_metric(x) for x in block[:6])
    return _clean(str(block), 80)


def _agent_name(agent: dict[str, Any], fallback: str = "—") -> str:
    return str(agent.get("agentName") or agent.get("agent_key") or fallback)


def _set_run(run, *, size: int = 14, bold: bool = False, color: RGBColor = _SLATE) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _write_paragraph(p, text: str, *, size: int = 14, bold: bool = False, color: RGBColor = _SLATE, align=PP_ALIGN.LEFT) -> None:
    """Replace paragraph text with a single styled run (avoids empty/duplicate runs)."""
    p.alignment = align
    # Clear existing runs cleanly
    p.clear()
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, bold=bold, color=color)


def _add_textbox(slide, left, top, width, height, text: str, *, size=14, bold=False, color=_SLATE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf.vertical_anchor = anchor
    except Exception:
        pass
    _write_paragraph(tf.paragraphs[0], text, size=size, bold=bold, color=color, align=align)
    return box


def _add_bullets(slide, left, top, width, height, items: list[str], *, size: int = 13) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(6)
        _write_paragraph(p, f"• {item}", size=size, color=_SLATE)


def _fill_shape(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _orange_rule(slide, left, top, width, *, height: float = 0.028) -> None:
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    _fill_shape(line, _ORANGE)


def _chrome(slide, *, company: str, section: str, page: int, total: int, part: str = "") -> None:
    """Live DiligenceIQ-style header / footer chrome (editable shapes)."""
    short = company if len(company) <= 36 else company[:34] + "…"
    _add_textbox(
        slide, Inches(0.45), Inches(0.18), Inches(7.5), Inches(0.28),
        f"MARKET INTEL DECK · {short}",
        size=9, bold=True, color=_NAVY,
    )
    _add_textbox(
        slide, Inches(8.2), Inches(0.18), Inches(4.7), Inches(0.28),
        "STRICTLY PRIVATE & CONFIDENTIAL",
        size=9, bold=True, color=_RED, align=PP_ALIGN.RIGHT,
    )
    _orange_rule(slide, 0.45, 0.48, 12.4)

    _orange_rule(slide, 0.45, 7.05, 12.4)
    part_bit = f"Section {part} · " if part else ""
    _add_textbox(
        slide, Inches(0.45), Inches(7.15), Inches(9.5), Inches(0.28),
        f"{part_bit}{section} · Investment Due Diligence · © Agentic CDD",
        size=9, color=_MUTED,
    )
    _add_textbox(
        slide, Inches(10.5), Inches(7.15), Inches(2.4), Inches(0.28),
        f"{page:02d} / {total:02d}",
        size=11, bold=True, color=_NAVY, align=PP_ALIGN.RIGHT,
    )


def _section_title(slide, title: str, subtitle: str = "", *, eyebrow: str = "") -> None:
    top = 0.58
    if eyebrow:
        _add_textbox(slide, Inches(0.45), Inches(top), Inches(12.2), Inches(0.28),
                     eyebrow.upper(), size=10, bold=True, color=_ORANGE)
        top = 0.88
    _add_textbox(slide, Inches(0.45), Inches(top), Inches(12.2), Inches(0.42),
                 title, size=24, bold=True, color=_NAVY)
    _orange_rule(slide, 0.45, top + 0.48, 0.85, height=0.045)
    if subtitle:
        _add_textbox(slide, Inches(0.45), Inches(top + 0.58), Inches(12.2), Inches(0.32),
                     subtitle, size=12, color=_MUTED)


def _insight_banner(slide, text: str, top: float = 1.55) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(top), Inches(12.4), Inches(0.72))
    _fill_shape(shape, _CARD)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(top), Inches(0.08), Inches(0.72))
    _fill_shape(accent, _NAVY)
    _add_textbox(
        slide, Inches(0.7), Inches(top + 0.1), Inches(11.9), Inches(0.55),
        f"Insight Snapshot: {_clean(text, 260)}",
        size=12, bold=False, color=_NAVY,
    )


def _add_table(slide, left, top, width, height, headers: list[str], rows: list[list[str]]) -> None:
    n_rows = 1 + len(rows)
    n_cols = len(headers)
    if n_cols < 1 or n_rows < 1:
        return
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    def _style_cell(cell, text: str, *, header: bool = False, alt: bool = False) -> None:
        # Assign text then restyle the single paragraph run python-pptx creates.
        cell.text = text if text is not None else ""
        p = cell.text_frame.paragraphs[0]
        # Keep one run only
        while len(p.runs) > 1:
            r = p.runs[-1]._r
            r.getparent().remove(r)
        if not p.runs:
            run = p.add_run()
            run.text = cell.text
        else:
            run = p.runs[0]
        _set_run(
            run,
            size=11 if header else 10,
            bold=header,
            color=_WHITE if header else _SLATE,
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = _NAVY if header else (_LIGHT if alt else _WHITE)

    for ci, h in enumerate(headers):
        _style_cell(table.cell(0, ci), str(h), header=True)
    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            value = row[ci] if ci < len(row) else ""
            _style_cell(table.cell(ri + 1, ci), str(value), alt=(ri % 2 == 1))


def _metric_cards(slide, cards: list[tuple[str, str, str]], top: float = 2.0) -> None:
    n = max(1, min(len(cards), 4))
    gap = 0.25
    usable = 12.4 - gap * (n - 1)
    w = usable / n
    for i, (label, value, note) in enumerate(cards):
        left = 0.45 + i * (w + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(1.55))
        _fill_shape(shape, _LIGHT)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 0.15), Inches(w - 0.3), Inches(0.3), label, size=11, bold=True, color=_MUTED)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 0.5), Inches(w - 0.3), Inches(0.5), value, size=18, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(left + 0.15), Inches(top + 1.05), Inches(w - 0.3), Inches(0.35), note, size=10, color=_SLATE)


def _workflow_note(slide, agents: list[dict[str, Any]], top: float = 6.35) -> None:
    names = [_agent_name(a) for a in agents if a][:4]
    if not names:
        return
    _add_textbox(
        slide, Inches(0.45), Inches(top), Inches(12.4), Inches(0.35),
        "Workflow: " + " · ".join(names) + "  ·  Not re-derived in deck UI",
        size=10, color=_MUTED,
    )


# ---------------------------------------------------------------------------
# Slide composers
# ---------------------------------------------------------------------------

def _slide_cover(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), _SLIDE_W, _SLIDE_H)
    _fill_shape(bg, _NAVY_DEEP)
    # Live-style orange corner block
    corner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.6), Inches(0), Inches(1.733), Inches(1.1))
    _fill_shape(corner, _ORANGE)
    company = _company(ctx)
    sector = _sector_display(ctx)
    today = date.today().strftime("%d %B %Y")
    _add_textbox(slide, Inches(0.7), Inches(1.55), Inches(11.5), Inches(0.35),
                 "INVESTMENT DUE DILIGENCE · MARKET INTELLIGENCE", size=11, bold=True, color=_TAN)
    _add_textbox(slide, Inches(0.7), Inches(2.15), Inches(11.5), Inches(0.75),
                 "Market Intel Deck", size=40, bold=True, color=_WHITE)
    _add_textbox(slide, Inches(0.7), Inches(3.0), Inches(11.5), Inches(0.4),
                 f"Investment Due Diligence — {company}", size=16, color=_WHITE)
    _add_textbox(slide, Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.35),
                 "Positioning matrices, market maps and competitive battlecards", size=13, color=_LIGHT)
    _orange_rule(slide, 0.7, 4.05, 4.5, height=0.035)

    meta = [
        ("DOCUMENT", "Market Intelligence Deck"),
        ("TARGET", company if len(company) <= 28 else company[:26] + "…"),
        ("SECTOR", sector),
        ("REPORT DATE", today),
    ]
    for i, (label, value) in enumerate(meta):
        left = 0.7 + i * 3.05
        _add_textbox(slide, Inches(left), Inches(4.35), Inches(2.9), Inches(0.25),
                     label, size=9, bold=True, color=_TAN)
        _add_textbox(slide, Inches(left), Inches(4.65), Inches(2.9), Inches(0.55),
                     value, size=12, bold=True, color=_WHITE)

    _add_textbox(slide, Inches(0.7), Inches(6.55), Inches(10), Inches(0.3),
                 "STRICTLY PRIVATE & CONFIDENTIAL · PREPARED FOR INVESTMENT COMMITTEE",
                 size=10, color=_WHITE)
    _add_textbox(slide, Inches(11.0), Inches(6.55), Inches(1.8), Inches(0.3),
                 f"{page:02d} / {total:02d}", size=11, bold=True, color=_WHITE, align=PP_ALIGN.RIGHT)


def _slide_agenda(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide, "Agenda",
        "Four IC-facing chapters mapped from the 17-section storyline.",
        eyebrow="Document overview",
    )
    parts = [
        ("01", "Key Findings", "Executive snapshot of the deal-level conclusions"),
        ("02", "Market Analysis", "Market definition, sizing, growth, pricing and demand drivers"),
        ("03", "Competitive Landscape", "Competitors, positioning, differentiation and SWOT analysis"),
        ("04", "Summary", "Final position and primary risks for the Investment Committee"),
    ]
    for i, (num, title, blurb) in enumerate(parts):
        top = 1.85 + i * 1.15
        _add_textbox(slide, Inches(0.55), Inches(top), Inches(1.3), Inches(0.55),
                     num, size=28, bold=True, color=_ORANGE)
        _add_textbox(slide, Inches(2.0), Inches(top + 0.05), Inches(10.5), Inches(0.35),
                     title, size=18, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(2.0), Inches(top + 0.45), Inches(10.5), Inches(0.3),
                     blurb, size=12, color=_MUTED)
        if i < len(parts) - 1:
            rule = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(2.0), Inches(top + 0.95), Inches(10.5), Inches(0.012)
            )
            _fill_shape(rule, RGBColor(0xE5, 0xE7, 0xEB))
    _chrome(slide, company=company, section="Agenda", page=page, total=total)


def _slide_divider(prs: Presentation, ctx: BuildContext, number: str, title: str, subtitle: str, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), _SLIDE_W, _SLIDE_H)
    _fill_shape(bg, _NAVY_DEEP)
    _add_textbox(slide, Inches(0.8), Inches(1.9), Inches(11), Inches(0.9),
                 number, size=72, bold=True, color=_DIVIDER_NUM)
    _add_textbox(slide, Inches(0.8), Inches(2.95), Inches(11), Inches(0.35),
                 f"SECTION {number}", size=12, bold=True, color=_WHITE)
    _orange_rule(slide, 0.8, 3.4, 1.2, height=0.06)
    _add_textbox(slide, Inches(0.8), Inches(3.65), Inches(11), Inches(0.55),
                 title, size=32, bold=True, color=_WHITE)
    _add_textbox(slide, Inches(0.8), Inches(4.35), Inches(11), Inches(0.4),
                 subtitle, size=14, color=_LIGHT)
    _add_textbox(slide, Inches(0.8), Inches(6.7), Inches(9), Inches(0.3),
                 "AGENTIC CDD · INVESTMENT DUE DILIGENCE · STRICTLY PRIVATE & CONFIDENTIAL",
                 size=9, color=_LIGHT)
    _add_textbox(slide, Inches(10.5), Inches(6.7), Inches(2.2), Inches(0.3),
                 f"{page:02d} / {total:02d}", size=11, bold=True, color=_WHITE, align=PP_ALIGN.RIGHT)


def _key_finding_cards(ctx: BuildContext) -> list[tuple[str, str]]:
    """Build the live-style 6-card Key Findings grid from agent specs."""
    mdef = _resolve(ctx, "market_definition")
    vol = _resolve(ctx, "market_volume_and_growth")
    strat = _resolve(ctx, "strategic_direction")
    diff = _resolve(ctx, "competitive_differentiation")
    comp = _resolve(ctx, "competitor_identification")
    share = _resolve(ctx, "market_share_strategy")

    framing = _spec(mdef).get("market_framing") or _spec(mdef).get("document")
    if isinstance(framing, list):
        framing = "; ".join(str(x) for x in framing[:2] if x)
    in_scope = _clean(str(framing or ""), 140) or (_findings(mdef, 1) or ["Primary addressable market from market_definition."])[0]

    geos = _as_list(_spec(mdef).get("geographies"), 1)
    if not geos:
        perim = _spec(mdef).get("perimeter") if isinstance(_spec(mdef).get("perimeter"), dict) else {}
        geo = perim.get("geography") if isinstance(perim, dict) else None
        if isinstance(geo, str) and geo.strip() and "Information request" not in geo:
            geos = [_clean(geo, 160)]
    segs = _as_list(_spec(mdef).get("segments"), 1)
    out_scope = (
        f"Adjacent / non-core segments outside {_company(ctx)}'s primary {_sector_display(ctx)} focus are excluded."
        if not geos else
        f"Outside primary geography focus ({geos[0][:120]})."
    )

    drivers = _as_list(_spec(strat).get("investment_drivers"), 1) or _findings(strat, 1)
    rationale = drivers[0] if drivers else f"Focus on high-growth {_sector_display(ctx)} with structural demand."

    verdict = (_findings(vol, 1) or _as_list(_spec(vol).get("growth_notes"), 1) or
               ["Market growth supported by volume and penetration trajectories."])[0]

    strength = (_as_list(_spec(diff).get("differentiators"), 1) or _findings(share, 1) or
                _as_list(_spec(share).get("strategy_notes"), 1) or
                ["Differentiated position from competitive / share agents."])[0]

    threat = (_as_list(_spec(diff).get("moat_signals"), 1) or _findings(comp, 1) or
              ["Competition from established and emerging players."])[0]

    return [
        ("IN-SCOPE DRIVERS", in_scope),
        ("OUT-OF-SCOPE", _clean(out_scope, 140)),
        ("ENGAGEMENT RATIONALE", _clean(str(rationale), 140)),
        ("VERDICT", _clean(str(verdict), 140)),
        ("RELATIVE STRENGTH", _clean(str(strength), 140)),
        ("PRIMARY COMP. THREAT", _clean(str(threat), 140)),
    ]


def _slide_key_findings(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide,
        f"Key Findings — {company if len(company) < 40 else 'Target'}",
        "Executive snapshot of in-scope drivers, verdict and primary competitive threat",
        eyebrow="Section 01",
    )
    cards = _key_finding_cards(ctx)
    # 2×3 grid
    for i, (label, body) in enumerate(cards[:6]):
        col, row = i % 3, i // 3
        left = 0.45 + col * 4.2
        top = 1.85 + row * 2.35
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(4.0), Inches(2.15))
        _fill_shape(card, _WHITE)
        card.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(2.15))
        _fill_shape(accent, _NAVY)
        _add_textbox(slide, Inches(left + 0.25), Inches(top + 0.15), Inches(0.4), Inches(0.35),
                     str(i + 1), size=16, bold=True, color=_ORANGE)
        _add_textbox(slide, Inches(left + 0.65), Inches(top + 0.2), Inches(3.1), Inches(0.3),
                     label, size=11, bold=True, color=_NAVY)
        _add_textbox(slide, Inches(left + 0.25), Inches(top + 0.65), Inches(3.55), Inches(1.3),
                     body, size=11, color=_SLATE)
    _chrome(slide, company=company, section="Key Findings", page=page, total=total, part="01")


def _slide_key_findings_workflow(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    """Live slide 04: diligence agents' read under Key Findings."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide,
        "Key Findings — the diligence agents' read",
        "Conclusions carried over from this deal's completed workflow analysis.",
        eyebrow="Workflow analysis",
    )
    agents = [
        _resolve(ctx, "market_definition"),
        _resolve(ctx, "market_volume_and_growth"),
        _resolve(ctx, "competitive_differentiation"),
        _resolve(ctx, "swot_analysis"),
        _resolve(ctx, "competitor_identification"),
        _resolve(ctx, "strategic_direction"),
    ]
    bullets: list[str] = []
    for a in agents:
        if not a:
            continue
        name = _agent_name(a)
        findings = _findings(a, 1)
        if findings:
            bullets.append(f"{findings[0]} (from the {name} agent)")
        if len(bullets) >= 5:
            break
    if not bullets:
        bullets = ["Workflow outputs available but no narrative findings extracted."]

    _add_textbox(slide, Inches(0.55), Inches(1.75), Inches(12), Inches(0.3),
                 "DILIGENCE WORKFLOW READ", size=12, bold=True, color=_ORANGE)
    # Orange-dash style bullets
    box = slide.shapes.add_textbox(Inches(0.55), Inches(2.15), Inches(12.2), Inches(4.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets[:5]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        _write_paragraph(p, f"—  {item}", size=13, color=_SLATE)

    sources: list[str] = []
    for a in agents:
        for s in list(a.get("sources") or [])[:2]:
            if isinstance(s, str) and s not in sources:
                sources.append(s)
    if sources:
        _add_textbox(slide, Inches(0.55), Inches(6.45), Inches(12.2), Inches(0.35),
                     "Source: " + ", ".join(sources[:5]), size=9, color=_MUTED)
    _chrome(slide, company=company, section="Key Findings", page=page, total=total, part="01")


def _slide_perimeter(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide,
        "2.1 MARKET DEFINITION: Market perimeter & definition",
        f"Defining the boundaries of where {company} operates and competes.",
        eyebrow="Section 02",
    )
    mdef = _resolve(ctx, "market_definition")
    buy = _resolve(ctx, "buying_behavior")
    seg = _resolve(ctx, "customer_segmentation")
    bg = _resolve(ctx, "company_background")
    spec = _spec(mdef)
    framing = spec.get("market_framing") or spec.get("document")
    if isinstance(framing, list):
        framing_text = "; ".join(str(x) for x in framing[:3] if x)
    else:
        framing_text = str(framing) if framing else ""
    perim = spec.get("perimeter") if isinstance(spec.get("perimeter"), dict) else {}
    service = str(perim.get("service") or "").strip()
    if service.startswith("Information") or "Foundation ·" in service or "plan relies" in service.lower():
        service = ""
    # Prefer company-background sells when market_definition only has a generic DOC cite
    bg_offers = _spec(bg).get("offers") if isinstance(_spec(bg).get("offers"), dict) else {}
    bg_sells = str(bg_offers.get("sells") or "").strip()
    if bg_sells and (
        not service
        or re.search(r"(?i)DOC:\s*data room financials", service)
    ):
        cites = ""
        bg_sources = bg.get("sources") if isinstance(bg.get("sources"), list) else []
        if bg_sources:
            n = min(3, len(bg_sources))
            cites = " (" + ", ".join(f"DOC: [{i}]" for i in range(1, n + 1)) + ")"
        service = f"{bg_sells}{cites}"
    elif service:
        service = _clean(service, 280, sources=mdef.get("sources") if isinstance(mdef.get("sources"), list) else None)
    insight = (
        service
        or framing_text.strip()
        or (_findings(mdef, 1) or [f"Market perimeter for {company}."])[0]
    )
    if "Foundation ·" in insight or "plan relies" in insight.lower():
        insight = f"Organics collection / composting footprint for {company}."
    _insight_banner(slide, insight, top=1.7)

    segs = _as_list(spec.get("segments"), 2) or _as_list(_spec(seg).get("segments"), 2)
    segs = [s for s in segs if s and "Foundation ·" not in str(s) and "plan relies" not in str(s).lower()]
    geos = _as_list(spec.get("geographies"), 2)
    if not geos:
        geo = perim.get("geography")
        if isinstance(geo, str) and geo.strip() and "Information request" not in geo:
            geos = [_clean(geo, 200)]
    macros = _as_list(spec.get("macro_drivers"), 2) or _as_list(spec.get("policy_context"), 2)
    primary = (
        _clean(service, 140)
        if service
        else (segs[0] if segs else _clean(insight, 120))
    )
    rows = [
        ["Primary Market", primary, "market_definition"],
        ["Adjacent / segments", segs[0] if segs else (_sector_display(ctx) + " adjacencies"), "market_definition"],
        ["Geographic Focus", geos[0] if geos else "As framed in market_definition / segmentation", "market_definition"],
        ["In-Scope Drivers", macros[0] if macros else (_findings(buy, 1) or ["Structural demand drivers from agents."])[0], "market_definition"],
        ["Out-of-Scope", f"Non-core segments outside {_sector_display(ctx)} primary focus", "market_definition"],
    ]
    _add_table(
        slide, Inches(0.45), Inches(2.6), Inches(12.4), Inches(3.6),
        ["Parameter", "Scope", "Source"],
        rows,
    )
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [mdef, buy, seg], top=6.55)


def _slide_tam(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.2 TAM / SAM / SOM estimation", "Structured sizing from market_volume_and_growth.")
    vol = _resolve(ctx, "market_volume_and_growth")
    vspec = _spec(vol)
    cards = []
    for key, label in (("tam", "TAM"), ("sam", "SAM"), ("som", "SOM")):
        block = vspec.get(key)
        note = "Combined Residential+Commercial TOTAL · Current"
        if isinstance(block, dict):
            raw = str(block.get("raw") or "")
            if "TOTAL" in raw.upper() or "Combined" in raw:
                note = "Combined Residential+Commercial TOTAL · Current"
            elif raw:
                note = _clean(raw, 60)
        cards.append((label, _fmt_metric(block), note))
    if not any(vspec.get(k) is not None for k in ("tam", "sam", "som")):
        cards = [("TAM", "—", "Not populated"), ("SAM", "—", "Not populated"), ("SOM", "—", "Not populated")]
    _metric_cards(slide, cards, top=1.85)

    # Shares
    ratios: list[str] = []
    tam_v = vspec.get("tam", {}).get("value") if isinstance(vspec.get("tam"), dict) else None
    sam_v = vspec.get("sam", {}).get("value") if isinstance(vspec.get("sam"), dict) else None
    som_v = vspec.get("som", {}).get("value") if isinstance(vspec.get("som"), dict) else None
    if isinstance(tam_v, (int, float)) and not isinstance(tam_v, bool) and tam_v and isinstance(sam_v, (int, float)) and not isinstance(sam_v, bool):
        ratios.append(f"SAM / TAM share: {float(sam_v) / float(tam_v):.1%}")
    if isinstance(sam_v, (int, float)) and not isinstance(sam_v, bool) and sam_v and isinstance(som_v, (int, float)) and not isinstance(som_v, bool):
        ratios.append(f"SOM / SAM capture: {float(som_v) / float(sam_v):.1%}")
    cagr = vspec.get("cagr_pct")
    if cagr is not None:
        ratios.append(f"CAGR observations: {_fmt_metric(cagr)}")
    pen = vspec.get("penetration_pct")
    if pen is not None:
        ratios.append(f"Penetration: {_fmt_metric(pen)}")
    notes = _as_list(vspec.get("growth_notes"), 4) or _as_list(vspec.get("unit_volume_notes"), 4) or _findings(vol, 4)
    _add_textbox(slide, Inches(0.55), Inches(3.15), Inches(12), Inches(0.3), "Derived shares & growth", size=13, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(0.55), Inches(3.5), Inches(12.2), Inches(2.5), (ratios + notes)[:8] or ["No sizing notes available."], size=13)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [vol, _resolve(ctx, "market_definition")])


def _slide_taxonomy(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.3 Market taxonomy, segmentation & boundaries", "")
    mdef = _resolve(ctx, "market_definition")
    seg = _resolve(ctx, "customer_segmentation")
    buy = _resolve(ctx, "buying_behavior")
    rows: list[list[str]] = []
    segments = _spec(seg).get("segments")
    if isinstance(segments, list):
        for s in segments[:8]:
            if isinstance(s, dict):
                rows.append([
                    str(s.get("name") or "—"),
                    f"{s.get('share_pct', '—')}%",
                    str(s.get("avg_age") or "—"),
                    _clean(str(s.get("use_case") or s.get("key_driver") or "—"), 80),
                ])
    if rows:
        _add_table(slide, Inches(0.45), Inches(1.25), Inches(12.4), Inches(3.6),
                   ["Segment", "Share", "Avg age", "Notes"], rows)
    else:
        bits = _as_list(_spec(mdef).get("segments"), 6) or _findings(seg, 6) or _findings(mdef, 5)
        _add_bullets(slide, Inches(0.55), Inches(1.4), Inches(12), Inches(4.5), bits or ["Segmentation not populated."], size=13)
    hhi = _spec(buy).get("segment_hhi")
    if hhi is not None:
        _add_textbox(slide, Inches(0.55), Inches(5.9), Inches(12), Inches(0.3), f"Segment HHI: {hhi}", size=12, bold=True, color=_NAVY)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [mdef, seg, buy])


def _slide_trajectory(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.4 Historical trajectory & growth rate", "")
    vol = _resolve(ctx, "market_volume_and_growth")
    growth = _resolve(ctx, "growth_opportunities")
    hist = _resolve(ctx, "historical_performance")
    share = _resolve(ctx, "market_share_strategy")
    insight = (_findings(vol, 1) or _findings(growth, 1) or [f"Growth trajectory for {company}."])[0]
    _insight_banner(slide, insight)
    left = _as_list(_spec(vol).get("growth_notes"), 5) or _findings(vol, 5)
    right = _as_list(_spec(growth).get("growth_levers"), 5) or _as_list(_spec(growth).get("market_growth"), 4) or _findings(growth, 4)
    right += _as_list(_spec(share).get("share_trends"), 3)
    _add_textbox(slide, Inches(0.55), Inches(2.0), Inches(6), Inches(0.3), "Volume / growth notes", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(0.55), Inches(2.35), Inches(6), Inches(2.2), left[:5] or ["—"], size=12)
    _add_textbox(slide, Inches(6.9), Inches(2.0), Inches(6), Inches(0.3), "Growth levers / share", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(6.9), Inches(2.35), Inches(5.8), Inches(2.2), right[:5] or _findings(hist, 3) or ["—"], size=12)

    # G1 — company financial history from released databook only (never silent agent finals).
    try:
        from agetic_cdd_api.services_databook_consume import released_pl_display_rows

        hist_pl = _spec(hist).get("pl_lines") if isinstance(_spec(hist).get("pl_lines"), list) else []
        fin_rows, _, footnote = released_pl_display_rows(ctx.deal_slug, hist_pl, limit=5)
        if fin_rows:
            _add_textbox(
                slide,
                Inches(0.55),
                Inches(4.55),
                Inches(12),
                Inches(0.25),
                "Company financials (released databook)",
                size=11,
                bold=True,
                color=_NAVY,
            )
            table_rows = [[r[0], r[1][:48], r[3]] for r in fin_rows]
            _add_table(
                slide,
                Inches(0.45),
                Inches(4.85),
                Inches(12.4),
                Inches(1.35),
                ["Line", "FY values", "Status"],
                table_rows,
            )
            if footnote:
                _add_textbox(
                    slide,
                    Inches(0.55),
                    Inches(6.25),
                    Inches(12),
                    Inches(0.25),
                    footnote,
                    size=10,
                    color=_NAVY,
                )
    except Exception:
        pass

    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [vol, growth, hist, share])


def _slide_cagr(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.5 CAGR benchmarking & market lifecycle", "")
    vol = _resolve(ctx, "market_volume_and_growth")
    vspec = _spec(vol)
    cagr = vspec.get("cagr_pct")
    cards = [("CAGR %", _fmt_metric(cagr) if cagr is not None else "—", "market_volume_and_growth")]
    if vspec.get("penetration_pct") is not None:
        cards.append(("Penetration", _fmt_metric(vspec.get("penetration_pct")), "observed / projected"))
    cards.append(("Lifecycle read", "Growth / expansion" if cagr else "—", "from CAGR & notes"))
    _metric_cards(slide, cards[:3], top=1.25)
    bits = _as_list(vspec.get("unit_volume_notes"), 5) or _findings(vol, 5)
    bits += _findings(_resolve(ctx, "growth_opportunities"), 3)
    _add_bullets(slide, Inches(0.55), Inches(3.2), Inches(12.2), Inches(2.8), bits[:8] or ["No CAGR notes available."], size=13)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [vol])


def _slide_pricing(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.6 Pricing power, competitive position & margin", "")
    pricing = _resolve(ctx, "market_pricing")
    diff = _resolve(ctx, "competitive_differentiation")
    pspec = _spec(pricing)
    points = pspec.get("price_points")
    rows: list[list[str]] = []
    if isinstance(points, list):
        for pt in points[:6]:
            if isinstance(pt, dict):
                rows.append([
                    str(pt.get("label") or "—"),
                    _fmt_metric(pt),
                    str(pt.get("unit") or "—"),
                    _clean(str(pt.get("raw") or ""), 70),
                ])
    if rows:
        _add_table(slide, Inches(0.45), Inches(1.2), Inches(12.4), Inches(2.4),
                   ["Price point", "Value", "Unit", "Source raw"], rows)
    bullets = _as_list(pspec.get("pricing_notes"), 4) + _as_list(pspec.get("power_signals"), 3)
    bullets += _as_list(_spec(diff).get("differentiators"), 3)
    top = 3.9 if rows else 1.4
    _add_bullets(slide, Inches(0.55), Inches(top), Inches(12.2), Inches(2.8), bullets[:8] or _findings(pricing, 5) or ["Pricing notes not populated."], size=13)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [pricing, diff])


def _slide_demand(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.7 Macro & structural demand drivers", "")
    demand = _resolve(ctx, "demand_drivers")
    risk = _resolve(ctx, "market_risk")
    dspec = _spec(demand)
    drivers = _as_list(dspec.get("demand_drivers"), 6) or _findings(demand, 6)
    regional = _as_list(dspec.get("regional_signals"), 4)
    cycles = _as_list(dspec.get("cycle_risks"), 4) or _findings(risk, 3)
    insight = drivers[0] if drivers else f"Demand backdrop for {company}."
    _insight_banner(slide, insight)
    _add_textbox(slide, Inches(0.55), Inches(2.0), Inches(6), Inches(0.3), "Structural drivers", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(0.55), Inches(2.35), Inches(6), Inches(3.5), drivers[:6], size=12)
    _add_textbox(slide, Inches(6.9), Inches(2.0), Inches(6), Inches(0.3), "Regional / cycle risks", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(6.9), Inches(2.35), Inches(5.8), Inches(3.5), (regional + cycles)[:6] or ["—"], size=12)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [demand, risk])


def _slide_sustainability(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "2.8 Sustainability, cyclicality & demand verdict", "")
    demand = _resolve(ctx, "demand_drivers")
    esg = _resolve(ctx, "esg_and_sustainability")
    rev = _resolve(ctx, "revenue_quality")
    bits = _as_list(_spec(demand).get("cycle_risks"), 4) + _findings(demand, 3)
    bits += _as_list(_spec(esg).get("themes"), 4) or _findings(esg, 3)
    bits += _findings(rev, 3)
    _insight_banner(slide, bits[0] if bits else f"Demand durability read for {company}.")
    _add_bullets(slide, Inches(0.55), Inches(2.0), Inches(12.2), Inches(4.0), bits[:9] or ["No sustainability / cyclicality notes populated."], size=13)
    _chrome(slide, company=company, section="Market Analysis", page=page, total=total, part="02")
    _workflow_note(slide, [demand, esg, rev])


def _slide_competitors(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.1 Direct & adjacent competitor universe", "")
    comp = _resolve(ctx, "competitor_identification")
    cspec = _spec(comp)
    competitors = cspec.get("competitors")
    rows: list[list[str]] = []
    if isinstance(competitors, list):
        for c in competitors[:8]:
            if isinstance(c, dict):
                rows.append([
                    str(c.get("name") or "—"),
                    f"{c.get('market_share_pct', '—')}%",
                    str(c.get("units_000s") or "—"),
                    str(c.get("flagship_price_inr") or "—"),
                    str(c.get("trend") or "—"),
                    str(c.get("nps") or "—"),
                ])
    if rows:
        _add_table(slide, Inches(0.35), Inches(1.2), Inches(12.6), Inches(4.2),
                   ["Competitor", "Share", "Units (000s)", "Flagship INR", "Trend", "NPS"], rows)
    else:
        _add_bullets(slide, Inches(0.55), Inches(1.4), Inches(12), Inches(4.5),
                     _as_list(cspec.get("positioning_notes"), 6) or _findings(comp, 6) or ["Competitor set not populated."], size=13)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [comp])


def _slide_positioning(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.2 Competitive positioning matrix", "Share vs price / NPS proxy grid from competitor_identification.")
    comp = _resolve(ctx, "competitor_identification")
    competitors = _spec(comp).get("competitors")
    rows: list[list[str]] = []
    if isinstance(competitors, list):
        for c in competitors[:8]:
            if isinstance(c, dict):
                share = c.get("market_share_pct")
                price = c.get("flagship_price_inr")
                nps = c.get("nps")
                soft = c.get("software")
                rows.append([
                    str(c.get("name") or "—"),
                    f"{share}%" if share is not None else "—",
                    str(price or "—"),
                    str(nps or "—"),
                    str(soft or "—"),
                    str(c.get("trend") or "—"),
                ])
    if rows:
        _add_table(slide, Inches(0.35), Inches(1.2), Inches(12.6), Inches(4.4),
                   ["Player", "Share", "Flagship", "NPS", "Software", "Trend"], rows)
    else:
        _add_bullets(slide, Inches(0.55), Inches(1.4), Inches(12), Inches(4.5),
                     _findings(comp, 6) or ["Positioning matrix inputs missing."], size=13)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [comp, _resolve(ctx, "competitive_differentiation")])


def _slide_capability(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.3 Competitive capability matrix & barriers", "")
    diff = _resolve(ctx, "competitive_differentiation")
    risk = _resolve(ctx, "market_risk")
    dspec = _spec(diff)
    left = _as_list(dspec.get("differentiators"), 5) or _findings(diff, 4)
    right = _as_list(dspec.get("feature_gaps"), 4) + _as_list(dspec.get("moat_signals"), 3)
    right += _findings(risk, 3)
    _add_textbox(slide, Inches(0.55), Inches(1.25), Inches(6), Inches(0.3), "Capabilities / differentiators", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(0.55), Inches(1.6), Inches(6), Inches(4.3), left[:7] or ["—"], size=12)
    _add_textbox(slide, Inches(6.9), Inches(1.25), Inches(6), Inches(0.3), "Gaps / barriers / moat signals", size=12, bold=True, color=_NAVY)
    _add_bullets(slide, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.3), right[:7] or ["—"], size=12)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [diff, risk])


def _slide_moat(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.4 Core strategic differentiators & moat", "")
    diff = _resolve(ctx, "competitive_differentiation")
    ip = _resolve(ctx, "ip_and_technology")
    strat = _resolve(ctx, "strategic_direction")
    bits = _as_list(_spec(diff).get("moat_signals"), 4) + _as_list(_spec(diff).get("differentiators"), 4)
    bits += _findings(ip, 3) + _as_list(_spec(strat).get("investment_drivers"), 3)
    _insight_banner(slide, bits[0] if bits else f"Moat read for {company}.")
    _add_bullets(slide, Inches(0.55), Inches(2.0), Inches(12.2), Inches(4.0), bits[:9] or ["Moat signals not populated."], size=13)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [diff, ip, strat])


def _slide_battlecards(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.5 Competitive battlecards", "Top competitors from competitor_identification.")
    comp = _resolve(ctx, "competitor_identification")
    share = _resolve(ctx, "market_share_strategy")
    competitors = _spec(comp).get("competitors")
    cards: list[dict[str, Any]] = []
    if isinstance(competitors, list):
        cards = [c for c in competitors if isinstance(c, dict)][:3]
    wins = _as_list(_spec(share).get("wins"), 2)
    losses = _as_list(_spec(share).get("losses"), 2)
    if not cards:
        _add_bullets(slide, Inches(0.55), Inches(1.4), Inches(12), Inches(4.5),
                     _findings(comp, 5) or ["Battlecard inputs missing."], size=13)
    else:
        w = 3.9
        for i, c in enumerate(cards):
            left = 0.45 + i * (w + 0.25)
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.25), Inches(w), Inches(4.7))
            _fill_shape(shape, _LIGHT)
            name = str(c.get("name") or f"Competitor {i+1}")
            _add_textbox(slide, Inches(left + 0.15), Inches(1.4), Inches(w - 0.3), Inches(0.4), name, size=16, bold=True, color=_NAVY)
            lines = [
                f"Share: {c.get('market_share_pct', '—')}%",
                f"Units (000s): {c.get('units_000s', '—')}",
                f"Flagship: {c.get('flagship_price_inr', '—')}",
                f"Trend: {c.get('trend', '—')}",
                f"NPS: {c.get('nps', '—')}",
                f"Software: {c.get('software', '—')}",
            ]
            if i == 0 and wins:
                lines.append(f"Win note: {wins[0]}")
            if i == 1 and losses:
                lines.append(f"Loss note: {losses[0]}")
            _add_bullets(slide, Inches(left + 0.15), Inches(1.95), Inches(w - 0.3), Inches(3.7), lines, size=11)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [comp, share, _resolve(ctx, "competitive_differentiation")])


def _slide_diff_verdict(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.6 Differentiation verdict & market share", "")
    share = _resolve(ctx, "market_share_strategy")
    diff = _resolve(ctx, "competitive_differentiation")
    growth = _resolve(ctx, "growth_opportunities")
    sspec = _spec(share)
    insight = (_findings(share, 1) or _as_list(sspec.get("strategy_notes"), 1) or [f"Share strategy for {company}."])[0]
    _insight_banner(slide, insight)
    bits = _as_list(sspec.get("strategy_notes"), 4) + _as_list(sspec.get("share_trends"), 3)
    bits += _as_list(sspec.get("wins"), 2) + _as_list(sspec.get("losses"), 2)
    bits += _as_list(_spec(diff).get("moat_signals"), 2) + _findings(growth, 2)
    _add_bullets(slide, Inches(0.55), Inches(2.0), Inches(12.2), Inches(4.0), bits[:9] or _findings(share, 6) or ["—"], size=13)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [share, diff, growth])


def _slide_swot(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(slide, "3.7 SWOT summary matrix", "")
    swot = _resolve(ctx, "swot_analysis")
    sspec = _spec(swot)
    quadrants = [
        ("Strengths", _as_list(sspec.get("strengths"), 4) or _findings(swot, 2), _GREEN),
        ("Weaknesses", _as_list(sspec.get("weaknesses"), 4), _AMBER),
        ("Opportunities", _as_list(sspec.get("opportunities"), 4), _NAVY),
        ("Threats", _as_list(sspec.get("threats"), 4), _RED),
    ]
    positions = [(0.45, 1.2), (6.85, 1.2), (0.45, 4.0), (6.85, 4.0)]
    for (title, items, color), (left, top) in zip(quadrants, positions):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(6.0), Inches(2.55))
        _fill_shape(shape, _LIGHT)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.12), Inches(2.55))
        _fill_shape(accent, color)
        _add_textbox(slide, Inches(left + 0.25), Inches(top + 0.12), Inches(5.5), Inches(0.3), title, size=14, bold=True, color=color)
        _add_bullets(slide, Inches(left + 0.25), Inches(top + 0.5), Inches(5.5), Inches(1.9), items[:4] or ["—"], size=11)
    _chrome(slide, company=company, section="Competitive Landscape", page=page, total=total, part="03")
    _workflow_note(slide, [swot], top=6.55)


def _slide_summary(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide,
        "Market Intel Deck — Summary",
        "Final consolidated position for the Investment Committee.",
        eyebrow="Closing summary",
    )
    cards = _key_finding_cards(ctx)
    share = _resolve(ctx, "market_share_strategy")
    bg = _resolve(ctx, "company_background")
    rec_bits = _findings(share, 1) or _findings(bg, 1) or [
        f"Underwrite {_sector_display(ctx)} ceiling and competitive intensity; pair with IC Memo for go/no-go."
    ]

    deal_rows = [
        ("Market Definition", cards[0][1]),
        ("Market Verdict", cards[3][1]),
        ("Relative Strength", cards[4][1]),
        ("Primary Comp. Threat", cards[5][1]),
        ("Recommendation", _clean(rec_bits[0], 160)),
    ]
    _add_textbox(slide, Inches(0.55), Inches(1.75), Inches(12), Inches(0.3),
                 "DEAL POSITION", size=12, bold=True, color=_ORANGE)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(2.15), Inches(12.2), Inches(3.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (label, body) in enumerate(deal_rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        p.clear()
        run_l = p.add_run()
        run_l.text = f"—  {label}: "
        _set_run(run_l, size=13, bold=True, color=_NAVY)
        run_b = p.add_run()
        run_b.text = body
        _set_run(run_b, size=13, bold=False, color=_SLATE)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(5.85), Inches(12.4), Inches(0.55))
    _fill_shape(bar, _NAVY)
    _add_textbox(
        slide, Inches(0.55), Inches(5.95), Inches(12.2), Inches(0.35),
        f"END OF MARKET INTEL DECK · {company.upper() if company.isascii() else company} · STRICTLY PRIVATE & CONFIDENTIAL",
        size=11, bold=True, color=_WHITE, align=PP_ALIGN.CENTER,
    )
    _chrome(slide, company=company, section="Summary", page=page, total=total, part="04")


def _slide_appendix(prs: Presentation, ctx: BuildContext, page: int, total: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    company = _company(ctx)
    _section_title(
        slide,
        "Appendix — Methodology & sources",
        "Market figures from agents; company financials from released databook only (G1).",
    )
    rows: list[list[str]] = []
    keys = [
        "market_definition", "market_volume_and_growth", "market_pricing", "demand_drivers",
        "competitor_identification", "competitive_differentiation", "market_share_strategy",
        "swot_analysis", "strategic_direction", "customer_segmentation", "buying_behavior",
    ]
    for key in keys:
        agent = _resolve(ctx, key)
        if not agent:
            continue
        sources = agent.get("sources") or _spec(agent).get("sources") or []
        src = ", ".join(s for s in sources if isinstance(s, str))[:70] if sources else "—"
        code = _spec(agent).get("dd_code") or _spec(agent).get("role_code") or "—"
        rows.append([_agent_name(agent, key), str(code), key, src or "—"])
    if rows:
        _add_table(slide, Inches(0.35), Inches(1.2), Inches(12.6), Inches(3.6),
                   ["Agent", "Code", "Slug", "Sources"], rows[:8])
    try:
        from agetic_cdd_api.services_databook_consume import released_pl_display_rows

        hist = _resolve(ctx, "historical_performance")
        hist_pl = _spec(hist).get("pl_lines") if isinstance(_spec(hist).get("pl_lines"), list) else []
        fin_rows, _, footnote = released_pl_display_rows(ctx.deal_slug, hist_pl, limit=6)
        if fin_rows:
            _add_textbox(
                slide,
                Inches(0.45),
                Inches(4.95),
                Inches(12),
                Inches(0.25),
                "Released databook — company financial history",
                size=11,
                bold=True,
                color=_NAVY,
            )
            _add_table(
                slide,
                Inches(0.35),
                Inches(5.25),
                Inches(12.6),
                Inches(1.3),
                ["Line", "FY values", "Unit", "Status"],
                fin_rows,
            )
            if footnote:
                _add_textbox(
                    slide,
                    Inches(0.45),
                    Inches(6.55),
                    Inches(12),
                    Inches(0.2),
                    footnote,
                    size=10,
                    color=_NAVY,
                )
    except Exception:
        pass
    _chrome(slide, company=company, section="Appendix", page=page, total=total, part="04")
    _workflow_note(slide, [_resolve(ctx, k) for k in keys[:4]])


# Storyline title → composer (content sections only; cover/agenda/dividers handled in export)
_COMPOSERS: dict[str, Callable[..., None]] = {
    "Key Findings": _slide_key_findings,
    "Market perimeter & definition": _slide_perimeter,
    "TAM / SAM / SOM estimation": _slide_tam,
    "Market taxonomy, segmentation & boundaries": _slide_taxonomy,
    "Historical trajectory & growth rate": _slide_trajectory,
    "CAGR benchmarking & market lifecycle": _slide_cagr,
    "Pricing power, competitive position & margin": _slide_pricing,
    "Macro & structural demand drivers": _slide_demand,
    "Sustainability, cyclicality & demand verdict": _slide_sustainability,
    "Direct & adjacent competitor universe": _slide_competitors,
    "Competitive positioning matrix": _slide_positioning,
    "Competitive capability matrix & barriers": _slide_capability,
    "Core strategic differentiators & moat": _slide_moat,
    "Competitive battlecards": _slide_battlecards,
    "Differentiation verdict & market share": _slide_diff_verdict,
    "SWOT summary matrix": _slide_swot,
    "Market Intel Deck - Summary": _slide_summary,
}

_MARKET_TITLES = {
    "Market perimeter & definition",
    "TAM / SAM / SOM estimation",
    "Market taxonomy, segmentation & boundaries",
    "Historical trajectory & growth rate",
    "CAGR benchmarking & market lifecycle",
    "Pricing power, competitive position & margin",
    "Macro & structural demand drivers",
    "Sustainability, cyclicality & demand verdict",
}
_COMP_TITLES = {
    "Direct & adjacent competitor universe",
    "Competitive positioning matrix",
    "Competitive capability matrix & barriers",
    "Core strategic differentiators & moat",
    "Competitive battlecards",
    "Differentiation verdict & market share",
    "SWOT summary matrix",
}


def _estimate_slide_count(sections: list[StorylineSection]) -> int:
    """Cover + agenda + key findings pair + dividers + storyline + appendix ≈ live 25."""
    enabled = [s for s in sections if s.included]
    titles = [s.title for s in enabled]
    market_n = sum(1 for t in titles if t in _MARKET_TITLES)
    comp_n = sum(1 for t in titles if t in _COMP_TITLES)
    key = 2 if any(t == "Key Findings" for t in titles) else 0  # cards + workflow read
    summary = 1 if any("Summary" in t for t in titles) else 0
    # Part dividers: 01 Key Findings, 02 Market, 03 Competitive, 04 Summary
    dividers = (1 if key else 0) + (1 if market_n else 0) + (1 if comp_n else 0) + (1 if summary else 0)
    return 2 + key + dividers + market_n + comp_n + summary + 1  # cover/agenda + appendix


class MarketDeckBuilder(ReportBuilder):
    report_type = "market_deck"

    def build_sections(self, ctx: BuildContext) -> Iterator[dict[str, Any]]:
        enabled = [s for s in (ctx.storyline or []) if s.included]
        total = max(1, len(enabled) + 5)  # cover/agenda/dividers/appendix approx
        current = 0
        for label in ("cover", "agenda", "part dividers"):
            current += 1
            yield _evt("section", f"Generating {label}", i=current, total=total)
        for section in enabled:
            current += 1
            yield _evt(
                "section",
                f"Generating slide {current} of {total}: {section.title}",
                i=current,
                total=total,
            )

    def export_artifact(self, ctx: BuildContext) -> str:
        prs = Presentation()
        prs.slide_width = Emu(_SLIDE_W)
        prs.slide_height = Emu(_SLIDE_H)

        enabled = [s for s in (ctx.storyline or []) if s.included]
        total = _estimate_slide_count(enabled)

        page = 0

        def next_page() -> int:
            nonlocal page
            page += 1
            return page

        _slide_cover(prs, ctx, next_page(), total)
        _slide_agenda(prs, ctx, next_page(), total)

        key_divider_done = False
        market_divider_done = False
        comp_divider_done = False
        summary_divider_done = False
        key_workflow_done = False

        for section in enabled:
            title = section.title
            if title == "Key Findings" and not key_divider_done:
                _slide_divider(
                    prs, ctx, "01", "Key Findings",
                    "Executive snapshot of the deal-level conclusions",
                    next_page(), total,
                )
                key_divider_done = True
            if title in _MARKET_TITLES and not market_divider_done:
                _slide_divider(
                    prs, ctx, "02", "Market Analysis",
                    "Market definition, sizing, growth, pricing and demand drivers",
                    next_page(), total,
                )
                market_divider_done = True
            if title in _COMP_TITLES and not comp_divider_done:
                _slide_divider(
                    prs, ctx, "03", "Competitive Landscape",
                    "Competitors, positioning, differentiation and SWOT analysis",
                    next_page(), total,
                )
                comp_divider_done = True
            if "Summary" in title and not summary_divider_done:
                _slide_divider(
                    prs, ctx, "04", "Summary",
                    "Final position and primary risks for the Investment Committee",
                    next_page(), total,
                )
                summary_divider_done = True

            composer = _COMPOSERS.get(title)
            if composer:
                composer(prs, ctx, next_page(), total)
                if title == "Key Findings" and not key_workflow_done:
                    _slide_key_findings_workflow(prs, ctx, next_page(), total)
                    key_workflow_done = True
            else:
                # Fallback generic findings slide
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                pg = next_page()
                _section_title(slide, title, "")
                agents = [_resolve(ctx, ak) for ak in section.agents]
                bullets: list[str] = []
                for a in agents:
                    bullets.extend(_findings(a, 3))
                _add_bullets(slide, Inches(0.55), Inches(1.7), Inches(12), Inches(4.8),
                             bullets[:10] or ["No findings available for this section."], size=13)
                _chrome(slide, company=_company(ctx), section=title[:40], page=pg, total=total)

        _slide_appendix(prs, ctx, next_page(), total)

        out_dir = report_artifact_dir(ctx.deal_slug, "market_deck")
        company = _company(ctx)
        safe = re.sub(r"[^\w\s-]", "", company).strip().replace(" ", "_") or ctx.deal_slug
        filename = f"{safe}_Market_Intel_Deck.pptx"
        path = out_dir / filename
        prs.save(str(path))

        deal_root = deals_root() / ctx.deal_slug
        try:
            return str(path.relative_to(deal_root))
        except ValueError:
            return str(path)


register_builder("market_deck", MarketDeckBuilder)
