"""Extract text and tables from VDR files.

Phase 2 adds per-page metadata for the page register. Full OCR / vision remain
out of scope — scanned and image-only pages are surfaced as unread instead.
"""

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path
from typing import Any

MAX_TEXT_CHARS = 80_000
MAX_TABLE_ROWS = 80
MAX_TABLES = 8
MAX_XLSX_SHEETS = 40  # S1-6: each sheet is its own table (raised from MAX_TABLES)
MAX_PDF_PAGES = 40


def _clip(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    cleaned = " ".join(text.split())
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[: limit - 1].rstrip() + "…"


def extract_file(path: Path) -> dict:
    """Return {text, tables, pages, page_count, parser, error} for a document on disk."""
    suffix = path.suffix.lower().lstrip(".")
    if not path.is_file():
        return {
            "text": "",
            "tables": [],
            "pages": [],
            "page_count": 0,
            "parser": "missing",
            "error": "File not found",
        }
    try:
        if suffix in {"txt", "md", "log"}:
            return _extract_text_file(path)
        if suffix == "csv":
            return _extract_csv(path)
        if suffix == "pdf":
            return _extract_pdf(path)
        if suffix == "docx":
            return _extract_docx(path)
        if suffix in {"xlsx", "xls"}:
            return _extract_xlsx(path)
        if suffix == "pptx":
            return _extract_pptx(path)
        if suffix in {"png", "jpg", "jpeg", "gif", "webp", "tif", "tiff", "bmp"}:
            return {
                "text": "",
                "tables": [],
                "pages": [
                    {
                        "page": 1,
                        "kind": "image",
                        "char_count": 0,
                        "has_images": True,
                        "unread": True,
                        "ocr_status": "not_available",
                        "notes": ["Standalone image — OCR not available"],
                    }
                ],
                "page_count": 1,
                "parser": "image",
                "error": "Image file — OCR not available",
            }
        return {
            "text": "",
            "tables": [],
            "pages": [],
            "page_count": 0,
            "parser": "unsupported",
            "error": f"No extractor for .{suffix}" if suffix else "Unknown format",
        }
    except Exception as exc:
        return {
            "text": "",
            "tables": [],
            "pages": [],
            "page_count": 0,
            "parser": suffix or "unknown",
            "error": str(exc),
        }


def _extract_text_file(path: Path) -> dict:
    raw = path.read_text(encoding="utf-8", errors="replace")
    return {
        "text": _clip(raw),
        "tables": [],
        "pages": [
            {
                "page": 1,
                "kind": "text",
                "char_count": len(raw),
                "has_images": False,
                "unread": False,
                "ocr_status": "not_needed",
            }
        ],
        "page_count": 1,
        "parser": "text",
        "error": None,
    }


def _extract_csv(path: Path) -> dict:
    raw = path.read_text(encoding="utf-8", errors="replace")
    rows: list[list[str]] = []
    reader = csv.reader(StringIO(raw))
    for i, row in enumerate(reader):
        if i >= MAX_TABLE_ROWS:
            break
        rows.append([str(cell) for cell in row])
    header = " ".join(rows[0]) if rows else ""
    sample = " ".join(" ".join(row) for row in rows[:12])
    return {
        "text": _clip(f"{header} {sample}"),
        "tables": [{"name": path.name, "rows": rows}] if rows else [],
        "pages": [
            {
                "page": 1,
                "kind": "text",
                "char_count": len(raw),
                "has_images": False,
                "sheet_name": path.name,
                "unread": False,
                "ocr_status": "not_needed",
            }
        ],
        "page_count": 1,
        "parser": "csv",
        "error": None,
    }


def _page_has_images(page: Any) -> bool:
    """Best-effort image detection via pypdf resources / images API."""
    try:
        images = getattr(page, "images", None)
        if images is not None and len(list(images)) > 0:
            return True
    except Exception:
        pass
    try:
        resources = page.get("/Resources")
        if resources is None:
            return False
        resolved = resources.get_object() if hasattr(resources, "get_object") else resources
        xobj = resolved.get("/XObject") if resolved else None
        if xobj is None:
            return False
        xobj = xobj.get_object() if hasattr(xobj, "get_object") else xobj
        for _name, ref in xobj.items():
            try:
                obj = ref.get_object() if hasattr(ref, "get_object") else ref
                if str(obj.get("/Subtype", "")) == "/Image":
                    return True
            except Exception:
                continue
    except Exception:
        return False
    return False


def _classify_pdf_page(char_count: int, has_images: bool) -> dict[str, Any]:
    if char_count <= 12 and not has_images:
        return {
            "kind": "blank",
            "unread": True,
            "ocr_status": "not_needed",
            "notes": ["Blank or near-empty page"],
        }
    if char_count <= 12 and has_images:
        return {
            "kind": "image",
            "unread": True,
            "ocr_status": "not_available",
            "notes": ["Image-only page — OCR not available"],
        }
    if char_count <= 80 and has_images:
        return {
            "kind": "mixed",
            "unread": True,
            "ocr_status": "not_available",
            "notes": ["Sparse text with images — OCR not available"],
        }
    if char_count <= 80:
        return {
            "kind": "scan",
            "unread": True,
            "ocr_status": "not_available",
            "notes": ["Very little extractable text — possible scan"],
        }
    if has_images:
        return {"kind": "mixed", "unread": False, "ocr_status": "not_needed", "notes": []}
    return {"kind": "text", "unread": False, "ocr_status": "not_needed", "notes": []}


def _extract_pdf(path: Path) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    page_texts: list[str] = []
    pages_meta: list[dict[str, Any]] = []
    for idx, page in enumerate(reader.pages[:MAX_PDF_PAGES]):
        text = page.extract_text() or ""
        page_texts.append(text)
        has_images = _page_has_images(page)
        meta = _classify_pdf_page(len(text.strip()), has_images)
        pages_meta.append(
            {
                "page": idx + 1,
                "char_count": len(text.strip()),
                "has_images": has_images,
                **meta,
            }
        )
    text = _clip("\n".join(page_texts))
    error = None if text.strip() else "No extractable text (scanned PDF requires OCR)"
    if error and pages_meta:
        for p in pages_meta:
            if p.get("kind") == "blank":
                p["kind"] = "scan"
                p["unread"] = True
                p["ocr_status"] = "not_available"
                p["notes"] = ["No extractable text — scanned PDF requires OCR"]
    return {
        "text": text,
        "tables": [],
        "pages": pages_meta,
        "page_count": len(reader.pages),
        "parser": "pypdf",
        "error": error,
    }


def _extract_docx(path: Path) -> dict:
    from docx import Document

    doc = Document(str(path))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []
    for i, table in enumerate(doc.tables[:MAX_TABLES]):
        rows = []
        for row in table.rows[:MAX_TABLE_ROWS]:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"name": f"table_{i + 1}", "rows": rows})
    page_count = max(1, len(doc.paragraphs) // 40)
    return {
        "text": _clip("\n".join(paras)),
        "tables": tables,
        "pages": [
            {
                "page": i + 1,
                "kind": "text",
                "char_count": max(0, len("\n".join(paras)) // page_count),
                "has_images": False,
                "unread": False,
                "ocr_status": "not_needed",
            }
            for i in range(page_count)
        ],
        "page_count": page_count,
        "parser": "python-docx",
        "error": None,
    }


def _extract_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    sheet_count = len(wb.worksheets)
    tables = []
    pages_meta: list[dict[str, Any]] = []
    chunks: list[str] = []
    for sheet_idx, sheet in enumerate(wb.worksheets[:MAX_XLSX_SHEETS]):
        rows: list[list[str]] = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_TABLE_ROWS:
                break
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cells):
                rows.append(cells)
        page_no = sheet_idx + 1
        if rows:
            # S1-6: each sheet is its own table/doc unit
            tables.append({"name": sheet.title, "rows": rows, "sheet_index": page_no})
            chunks.append(sheet.title)
            chunks.append(" ".join(" ".join(r) for r in rows[:8]))
            char_count = sum(len(c) for r in rows for c in r)
            pages_meta.append(
                {
                    "page": page_no,
                    "kind": "text",
                    "char_count": char_count,
                    "has_images": False,
                    "sheet_name": sheet.title,
                    "unread": False,
                    "ocr_status": "not_needed",
                }
            )
        else:
            pages_meta.append(
                {
                    "page": page_no,
                    "kind": "blank",
                    "char_count": 0,
                    "has_images": False,
                    "sheet_name": sheet.title,
                    "unread": True,
                    "ocr_status": "not_needed",
                    "notes": ["Empty sheet"],
                }
            )
    wb.close()
    return {
        "text": _clip(" ".join(chunks)),
        "tables": tables,
        "pages": pages_meta,
        "page_count": sheet_count,
        "parser": "openpyxl",
        "error": None,
    }


def _extract_pptx(path: Path) -> dict:
    from pptx import Presentation

    pres = Presentation(str(path))
    texts: list[str] = []
    pages_meta: list[dict[str, Any]] = []
    for idx, slide in enumerate(pres.slides):
        slide_text: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text)
                texts.append(shape.text)
        char_count = len(" ".join(slide_text).strip())
        pages_meta.append(
            {
                "page": idx + 1,
                "kind": "text" if char_count > 12 else "blank",
                "char_count": char_count,
                "has_images": False,
                "unread": char_count <= 12,
                "ocr_status": "not_needed",
            }
        )
    return {
        "text": _clip("\n".join(texts)),
        "tables": [],
        "pages": pages_meta,
        "page_count": len(pres.slides),
        "parser": "python-pptx",
        "error": None,
    }
