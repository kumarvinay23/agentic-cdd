#!/usr/bin/env python3
"""Step 6: populate a COPY of the fixed template from content.json. Rendering only - no analysis.

Usage:
  python render_deck.py --content <work>/content.json --work <work> --out <deck.pptx>
                        [--template assets/cdd_template.pptx]

Contract with the template: slides are found by name (slide_map id); variable elements are the
shapes whose names match element names. Nothing else on a slide is touched, so masters, layouts,
theme, dimensions, branding and fixed text are preserved. Slides are only duplicated when
slide_map marks them `repeatable` and content.json supplies a list of pages for them.

Writes <work>/render_log.json with the font size chosen for every element and any overflow;
check_content.py and qa_deck.py read that log.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.common import SKILL_DIR, read_json, slide_map, write_json  # noqa: E402
from lib.fit import fit_size, wrap_lines  # noqa: E402

BRAND = json.loads((SKILL_DIR / "scripts" / "template" / "brand.json").read_text())
C = {k: RGBColor.from_string(v) for k, v in BRAND["colors"].items()}
PALETTE = [RGBColor.from_string(c) for c in BRAND["chart_palette"]]
FONT = BRAND["fonts"]["body"]
EMU_PER_PT = 12700

RAG = {
    "good": (RGBColor.from_string("E3F1E8"), C["rag_high"]),
    "med": (RGBColor.from_string("FBF0DC"), RGBColor.from_string("8A5A0B")),
    "bad": (RGBColor.from_string("F6E0DD"), C["rag_low"]),
    "na": (C["tint"], C["slate"]),
    "neutral": (C["tint"], C["ink"]),
}
WORD_RAG = {
    "high": "H", "very high": "H", "strong": "H", "yes": "H", "met": "good", "achieved": "good", "positive": "good",
    "proceed": "good", "compliant": "good", "accretive": "good", "exceeded": "good", "increasing": "neutral",
    "medium": "M", "moderate": "M", "med": "M", "mid": "M", "partly": "med", "partial": "med", "partially compliant": "med",
    "mixed": "med", "cautious": "med", "stable": "M",
    "low": "L", "weak": "L", "no": "L", "critical": "H", "missed": "bad", "negative": "bad", "reject": "bad", "non-compliant": "bad",
    "dilutive": "bad",
}


def rag_class(value: str, scale: str) -> str:
    v = (value or "").strip().lower()
    if not v or v.startswith(("n/a", "not provided", "insufficient", "unknown", "not comparable", "not rated", "not stated", "-", "–")):
        return "na"
    key = None
    for w in sorted(WORD_RAG, key=len, reverse=True):
        if v == w or v.startswith(w + " ") or v.startswith(w + ",") or v.startswith(w + " ("):
            key = WORD_RAG[w]
            break
    if key is None:
        return "neutral"
    if key in ("good", "med", "bad"):
        return key
    if scale == "neutral":
        return "neutral"
    if scale == "bad_high":
        return {"H": "bad", "M": "med", "L": "good"}[key]
    return {"H": "good", "M": "med", "L": "bad"}[key]


# ------------------------------------------------------------------ low-level helpers
def slide_name(slide):
    return slide._element.cSld.get("name")


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    return None


def delete_shape(sh):
    el = sh._element
    el.getparent().remove(el)


def geom(sh):
    return sh.left, sh.top, sh.width, sh.height


def add_runs(paragraph, text, size, color, bold=False, italic=False):
    """Supports **bold** spans inside text."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        r = paragraph.add_run()
        is_b = part.startswith("**") and part.endswith("**")
        r.text = part[2:-2] if is_b else part
        f = r.font
        f.name = FONT
        f.size = Pt(size)
        f.bold = bold or is_b
        f.italic = italic
        f.color.rgb = color


def plain(text):
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", str(text))


def set_text_keep_format(sh, text):
    """Replace the text of a template text shape, keeping the first run's formatting."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    rpr = copy.deepcopy(runs[0]._r.find(qn("a:rPr"))) if runs else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    r = p0.add_run()
    r.text = text
    if rpr is not None:
        r._r.insert(0, rpr)


def first_size(sh, default=12):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return default


def set_bullet(paragraph, char="•", indent_emu=Inches(0.17)):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent_emu)))
    pPr.set("indent", str(-int(indent_emu)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def clear_region_style(sh, fill=None):
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    sh.line.fill.background()


def cell_border(cell, color="D5DBE3", width_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)
        ln = etree.SubElement(tcPr, qn(tag))
        ln.set("w", str(int(width_pt * EMU_PER_PT)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", color)
    # schema order: borders must precede fill elements
    fills = [e for e in tcPr if e.tag in (qn("a:solidFill"), qn("a:noFill"))]
    for f in fills:
        tcPr.remove(f)
        tcPr.append(f)


def _free_slide_partname(prs):
    """A /ppt/slides/slideN.xml that no part in the package already uses.

    python-pptx derives the next partname from the parts it can reach, and dropping a slide
    (as coverage gating does) can leave that counter pointing at a name a still-reachable part
    holds - the package then serialises two entries with the same name and the .pptx will not
    open. Allocate explicitly instead.
    """
    from pptx.opc.packuri import PackURI
    used = {str(part.partname) for part in prs.part.package.iter_parts()}
    n = 1
    while f"/ppt/slides/slide{n}.xml" in used:
        n += 1
    return PackURI(f"/ppt/slides/slide{n}.xml")


def duplicate_slide(prs, src):
    layout = src.slide_layout
    new = prs.slides.add_slide(layout)
    new.part.partname = _free_slide_partname(prs)
    for sh in list(new.shapes):
        delete_shape(sh)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new.shapes._spTree.append(copy.deepcopy(el))
    # move the new slide directly after the source
    lst = prs.slides._sldIdLst
    ids = list(lst)
    new_id = ids[-1]
    src_idx = [i for i, s in enumerate(prs.slides) if s.slide_id == src.slide_id][0]
    lst.remove(new_id)
    lst.insert(src_idx + 1, new_id)
    return new


# ------------------------------------------------------------------ element renderers
class Log:
    def __init__(self):
        self.items = []

    def add(self, slide, element, **kw):
        self.items.append({"slide": slide, "element": element, **kw})


def render_text(sh, spec, sid, name, log, style="panel"):
    """Single text block (takeaway, lead, data notice, cover fields)."""
    text = spec["text"] if isinstance(spec, dict) else str(spec)
    w_pt = sh.width / EMU_PER_PT - 14
    h_pt = sh.height / EMU_PER_PT - 8
    pref, mn = (12, 10)
    bold = style == "takeaway"
    size, ok, need = fit_size([(plain(text), bold)], w_pt, h_pt, pref, mn)
    clear_region_style(sh, C["tint"])
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    add_runs(p, text, size, C["ink"], bold=bold)
    log.add(sid, name, kind="text", size=size, fits=ok, needed_pt=need)


def render_insufficient(sh, spec, sid, name, log):
    text = spec.get("text") or "Insufficient evidence"
    clear_region_style(sh, C["tint2"])
    sh.line.color.rgb = C["line"]
    sh.line.width = Pt(0.75)
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_runs(p, text, 11, C["slate"], italic=True)
    log.add(sid, name, kind="insufficient", size=11, fits=True)


def render_bullets(sh, spec, sid, name, log):
    items = spec.get("items", [])
    paras = []
    for it in items:
        t = it["text"] if isinstance(it, dict) else str(it)
        paras.append((plain(t), False))
    w_pt = sh.width / EMU_PER_PT - 16
    h_pt = sh.height / EMU_PER_PT - 10
    size, ok, need = fit_size(paras, w_pt, h_pt, spec.get("font_pt", 14), spec.get("font_min_pt", 10), space_after_pt=6, indent_pt=13)
    clear_region_style(sh, None)
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.06)
    for i, it in enumerate(items):
        t = it["text"] if isinstance(it, dict) else str(it)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        set_bullet(p)
        mark = "†" if isinstance(it, dict) and it.get("kind") == "calc" else ""
        add_runs(p, t + mark, size, C["text"])
    log.add(sid, name, kind="bullets", size=size, fits=ok, needed_pt=need, items=len(items))


def _col_widths(spec, ncols, total_w):
    if spec.get("col_widths"):
        fr = spec["col_widths"]
        s = sum(fr)
        return [int(total_w * f / s) for f in fr]
    lens = []
    for j in range(ncols):
        col = [str(spec["columns"][j])] + [plain(r[j]) if j < len(r) else "" for r in spec["rows"]]
        lens.append(max(6, min(60, max(len(c) for c in col))) ** 0.8)
    s = sum(lens)
    return [int(total_w * l / s) for l in lens]


def render_table(slide, sh, spec, sid, name, log):
    x, y, w, h = geom(sh)
    cols = spec["columns"]
    rows = spec["rows"]
    nc = len(cols)
    widths = _col_widths(spec, nc, w)
    pill_cols = {int(p["col"]) if isinstance(p, dict) else int(p): (p.get("scale", "good_high") if isinstance(p, dict) else "good_high")
                 for p in spec.get("pill_cols", [])}
    pref, mn = spec.get("font_pt", 10.5), spec.get("font_min_pt", 9)
    pad_pt = 7.0

    def table_height(size):
        hh = 0
        for ri, row in enumerate([cols] + rows):
            mx = 1
            for j in range(nc):
                txt = plain(row[j]) if j < len(row) else ""
                mx = max(mx, wrap_lines(txt, widths[j] / EMU_PER_PT - 10, size, bold=(ri == 0)))
            hh += mx * size * 1.2 + pad_pt
        return hh

    size = pref
    while size > mn and table_height(size) > h / EMU_PER_PT:
        size -= 0.5
    # sparse tables: grow the font (max 13pt) while the table stays within ~80% of the region
    mx_pt = spec.get("font_max_pt", 13)
    while size < mx_pt and table_height(size + 0.5) <= 0.8 * h / EMU_PER_PT:
        size += 0.5
    need = table_height(size)
    # then pad rows so the table fills the region (never more than +60% of its natural height)
    stretch = min(1.6, (h / EMU_PER_PT) / need) if need else 1.0
    fits = need <= h / EMU_PER_PT + 2
    delete_shape(sh)
    stretch = max(1.0, stretch)
    gf = slide.shapes.add_table(len(rows) + 1, nc, x, y, w, Emu(int(min(need * stretch, h / EMU_PER_PT) * EMU_PER_PT)))
    gf.name = name
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    sid_el = tblPr.find(qn("a:tableStyleId"))
    if sid_el is not None:
        sid_el.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for j, cw in enumerate(widths):
        tbl.columns[j].width = cw
    for ri, row in enumerate([cols] + rows):
        # row height = its own estimated need, so the table does not stretch past the region
        mx = 1
        for j in range(nc):
            txt = plain(row[j]) if j < len(row) else ""
            mx = max(mx, wrap_lines(txt, widths[j] / EMU_PER_PT - 10, size, bold=(ri == 0)))
        rh = mx * size * 1.2 + pad_pt
        if ri > 0:
            rh *= stretch
        tbl.rows[ri].height = Emu(int(rh * EMU_PER_PT))
        for j in range(nc):
            cell = tbl.cell(ri, j)
            txt = str(row[j]) if j < len(row) and row[j] is not None else ""
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.035)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if ri == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = C["ink"]
                add_runs(p, txt, size, C["white"], bold=True)
            else:
                if j in pill_cols:
                    cls = rag_class(plain(txt), pill_cols[j])
                    bg, fg = RAG[cls]
                    cell.fill.solid(); cell.fill.fore_color.rgb = bg
                    p.alignment = PP_ALIGN.CENTER
                    add_runs(p, txt, size, fg, bold=True)
                else:
                    cell.fill.solid(); cell.fill.fore_color.rgb = C["white"] if ri % 2 else C["tint2"]
                    first = j == 0 and spec.get("bold_first_col", True)
                    add_runs(p, txt, size, C["ink"] if first else C["text"], bold=first)
            cell_border(cell)
    log.add(sid, name, kind="table", size=size, fits=fits, needed_pt=round(need, 1),
            available_pt=round(h / EMU_PER_PT, 1), rows=len(rows))


CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked_column_100": XL_CHART_TYPE.COLUMN_STACKED_100, "line": XL_CHART_TYPE.LINE_MARKERS,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def _hide_label(point):
    """Remove one point's data label (a label on a sliver of a stacked bar only overlaps)."""
    from pptx.oxml.ns import qn
    dLbl = point.data_label._get_or_add_dLbl()
    for child in list(dLbl):
        if child.tag != qn("c:idx"):
            dLbl.remove(child)
    el = dLbl.makeelement(qn("c:delete"), {"val": "1"})
    dLbl.append(el)


def _fmt_num(v, spec):
    nf = spec.get("number_format", "")
    dec = 0 if "0.0" not in nf else (2 if "0.00" in nf else 1)
    pre = "$" if "$" in nf else ""
    suf = "M" if '"M"' in nf else ("%" if "%" in nf else "")
    sign = "+" if v > 0 else ("-" if v < 0 else "")
    return f"{sign}{pre}{abs(v):,.{dec}f}{suf}"


def render_waterfall(slide, sh, spec, sid, name, log):
    """Bridge chart: first/last (and any index in `totals`) are absolute bars, the rest are
    deltas floating on an invisible base. Values are the deltas / totals as stated in the spec."""
    x, y, w, h = geom(sh)
    delete_shape(sh)
    vals = [float(v) for v in spec["series"][0]["values"]]
    totals = set(spec.get("totals", [0, len(vals) - 1]))
    base, bar, run = [], [], 0.0
    for i, v in enumerate(vals):
        if i in totals:
            base.append(0.0); bar.append(v); run = v
        else:
            lo = run + min(v, 0)
            base.append(lo); bar.append(abs(v)); run += v
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    cd.add_series("base", base)
    cd.add_series(spec["series"][0].get("name", "value"), bar)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, x, y, w, h, cd)
    gf.name = name
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(10)
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 50
    va = ch.value_axis
    va.major_gridlines.format.line.color.rgb = C["line"]
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(9)
    if spec.get("number_format"):
        va.tick_labels.number_format = spec["number_format"]
        va.tick_labels.number_format_is_linked = False
    if spec.get("value_axis_title"):
        va.has_title = True
        va.axis_title.text_frame.text = spec["value_axis_title"]
        va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    ch.category_axis.tick_labels.font.size = Pt(9)
    b, s_ = plot.series[0], plot.series[1]
    b.format.fill.background()
    b.format.line.fill.background()
    for i, v in enumerate(vals):
        pt = s_.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = (C["ink"] if i in totals else
                                         (PALETTE[1 % len(PALETTE)] if v >= 0 else C["rag_low"]))
        tf = pt.data_label.text_frame
        tf.text = _fmt_num(v, spec) if i not in totals else _fmt_num(v, spec).lstrip("+")
        tf.paragraphs[0].runs[0].font.size = Pt(9)
        if i in totals:
            tf.paragraphs[0].runs[0].font.color.rgb = C["white"]
        pt.data_label.position = XL_LABEL_POSITION.INSIDE_END if abs(v) > 0.12 * max(map(abs, vals)) \
            else XL_LABEL_POSITION.INSIDE_BASE
    log.add(sid, name, kind="chart", chart_type="waterfall", points=len(vals), series=1)


def render_chart(slide, sh, spec, sid, name, log):
    if spec.get("chart_type") == "waterfall":
        return render_waterfall(slide, sh, spec, sid, name, log)
    x, y, w, h = geom(sh)
    delete_shape(sh)
    cd = CategoryChartData()
    cd.categories = spec["categories"]
    for s in spec["series"]:
        cd.add_series(s["name"], [None if v is None else float(v) for v in s["values"]])
    ctype = spec.get("chart_type", "column")
    gf = slide.shapes.add_chart(CHART_TYPES[ctype], x, y, w, h, cd)
    gf.name = name
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(10)
    ch.font.color.rgb = C["text"]
    multi = len(spec["series"]) > 1
    ch.has_legend = multi or ctype == "doughnut"
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(10)
    ch.has_title = False
    plot = ch.plots[0]
    if ctype != "doughnut":
        plot.gap_width = spec.get("gap_width", 60)
        if multi and ctype in ("column", "bar"):
            plot.overlap = -10
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = C["line"]
        va.major_gridlines.format.line.width = Pt(0.5)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(9)
        va.tick_labels.font.color.rgb = C["slate"]
        if spec.get("number_format"):
            va.tick_labels.number_format = spec["number_format"]
            va.tick_labels.number_format_is_linked = False
        if spec.get("value_axis_title"):
            va.has_title = True
            va.axis_title.text_frame.text = spec["value_axis_title"]
            va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            va.axis_title.text_frame.paragraphs[0].runs[0].font.color.rgb = C["slate"]
        if spec.get("value_axis_min") is not None:
            va.minimum_scale = spec["value_axis_min"]
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(10)
        ca.tick_labels.font.color.rgb = C["text"]
        ca.format.line.color.rgb = C["line"]
        ca.has_major_gridlines = False
    if spec.get("data_labels", True):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(9)
        dl.font.color.rgb = C["text"] if ctype not in ("stacked_column", "stacked_bar", "stacked_column_100", "doughnut") else C["white"]
        if spec.get("number_format"):
            dl.number_format = spec["number_format"]
            dl.number_format_is_linked = False
        if ctype in ("column", "bar"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    hl = spec.get("highlight")
    for i, s in enumerate(plot.series):
        fill = s.format.fill
        fill.solid()
        fill.fore_color.rgb = PALETTE[i % len(PALETTE)]
        if ctype == "line":
            s.format.line.color.rgb = PALETTE[i % len(PALETTE)]
            s.format.line.width = Pt(2.25)
        if hl is not None and not multi and ctype in ("column", "bar"):
            pt = s.points[int(hl)]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = C["amber"]
    if ctype in ("stacked_column", "stacked_bar") and spec.get("data_labels", True):
        tot = [sum(abs(float(sr["values"][j] or 0)) for sr in spec["series"]) for j in range(len(spec["categories"]))]
        mx = max(tot) if tot else 0
        for i, sr in enumerate(spec["series"]):
            for j, v in enumerate(sr["values"]):
                if v is not None and mx and abs(float(v)) < 0.06 * mx:
                    _hide_label(plot.series[i].points[j])
            # a per-point override creates series-level dLbls that default to 'show nothing';
            # restate the plot-level settings on the series so the other labels stay visible
            ser = plot.series[i]._element
            dls = ser.find(qn("c:dLbls"))
            if dls is not None:
                from pptx.chart.datalabel import DataLabels
                d = DataLabels(dls)
                d.show_value = True
                d.font.size = Pt(9)
                d.font.color.rgb = C["white"]
                if spec.get("number_format"):
                    d.number_format = spec["number_format"]
                    d.number_format_is_linked = False
    if ctype == "doughnut":
        for j, _ in enumerate(spec["categories"]):
            pt = plot.series[0].points[j]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = PALETTE[j % len(PALETTE)]
    log.add(sid, name, kind="chart", chart_type=ctype, points=len(spec["categories"]), series=len(spec["series"]))


def render_kpis(slide, spec, sid, log):
    tiles = spec.get("items", [])
    for i in range(1, 6):
        sh = shape_by_name(slide, f"kpi{i}")
        if sh is None:
            continue
        if i > len(tiles):
            delete_shape(sh)
            log.add(sid, f"kpi{i}", kind="kpi", removed=True)
            continue
        t = tiles[i - 1]
        clear_region_style(sh, C["tint"])
        tf = sh.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.12)
        val = str(t.get("value", ""))
        w_pt = sh.width / EMU_PER_PT - 18
        vsize = 24
        while vsize > 14 and wrap_lines(val, w_pt, vsize, True) > 1:
            vsize -= 1
        color = C["ink"]
        if t.get("rag"):
            color = RAG[t["rag"]][1]
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        add_runs(p, val + ("†" if t.get("kind") == "calc" else ""), vsize, color, bold=True)
        p2 = tf.add_paragraph()
        add_runs(p2, str(t.get("label", "")), 10, C["slate"], bold=True)
        if t.get("note"):
            p3 = tf.add_paragraph()
            add_runs(p3, str(t["note"]), 9, C["slate"], italic=True)
        need = vsize * 1.2 + 12 + (wrap_lines(str(t.get("note", "")), w_pt, 9) * 11 if t.get("note") else 0) + 10
        log.add(sid, f"kpi{i}", kind="kpi", size=vsize, fits=need <= sh.height / EMU_PER_PT)


def render_contents(slide, sh, sections, sid, log):
    x, y, w, h = geom(sh)
    delete_shape(sh)
    gf = slide.shapes.add_table(len(sections), 2, x, y, w, Emu(int(len(sections) * 0.46 * 914400)))
    gf.name = "contents_body"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    el = tblPr.find(qn("a:tableStyleId"))
    if el is not None:
        el.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    tbl.first_row = False
    tbl.columns[0].width = int(w * 0.85)
    tbl.columns[1].width = w - int(w * 0.85)
    for i, (num, name_, page) in enumerate(sections):
        for j, txt in enumerate([f"{num:02d}   {name_}", f"{page:02d}"]):
            c = tbl.cell(i, j)
            c.fill.background()
            tf = c.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if j else PP_ALIGN.LEFT
            add_runs(p, txt, 16, C["ink"], bold=(j == 0))
            tcPr = c._tc.get_or_add_tcPr()
            ln = etree.SubElement(tcPr, qn("a:lnB"))
            ln.set("w", "6350")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", BRAND["colors"]["line"])
    log.add(sid, "contents_body", kind="contents", rows=len(sections))


# ------------------------------------------------------------------ source lines & notes
def short_doc(d):
    d = re.sub(r"\.(xlsx|xls|csv|pdf|docx|pptx)$", "", d, flags=re.I)
    syn = bool(re.match(r"(?i)synthetic", d))
    d = re.sub(r"(?i)^synthetic_?\d*_?", "", d).replace("_", " ").strip()
    return ("synthetic " if syn else "") + d


def collect_refs(obj, acc):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k in ("refs", "row_refs") and isinstance(v, list):
                for r in v:
                    if isinstance(r, list):
                        acc.extend(r)
                    else:
                        acc.append(r)
            else:
                collect_refs(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            collect_refs(v, acc)
    return acc


def source_line(refs, facts, checks, has_calc):
    mods, docs, webs = [], [], []
    for r in refs:
        f = facts.get(r)
        if not f:
            continue
        if f["module_title"] not in mods:
            mods.append(f["module_title"])
        for d in f["citations"]["doc"]:
            sd = short_doc(d)
            if sd not in docs:
                docs.append(sd)
        for w in f["citations"]["web"]:
            res = checks.get("citation_index", {}).get("web", {}).get(str(w), {}).get("resolved_as") or []
            lab = f"[{w}] " + (re.sub(r"\s*\(.*$", "", res[0]) if res else "unresolved web source")
            if lab not in webs:
                webs.append(lab)
    parts = []
    if mods:
        parts.append("DiligenceIQ agent outputs: " + ", ".join(mods))
    if docs:
        parts.append("data room: " + ", ".join(docs))
    if webs:
        parts.append("external: " + "; ".join(webs))
    line = "Source: " + ("; ".join(parts) if parts else "DiligenceIQ analysis")
    if has_calc:
        line += ". † DiligenceIQ calculation (formula in notes)"
    if len(line) > 420:
        line = "Source: DiligenceIQ agent outputs: " + ", ".join(mods[:6]) + (f" (+{len(mods)-6} more)" if len(mods) > 6 else "") + \
               f"; {len(docs)} data-room documents; {len(webs)} external sources (full list in notes)" + \
               (". † DiligenceIQ calculation" if has_calc else "")
    return line


def notes_text(sid, sc, facts, calcs, exceptions):
    out = [f"EVIDENCE TRAIL - {sid}", "Kinds: FACT = sourced; CALC = calculated (†); VIEW = DiligenceIQ interpretation", ""]

    def walk(label, obj):
        if isinstance(obj, dict):
            refs = obj.get("refs") or []
            kind = obj.get("kind", "")
            if refs or kind:
                txt = obj.get("text") or obj.get("value") or ""
                out.append(f"[{label}] {kind.upper() or 'ITEM'}: {plain(txt)[:140]}")
                for r in refs:
                    f = facts.get(r)
                    out.append(f"    <- {r}: {f['text'][:160] if f else 'UNKNOWN REF'}")
            for k, v in obj.items():
                if k in ("items", "rows_meta", "elements"):
                    walk(label, v)
                elif isinstance(v, (dict, list)) and k not in ("refs", "row_refs", "series", "rows", "columns", "categories"):
                    walk(f"{label}.{k}" if label else k, v)
            if obj.get("row_refs"):
                for i, rr in enumerate(obj["row_refs"]):
                    out.append(f"[{label} row {i+1}] refs: {', '.join(rr) if isinstance(rr, list) else rr}")
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                walk(f"{label}[{i+1}]", v)

    walk("", {k: v for k, v in sc.items() if k != "notes"})
    rel_calcs = [c for c in calcs if sid in c.get("used_on", [])]
    if rel_calcs:
        out += ["", "CALCULATIONS"]
        for c in rel_calcs:
            ins = ", ".join(f"{k}={v.get('value')} ({v.get('ref')})" for k, v in c.get("inputs", {}).items())
            out.append(f"{c['id']}: {c['formula']} | inputs: {ins} | output: {c.get('output')} {c.get('unit','')}")
    rel_ex = [e for e in exceptions if sid in e.get("slides", [])]
    if rel_ex:
        out += ["", "EXCEPTIONS AFFECTING THIS SLIDE"]
        for e in rel_ex:
            out.append(f"[{e['severity']}] {e['issue']} -> {e.get('treatment','')}")
    if sc.get("notes"):
        out += ["", "ANALYST NOTES", sc["notes"]]
    return "\n".join(out)


# ------------------------------------------------------------------ main
def render_slide(prs, slide, sid, sc, sm_slide, facts, checks, calcs, exceptions, log):
    arch = sm_slide["archetype"]
    for key in ("title", "subtitle"):
        if key in sc and shape_by_name(slide, key) is not None:
            val = sc[key]["text"] if isinstance(sc[key], dict) else sc[key]
            sh = shape_by_name(slide, key)
            set_text_keep_format(sh, val)
            lim = {"title": (20, 18), "subtitle": (13, 12)}[key]
            size, ok, need = fit_size([(val, key == "title")], sh.width / EMU_PER_PT - 6, sh.height / EMU_PER_PT - 2, lim[0], lim[1])
            for r in sh.text_frame.paragraphs[0].runs:
                r.font.size = Pt(size)
            log.add(sid, key, kind="text", size=size, fits=ok, needed_pt=need)
    # optional panel-heading overrides (template supplies defaults, e.g. {"chart_head": "..."})
    for hk, hv in (sc.get("heads") or {}).items():
        hs = shape_by_name(slide, hk)
        if hs is not None and hk.endswith("_head"):
            set_text_keep_format(hs, hv)
    els = sc.get("elements", {})
    for el in sm_slide.get("elements", []):
        name = el["name"]
        spec = els.get(name)
        if el["type"] == "auto":
            continue
        if name == "kpis":
            if spec is None:
                spec = {"items": []}
            render_kpis(slide, spec, sid, log)
            continue
        sh = shape_by_name(slide, name)
        if sh is None:
            log.add(sid, name, error="SHAPE_NOT_IN_TEMPLATE")
            continue
        if spec is None:
            spec = {"type": "insufficient", "text": "Not provided in agent outputs"}
            log.add(sid, name, warning="ELEMENT_MISSING_IN_CONTENT")
        t = spec.get("type", el["type"])
        if arch == "cover":
            set_text_keep_format(sh, spec["text"] if isinstance(spec, dict) else spec)
            log.add(sid, name, kind="text", fits=True)
        elif t == "insufficient":
            render_insufficient(sh, spec, sid, name, log)
        elif t == "text":
            render_text(sh, spec, sid, name, log, style="takeaway" if name == "takeaway" else "panel")
        elif t == "bullets":
            render_bullets(sh, spec, sid, name, log)
        elif t == "table":
            render_table(slide, sh, spec, sid, name, log)
        elif t == "chart":
            render_chart(slide, sh, spec, sid, name, log)
        else:
            log.add(sid, name, error=f"UNKNOWN_TYPE {t}")
    if arch not in ("cover", "back_cover", "divider", "contents"):
        src = shape_by_name(slide, "source")
        if src is not None:
            refs = collect_refs(sc, [])
            has_calc = "†" in json.dumps(sc, ensure_ascii=False) or '"kind": "calc"' in json.dumps(sc)
            line = sc.get("source_override") or source_line(refs, facts, checks, has_calc)
            set_text_keep_format(src, line)
            size, ok, _ = fit_size([(line, False)], src.width / EMU_PER_PT, 24, 8, 7)
            for r in src.text_frame.paragraphs[0].runs:
                r.font.size = Pt(size)
            log.add(sid, "source", kind="text", size=size, fits=ok, chars=len(line))
    slide.notes_slide.notes_text_frame.text = notes_text(sid, sc, facts, calcs, exceptions)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--content", required=True)
    ap.add_argument("--work", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", default=str(SKILL_DIR / "assets" / "cdd_template.pptx"))
    a = ap.parse_args()
    work = Path(a.work)
    content = read_json(a.content)
    sm = slide_map()
    facts = {f["id"]: f for f in read_json(work / "evidence.json")["facts"]}
    checks = read_json(work / "auto_checks.json") if (work / "auto_checks.json").exists() else {}
    calcs = content.get("calculations", [])
    exceptions = content.get("exceptions", [])
    shutil.copy(a.template, a.out)
    prs = Presentation(a.out)
    log = Log()
    sm_by_id = {s["id"]: s for s in sm["slides"]}

    # 0) coverage gating: a slide the evidence cannot support is REMOVED, not filled with a
    #    placeholder. Every removal is already recorded in gates.json with its reason and
    #    surfaces in the exceptions appendix - silent omission is a defect (axiom A6).
    gates_path = work / "gates.json"
    dropped = {}
    if gates_path.exists():
        for g in read_json(gates_path):
            if not g.get("render"):
                dropped[g["id"]] = g
    forced = set(content.get("force_render", []))
    dropped = {k: v for k, v in dropped.items() if k not in forced}
    for w in content.get("withhold", []) or []:
        dropped[w["id"]] = {"id": w["id"], "exception": w.get("exception", "E13-WITHHELD"),
                            "reason": w.get("reason", "")}
    # a section survives only if it keeps a findings page; the back cover is not one
    alive = {sm_by_id[s["id"]]["section"] for s in sm["slides"]
             if s["id"] not in dropped and s["archetype"] not in ("divider", "back_cover")}
    for sid_, sd_ in sm_by_id.items():
        if sd_["archetype"] == "divider" and sd_["section"] not in alive:
            dropped.setdefault(sid_, {"id": sid_, "exception": "E1-NOT-PROVIDED",
                                      "reason": "every page in this section was withheld"})
    if dropped:
        # python-pptx has no delete; remove by index via the sldIdLst
        sldIdLst = prs.slides._sldIdLst
        present = [(idx, slide_name(sl)) for idx, sl in enumerate(prs.slides)]
        for idx, nm in reversed(present):
            if nm in dropped:
                rId = sldIdLst[idx].rId
                prs.part.drop_rel(rId)
                del sldIdLst[idx]
        print(f"[render] gated out {len(dropped)} slides: "
              f"{', '.join(sorted(dropped)[:6])}{' ...' if len(dropped) > 6 else ''}")

    by_name = {slide_name(s): s for s in prs.slides}
    sm_by_id = {k: v for k, v in sm_by_id.items() if k in by_name}

    # 1) structure first: duplicate repeatable slides that have multiple pages of content
    pages = {}
    for sid, sm_slide in sm_by_id.items():
        sc = content["slides"].get(sid)
        if isinstance(sc, list):
            if not sm_slide.get("repeatable") and len(sc) > 1:
                log.add(sid, "*", error="SLIDE_NOT_REPEATABLE - extra pages ignored")
                sc = sc[:1]
            base = by_name[sid]
            slides_ = [base]
            for k in range(1, len(sc)):
                dup = duplicate_slide(prs, slides_[-1])
                dup._element.cSld.set("name", f"{sid}__{k+1}")
                slides_.append(dup)
            pages[sid] = list(zip(slides_, sc))
        else:
            pages[sid] = [(by_name[sid], sc)]

    # 2) project-level text in layouts (footer project name)
    proj = content.get("project", {})
    for layout in prs.slide_layouts:
        for sh in layout.shapes:
            if sh.has_text_frame and "[Project name]" in sh.text_frame.text:
                set_text_keep_format(sh, proj.get("footer", proj.get("name", "")))

    # 3) contents numbering after structure is final
    order = [slide_name(s) for s in prs.slides]
    sec_order = [s for s in sm["sections"] if s["id"] != "front"]
    # number the sections that survived gating contiguously - a gap in the contents reads as a
    # rendering fault rather than as a withheld section, and the withholding is stated in the
    # exceptions appendix instead
    contents = []
    for sec in sec_order:
        div = [s["id"] for s in sm["slides"] if s["section"] == sec["id"] and s["archetype"] == "divider"]
        if div and div[0] in order:
            contents.append((len(contents) + 1, sec["name"], order.index(div[0]) + 1))

    # 3a) appendix letters follow the appendices that survived (A, B, ...), not the template's
    rename, k = {}, 0
    for i, (n, name, pg) in enumerate(contents):
        mm = re.match(r"^Appendix [A-Z]\b(.*)$", name)
        if mm:
            new = f"Appendix {chr(65 + k)}{mm.group(1)}"
            k += 1
            rename[name] = new
            contents[i] = (n, new, pg)
    if rename:
        for sl in prs.slides:
            for shp in sl.shapes:
                if shp.name in ("section_tag", "div_title") and shp.has_text_frame:
                    t = shp.text_frame.text
                    for old_, new_ in rename.items():
                        if t.strip().upper() == old_.upper():
                            set_text_keep_format(shp, new_.upper() if t.isupper() else new_)

    # 3b) dividers: contiguous numbers and a navigation list of the sections that survived
    blurbs = content.get("dividers") or {}
    live = [(n, name) for n, name, _ in contents]
    for sid, plist in pages.items():
        sm_slide = sm_by_id[sid]
        if sm_slide["archetype"] != "divider":
            continue
        slide = plist[0][0]
        sec = dict(next(x for x in sm["sections"] if x["id"] == sm_slide["section"]))
        sec["name"] = rename.get(sec["name"], sec["name"])
        num = next((n for n, name in live if name == sec["name"]), 0)
        sh = shape_by_name(slide, "div_number")
        if sh is not None:
            set_text_keep_format(sh, f"{num:02d}")
        if sec["id"] in blurbs and shape_by_name(slide, "div_blurb") is not None:
            set_text_keep_format(shape_by_name(slide, "div_blurb"), blurbs[sec["id"]])
        navs = sorted([x for x in slide.shapes if x.name.startswith("div_list_")],
                      key=lambda x: int(x.name.rsplit("_", 1)[1]))
        for i, shp in enumerate(navs):
            if i < len(live):
                n, name = live[i]
                set_text_keep_format(shp, f"{n:02d}  {name}")
                cur = name == sec["name"]
                for para in shp.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.bold = cur
                        r.font.color.rgb = RGBColor.from_string("F2C879" if cur else "B9C3D1")
            else:
                delete_shape(shp)
        log.add(sid, "div_number", kind="divider", number=num, nav=len(live))

    # 4) fill
    for sid, plist in pages.items():
        sm_slide = sm_by_id[sid]
        for slide, sc in plist:
            if sm_slide["archetype"] == "divider":
                continue
            if sm_slide["archetype"] == "contents":
                render_contents(slide, shape_by_name(slide, "contents_body"), contents, sid, log)
                continue
            if sc is None:
                sc = {"title": {"text": "Insufficient evidence to populate this slide"}, "elements": {}}
                log.add(sid, "*", warning="SLIDE_MISSING_IN_CONTENT")
            render_slide(prs, slide, sid, sc, sm_slide, facts, checks, calcs, exceptions, log)

    # the blueprint number is an authoring aid; it collides with the wordmark on the page
    for sl in prs.slides:
        for shp in list(sl.shapes):
            if shp.name == "slide_no_tag":
                delete_shape(shp)

    # reference pages (glossary) sit at the back, before the back cover, unless content asks otherwise
    if content.get("glossary_position", "back") == "back":
        ids_ = list(prs.slides._sldIdLst)
        gl = [el for el, sl in zip(ids_, prs.slides) if re.sub(r"__\d+$", "", slide_name(sl) or "") == "S007_glossary"]
        for el in gl:
            prs.slides._sldIdLst.remove(el)
            prs.slides._sldIdLst.append(el)

    # the back cover closes the deck, after the appendices
    ids = list(prs.slides._sldIdLst)
    for idx, sl in enumerate(prs.slides):
        if sm_by_id.get(slide_name(sl), {}).get("archetype") == "back_cover" and idx != len(ids) - 1:
            el = ids[idx]
            prs.slides._sldIdLst.remove(el)
            prs.slides._sldIdLst.append(el)
            break
    prs.save(a.out)
    overflow = [i for i in log.items if i.get("fits") is False]
    errors = [i for i in log.items if i.get("error")]
    write_json(work / "render_log.json", {"out": a.out, "items": log.items, "overflow": overflow,
                                          "errors": errors, "dropped": dropped})
    print(f"[render] {a.out}: {len(prs.slides)} slides; overflow {len(overflow)}; errors {len(errors)}")
    for o in overflow:
        print("  OVERFLOW", o)
    for e in errors:
        print("  ERROR", e)


if __name__ == "__main__":
    main()
