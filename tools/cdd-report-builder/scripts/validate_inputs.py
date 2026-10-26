#!/usr/bin/env python3
"""Step 2: validate input completeness and template compatibility before any drafting.

Usage: python validate_inputs.py --work <work> [--template assets/cdd_template.pptx]

Writes <work>/input_validation.json and .md. Exit code 2 = blocker (generation would be unreliable):
  * template does not carry every slide id / element shape required by config/slide_map.yaml
  * fewer than half of the `core` modules could be identified
Everything else (missing supporting modules, duplicates, missing sections, low OCR confidence,
synthetic-data markers) is recorded as an exception and generation continues.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.common import SKILL_DIR, read_json, registry, slide_map, write_json  # noqa: E402


def template_check(path, sm):
    prs = Presentation(path)
    names = {}
    for s in prs.slides:
        names[s._element.cSld.get("name")] = {sh.name for sh in s.shapes}
    problems = []
    for sd in sm["slides"]:
        if sd["id"] not in names:
            problems.append(f"slide {sd['id']} not found in template")
            continue
        shapes = names[sd["id"]]
        for el in sd.get("elements", []):
            n = el["name"]
            if n == "kpis":
                if not any(x.startswith("kpi") for x in shapes):
                    problems.append(f"{sd['id']}: no kpi tiles")
                continue
            if el["type"] != "auto" and n not in shapes:
                problems.append(f"{sd['id']}: shape '{n}' missing")
        # archetypes whose title is fixed by the template and which carry no subtitle/source line
        FIXED_TITLE = ("cover", "back_cover", "divider", "contents", "disclaimer", "text_panel")
        if sd["archetype"] not in FIXED_TITLE:
            for req in ("title", "subtitle", "source"):
                if req not in shapes:
                    problems.append(f"{sd['id']}: shape '{req}' missing")
    w, h = prs.slide_width / 914400, prs.slide_height / 914400
    return problems, {"slides": len(prs.slides), "width_in": round(w, 3), "height_in": round(h, 3)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True)
    ap.add_argument("--template", default=str(SKILL_DIR / "assets" / "cdd_template.pptx"))
    a = ap.parse_args()
    work = Path(a.work)
    reg = registry()
    sm = slide_map()
    man = read_json(work / "manifest.json")
    checks = read_json(work / "auto_checks.json") if (work / "auto_checks.json").exists() else {}
    issues, blockers = [], []
    core = [m for m, d in reg["modules"].items() if d.get("criticality") == "core"]
    found = set(man["modules"])
    miss_core = [m for m in core if m not in found]
    miss_sup = [m for m in reg["modules"] if m not in found and m not in core]
    if len(found & set(core)) < len(core) / 2:
        blockers.append(f"only {len(found & set(core))}/{len(core)} core modules identified")
    for m in miss_core:
        issues.append({"severity": "HIGH", "code": "MODULE_MISSING", "issue": f"Core module missing: {reg['modules'][m]['title']}"})
    for m in miss_sup:
        issues.append({"severity": "MEDIUM", "code": "MODULE_MISSING", "issue": f"Supporting module missing: {reg['modules'][m]['title']}"})
    for rec in man.get("unidentified", []):
        issues.append({"severity": "MEDIUM", "code": "UNIDENTIFIED_FILE", "issue": f"File not recognised as a module: {Path(rec['file']).name} (best guess {rec['candidates'][0]})"})
    for mid, m in man["modules"].items():
        for o in m.get("other_versions", []):
            issues.append({"severity": "LOW", "code": "DUPLICATE_INPUT",
                           "issue": f"{m['title']}: additional file {Path(o['file']).name} "
                                    f"({'identical content' if o['identical_content'] else 'DIFFERENT content - older version ignored'}); using {Path(m['chosen']).name}"})
        exp = {s["n"] for s in reg["modules"][mid]["sections"]}
        miss = sorted(exp - set(m["sections_found"]))
        if miss:
            issues.append({"severity": "MEDIUM", "code": "SECTION_MISSING", "issue": f"{m['title']}: sections {miss} not found"})
        if m.get("ocr_mean_conf") and m["ocr_mean_conf"] < 85:
            issues.append({"severity": "MEDIUM", "code": "LOW_OCR_CONF", "issue": f"{m['title']}: mean OCR confidence {m['ocr_mean_conf']}"})
    if checks.get("synthetic", {}).get("detected"):
        issues.append({"severity": "CRITICAL", "code": "SYNTHETIC_DATA",
                       "issue": f"Inputs cite synthetic / machine-generated test data ({checks['synthetic']['count']} evidence items). Output must be labelled illustrative."})
    tprob, tinfo = template_check(a.template, sm)
    if tprob:
        blockers += tprob
    res = {"modules_found": sorted(found), "missing_core": miss_core, "missing_supporting": miss_sup,
           "template": tinfo, "issues": issues, "blockers": blockers}
    write_json(work / "input_validation.json", res)
    md = ["# Input validation", "", f"Modules identified: {len(found)}/{len(reg['modules'])}",
          f"Template: {tinfo}", "", "## Blockers"] + ([f"- {b}" for b in blockers] or ["- none"]) + ["", "## Issues"]
    md += [f"- [{i['severity']}] {i['code']}: {i['issue']}" for i in issues] or ["- none"]
    (work / "input_validation.md").write_text("\n".join(md), encoding="utf-8")
    print(f"[validate] modules {len(found)}; blockers {len(blockers)}; issues {len(issues)}")
    for b in blockers:
        print("  BLOCKER", b)
    sys.exit(2 if blockers else 0)


if __name__ == "__main__":
    main()
