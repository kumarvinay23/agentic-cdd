#!/usr/bin/env python3
"""Step 7-9 (deck side): structural QA, rendering for visual inspection, exception report.

Usage: python qa_deck.py --deck <deck.pptx> --work <work> [--content <work>/content.json] [--dpi 110]

  * structural checks on the .pptx: leftover template instructions / banned text, shapes outside the
    slide, overlapping content shapes, empty variable shapes, slide count/order vs slide_map
  * renders every slide to PNG (LibreOffice -> PDF -> pdftoppm) and builds contact sheets
    (<work>/qa/sheet-N.png) - inspect every one; fix and re-run until clean
  * writes <work>/qa/qa_report.json and <work>/exception_report.md (the concise report delivered
    with the deck: input issues + content exceptions + unresolved check warnings)
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.common import load_yaml, read_json, slide_map, write_json  # noqa: E402

SEV_ORDER = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "INFO": 4}
INSTR = re.compile(r"\[(TEXT|BULLETS|TABLE|CHART|KPI|AUTO)\b|\[Action title|\[Supporting message|\[Source line|\[Project name\]|<target name>")


def all_text(shape):
    out = []
    if shape.has_text_frame:
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for r in shape.table.rows:
            for c in r.cells:
                out.append(c.text_frame.text)
    return "\n".join(out)


def to_pdf(deck, outdir):
    wrapper = "/mnt/skills/public/pptx/scripts/office/soffice.py"
    if os.path.exists(wrapper):
        cmd = [sys.executable, wrapper, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(deck)]
    else:
        prof = tempfile.mkdtemp()
        cmd = ["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(deck)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=600)
    pdf = Path(outdir) / (Path(deck).stem + ".pdf")
    return pdf


def contact_sheets(pngs, outdir, cols=3, rows=3, width=640):
    sheets = []
    per = cols * rows
    for k in range(0, len(pngs), per):
        batch = pngs[k:k + per]
        ims = [Image.open(p) for p in batch]
        h = int(ims[0].height * width / ims[0].width)
        sheet = Image.new("RGB", (cols * (width + 12) + 12, rows * (h + 34) + 12), "white")
        d = ImageDraw.Draw(sheet)
        for i, (p, im) in enumerate(zip(batch, ims)):
            x = 12 + (i % cols) * (width + 12)
            y = 12 + (i // cols) * (h + 34)
            sheet.paste(im.resize((width, h)), (x, y + 22))
            d.text((x, y + 4), Path(p).stem, fill="black")
        path = Path(outdir) / f"sheet-{k // per + 1}.png"
        sheet.save(path)
        sheets.append(str(path))
    return sheets


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--work", required=True)
    ap.add_argument("--content")
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--no-render", action="store_true")
    a = ap.parse_args()
    work = Path(a.work)
    qa = work / "qa"
    qa.mkdir(exist_ok=True)
    sm = slide_map()
    banned_cfg = load_yaml("banned_terms.yaml")
    banned = [re.compile(p, re.I) for p in banned_cfg.get("patterns", [])]
    banned += [re.compile(r"\b" + re.escape(t) + r"\b", re.I) for t in banned_cfg.get("reference_report_terms", []) + banned_cfg.get("project_terms", [])]
    prs = Presentation(a.deck)
    W, H = prs.slide_width, prs.slide_height
    findings = []

    order = [s._element.cSld.get("name") for s in prs.slides]
    base_order = [re.sub(r"__\d+$", "", n or "") for n in order]
    # Coverage gating removes slides by design; the expected sequence is the template order
    # minus whatever gates.json withheld, in the same order.
    gated = set()
    gp = Path(a.work) / "gates.json"
    if gp.exists():
        gated = {g["id"] for g in json.loads(gp.read_text()) if not g.get("render")}
    try:
        cj = json.loads(Path(a.content).read_text()) if a.content else {}
    except Exception:  # noqa: BLE001
        cj = {}
    gated |= {w["id"] for w in cj.get("withhold", []) or []}
    alive = {s["section"] for s in sm["slides"] if s["id"] not in gated
             and s["archetype"] not in ("divider", "back_cover")}
    gated |= {s["id"] for s in sm["slides"] if s["archetype"] == "divider" and s["section"] not in alive}
    back = cj.get("glossary_position", "back") == "back"
    tail_ids = {"S007_glossary"} if back else set()
    expected = [s["id"] for s in sm["slides"] if s["id"] not in gated and s["archetype"] != "back_cover"
                and s["id"] not in tail_ids]
    expected += [s["id"] for s in sm["slides"] if s["id"] in tail_ids and s["id"] not in gated]
    expected += [s["id"] for s in sm["slides"] if s["id"] not in gated and s["archetype"] == "back_cover"]
    dedup = [x for i, x in enumerate(base_order) if i == 0 or x != base_order[i - 1]]
    if dedup != expected:
        findings.append({"severity": "HIGH", "slide": "-", "issue": "slide order/count differs from template sequence"})
    for idx, s in enumerate(prs.slides, 1):
        name = s._element.cSld.get("name")
        shapes = list(s.shapes)
        boxes = []
        for sh in shapes:
            txt = all_text(sh)
            if INSTR.search(txt):
                findings.append({"severity": "HIGH", "slide": f"{idx} {name}", "issue": f"template instruction left in '{sh.name}': {txt[:80]}"})
            for rx in banned:
                m = rx.search(txt)
                if m:
                    findings.append({"severity": "HIGH", "slide": f"{idx} {name}", "issue": f"banned text '{m.group(0)}' in '{sh.name}'"})
            if sh.left is not None and (sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 5 or sh.top + sh.height > H + 5):
                findings.append({"severity": "HIGH", "slide": f"{idx} {name}", "issue": f"'{sh.name}' extends beyond the slide"})
            if sh.name in ("title", "subtitle", "source") and not txt.strip():
                findings.append({"severity": "MEDIUM", "slide": f"{idx} {name}", "issue": f"'{sh.name}' is empty"})
            if sh.name not in ("title", "subtitle", "section_tag", "source") and not sh.name.startswith(("motif", "div_", "cover_")):
                boxes.append((sh.name, sh.left, sh.top, sh.left + sh.width, sh.top + sh.height))
        # overlapping content shapes (heads sit above their panels by design, so only same-row regions count)
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                n1, a1, b1, c1, d1 = boxes[i]
                n2, a2, b2, c2, d2 = boxes[j]
                ix = max(0, min(c1, c2) - max(a1, a2))
                iy = max(0, min(d1, d2) - max(b1, b2))
                area = min((c1 - a1) * (d1 - b1), (c2 - a2) * (d2 - b2)) or 1
                if ix * iy / area > 0.03:
                    findings.append({"severity": "HIGH", "slide": f"{idx} {name}", "issue": f"'{n1}' overlaps '{n2}'"})
        if s.has_notes_slide and INSTR.search(s.notes_slide.notes_text_frame.text or ""):
            findings.append({"severity": "LOW", "slide": f"{idx} {name}", "issue": "instruction marker in notes"})

    pngs, sheets = [], []
    if not a.no_render:
        for f in glob.glob(str(qa / "slide-*.png")):
            os.remove(f)
        pdf = to_pdf(a.deck, qa)
        subprocess.run(["pdftoppm", "-r", str(a.dpi), "-png", str(pdf), str(qa / "slide")], check=True)
        pngs = sorted(glob.glob(str(qa / "slide-*.png")))
        sheets = contact_sheets(pngs, qa)

    # ---- exception report
    rows = []
    if (work / "input_validation.json").exists():
        for i in read_json(work / "input_validation.json")["issues"]:
            rows.append({"severity": i["severity"], "issue": i["issue"], "slides": "-", "source": "input validation", "treatment": "Recorded; affected content marked"})
    content = read_json(a.content) if a.content and Path(a.content).exists() else {}
    for e in content.get("exceptions", []):
        rows.append({"severity": e["severity"], "issue": e["issue"], "slides": ", ".join(e.get("slides", [])) or "-",
                     "source": e.get("source", ""), "treatment": e.get("treatment", "")})
    if (work / "check_report.json").exists():
        cr = read_json(work / "check_report.json")
        for er in cr.get("errors", []):
            rows.append({"severity": "HIGH", "issue": f"Unresolved content check: {er['msg']}", "slides": er["where"], "source": "check_content", "treatment": "Must be fixed before issue"})
        wc = Counter(re.sub(r" [A-Z]{3}-\S+.*$", "", w["msg"]) for w in cr.get("warnings", []))
        for msg, n in wc.most_common():
            rows.append({"severity": "INFO", "issue": f"{msg} ({n} occurrences)", "slides": "various", "source": "check_content",
                         "treatment": "Figures verified visually where used; see notes"})
    for f in findings:
        rows.append({"severity": f["severity"], "issue": f["issue"], "slides": f["slide"], "source": "qa_deck", "treatment": "Fix and re-render"})
    rows.sort(key=lambda r: SEV_ORDER.get(r["severity"], 9))
    md = ["# Exception report", "", f"Deck: `{Path(a.deck).name}`  |  slides: {len(prs.slides)}", "",
          "| Severity | Issue | Slide(s) | Source | Treatment |", "|---|---|---|---|---|"]
    for r in rows:
        md.append(f"| {r['severity']} | {r['issue'].replace('|', '/')} | {r['slides']} | {r['source'].replace('|','/')} | {r['treatment'].replace('|','/')} |")
    (work / "exception_report.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    write_json(qa / "qa_report.json", {"findings": findings, "pngs": pngs, "sheets": sheets})
    print(f"[qa] structural findings {len(findings)}; rendered {len(pngs)} slides; sheets: {sheets}")
    for f in findings:
        print("  ", f)


if __name__ == "__main__":
    main()
